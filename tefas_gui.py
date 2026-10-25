#!/usr/bin/env python3
"""
TEFAS Fon Analiz — Streamlit Arayüz
====================================
Mevcut tefas_core motorunu kullanır.
Çalıştır: streamlit run tefas_gui.py
"""

from __future__ import annotations

import io
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st

# ── Core motor ──────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

import tefas_core as core  # noqa: E402
import tefas_filtre  # noqa: E402
import tefas_ek as TE  # noqa: E402
import tefas_portfoyum  # noqa: E402
import tefas_bist  # noqa: E402
import tefas_panel  # noqa: E402
# ── Sayfa ayarı ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="M.Durmaz · TEFAS Fon Analiz",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── CSS ─────────────────────────────────────────────────────────────────────
st.markdown(
    """
<style>
  /* ═══════════════════════════════════════════
     DESIGN TOKENS (tek kaynak — tüm sayfalar)
     ═══════════════════════════════════════════ */
  :root {
    --bg: #0b1220;
    --bg2: #0f172a;
    --surface: #1e293b;
    --surface2: #162032;
    --line: #334155;
    --line2: #1e293b;
    --text: #f1f5f9;
    --muted: #94a3b8;
    --muted2: #64748b;
    --accent: #3b82f6;
    --accent2: #0ea5e9;
    --accent-dim: #1d4ed8;
    --pos: #3ddc97;
    --pos-bg: #0f3d32;
    --neg: #ff6b6b;
    --neg-bg: #3d1a1a;
    --warn: #fcd34d;
    --radius: 14px;
    --radius-sm: 10px;
    --shadow: 0 8px 28px rgba(0,0,0,.28);
  }

  .block-container {
    padding-top: 1.1rem; padding-bottom: 3rem; max-width: 1280px;
  }
  html, body, [class*="css"] {
    font-family: "Segoe UI", "Inter", system-ui, -apple-system, sans-serif !important;
  }
  h1, h2, h3 { letter-spacing: -0.02em; color: var(--text); }

  /* Sidebar */
  [data-testid="stSidebar"] {
    background: linear-gradient(180deg, #0b1220 0%, #0f172a 55%, #111827 100%);
    border-right: 1px solid var(--line2);
  }
  [data-testid="stSidebar"] * { color: #e2e8f0 !important; }
  [data-testid="stSidebar"] .stRadio label { font-size: 0.9rem; padding: 4px 2px; }
  div[data-testid="stSidebar"] .stButton > button {
    width: 100%; border-radius: var(--radius-sm);
    background: var(--surface); border: 1px solid var(--line); color: #e2e8f0;
  }
  div[data-testid="stSidebar"] .stButton > button:hover {
    border-color: var(--accent); color: #93c5fd;
  }

  /* Hero */
  .hero {
    background: linear-gradient(125deg, #0f2744 0%, var(--accent-dim) 48%, #0e7490 100%);
    color: white; border-radius: 18px; padding: 24px 28px; margin-bottom: 18px;
    box-shadow: var(--shadow); border: 1px solid rgba(255,255,255,0.08);
  }
  .hero h1 { color: white !important; margin: 0 0 8px; font-size: 1.55rem; }
  .hero p { margin: 0; opacity: 0.92; font-size: 0.95rem; line-height: 1.5; }

  /* Cards */
  .flow-step, .token-card {
    background: linear-gradient(145deg, var(--surface), var(--bg2));
    border: 1px solid var(--line); border-radius: var(--radius);
    padding: 16px 18px; color: var(--text); box-shadow: var(--shadow);
  }
  .flow-step:hover, .token-card:hover { border-color: var(--accent); }
  .flow-step b { color: #60a5fa; }

  .badge {
    display:inline-block; padding:3px 10px; border-radius:999px;
    font-size:0.75rem; font-weight:600; margin-right:4px;
  }
  .badge-ok  { background: var(--pos-bg); color: var(--pos); }
  .badge-warn{ background:#78350f; color: var(--warn); }
  .badge-bad { background: var(--neg-bg); color: var(--neg); }

  /* Metrics */
  [data-testid="stMetric"] {
    background: linear-gradient(160deg, #1e293bcc, #0f172acc);
    border: 1px solid var(--line); border-radius: var(--radius);
    padding: 12px 14px; box-shadow: 0 4px 16px rgba(0,0,0,.18);
  }
  [data-testid="stMetricLabel"] { color: var(--muted) !important; font-size: 0.8rem !important; }
  [data-testid="stMetricValue"] { color: var(--text) !important; font-weight: 700 !important; }

  /* Tabs */
  .stTabs [data-baseweb="tab-list"] {
    gap: 6px; background: transparent; border-bottom: 1px solid var(--line2);
  }
  .stTabs [data-baseweb="tab"] {
    border-radius: 10px 10px 0 0; padding: 8px 14px;
    background: var(--bg2); border: 1px solid var(--line2); color: var(--muted);
  }
  .stTabs [aria-selected="true"] {
    background: var(--surface) !important; color: #93c5fd !important;
    border-color: var(--accent) !important;
  }

  /* Buttons */
  button[data-testid="baseButton-primary"],
  .stButton > button[kind="primary"] {
    border-radius: var(--radius-sm) !important;
    background: linear-gradient(135deg, var(--accent), var(--accent2)) !important;
    border: none !important; font-weight: 600 !important;
    box-shadow: 0 4px 14px rgba(59,130,246,.35) !important;
  }
  .stButton > button { border-radius: var(--radius-sm); }

  /* Data / expander / charts */
  [data-testid="stDataFrame"] {
    border: 1px solid var(--line2); border-radius: 12px; overflow: hidden;
  }
  [data-testid="stExpander"] {
    border: 1px solid var(--line2) !important; border-radius: 12px !important;
    background: #0f172a88;
  }
  [data-testid="stPlotlyChart"] {
    border: 1px solid var(--line2); border-radius: var(--radius);
    padding: 6px; background: var(--bg);
  }
  [data-testid="stImage"] img { border-radius: 12px; }

  .stProgress > div > div > div > div {
    background: linear-gradient(90deg, var(--pos), var(--accent));
  }

  .sepet-box {
    background: var(--surface); border: 1px solid var(--line);
    border-radius: 12px; padding: 12px 14px; margin-top: 8px;
  }
  .sepet-badge {
    display:inline-block;
    background: linear-gradient(135deg, var(--accent-dim), var(--accent2));
    color:#fff; border-radius:999px; padding:3px 10px;
    font-size:0.75rem; font-weight:700;
  }
  .disclaimer {
    margin-top: 2.2rem; padding: 14px 18px; border-radius: 12px;
    background: var(--bg2); border: 1px solid var(--line); color: var(--muted);
    font-size: 0.8rem; line-height: 1.5;
  }

  /* Portföyüm tokens (fy-*) — aynı palet */
  .fy-card {
    background: linear-gradient(145deg, var(--surface), var(--bg2));
    border: 1px solid var(--line); border-radius: var(--radius);
    padding: 14px 16px; margin-bottom: 10px; box-shadow: var(--shadow);
  }
  .fy-card .lbl {
    font-size: 0.72rem; letter-spacing: .06em; text-transform: uppercase;
    color: var(--muted2); margin-bottom: 4px;
  }
  .fy-card .val { font-size: 1.45rem; font-weight: 700; color: var(--text); }
  .fy-card .sub { font-size: 0.85rem; color: var(--muted); margin-top: 2px; }
  .fy-pos { color: var(--pos) !important; }
  .fy-neg { color: var(--neg) !important; }
  .fy-hero {
    background: linear-gradient(120deg, #0f2744, #163a2a 40%, #0e4d6b);
    border: 1px solid #1e3d32; border-radius: 16px;
    padding: 18px 20px; margin-bottom: 14px; box-shadow: var(--shadow);
  }
  .fy-hero h2 { margin: 0 0 4px; color: var(--text); font-size: 1.25rem; }
  .fy-hero p { margin: 0; color: var(--muted); font-size: 0.9rem; }

  /* Benchmark panel */
  .bm-panel {
    background: linear-gradient(160deg, #0f172a, #132337);
    border: 1px solid var(--line); border-radius: 16px;
    padding: 18px 20px; margin: 12px 0 18px; box-shadow: var(--shadow);
  }
  .bm-panel h3 {
    margin: 0 0 12px; font-size: 1.05rem; color: var(--text);
    display: flex; align-items: center; gap: 8px;
  }
  .bm-chip {
    display: inline-flex; align-items: center; gap: 6px;
    background: var(--surface); border: 1px solid var(--line);
    border-radius: 999px; padding: 4px 12px; font-size: 0.8rem; color: var(--muted);
    margin-right: 6px; margin-bottom: 6px;
  }
  .bm-chip strong { color: var(--text); }
  .rb-bar-wrap { margin: 8px 0 12px; }
  .rb-label {
    display: flex; justify-content: space-between; font-size: 0.85rem;
    color: var(--muted); margin-bottom: 4px;
  }
  .rb-label b { color: var(--text); }


  .sticky-strip {
    position: sticky; top: 0; z-index: 20;
    display: flex; flex-wrap: wrap; gap: 8px; align-items: center;
    background: linear-gradient(90deg, #0f172aee, #132337ee);
    border: 1px solid var(--line); border-radius: 14px;
    padding: 10px 14px; margin: 8px 0 14px;
    backdrop-filter: blur(8px); box-shadow: var(--shadow);
  }
  .sticky-item {
    min-width: 90px; padding: 4px 10px; border-right: 1px solid var(--line2);
  }
  .sticky-item:last-child { border-right: none; }
  .sticky-strip .sk {
    font-size: 0.7rem; text-transform: uppercase; letter-spacing: .05em; color: var(--muted2);
  }
  .sticky-strip .sv { font-size: 0.98rem; font-weight: 700; color: var(--text); }

  ::-webkit-scrollbar { width: 8px; height: 8px; }
  ::-webkit-scrollbar-thumb { background: var(--line); border-radius: 8px; }
  ::-webkit-scrollbar-track { background: var(--bg2); }
</style>
""",
    unsafe_allow_html=True,
)



# ── Session state ───────────────────────────────────────────────────────────
def _init_state():
    defaults = {
        "cfg": core.load_config(),
        "fon_tipi": "YAT",
        "analiz": None,
        "son_tarama": None,
        "son_tarama_meta": {},
        "son_siralama": None,
        "son_portfoy": None,
        "son_karsilastir": None,
        "detay_sonuc": None,
        "opt_sonuc": None,
        "akilli_portfoy": None,
        "akilli_portfoy_sira": None,
        "hisse_sonuc": None,
        "karma_sonuc": None,
        "hisse_liste": "THYAO,GARAN,ASELS",
        "portfoy_sepet": [],
        "sepet_analiz": None,
        "sayfa": "🏠 Ana Sayfa",
        "sepet_auto_analiz": False,
        "app_version": "2.1.0",
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


_init_state()


def get_analiz() -> core.FonAnaliz:
    """FonAnaliz motorunu session'da tutar; tip/cfg değişince yeniler."""
    a = st.session_state.get("analiz")
    tip = st.session_state.get("fon_tipi", "YAT")
    cfg = st.session_state.get("cfg") or core.load_config()
    need_new = a is None
    if not need_new:
        try:
            need_new = getattr(a.api, "fon_tipi", None) != tip
        except Exception:
            need_new = True
    if need_new:
        st.session_state.analiz = core.FonAnaliz(fon_tipi=tip, cfg=cfg)
    else:
        try:
            st.session_state.analiz.cfg = cfg
            st.session_state.analiz.rf = float(cfg.get("risk_free", 0.40))
        except Exception:
            st.session_state.analiz = core.FonAnaliz(fon_tipi=tip, cfg=cfg)
    return st.session_state.analiz



def sepet_ekle(kodlar, ad_map=None, kaynak="manuel"):
    """Fon kodlarını portföy sepetine ekler (tekrarsız, normalize)."""
    ad_map = ad_map or {}
    sepet = list(st.session_state.get("portfoy_sepet") or [])
    mevcut = {str(x.get("kod")).upper() for x in sepet}
    eklendi = 0
    for k in kodlar:
        k = str(k).strip().upper().replace(" ", "")
        if not k or k in mevcut:
            continue
        # ad_map anahtarları da normalize
        ad = ad_map.get(k) or ad_map.get(str(k)) or k
        sepet.append({"kod": k, "ad": ad, "kaynak": kaynak})
        mevcut.add(k)
        eklendi += 1
    st.session_state.portfoy_sepet = sepet
    return eklendi


def sepet_analiz_et(weights: dict | None = None):
    """
    Sepet derin raporu:
    - fon bazlı metrik tablosu
    - eşit veya özel ağırlıklı birleşik portföy
    - korelasyon matrisi, çeşitlendirme skoru
    - kategori / risk dağılımı
    - en iyi / en zayıf fonlar
    - sade dilde genel değerlendirme
    """
    sepet = st.session_state.get("portfoy_sepet") or []
    if not sepet:
        return {"error": "Sepet boş"}
    analiz = get_analiz()
    rows = []
    series = {}
    dagilim_list = []

    for item in sepet:
        kod = item["kod"]
        try:
            d = analiz.fon_detay(kod)
            m = d["metrikler"]
            o, p = d["ozet"], d["profil"]
            ad = o.get("fonUnvan") or p.get("fonUnvan") or item.get("ad") or kod
            item["ad"] = ad
            row = {
                "kod": kod, "ad": (ad or "")[:55], "kaynak": item.get("kaynak"),
                "getiri_1a": m.get("getiri_1a"), "getiri_3a": m.get("getiri_3a"),
                "getiri_6a": m.get("getiri_6a"), "getiri_1y": m.get("getiri_1y"),
                "getiri_3y": m.get("getiri_3y"),
                "reel_1y": m.get("reel_1y"), "vs_mevduat": m.get("vs_mevduat"),
                "vs_bist100": m.get("vs_bist100"),
                "sharpe": m.get("sharpe"), "sortino": m.get("sortino"),
                "calmar": m.get("calmar"), "cagr": m.get("cagr"),
                "volatilite": m.get("volatilite"), "max_drawdown": m.get("max_drawdown"),
                "kazanma_orani": m.get("kazanma_orani"),
                "en_kotu_gun": m.get("en_kotu_gun"), "en_iyi_gun": m.get("en_iyi_gun"),
                "var95": m.get("var95"), "cvar95": m.get("cvar95"),
                "risk_degeri": p.get("riskDegeri"),
                "kategori": o.get("fonKategori") or p.get("fonKategori"),
                "fiyat": m.get("fiyat"),
            }
            # dağılım özeti
            try:
                dag = analiz.api.dagilimlar()
                if dag is not None and not dag.empty:
                    rr = dag[dag["kod"].astype(str).str.upper() == kod]
                    if not rr.empty:
                        oz = analiz.api.dagilim_ozet(rr.iloc[0])
                        for kk, vv in (oz or {}).items():
                            row[f"d_{kk}"] = vv
                        dagilim_list.append((kod, oz or {}))
            except Exception:
                pass
            rows.append(row)
            fd = d.get("fiyat_df")
            if fd is None or len(fd) < 60:
                try:
                    fd = analiz.api.fiyat(kod, periyod=36)
                except Exception:
                    fd = None
            if fd is not None and not fd.empty:
                s = fd.set_index("tarih")["fiyat"].astype(float).sort_index()
                s = s[~s.index.duplicated(keep="last")]
                series[kod] = s
        except Exception as e:
            rows.append({"kod": kod, "ad": item.get("ad"), "hata": str(e), "kaynak": item.get("kaynak")})

    df = pd.DataFrame(rows)
    if "hata" in df.columns:
        ok = df[df["hata"].isna()].copy()
    else:
        ok = df.copy()

    # Ağırlıklar
    kodlar_ok = ok["kod"].astype(str).tolist() if not ok.empty else []
    if weights:
        w_raw = {str(k).upper(): float(v) for k, v in weights.items() if str(k).upper() in kodlar_ok and float(v) > 0}
    else:
        w_raw = {k: 1.0 for k in kodlar_ok}
    tot = sum(w_raw.values()) or 1.0
    w_norm = {k: v / tot for k, v in w_raw.items()}

    # Birleşik portföy serisi
    port_m = None
    corr_df = None
    diversify_score = None
    port_ret = None
    if len(series) >= 1 and w_norm:
        use_cols = [c for c in series if c in w_norm]
        if use_cols:
            dfp = pd.DataFrame({c: series[c] for c in use_cols}).dropna(how="all").ffill().dropna()
            if len(dfp) >= 30:
                rets = dfp.pct_change().dropna()
                if len(use_cols) >= 2 and len(rets) >= 40:
                    try:
                        corr_df = rets[use_cols].corr()
                        # çeşitlendirme: ortalama |korelasyon| ne kadar düşükse o kadar iyi
                        mask = np.triu(np.ones(corr_df.shape), k=1).astype(bool)
                        avg_abs = float(np.abs(corr_df.where(mask)).stack().mean())
                        diversify_score = round(max(0, min(100, (1 - avg_abs) * 100)), 1)
                    except Exception:
                        corr_df = None
                ww = np.array([w_norm[c] for c in use_cols])
                ww = ww / ww.sum()
                port_ret = (rets[use_cols] * ww).sum(axis=1)
                port_px = (1 + port_ret).cumprod() * 100
                fd_port = pd.DataFrame({"tarih": port_px.index, "fiyat": port_px.values})
                port_m = core.full_metrics(fd_port, rf=analiz.rf)
                try:
                    port_m = core.attach_benchmark(port_m, core.fetch_benchmarks(st.session_state.cfg))
                except Exception:
                    pass
                for i, c in enumerate(use_cols):
                    w_norm[c] = float(ww[i])

    # Birleşik varlık dağılımı (ağırlıklı)
    birlesik_dag = {}
    if dagilim_list and w_norm:
        keys = set()
        for _, oz in dagilim_list:
            keys.update(oz.keys())
        for k in keys:
            s = 0.0
            for kod, oz in dagilim_list:
                s += float(oz.get(k) or 0) * float(w_norm.get(kod) or 0)
            if s > 0.05:
                birlesik_dag[k] = round(s, 1)

    # En iyi / en zayıf
    liderler = {}
    if not ok.empty:
        for col, label, asc in [
            ("getiri_1y", "en_yuksek_1y", False),
            ("sharpe", "en_iyi_sharpe", False),
            ("max_drawdown", "en_dusuk_dd", False),  # max_dd is negative, higher is better
            ("volatilite", "en_dusuk_vol", True),
        ]:
            if col in ok.columns and ok[col].notna().any():
                sub = ok.dropna(subset=[col])
                if sub.empty:
                    continue
                r = sub.sort_values(col, ascending=asc).iloc[0]
                liderler[label] = {"kod": r["kod"], "ad": r.get("ad"), "deger": r[col]}

    # Kategori dağılımı
    kat_dag = {}
    if not ok.empty and "kategori" in ok.columns:
        for _, r in ok.iterrows():
            k = str(r.get("kategori") or "Diğer")[:40]
            kat_dag[k] = kat_dag.get(k, 0) + float(w_norm.get(str(r["kod"]), 0)) * 100

    # Ağırlıklı ortalama tehlike + getiri (tüm portföy)
    def _wavg(col):
        if ok.empty or col not in ok.columns:
            return None
        s = pd.to_numeric(ok[col], errors="coerce")
        ww = ok["kod"].astype(str).map(lambda k: float(w_norm.get(k, 0)))
        mask = s.notna() & (ww > 0)
        if not mask.any():
            return None
        return float((s[mask] * ww[mask]).sum() / ww[mask].sum())

    def _simple_avg(col):
        if ok.empty or col not in ok.columns:
            return None
        s = pd.to_numeric(ok[col], errors="coerce")
        return float(s.mean()) if s.notna().any() else None

    ort_tehlike = _wavg("risk_degeri")  # SPK 1–7
    ort_getiri_1a = _wavg("getiri_1a")
    ort_getiri_3a = _wavg("getiri_3a")
    ort_getiri_6a = _wavg("getiri_6a")
    ort_getiri_1y = _wavg("getiri_1y")
    # Birleşik portföy getirisi varsa onu da tercih et
    if port_m:
        if port_m.get("getiri_1y") is not None:
            ort_getiri_1y = port_m.get("getiri_1y")
        if port_m.get("getiri_3a") is not None:
            ort_getiri_3a = port_m.get("getiri_3a")
        if port_m.get("getiri_1a") is not None:
            ort_getiri_1a = port_m.get("getiri_1a")
        if port_m.get("getiri_6a") is not None:
            ort_getiri_6a = port_m.get("getiri_6a")

    def _tehlike_yazi(v):
        if v is None:
            return "—"
        try:
            r = float(v)
        except Exception:
            return "—"
        etiket = (
            "Çok düşük" if r < 1.5 else
            "Düşük" if r < 2.5 else
            "Orta-düşük" if r < 3.5 else
            "Orta" if r < 4.5 else
            "Orta-yüksek" if r < 5.5 else
            "Yüksek" if r < 6.5 else
            "Çok yüksek"
        )
        return f"{r:.1f} · {etiket}"

    # Özet markdown
    lines = [f"### Portföy sepeti derin rapor ({len(sepet)} fon)", ""]
    lines.append("#### Tüm portföy — ortalama tehlike & getiri")
    lines.append(f"- **Ortalama tehlike seviyesi (ağırlıklı):** {_tehlike_yazi(ort_tehlike)} (SPK 1–7)")
    if ort_getiri_1a is not None:
        lines.append(f"- **Ortalama getiri 1A:** **%{ort_getiri_1a:.1f}**")
    if ort_getiri_3a is not None:
        lines.append(f"- **Ortalama getiri 3A:** **%{ort_getiri_3a:.1f}**")
    if ort_getiri_6a is not None:
        lines.append(f"- **Ortalama getiri 6A:** **%{ort_getiri_6a:.1f}**")
    if ort_getiri_1y is not None:
        lines.append(f"- **Ortalama getiri 1Y:** **%{ort_getiri_1y:.1f}**")
    lines.append("")
    if not ok.empty:
        a1y, ash, add, avol = _simple_avg("getiri_1y"), _simple_avg("sharpe"), _simple_avg("max_drawdown"), _simple_avg("volatilite")
        lines.append("#### Sepet ortalamaları (fon bazlı, eşit ağırlık)")
        if a1y is not None:
            lines.append(f"- Ortalama 1Y getiri: **%{a1y:.1f}**")
        if ash is not None:
            lines.append(f"- Ortalama Sharpe: **{ash:.2f}**")
        if add is not None:
            lines.append(f"- Ortalama MaxDD: **%{add:.1f}**")
        if avol is not None:
            lines.append(f"- Ortalama oynaklık: **%{avol:.1f}**")
        lines.append("")
    if liderler:
        lines.append("#### Öne çıkanlar")
        if "en_yuksek_1y" in liderler:
            x = liderler["en_yuksek_1y"]
            lines.append(f"- En yüksek 1Y: **{x['kod']}** ({x['ad']}) · %{x['deger']:.1f}")
        if "en_iyi_sharpe" in liderler:
            x = liderler["en_iyi_sharpe"]
            lines.append(f"- En iyi Sharpe: **{x['kod']}** · {x['deger']:.2f}")
        if "en_dusuk_dd" in liderler:
            x = liderler["en_dusuk_dd"]
            lines.append(f"- En sınırlı düşüş: **{x['kod']}** · %{x['deger']:.1f}")
        if "en_dusuk_vol" in liderler:
            x = liderler["en_dusuk_vol"]
            lines.append(f"- En düşük vol: **{x['kod']}** · %{x['deger']:.1f}")
        lines.append("")
    if diversify_score is not None:
        yorum = (
            "fonlar birbirinden görece bağımsız hareket ediyor"
            if diversify_score >= 55 else
            "fonlar orta düzeyde birlikte hareket ediyor"
            if diversify_score >= 35 else
            "fonlar yüksek korelasyonlu — çeşitlendirme zayıf olabilir"
        )
        lines.append(f"#### Çeşitlendirme skoru: **{diversify_score}/100**")
        lines.append(f"- {yorum} (ortalama mutlak korelasyona dayalı)")
        lines.append("")
    if kat_dag:
        lines.append("#### Kategori ağırlıkları")
        for k, v in sorted(kat_dag.items(), key=lambda x: -x[1]):
            lines.append(f"- {k}: **%{v:.0f}**")
        lines.append("")
    if birlesik_dag:
        lines.append("#### Birleşik varlık kırılımı (ağırlıklı)")
        labels = {"hisse": "Hisse", "tahvil": "Tahvil", "repo": "Repo", "mevduat": "Mevduat",
                  "altin": "Altın", "doviz": "Döviz", "fon": "Fon", "diger": "Diğer"}
        for k, v in sorted(birlesik_dag.items(), key=lambda x: -x[1]):
            lines.append(f"- {labels.get(k, k)}: **%{v:.1f}**")
        lines.append("")
    if port_m:
        lines.append("#### Birleşik portföy (seçilen ağırlıklar)")
        lines.append(
            core.turkce_degerlendirme(
                sharpe=port_m.get("sharpe"), sortino=port_m.get("sortino"),
                max_dd=port_m.get("max_drawdown"), vol=port_m.get("volatilite"),
                cagr=port_m.get("cagr"), calmar=port_m.get("calmar"),
                getiri_1y=port_m.get("getiri_1y"), reel_1y=port_m.get("reel_1y"),
                vs_mevduat=port_m.get("vs_mevduat"), vs_bist100=port_m.get("vs_bist100"),
                kazanma_orani=port_m.get("kazanma_orani"),
                en_kotu_gun=port_m.get("en_kotu_gun"),
                plain=True,
            )
        )
        lines.append("")
    lines.append("#### Fon listesi")
    for _, r in df.iterrows():
        if pd.notna(r.get("hata")) and str(r.get("hata")) not in ("", "nan", "None"):
            lines.append(f"- `{r['kod']}`: ⚠ {r['hata']}")
        else:
            wpct = w_norm.get(str(r["kod"]), 0) * 100
            lines.append(
                f"- `{r['kod']}` %{wpct:.0f} · {r.get('ad') or ''} — "
                f"1Y %{fmt_num(r.get('getiri_1y'))} · Sh {fmt_num(r.get('sharpe'), 2)} · "
                f"DD {fmt_num(r.get('max_drawdown'))}% · vol {fmt_num(r.get('volatilite'))}%"
            )
    lines.append("")
    lines.append("*Geçmiş performans geleceği garanti etmez; yatırım tavsiyesi değildir.*")

    # Portföy fiyat serisi (HTML / quantstats için)
    port_fd = None
    if port_ret is not None and len(port_ret) >= 20:
        try:
            port_px = (1 + port_ret).cumprod() * 100
            port_fd = pd.DataFrame({"tarih": port_px.index, "fiyat": port_px.values})
        except Exception:
            port_fd = None

    # Çeşitlendirme önerileri (aday = son tarama/sıralama)
    oneriler = []
    try:
        aday = None
        for key in ("son_siralama", "son_tarama"):
            cand = st.session_state.get(key)
            if cand is not None and hasattr(cand, "empty") and not cand.empty:
                aday = cand
                break
        if series and len(series) >= 1:
            oneriler = core.cesitlendirme_onerileri(
                analiz, series, aday_df=aday, max_aday=35, top_n=8,
                mevcut_kodlar=set(series.keys()),
            )
            if oneriler:
                lines.append("")
                lines.append("#### Çeşitlendirme önerileri (TEFAS açık · tehlike + getiri)")
                for o in oneriler[:6]:
                    lines.append(
                        f"- **{o['kod']}** — {o.get('ad','')} · "
                        f"tehlike: {o.get('tehlike_seviyesi') or '—'} · "
                        f"1Y %{o.get('getiri_1y') if o.get('getiri_1y') is not None else '—'} · "
                        f"3A %{o.get('getiri_3a') if o.get('getiri_3a') is not None else '—'} · "
                        f"ort.ρ {o.get('ortalama_korelasyon')} · "
                        f"çeşitl. {o.get('beklenen_cesitlilik')}/100 · {o.get('neden')}"
                    )
    except Exception as e:
        oneriler = []
        lines.append(f"\n_Çeşitlendirme önerisi hesaplanamadı: {e}_")

    return {
        "df": df,
        "ok": ok,
        "port_metrikler": port_m,
        "ozet_md": "\n".join(lines),
        "weights": w_norm,
        "corr": corr_df,
        "diversify_score": diversify_score,
        "liderler": liderler,
        "kategori_dag": kat_dag,
        "birlesik_dag": birlesik_dag,
        "port_ret": port_ret,
        "port_fd": port_fd,
        "series": series,
        "oneriler": oneriler,
        "n_fon": len(kodlar_ok),
        "ort_tehlike": ort_tehlike,
        "ort_tehlike_yazi": _tehlike_yazi(ort_tehlike),
        "ort_getiri_1a": ort_getiri_1a,
        "ort_getiri_3a": ort_getiri_3a,
        "ort_getiri_6a": ort_getiri_6a,
        "ort_getiri_1y": ort_getiri_1y,
    }



def tr_num(v, digits: int = 2, *, suffix: str = "", force_sign: bool = False) -> str:
    """
    Tüm rapor / tablo sayıları için TR biçim:
    binlik ayırıcı nokta, ondalık ayırıcı virgül → 1.234.567,89
    """
    if v is None:
        return "—"
    try:
        if isinstance(v, float) and (np.isnan(v) or not np.isfinite(v)):
            return "—"
    except Exception:
        pass
    try:
        x = float(v)
    except (TypeError, ValueError):
        return "—"
    try:
        if digits <= 0:
            s = f"{abs(x):,.0f}"
        else:
            s = f"{abs(x):,.{int(digits)}f}"
        # en-US 1,234,567.89 → TR 1.234.567,89
        s = s.replace(",", "X").replace(".", ",").replace("X", ".")
        if x < 0:
            s = "-" + s
        elif force_sign and x > 0:
            s = "+" + s
        return f"{s}{suffix}"
    except Exception:
        return "—"


def fmt_num(v, digits=1, suffix=""):
    """Genel sayı — ondalık/binlik TR ayırıcı."""
    return tr_num(v, digits=digits, suffix=suffix)


def _fmt_tl(v):
    """TL tutar — TR ayırıcı + mlr/mn kısaltma."""
    if v is None:
        return "—"
    try:
        x = float(v)
        if abs(x) >= 1e9:
            return f"{tr_num(x/1e9, 2)} mlr"
        if abs(x) >= 1e6:
            return f"{tr_num(x/1e6, 1)} mn"
        return tr_num(x, 0 if abs(x) >= 100 else 2)
    except (TypeError, ValueError):
        return "—"

# Eski / eksik deploy'larda NameError olmasın diye alias
fmt_tl = _fmt_tl
tr_fmt = tr_num



def tarihsel_ucuzluk_analizi(sembol: str, fd: pd.DataFrame | None = None, deg: dict | None = None) -> dict:
    """
    Tarihsel ucuzluk (Filtre Stüdyosu / BIST Değerleme uyumu).
    - Fiyat: 1Y ve 3Y persentil (0=en düşük fiyat=ucuz seviye, 100=zirve)
    - F/K, PD/DD: mevcut çarpan + 52h fiyat konumuna göre yorum
    Skor 0–100: yüksek = tarihsel olarak daha ucuz görünüm.
    """
    deg = deg or {}
    out = {
        "skor": None,
        "fiyat_persentil_1y": None,
        "fiyat_persentil_3y": None,
        "fk": None,
        "pddd": None,
        "yorum": [],
        "ucuz_mu": None,
    }
    try:
        px = None
        if fd is not None and not getattr(fd, "empty", True) and "fiyat" in fd.columns:
            s = fd.copy()
            if "tarih" in s.columns:
                s["tarih"] = pd.to_datetime(s["tarih"])
                s = s.sort_values("tarih")
            px = pd.to_numeric(s["fiyat"], errors="coerce").dropna()
        if px is None or len(px) < 30:
            # yfinance dene
            try:
                import yfinance as yf
                sym = sembol if str(sembol).endswith(".IS") else f"{str(sembol).upper().replace('.IS','')}.IS"
                hist = yf.Ticker(sym).history(period="3y", auto_adjust=True)
                if hist is not None and not hist.empty:
                    px = hist["Close"].dropna()
            except Exception:
                pass
        if px is None or len(px) < 20:
            out["yorum"].append("Yetersiz fiyat serisi — tarihsel ucuzluk hesaplanamadı.")
            return out

        last = float(px.iloc[-1])

        def _pct_rank(series):
            if series is None or len(series) < 5:
                return None
            return float((series <= last).mean() * 100.0)

        p1 = _pct_rank(px.tail(min(252, len(px))))
        p3 = _pct_rank(px.tail(min(756, len(px))))
        out["fiyat_persentil_1y"] = round(p1, 1) if p1 is not None else None
        out["fiyat_persentil_3y"] = round(p3, 1) if p3 is not None else None

        fk = deg.get("F/K")
        pddd = deg.get("PD/DD")
        try:
            fk = float(fk) if fk is not None else None
        except (TypeError, ValueError):
            fk = None
        try:
            pddd = float(pddd) if pddd is not None else None
        except (TypeError, ValueError):
            pddd = None
        out["fk"] = fk
        out["pddd"] = pddd

        # Skor: düşük persentil = ucuz fiyat seviyesi → yüksek skor
        # F/K düşük ve PD/DD < 1 bonus
        skor = 50.0
        if p1 is not None:
            skor = 100.0 - p1  # 0. persentil → 100 puan
            out["yorum"].append(
                f"1Y fiyat persentili **%{p1:.0f}** "
                f"({'ucuz bölge' if p1 <= 30 else ('pahalı bölge' if p1 >= 70 else 'orta bant')})."
            )
        if p3 is not None:
            out["yorum"].append(f"3Y fiyat persentili **%{p3:.0f}**.")
            skor = 0.6 * skor + 0.4 * (100.0 - p3)
        if fk is not None:
            if fk < 8:
                skor += 8
                out["yorum"].append(f"F/K **{fk:.1f}** düşük çarpan (ucuzluk lehine).")
            elif fk > 20:
                skor -= 8
                out["yorum"].append(f"F/K **{fk:.1f}** yüksek — tarihsel ucuz görünümü zayıflatır.")
        if pddd is not None:
            if pddd < 1:
                skor += 6
                out["yorum"].append(f"PD/DD **{pddd:.2f}** defter altı.")
            elif pddd > 2.5:
                skor -= 5
                out["yorum"].append(f"PD/DD **{pddd:.2f}** primli.")

        skor = max(0, min(100, round(skor, 1)))
        out["skor"] = skor
        out["ucuz_mu"] = bool(skor >= 60)
        if skor >= 70:
            out["yorum"].insert(0, f"**Tarihsel ucuzluk skoru: {skor}/100** — görece ucuz bölge.")
        elif skor >= 40:
            out["yorum"].insert(0, f"**Tarihsel ucuzluk skoru: {skor}/100** — nötr.")
        else:
            out["yorum"].insert(0, f"**Tarihsel ucuzluk skoru: {skor}/100** — pahalı / zirveye yakın seviye.")
    except Exception as e:
        out["yorum"].append(f"Hesap hatası: {e}")
    return out







def mali_yoy_karsilastirma_df(df: pd.DataFrame, col_labels: list | None = None) -> pd.DataFrame:
    """
    Güncel dönem vs önceki yıl aynı dönem (ör. 2026/9 vs 2025/9).
    Kolonlar: kod, kalem, güncel, YoY dönem, fark %, güncel biçimli.
    """
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return pd.DataFrame()
    cols = list(col_labels) if col_labels else [c for c in df.columns if c not in ("kod", "kalem")]
    if not cols:
        return pd.DataFrame()
    cur = cols[0]
    # aynı çeyrek önceki yıl
    yoy_col = None
    try:
        y, p = cur.split("/")
        y, p = int(y), int(p)
        cand = f"{y-1}/{p}"
        if cand in df.columns:
            yoy_col = cand
    except Exception:
        pass
    if yoy_col is None and len(cols) > 1:
        yoy_col = cols[1]
    rows = []
    for _, r in df.iterrows():
        a = r.get(cur)
        b = r.get(yoy_col) if yoy_col else None
        try:
            af = float(a) if a is not None and a == a else None
        except (TypeError, ValueError):
            af = None
        try:
            bf = float(b) if b is not None and b == b else None
        except (TypeError, ValueError):
            bf = None
        pct = None
        if af is not None and bf not in (None, 0):
            pct = round((af / bf - 1.0) * 100.0, 1)
        rows.append({
            "kod": r.get("kod"),
            "kalem": r.get("kalem"),
            cur: af,
            (yoy_col or "YoY dönem"): bf,
            "YoY %": pct,
        })
    return pd.DataFrame(rows)


def reel_satis_karlilik(isy: dict | None, tufe: float | None = None) -> dict:
    """
    Nominal YoY satış/kâr − TÜFE = reel büyüme.
    """
    isy = isy or {}
    oz = isy.get("ozet") or {}
    if tufe is None:
        try:
            tufe = float((st.session_state.get("cfg") or {}).get("tufe_yillik") or 0) or None
        except Exception:
            tufe = None
    if tufe is None:
        tufe = 40.0
        src = "varsayılan"
    else:
        src = "Ayarlar"
    out = {"tufe": tufe, "tufe_kaynak": src, "donem": isy.get("gunce_donem"), "satirlar": []}

    def _item(key, label):
        it = oz.get(key) or {}
        if not isinstance(it, dict):
            return
        nom = it.get("degisim_%")
        try:
            nom = float(nom) if nom is not None else None
        except (TypeError, ValueError):
            nom = None
        reel = round(nom - float(tufe), 1) if nom is not None else None
        out["satirlar"].append({
            "Kalem": label,
            "Nominal YoY %": nom,
            "TÜFE %": tufe,
            "Reel % (nom−TÜFE)": reel,
            "Yorum": (
                "Reel büyüme" if reel is not None and reel > 0 else
                ("Enflasyon altı" if reel is not None else "—")
            ),
            "Tutar": it.get("deger"),
            "YoY tutar": it.get("onceki_yoy"),
        })

    _item("hasilat", "Satışlar")
    _item("brut_kar", "Brüt kâr")
    _item("esas_faaliyet_kari", "Faaliyet kârı")
    _item("net_donem_kari", "Net kâr")
    return out


def format_mali_df(df: pd.DataFrame, *, birim: str = "mn") -> pd.DataFrame:
    """
    Bilanço / gelir tablosu hücre biçimi.
    - None / NaN → "—"
    - Büyük tutarlar: mn (milyon) veya mlr (milyar) + TR binlik
    - Küçük sayılar: ondalıklı
    kod / kalem kolonları aynen kalır.
    """
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return df
    out = df.copy()
    skip = {"kod", "kalem", "itemCode", "itemDesc", "Kalem", "Kod"}

    def _cell(v):
        if v is None:
            return "—"
        try:
            if isinstance(v, float) and (v != v):  # NaN
                return "—"
            if isinstance(v, str):
                s = v.strip()
                if s.lower() in ("none", "nan", "", "-"):
                    return "—"
                try:
                    v = float(s.replace(".", "").replace(",", ".")) if s.count(",") == 1 and s.count(".") > 1 else float(s.replace(",", "."))
                except ValueError:
                    return s
            x = float(v)
        except (TypeError, ValueError):
            return "—" if str(v).lower() in ("none", "nan") else str(v)
        if x == 0:
            return "0"
        ax = abs(x)
        sign = "-" if x < 0 else ""
        if birim == "mlr" or ax >= 1e9:
            return f"{sign}{ax/1e9:,.2f} mlr".replace(",", "X").replace(".", ",").replace("X", ".")
        if birim == "mn" or ax >= 1e6:
            return f"{sign}{ax/1e6:,.2f} mn".replace(",", "X").replace(".", ",").replace("X", ".")
        if ax >= 1:
            return f"{sign}{ax:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{sign}{ax:.4f}".replace(".", ",")

    for c in out.columns:
        if str(c) in skip:
            continue
        # dönem kolonları veya sayısal
        try:
            sample = out[c].dropna().head(3)
            if len(sample) == 0:
                out[c] = out[c].apply(lambda v: "—" if v is None or (isinstance(v, float) and v != v) else v)
                continue
        except Exception:
            pass
        out[c] = out[c].map(_cell)
    return out


def isyatirim_carpanlar(kod: str, *, force: bool = False) -> dict:
    """
    İş Yatırım şirket kartından F/K, PD/DD, FD/FAVÖK, fiyat, piyasa değeri.
    Yahoo kullanılmaz — güncel BIST çarpanları için tek kaynak.
    """
    import urllib.request
    import re

    kod = str(kod).upper().replace(".IS", "").strip()
    bugun = datetime.now().strftime("%Y-%m-%d")
    cache = st.session_state.setdefault("_isy_carpan_cache", {})
    if not force and kod in cache and cache[kod].get("_cache_day") == bugun and not cache[kod].get("error"):
        return cache[kod]

    def parse_tr(s: str):
        s = (s or "").strip().replace("\xa0", "").replace(" ", "")
        if not s or s in ("-", "—", "N/A", "n/a"):
            return None
        # 48.213,2 → 48213.2 ; 9,8 → 9.8
        if re.match(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$", s):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None

    def parse_mcap(raw: str):
        if not raw:
            return None
        t = raw.strip().replace("\xa0", " ")
        mult = 1.0
        low = t.lower()
        if "mntl" in low or "mn tl" in low or "mn.tl" in low:
            mult = 1e6
        elif "mlrtl" in low or "mlr tl" in low or "mlr.tl" in low:
            mult = 1e9
        elif "bntl" in low or "bn tl" in low:
            mult = 1e9
        t2 = re.sub(r"[a-zA-ZğüşıöçĞÜŞİÖÇ.]", "", t).strip()
        # leftover like 48.213,2
        val = parse_tr(t2.replace(" ", ""))
        if val is None:
            # try digits only pattern
            m = re.search(r"([\d.,]+)", t)
            if m:
                val = parse_tr(m.group(1))
        return val * mult if val is not None else None

    url = f"https://www.isyatirim.com.tr/tr-tr/analiz/hisse/Sayfalar/sirket-karti.aspx?hisse={kod}"
    try:
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; TEFAS-GUI/2.4)",
                "Accept": "text/html,application/xhtml+xml",
            },
        )
        with urllib.request.urlopen(req, timeout=25) as resp:
            html = resp.read().decode("utf-8", errors="replace")
    except Exception as e:
        out = {"error": str(e), "kod": kod, "_cache_day": bugun}
        cache[kod] = out
        return out

    text = re.sub(r"\s+", " ", html)
    out = {
        "kod": kod,
        "kaynak": "İş Yatırım şirket kartı",
        "url": url,
        "_cache_day": bugun,
    }
    patterns = {
        "F/K": [r"<th>\s*F/K\s*</th>\s*<td>\s*([^<]+?)\s*</td>"],
        "PD/DD": [r"<th>\s*PD/DD\s*</th>\s*<td>\s*([^<]+?)\s*</td>"],
        "FD/FAVÖK": [r"<th>\s*FD/FAVÖK\s*</th>\s*<td>\s*([^<]+?)\s*</td>"],
        "FD/Satışlar": [r"<th>\s*FD/Satışlar\s*</th>\s*<td>\s*([^<]+?)\s*</td>"],
    }
    for key, pats in patterns.items():
        for p in pats:
            m = re.search(p, text, re.I)
            if m:
                out[key] = parse_tr(m.group(1))
                break
    m = re.search(r"Piyasa Değeri\s*</th>\s*<td>\s*([^<]+?)\s*</td>", text, re.I)
    if m:
        out["piyasa_degeri_raw"] = m.group(1).strip()
        out["piyasa_degeri"] = parse_mcap(m.group(1))
    for lab, key in [
        (r"Son Fiyat\s*</th>\s*<td>\s*([^<]+)", "fiyat"),
        (r"Kapanış\s*</th>\s*<td>\s*([^<]+)", "fiyat"),
    ]:
        m = re.search(lab, text, re.I)
        if m and out.get("fiyat") is None:
            out["fiyat"] = parse_tr(m.group(1))
    cache[kod] = out
    return out


def deg_from_isyatirim(kod: str, force: bool = False) -> dict:
    """Çarpan dict — yalnızca İş Yatırım. Yahoo yok."""
    c = isyatirim_carpanlar(kod, force=force)
    if c.get("error"):
        return {"_error": c.get("error"), "_kaynak": "İş Yatırım"}
    deg = {
        "F/K": c.get("F/K"),
        "PD/DD": c.get("PD/DD"),
        "FD/FAVÖK": c.get("FD/FAVÖK"),
        "FD/Satış": c.get("FD/Satışlar"),
        "fiyat": c.get("fiyat"),
        "piyasa_değeri": c.get("piyasa_degeri"),
        "_kaynak": "İş Yatırım şirket kartı",
        "_url": c.get("url"),
    }
    return deg


def bilanco_kalite_skoru(isy: dict | None, deg: dict | None = None) -> dict:
    """Bilanço kalite skoru 0–100. Yatırım tavsiyesi değildir."""
    isy = isy or {}
    deg = deg or {}
    oz = isy.get("ozet") or {}
    oran = isy.get("oranlar") or {}
    artilar, eksiler = [], []
    skor = 55.0

    def _f(v):
        try:
            if v is None:
                return None
            x = float(v)
            return x if x == x else None
        except (TypeError, ValueError):
            return None

    def _oz(key, field="deger"):
        item = oz.get(key) or {}
        return _f(item.get(field)) if isinstance(item, dict) else None

    net_chg = _oz("net_donem_kari", "degisim_%")
    has_chg = _oz("hasilat", "degisim_%")
    oz_chg = _oz("ozkaynaklar", "degisim_%")
    ef_chg = _oz("esas_faaliyet_kari", "degisim_%")
    nakit = _oz("nakit")
    kv = _oz("kisa_vadeli_yukumluluk")
    ozv = _oz("ozkaynaklar")
    tv = _oz("toplam_varliklar")
    roe = _f(oran.get("ROE_%")) or _f(deg.get("ROE"))
    marj = _f(oran.get("net_kar_marji_%")) or _f(deg.get("net_kar_marjı"))
    borc = _f(oran.get("borc_ozkaynak")) or _f(deg.get("borç/özkaynak"))
    if borc is not None and borc > 5:
        borc = borc / 100.0
    cari = _f(oran.get("cari_oran")) or _f(deg.get("cari_oran"))

    if net_chg is not None:
        if net_chg >= 25:
            skor += 10; artilar.append(f"Net kâr YoY +%{net_chg:.0f}")
        elif net_chg <= -20:
            skor -= 12; eksiler.append(f"Net kâr YoY %{net_chg:.0f}")
    if has_chg is not None:
        if has_chg >= 10:
            skor += 6; artilar.append(f"Hasılat YoY +%{has_chg:.0f}")
        elif has_chg < 0:
            skor -= 6; eksiler.append(f"Hasılat YoY %{has_chg:.0f}")
    if oz_chg is not None and oz_chg >= 5:
        skor += 5; artilar.append(f"Özkaynak +%{oz_chg:.0f}")
    if ef_chg is not None:
        if ef_chg >= 15:
            skor += 5; artilar.append(f"Esas faaliyet +%{ef_chg:.0f}")
        elif ef_chg < -10:
            skor -= 5; eksiler.append(f"Esas faaliyet %{ef_chg:.0f}")
    if roe is not None:
        if roe >= 15:
            skor += 8; artilar.append(f"ROE %{roe:.1f}")
        elif roe < 5:
            skor -= 8; eksiler.append(f"ROE %{roe:.1f}")
    if marj is not None:
        if marj >= 10:
            skor += 5; artilar.append(f"Net marj %{marj:.1f}")
        elif marj < 3:
            skor -= 5; eksiler.append(f"Net marj %{marj:.1f}")
    if borc is not None:
        if borc <= 0.5:
            skor += 7; artilar.append(f"Düşük kaldıraç {borc:.2f}")
        elif borc >= 1.5:
            skor -= 10; eksiler.append(f"Yüksek kaldıraç {borc:.2f}")
    if cari is not None:
        if cari >= 1.5:
            skor += 5; artilar.append(f"Cari oran {cari:.2f}")
        elif cari < 1:
            skor -= 8; eksiler.append(f"Cari oran {cari:.2f}")
    if nakit is not None and kv:
        rn = nakit / kv if kv else None
        if rn is not None:
            if rn >= 0.3:
                skor += 4; artilar.append("Nakit/KV yükümlülük makul")
            elif rn < 0.1:
                skor -= 4; eksiler.append("Nakit/KV yükümlülük düşük")
    if tv and ozv:
        op = ozv / tv
        if op >= 0.5:
            skor += 4; artilar.append(f"Özkaynak payı %{op*100:.0f}")
        elif op < 0.3:
            skor -= 4; eksiler.append(f"Özkaynak payı %{op*100:.0f}")

    skor = max(0, min(100, round(skor, 1)))
    if skor >= 70:
        seviye, yorum = "İyi", "Bilanço kalitesi görece güçlü görünüyor."
    elif skor >= 45:
        seviye, yorum = "Orta", "Bilanço kalitesi karışık / nötr."
    else:
        seviye, yorum = "Zayıf", "Bilanço kalitesi temkinli — dipnot ve sektör ile doğrulayın."
    return {
        "skor": skor, "seviye": seviye, "yorum": yorum,
        "artilar": artilar, "eksiler": eksiler,
        "donem": isy.get("gunce_donem"), "kaynak": isy.get("kaynak") or "—",
    }


def model_sinyal_skoru(
    *,
    m: dict | None = None,
    deg: dict | None = None,
    tek: dict | None = None,
    tem: dict | None = None,
    kalite: dict | None = None,
    tu: dict | None = None,
) -> dict:
    """
    Kural tabanlı model sinyal skoru (−100…+100).
    YATIRIM TAVSİYESİ VEYA YAPAY ZEKÂ AL/SAT EMRİ DEĞİLDİR.
    """
    m, deg, tek = m or {}, deg or {}, tek or {}
    tem, kalite, tu = tem or {}, kalite or {}, tu or {}
    bilesen = {}
    toplam_w = 0.0
    skor = 0.0

    def _f(v):
        try:
            if v is None:
                return None
            x = float(v)
            return x if x == x else None
        except (TypeError, ValueError):
            return None

    def _add(name, val, w):
        nonlocal skor, toplam_w
        if val is None:
            return
        bilesen[name] = round(val, 1)
        skor += val * w
        toplam_w += w

    ts = _f(tem.get("skor"))
    if ts is not None:
        _add("temel", (ts - 50) * 2, 0.25)
    ks = _f(kalite.get("skor"))
    if ks is not None:
        _add("bilanco_kalite", (ks - 50) * 2, 0.20)
    us = _f(tu.get("skor"))
    if us is not None:
        _add("tarihsel_ucuzluk", (us - 50) * 2, 0.15)
    sh = _f(m.get("sharpe"))
    if sh is not None:
        _add("sharpe", max(-100, min(100, sh * 40)), 0.10)
    g1y = _f(m.get("getiri_1y"))
    if g1y is not None:
        _add("getiri_1y", max(-100, min(100, g1y)), 0.10)
    dd = _f(m.get("max_drawdown"))
    if dd is not None:
        _add("drawdown", max(-100, min(100, 50 + dd)), 0.08)
    rsi = _f(tek.get("rsi14"))
    if rsi is not None:
        if rsi <= 30:
            _add("rsi", 40, 0.07)
        elif rsi >= 70:
            _add("rsi", -40, 0.07)
        else:
            _add("rsi", (50 - rsi) * 0.8, 0.07)
    fk = _f(deg.get("F/K"))
    if fk is not None:
        if fk < 8:
            _add("fk", 30, 0.05)
        elif fk > 25:
            _add("fk", -30, 0.05)
        else:
            _add("fk", (15 - fk) * 2, 0.05)

    final = (skor / toplam_w) if toplam_w > 0 else 0.0
    final = max(-100, min(100, round(final, 1)))
    if final >= 35:
        etiket = "Model: pozitif eğilim"
    elif final <= -35:
        etiket = "Model: negatif eğilim"
    else:
        etiket = "Model: nötr / belirsiz"
    return {
        "skor": final,
        "etiket": etiket,
        "bilesenler": bilesen,
        "disclaimer": (
            "Bu çıktı kural tabanlı bir model skorudur; yapay zekâ al/sat emri veya "
            "yatırım tavsiyesi değildir. Sermaye kaybı riski vardır. Karar size aittir."
        ),
    }





def fintables_tarzi_karne(isy: dict | None, deg: dict | None = None, m: dict | None = None) -> dict:
    """
    Fintables Karne benzeri: kârlılık, büyüme, borçluluk, değerleme (0–100 alt skorlar).
    Kaynak: İş Yatırım mali tablo + şirket kartı çarpanları.
    """
    isy = isy or {}
    deg = deg or {}
    m = m or {}
    oz = isy.get("ozet") or {}
    oran = isy.get("oranlar") or {}

    def _f(v):
        try:
            if v is None:
                return None
            x = float(v)
            return x if x == x else None
        except (TypeError, ValueError):
            return None

    def _oz(key, field="deger"):
        item = oz.get(key) or {}
        return _f(item.get(field)) if isinstance(item, dict) else None

    def _oz_yoy(key):
        item = oz.get(key) or {}
        return _f(item.get("degisim_%")) if isinstance(item, dict) else None

    roe = _f(oran.get("ROE_%")) or _f(deg.get("ROE"))
    marj = _f(oran.get("net_kar_marji_%")) or _f(deg.get("net_kar_marjı"))
    borc = _f(oran.get("borc_ozkaynak")) or _f(deg.get("borç/özkaynak"))
    if borc is not None and borc > 5:
        borc = borc / 100.0
    cari = _f(oran.get("cari_oran")) or _f(deg.get("cari_oran"))
    fk = _f(deg.get("F/K"))
    pddd = _f(deg.get("PD/DD"))
    sat_yoy = _oz_yoy("hasilat")
    net_yoy = _oz_yoy("net_donem_kari")

    # Kârlılık
    k = 50.0
    if roe is not None:
        k += 20 if roe >= 20 else (10 if roe >= 12 else (-10 if roe < 5 else 0))
    if marj is not None:
        k += 15 if marj >= 12 else (8 if marj >= 6 else (-10 if marj < 2 else 0))
    k = max(0, min(100, round(k, 1)))

    # Büyüme
    b = 50.0
    if sat_yoy is not None:
        b += 20 if sat_yoy >= 30 else (10 if sat_yoy >= 15 else (-15 if sat_yoy < 0 else 0))
    if net_yoy is not None:
        b += 15 if net_yoy >= 25 else (8 if net_yoy >= 10 else (-15 if net_yoy < -10 else 0))
    b = max(0, min(100, round(b, 1)))

    # Borçluluk (yüksek skor = sağlıklı / düşük borç)
    d = 55.0
    if borc is not None:
        d += 25 if borc <= 0.4 else (10 if borc <= 0.8 else (-20 if borc >= 1.5 else -5))
    if cari is not None:
        d += 15 if cari >= 1.5 else (-15 if cari < 1 else 0)
    d = max(0, min(100, round(d, 1)))

    # Değerleme (yüksek = daha ucuz görünüm)
    v = 50.0
    if fk is not None and fk > 0:
        v += 25 if fk < 8 else (10 if fk <= 12 else (-15 if fk > 25 else 0))
    if pddd is not None and pddd > 0:
        v += 15 if pddd < 1 else (5 if pddd < 1.5 else (-10 if pddd > 3 else 0))
    v = max(0, min(100, round(v, 1)))

    genel = round((k + b + d + v) / 4.0, 1)
    if genel >= 70:
        harf, yorum = "A", "Karne genel olarak güçlü — yine de sektör ve dipnot kontrol edin."
    elif genel >= 55:
        harf, yorum = "B", "Karne orta-üstü; büyümeyi ve borcu birlikte okuyun."
    elif genel >= 40:
        harf, yorum = "C", "Karne orta; zayıf alanları (kârlılık/borç/değer) inceleyin."
    else:
        harf, yorum = "D", "Karne zayıf — yüksek temkin."
    return {
        "karlilik": k,
        "buyume": b,
        "borcluluk": d,
        "degerleme": v,
        "genel": genel,
        "harf": harf,
        "yorum": yorum,
        "metrikler": {
            "ROE": roe, "net_marj": marj, "borc_oz": borc, "cari": cari,
            "F/K": fk, "PD/DD": pddd, "satis_yoy": sat_yoy, "net_yoy": net_yoy,
        },
        "kaynak": "İş Yatırım + model karne (Fintables tarzı özet)",
    }


def birlesik_profil_skoru(
    kalite: dict | None = None,
    sinyal: dict | None = None,
    tu: dict | None = None,
) -> dict:
    """
    Hazır profil birleşik skor — eşit ağırlık:
      1) Bilanço kalite (0–100)
      2) Model sinyal normalize (0–100)
      3) Tarihsel ucuzluk (0–100)
    Ortalama yüksek → düşük sıralanır. Yatırım tavsiyesi değildir.
    """
    kalite = kalite or {}
    sinyal = sinyal or {}
    tu = tu or {}

    def _f(v):
        try:
            if v is None:
                return None
            x = float(v)
            return x if x == x else None
        except (TypeError, ValueError):
            return None

    k = _f(kalite.get("skor"))
    s = _f(sinyal.get("skor"))
    u = _f(tu.get("skor"))
    # sinyal −100..100 → 0..100
    s_n = None if s is None else max(0.0, min(100.0, (s + 100.0) / 2.0))
    parts = []
    labels = []
    if k is not None:
        parts.append(k)
        labels.append("kalite")
    if s_n is not None:
        parts.append(s_n)
        labels.append("sinyal")
    if u is not None:
        parts.append(u)
        labels.append("ucuzluk")
    # eşit ağırlık
    n = len(parts)
    ort = round(sum(parts) / n, 1) if n else None
    agirlik = round(100.0 / n, 1) if n else None
    if ort is None:
        seviye = "—"
    elif ort >= 70:
        seviye = "Yüksek"
    elif ort >= 45:
        seviye = "Orta"
    else:
        seviye = "Düşük"
    return {
        "kalite": k,
        "sinyal": s,
        "sinyal_norm": round(s_n, 1) if s_n is not None else None,
        "ucuzluk": u,
        "ortalama": ort,
        "seviye": seviye,
        "bilesen_sayisi": n,
        "esit_agirlik_%": agirlik,
        "kullanilan": labels,
        "not": (
            f"Eşit ağırlık ({agirlik}% × {n} bileşen): "
            f"kalite={k if k is not None else '—'}, "
            f"sinyal_norm={round(s_n,1) if s_n is not None else '—'}, "
            f"ucuzluk={u if u is not None else '—'}. "
            f"Birleşik ortalama={ort if ort is not None else '—'}. "
            "Sıralama yüksek → düşük. Yatırım tavsiyesi değildir."
        ),
    }


# Geriye uyumluluk
def birlesik_kalite_sinyal_skoru(kalite: dict | None, sinyal: dict | None, tu: dict | None = None) -> dict:
    return birlesik_profil_skoru(kalite=kalite, sinyal=sinyal, tu=tu)



def empty_state(baslik: str, aciklama: str, ipucu: str = "") -> None:
    """Görsel boş durum kutusu."""
    ip = f'<div style="color:#64748b;font-size:0.8rem">{ipucu}</div>' if ipucu else ""
    st.markdown(
        (
            '<div style="text-align:center;padding:28px 20px;margin:12px 0;'
            'background:linear-gradient(160deg,#0f172a,#132337);'
            'border:1px dashed #334155;border-radius:16px;">'
            '<div style="font-size:1.6rem;margin-bottom:8px">📭</div>'
            f'<div style="font-size:1.05rem;font-weight:650;color:#f1f5f9;margin-bottom:6px">{baslik}</div>'
            f'<div style="color:#94a3b8;font-size:0.9rem;max-width:420px;margin:0 auto 8px">{aciklama}</div>'
            f'{ip}</div>'
        ),
        unsafe_allow_html=True,
    )


def badge_fk(fk) -> str:
    """F/K renk eşiği: <10 yeşil, 10-20 sarı, >20 kırmızı."""
    if fk is None:
        return "—"
    try:
        x = float(fk)
    except (TypeError, ValueError):
        return "—"
    if x < 10:
        cls, lbl = "badge-ok", "ucuz band"
    elif x <= 20:
        cls, lbl = "badge-warn", "orta"
    else:
        cls, lbl = "badge-bad", "pahalı band"
    return f'<span class="badge {cls}">F/K {tr_num(x, 1)} · {lbl}</span>'


def badge_skor(v, *, iyi=70, orta=45) -> str:
    if v is None:
        return "—"
    try:
        x = float(v)
    except (TypeError, ValueError):
        return "—"
    if x >= iyi:
        cls = "badge-ok"
    elif x >= orta:
        cls = "badge-warn"
    else:
        cls = "badge-bad"
    return f'<span class="badge {cls}">{tr_num(x, 0)}</span>'


def bilanco_momentum(isy: dict | None) -> dict:
    """Güncel vs YoY aynı dönem: satış / net kâr / borç momentumu."""
    isy = isy or {}
    oz = isy.get("ozet") or {}
    out = {"donem": isy.get("gunce_donem"), "satirlar": [], "skor": 50, "ozet": ""}
    score = 50
    labels = [
        ("hasilat", "Satışlar"),
        ("brut_kar", "Brüt kâr"),
        ("esas_faaliyet_kari", "Faaliyet kârı"),
        ("net_donem_kari", "Net kâr"),
        ("ozkaynaklar", "Özkaynak"),
        ("toplam_yukumluluk", "Toplam yükümlülük"),
    ]
    for key, lab in labels:
        it = oz.get(key) or {}
        if not isinstance(it, dict):
            continue
        d = it.get("degisim_%")
        try:
            d = float(d) if d is not None else None
        except (TypeError, ValueError):
            d = None
        yon = "—"
        if d is not None:
            if key == "toplam_yukumluluk":
                if d <= 0:
                    yon, score = "olumlu", score + 6
                elif d > 15:
                    yon, score = "olumsuz", score - 8
                else:
                    yon = "nötr"
            else:
                if d >= 10:
                    yon, score = "olumlu", score + 8
                elif d >= 0:
                    yon, score = "hafif+", score + 3
                elif d > -10:
                    yon, score = "hafif−", score - 3
                else:
                    yon, score = "olumsuz", score - 8
        out["satirlar"].append({
            "Kalem": lab,
            "Güncel": it.get("deger"),
            "YoY dönem": it.get("onceki_yoy"),
            "YoY %": d,
            "Yön": yon,
        })
    score = max(0, min(100, score))
    out["skor"] = score
    if score >= 70:
        out["ozet"] = "Bilanço momentumu güçlü (çoğu kalem YoY iyileşiyor)."
    elif score >= 45:
        out["ozet"] = "Bilanço momentumu karışık — kalem bazında ayırın."
    else:
        out["ozet"] = "Bilanço momentumu zayıf (kâr/satış baskısı veya borç artışı)."
    return out



def mali_tablo_otomatik_yorum(isy: dict | None, deg: dict | None = None) -> dict:
    """
    Yeni / güncel bilanço için otomatik yorum:
    kârlılık, büyüme, borç, nakit, çarpan, momentum.
    """
    isy = isy or {}
    deg = deg or {}
    oz = isy.get("ozet") or {}
    oran = isy.get("oranlar") or {}
    donem = isy.get("gunce_donem") or "—"
    olumlu, olumsuz, notr, aksiyon = [], [], [], []

    def _g(key):
        it = oz.get(key) or {}
        return it if isinstance(it, dict) else {}

    def _f(v):
        try:
            return float(v) if v is not None else None
        except (TypeError, ValueError):
            return None

    # Satış / kâr büyüme
    for key, lab in (
        ("hasilat", "Satışlar"),
        ("net_donem_kari", "Net kâr"),
        ("esas_faaliyet_kari", "Faaliyet kârı"),
        ("brut_kar", "Brüt kâr"),
    ):
        d = _f(_g(key).get("degisim_%"))
        if d is None:
            continue
        if d >= 20:
            olumlu.append(f"**{lab}** YoY **%{tr_num(d,1)}** — güçlü büyüme ({donem}).")
        elif d >= 0:
            notr.append(f"**{lab}** YoY **%{tr_num(d,1)}** — enflasyon altında kalabilir.")
        elif d > -15:
            olumsuz.append(f"**{lab}** YoY **%{tr_num(d,1)}** — hafif gerileme.")
        else:
            olumsuz.append(f"**{lab}** YoY **%{tr_num(d,1)}** — belirgin baskı.")

    roe = _f(oran.get("ROE_%")) or _f(deg.get("ROE"))
    if roe is not None:
        if roe >= 20:
            olumlu.append(f"**ROE %{tr_num(roe,1)}** yüksek — özkaynak verimi güçlü.")
        elif roe >= 10:
            notr.append(f"**ROE %{tr_num(roe,1)}** makul band.")
        elif roe > 0:
            olumsuz.append(f"**ROE %{tr_num(roe,1)}** düşük.")
        else:
            olumsuz.append(f"**ROE %{tr_num(roe,1)}** negatif — zarar dönemi.")

    marj = _f(oran.get("net_kar_marji_%"))
    if marj is not None:
        if marj >= 15:
            olumlu.append(f"Net marj **%{tr_num(marj,1)}**.")
        elif marj < 5:
            olumsuz.append(f"Net marj **%{tr_num(marj,1)}** dar.")

    borc = _f(oran.get("borc_ozkaynak")) or _f(deg.get("borç/özkaynak"))
    if borc is not None:
        if borc <= 0.5:
            olumlu.append(f"Borç/özkaynak **{tr_num(borc,2)}** düşük kaldıraç.")
        elif borc <= 1.2:
            notr.append(f"Borç/özkaynak **{tr_num(borc,2)}** yönetilebilir.")
        else:
            olumsuz.append(f"Borç/özkaynak **{tr_num(borc,2)}** yüksek — faiz/refinansman riski.")
            aksiyon.append("Borç vadesi ve faiz yükünü kontrol edin.")

    cari = _f(oran.get("cari_oran"))
    if cari is not None:
        if cari >= 1.5:
            olumlu.append(f"Cari oran **{tr_num(cari,2)}** — kısa vade likidite rahat.")
        elif cari < 1.0:
            olumsuz.append(f"Cari oran **{tr_num(cari,2)}** — kısa vadeli baskı.")
            aksiyon.append("KV yükümlülük / nakit dengesini izleyin.")

    nakit = _f(_g("nakit").get("deger"))
    # nakit değişim
    nd = _f(_g("nakit").get("degisim_%"))
    if nd is not None:
        if nd >= 10:
            olumlu.append(f"Nakit pozisyonu YoY **%{tr_num(nd,1)}** artmış.")
        elif nd <= -20:
            olumsuz.append(f"Nakit YoY **%{tr_num(nd,1)}** gerilemiş.")

    fk = _f(deg.get("F/K"))
    if fk is not None:
        if fk < 8:
            olumlu.append(f"F/K **{tr_num(fk,1)}** düşük band (ucuz görünüm — sektör ile doğrulayın).")
        elif fk > 25:
            olumsuz.append(f"F/K **{tr_num(fk,1)}** yüksek band — büyüme beklentisi fiyatlanmış olabilir.")
            aksiyon.append("Sektör F/K medyanı ile kıyaslayın.")
        else:
            notr.append(f"F/K **{tr_num(fk,1)}** orta band.")

    pddd = _f(deg.get("PD/DD"))
    if pddd is not None:
        if pddd < 1:
            olumlu.append(f"PD/DD **{tr_num(pddd,2)}** defter altı.")
        elif pddd > 3:
            olumsuz.append(f"PD/DD **{tr_num(pddd,2)}** primli.")

    mom = bilanco_momentum(isy)
    if mom.get("skor", 50) >= 70:
        olumlu.append(f"Bilanço momentum skoru **{mom['skor']}/100** — {mom.get('ozet','')}")
    elif mom.get("skor", 50) < 45:
        olumsuz.append(f"Bilanço momentum skoru **{mom['skor']}/100** — {mom.get('ozet','')}")
        aksiyon.append("Zarar eden kalemleri ve borç değişimini derinlemesine bakın.")

    # Özet paragraf
    n_ol = len(olumlu)
    n_olms = len(olumsuz)
    if n_ol >= n_olms + 2:
        sonuc = f"**{donem}** bilançosu genel olarak **olumlu** okunuyor ({n_ol} artı / {n_olms} eksi)."
    elif n_olms >= n_ol + 2:
        sonuc = f"**{donem}** bilançosu **temkinli** okunmalı ({n_olms} eksi / {n_ol} artı)."
    else:
        sonuc = f"**{donem}** bilançosu **karışık** — kalem bazında ayırım gerekli."

    if not aksiyon:
        aksiyon.append("Sonraki çeyrek satış/marj trendini ve net borç değişimini takip edin.")

    return {
        "donem": donem,
        "sonuc": sonuc,
        "olumlu": olumlu,
        "olumsuz": olumsuz,
        "notr": notr,
        "aksiyon": aksiyon,
        "momentum_skor": mom.get("skor"),
    }


def render_otomatik_bilanco_yorum(yorum: dict, *, baslik: str = "Otomatik bilanço yorumu") -> None:
    if not yorum:
        return
    st.markdown(f"##### 📝 {baslik}")
    st.info(yorum.get("sonuc") or "")
    c1, c2 = st.columns(2)
    with c1:
        if yorum.get("olumlu"):
            st.success("**Olumlu**\n\n" + "\n".join(f"- {x}" for x in yorum["olumlu"][:8]))
        if yorum.get("notr"):
            st.caption("Nötr: " + " · ".join(yorum["notr"][:4]))
    with c2:
        if yorum.get("olumsuz"):
            st.warning("**Olumsuz / risk**\n\n" + "\n".join(f"- {x}" for x in yorum["olumsuz"][:8]))
        if yorum.get("aksiyon"):
            st.markdown("**İzleme**\n\n" + "\n".join(f"- {x}" for x in yorum["aksiyon"][:5]))
    st.caption(f"Dönem: **{yorum.get('donem')}** · momentum {yorum.get('momentum_skor', '—')}/100 · otomatik kural motoru — tavsiye değildir.")


def hazir_profil_filtrele(df: pd.DataFrame, profil: str) -> pd.DataFrame:
    """Dengeli / Değer / Kalite / Momentum / Taze bilanço."""
    if df is None or getattr(df, "empty", True):
        return df
    d = df.copy()
    profil = (profil or "Dengeli").strip()

    def _num(s):
        return pd.to_numeric(s, errors="coerce")

    for col, alias in (
        ("F/K", "_fk"), ("PD/DD", "_pddd"), ("kalite_skor", "_kal"),
        ("ucuzluk_skor", "_uc"), ("model_sinyal", "_sig"), ("birlesik_ortalama", "_bir"),
    ):
        if col in d.columns:
            d[alias] = _num(d[col])
    if "bilanco_donem" in d.columns:
        d["_don"] = d["bilanco_donem"].astype(str)

    if profil == "Değer":
        if "_fk" in d.columns:
            d = d[(d["_fk"].isna()) | (d["_fk"] <= 15)]
        if "_uc" in d.columns:
            d = d.sort_values(["_uc", "_fk"] if "_fk" in d.columns else ["_uc"], ascending=[False, True] if "_fk" in d.columns else [False], na_position="last")
        elif "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")
    elif profil == "Kalite":
        if "_kal" in d.columns:
            d = d[(d["_kal"].isna()) | (d["_kal"] >= 55)].sort_values("_kal", ascending=False, na_position="last")
        elif "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")
    elif profil == "Momentum":
        if "_sig" in d.columns:
            d = d.sort_values("_sig", ascending=False, na_position="last")
        elif "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")
    elif profil == "Taze bilanço":
        if "_don" in d.columns:
            y = datetime.now().year
            d["_taze"] = d["_don"].str.contains(str(y), na=False).astype(int)
            cols = ["_taze"] + (["_bir"] if "_bir" in d.columns else [])
            asc = [False] * len(cols)
            d = d.sort_values(cols, ascending=asc, na_position="last")
        elif "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")
    elif profil == "Düşük kaldıraç":
        # PD/DD ve kalite; borç kolonu yoksa kalite+birleşik
        if "_pddd" in d.columns:
            d = d[(d["_pddd"].isna()) | (d["_pddd"] <= 2.5)]
        if "_kal" in d.columns:
            d = d.sort_values(["_kal", "_bir"] if "_bir" in d.columns else ["_kal"], ascending=[False, False] if "_bir" in d.columns else [False], na_position="last")
        elif "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")
    else:
        if "_bir" in d.columns:
            d = d.sort_values("_bir", ascending=False, na_position="last")

    return d.drop(columns=[c for c in d.columns if c.startswith("_")], errors="ignore")


def hisse_birlesik_satir(kod: str, *, force_isy: bool = False) -> dict:
    """Tek hisse: kalite + sinyal + tarihsel ucuzluk + eşit ağırlıklı ortalama."""
    kod = str(kod).upper().replace(".IS", "").strip()
    row = {
        "kod": kod, "ad": kod, "bilanco_donem": None,
        "kalite_skor": None, "model_sinyal": None, "sinyal_norm": None,
        "ucuzluk_skor": None, "birlesik_ortalama": None, "seviye": None,
        "F/K": None, "PD/DD": None, "ROE": None, "getiri_1y": None,
        "fiyat_persentil_1y": None,
    }
    deg, m = {}, {}
    fd = None
    try:
        import yfinance as yf
        t = yf.Ticker(f"{kod}.IS")
        hist = t.history(period="3y", auto_adjust=True)
        if hist is not None and not hist.empty:
            hist2 = hist.reset_index()
            dc = "Date" if "Date" in hist2.columns else hist2.columns[0]
            fd = hist2.rename(columns={dc: "tarih", "Close": "fiyat"})[["tarih", "fiyat"]]
            px = hist["Close"].dropna()
            if len(px) > 2:
                # 1Y getiri
                px1 = px.tail(min(252, len(px)))
                m["getiri_1y"] = float((px1.iloc[-1] / px1.iloc[0] - 1) * 100)
                row["getiri_1y"] = m["getiri_1y"]
            rets = px.pct_change().dropna()
            if len(rets) > 10:
                m["volatilite"] = float(rets.std() * (252 ** 0.5) * 100)
                cum = (1 + rets).cumprod()
                m["max_drawdown"] = float((cum / cum.cummax() - 1).min() * 100)
        try:
            info = dict(t.info or {})
        except Exception:
            info = {}
        row["ad"] = (info.get("longName") or info.get("shortName") or kod)[:48]
        # Yahoo trailingPE KULLANILMAZ
        if info.get("returnOnEquity") is not None:
            row["ROE"] = float(info["returnOnEquity"]) * 100
            deg["ROE"] = row["ROE"]
    except Exception:
        pass
    try:
        _ic = deg_from_isyatirim(kod)
        if not _ic.get("_error"):
            for _ck in ("F/K", "PD/DD", "FD/FAVÖK"):
                if _ic.get(_ck) is not None:
                    deg[_ck] = _ic[_ck]
                    row[_ck] = _ic[_ck]
            if _ic.get("fiyat") is not None:
                row["fiyat"] = _ic["fiyat"]
    except Exception:
        pass
    isy = {}
    try:
        isy = isyatirim_mali_tablo(kod, force=force_isy)
        if isy and not isy.get("error"):
            row["bilanco_donem"] = isy.get("gunce_donem")
            orans = isy.get("oranlar") or {}
            if orans.get("ROE_%") is not None:
                row["ROE"] = orans["ROE_%"]
                deg["ROE"] = row["ROE"]
    except Exception:
        isy = {}
    kal = bilanco_kalite_skoru(isy if isinstance(isy, dict) else {}, deg)
    tu = tarihsel_ucuzluk_analizi(kod, fd=fd, deg=deg)
    sig = model_sinyal_skoru(m=m, deg=deg, kalite=kal, tu=tu)
    bir = birlesik_profil_skoru(kalite=kal, sinyal=sig, tu=tu)
    row["kalite_skor"] = bir.get("kalite")
    row["model_sinyal"] = bir.get("sinyal")
    row["sinyal_norm"] = bir.get("sinyal_norm")
    row["ucuzluk_skor"] = bir.get("ucuzluk")
    row["fiyat_persentil_1y"] = tu.get("fiyat_persentil_1y")
    row["birlesik_ortalama"] = bir.get("ortalama")
    row["seviye"] = bir.get("seviye")
    row["_kal"] = kal
    row["_sig"] = sig
    row["_tu"] = tu
    row["_bir"] = bir
    return row



def _html_escape(s) -> str:
    return (
        str(s if s is not None else "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _svg_bar_chart(labels: list, values: list, title: str = "", width: int = 640, height: int = 220) -> str:
    """Basit yatay-olmayan dikey bar SVG (bağımlılıksız)."""
    if not labels or not values or len(labels) != len(values):
        return ""
    nums = []
    for v in values:
        try:
            nums.append(float(v) if v is not None else 0.0)
        except (TypeError, ValueError):
            nums.append(0.0)
    if not any(nums):
        return f"<p class='muted'>{_html_escape(title)}: veri yok</p>"
    pad_l, pad_r, pad_t, pad_b = 48, 16, 28, 40
    w = width - pad_l - pad_r
    h = height - pad_t - pad_b
    mx = max(abs(x) for x in nums) or 1.0
    n = len(nums)
    gap = 8
    bw = max(12, (w - gap * (n + 1)) / n)
    bars = []
    zero_y = pad_t + h  # baseline bottom for positive-only; if negatives, mid
    has_neg = any(x < 0 for x in nums)
    if has_neg:
        zero_y = pad_t + h * (mx / (2 * mx))
    for i, (lab, val) in enumerate(zip(labels, nums)):
        x = pad_l + gap + i * (bw + gap)
        if has_neg:
            bar_h = abs(val) / (2 * mx) * h
            y = zero_y - bar_h if val >= 0 else zero_y
        else:
            bar_h = (val / mx) * h if mx else 0
            y = pad_t + h - bar_h
        color = "#1B4F72" if val >= 0 else "#C0392B"
        bars.append(
            f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw:.1f}" height="{max(bar_h,0):.1f}" fill="{color}" rx="3"/>'
            f'<text x="{x + bw/2:.1f}" y="{height - 12}" text-anchor="middle" font-size="11" fill="#333">{_html_escape(lab)}</text>'
            f'<text x="{x + bw/2:.1f}" y="{y - 4:.1f}" text-anchor="middle" font-size="10" fill="#555">{_fmt_tl(val)}</text>'
        )
    return (
        f'<div class="chart-wrap"><div class="chart-title">{_html_escape(title)}</div>'
        f'<svg viewBox="0 0 {width} {height}" width="100%" height="{height}">'
        f'{"".join(bars)}</svg></div>'
    )


def rapor_paketi_html(
    *,
    baslik: str,
    kod: str,
    ad: str = "",
    donem: str = "",
    bloklar: list | None = None,
    meta: dict | None = None,
    tables: list | None = None,
    charts: list | None = None,
    kpi: dict | None = None,
    fk_yorum: str = "",
) -> str:
    """
    Kurumsal HTML rapor.
    tables: [(baslik, headers, rows)]
    charts: [(baslik, labels, values)]
    kpi: {etiket: deger}
    """
    bloklar = bloklar or []
    meta = meta or {}
    tables = tables or []
    charts = charts or []
    kpi = kpi or {}
    ts = datetime.now().strftime("%d.%m.%Y %H:%M")

    kpi_html = ""
    if kpi:
        cards = []
        for k, v in kpi.items():
            cards.append(
                f'<div class="kpi"><div class="kpi-l">{_html_escape(k)}</div>'
                f'<div class="kpi-v">{_html_escape(v)}</div></div>'
            )
        kpi_html = f'<div class="kpi-row">{"".join(cards)}</div>'

    tbl_html = []
    for title, headers, rows in tables:
        th = "".join(f"<th>{_html_escape(h)}</th>" for h in headers)
        body = []
        for r in rows:
            tds = "".join(f"<td>{_html_escape(c)}</td>" for c in r)
            body.append(f"<tr>{tds}</tr>")
        tbl_html.append(
            f'<section class="sec"><h2>{_html_escape(title)}</h2>'
            f'<table class="data"><thead><tr>{th}</tr></thead>'
            f'<tbody>{"".join(body)}</tbody></table></section>'
        )

    chart_html = []
    for title, labels, values in charts:
        chart_html.append(_svg_bar_chart(labels, values, title=title))

    blk_html = []
    for h, b in bloklar:
        blk_html.append(
            f'<section class="sec"><h2>{_html_escape(h)}</h2>'
            f'<div class="blk">{_html_escape(b)}</div></section>'
        )

    meta_rows = "".join(
        f"<tr><td>{_html_escape(k)}</td><td>{_html_escape(v)}</td></tr>"
        for k, v in meta.items() if v is not None
    )
    fk_box = ""
    if fk_yorum:
        fk_box = f'<div class="callout">{_html_escape(fk_yorum)}</div>'

    return f"""<!DOCTYPE html>
<html lang="tr">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>{_html_escape(baslik)} — {_html_escape(kod)}</title>
<style>
:root {{
  --navy: #0f2744;
  --navy2: #1B4F72;
  --line: #d5dde8;
  --bg: #f4f6f9;
  --card: #ffffff;
  --muted: #5c6b7a;
  --warn: #fff8e6;
  --warn-b: #e8c96a;
}}
* {{ box-sizing: border-box; }}
body {{
  margin: 0; font-family: "Segoe UI", system-ui, -apple-system, sans-serif;
  background: var(--bg); color: #1a2332; line-height: 1.45;
}}
.page {{ max-width: 920px; margin: 0 auto; padding: 28px 20px 48px; }}
.header {{
  background: linear-gradient(135deg, var(--navy), var(--navy2));
  color: #fff; border-radius: 12px; padding: 22px 24px; margin-bottom: 18px;
}}
.header h1 {{ margin: 0 0 6px; font-size: 1.35rem; font-weight: 650; }}
.header .sub {{ opacity: 0.9; font-size: 0.95rem; }}
.badge {{
  display: inline-block; background: rgba(255,255,255,0.15);
  padding: 3px 10px; border-radius: 999px; font-size: 0.8rem; margin-top: 8px;
}}
.warn {{
  background: var(--warn); border: 1px solid var(--warn-b);
  border-radius: 8px; padding: 12px 14px; margin: 14px 0; font-size: 0.9rem;
}}
.callout {{
  background: #eef5fb; border-left: 4px solid var(--navy2);
  padding: 12px 14px; border-radius: 0 8px 8px 0; margin: 12px 0; font-size: 0.92rem;
}}
.kpi-row {{
  display: grid; grid-template-columns: repeat(auto-fit, minmax(140px, 1fr));
  gap: 10px; margin: 14px 0 18px;
}}
.kpi {{
  background: var(--card); border: 1px solid var(--line); border-radius: 10px;
  padding: 12px 14px;
}}
.kpi-l {{ font-size: 0.75rem; color: var(--muted); text-transform: uppercase; letter-spacing: .03em; }}
.kpi-v {{ font-size: 1.15rem; font-weight: 650; margin-top: 4px; color: var(--navy); }}
.sec {{
  background: var(--card); border: 1px solid var(--line); border-radius: 10px;
  padding: 16px 18px; margin: 14px 0;
}}
.sec h2 {{
  margin: 0 0 12px; font-size: 1.05rem; color: var(--navy);
  border-bottom: 1px solid var(--line); padding-bottom: 8px;
}}
table.data {{ width: 100%; border-collapse: collapse; font-size: 0.9rem; }}
table.data th {{
  text-align: left; background: #eef2f7; color: var(--navy);
  padding: 8px 10px; border-bottom: 1px solid var(--line); font-weight: 600;
}}
table.data td {{
  padding: 8px 10px; border-bottom: 1px solid #eef1f5;
}}
table.data tr:last-child td {{ border-bottom: none; }}
.blk {{ white-space: pre-wrap; font-size: 0.92rem; color: #2c3a4a; }}
.chart-wrap {{ margin: 12px 0 8px; }}
.chart-title {{ font-size: 0.9rem; font-weight: 600; color: var(--navy); margin-bottom: 6px; }}
.muted {{ color: var(--muted); font-size: 0.85rem; }}
.footer {{
  margin-top: 24px; padding-top: 12px; border-top: 1px solid var(--line);
  font-size: 0.8rem; color: var(--muted);
}}
@media print {{
  body {{ background: #fff; }}
  .page {{ padding: 0; }}
  .header {{ border-radius: 0; }}
}}
</style>
</head>
<body>
<div class="page">
  <div class="header">
    <h1>{_html_escape(baslik)}</h1>
    <div class="sub"><b>{_html_escape(kod)}</b> · {_html_escape(ad)} · Bilanço dönemi: <b>{_html_escape(donem or "—")}</b></div>
    <div class="badge">Üretim: {_html_escape(ts)} · M.Durmaz Analiz</div>
  </div>
  <div class="warn">
    <b>Yasal uyarı:</b> Bu rapor bilgilendirme amaçlıdır; yatırım tavsiyesi değildir.
    Veriler TEFAS, İş Yatırım ve piyasa kaynaklarından derlenmiştir; gecikme / hata olabilir.
  </div>
  {fk_box}
  {kpi_html}
  <section class="sec">
    <h2>Rapor özeti</h2>
    <table class="data"><thead><tr><th>Alan</th><th>Değer</th></tr></thead>
    <tbody>{meta_rows}</tbody></table>
  </section>
  {"".join(tbl_html)}
  {('<section class="sec"><h2>Grafikler</h2>' + "".join(chart_html) + "</section>") if chart_html else ""}
  {"".join(blk_html)}
  <div class="footer">
    M.Durmaz · TEFAS / BIST rapor paketi · F/K ve çarpanlar piyasa verisine bağlıdır ·
    Nakit / PD oranı dönem sonu nakit ÷ piyasa değeri
  </div>
</div>
</body>
</html>"""


# ── Bilanço bildirimleri (FON / HİSSE ayrı) ─────────────────────────────────
def _bilanco_state_path() -> Path:
    try:
        return Path(core.CACHE_DIR) / "bilanco_bildirim_state.json"
    except Exception:
        return ROOT / "bilanco_bildirim_state.json"


def _load_bilanco_state() -> dict:
    p = _bilanco_state_path()
    if not p.exists():
        return {"fon": {}, "hisse": {}, "hisse_metrik_onceki": {}, "fon_metrik_onceki": {}}
    try:
        import json
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {"fon": {}, "hisse": {}, "hisse_metrik_onceki": {}, "fon_metrik_onceki": {}}


def _save_bilanco_state(state: dict) -> None:
    try:
        import json
        p = _bilanco_state_path()
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(state, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    except Exception:
        pass


def _beklenen_bilanco_donemi(ref: datetime | None = None) -> dict:
    """
    Takvime göre *beklenen* son bilanço dönemi (KAP yayın gecikmesi dahil).
    Q1 ~Mayıs, Q2 ~Ağustos, Q3 ~Ekim sonu, Yıllık ~Şubat–Mart.
    Asıl kaynak her zaman core.bilanco_donem_durumu / canlı probe'dur;
    bu fonksiyon yedek takvimdir.
    """
    ref = ref or datetime.now()
    y, m = ref.year, ref.month
    if m <= 2:
        # Yıllık tablolar henüz gelmemiş olabilir → önceki yıl Q3
        return {"yil": y - 1, "donem": 9, "etiket": f"{y-1}/9", "tur": "ara", "aciklama": "9 aylık"}
    if m <= 4:
        return {"yil": y - 1, "donem": 12, "etiket": f"{y-1}/12", "tur": "yillik", "aciklama": "yıllık"}
    if m <= 7:
        return {"yil": y, "donem": 3, "etiket": f"{y}/3", "tur": "ara", "aciklama": "3 aylık"}
    if m <= 9:
        return {"yil": y, "donem": 6, "etiket": f"{y}/6", "tur": "ara", "aciklama": "6 aylık"}
    # Ekim–Aralık: 9 aylık (Q3) sezonu
    return {"yil": y, "donem": 9, "etiket": f"{y}/9", "tur": "ara", "aciklama": "9 aylık"}


def son_bilanco_donemi(*, probe: bool = False, force_refresh: bool = False) -> dict:
    """
    Analiz ve raporların bağlanacağı *güncel* bilanço dönemi.
    Öncelik: core canlı durum > kayıtlı state > takvim beklentisi.
    Yeni dönem gelince state güncellenir; rapor meta'sına yazılır.
    """
    state = _load_bilanco_state()
    bek = _beklenen_bilanco_donemi()
    result = {
        "yil": bek["yil"],
        "donem": bek["donem"],
        "etiket": bek["etiket"],
        "aciklama": bek.get("aciklama") or f"{bek['donem']} aylık",
        "kaynak": "takvim",
        "durum": "bekleniyor",
        "yeni_donem": False,
        "onceki": None,
    }
    bil = {}
    try:
        bil = core.bilanco_donem_durumu(probe=probe) or {}
    except Exception as e:
        bil = {"mesaj": str(e)}

    if bil.get("yil") and bil.get("donem"):
        result["yil"] = int(bil["yil"])
        result["donem"] = int(bil["donem"])
        result["etiket"] = f"{result['yil']}/{result['donem']}"
        result["aciklama"] = (
            "yıllık" if result["donem"] == 12 else f"{result['donem']} aylık"
        )
        result["kaynak"] = "core"
        result["durum"] = bil.get("durum") or "geldi"
        if bil.get("mesaj"):
            result["mesaj"] = bil["mesaj"]
        if "probe" in bil:
            result["probe"] = bil.get("probe")
    else:
        # Core yoksa: kayıtlı son dönem, takvimden eskiyse takvimi kullan
        prev = state.get("aktif_donem") or state.get("hisse") or {}
        if prev.get("yil") and prev.get("donem"):
            py, pd_ = int(prev["yil"]), int(prev["donem"])
            # takvim daha yeniyse takvime geç
            if (bek["yil"], bek["donem"]) > (py, pd_):
                result["kaynak"] = "takvim"
            else:
                result["yil"], result["donem"] = py, pd_
                result["etiket"] = f"{py}/{pd_}"
                result["aciklama"] = "yıllık" if pd_ == 12 else f"{pd_} aylık"
                result["kaynak"] = "state"
                result["durum"] = prev.get("durum") or "geldi"
        if bil.get("mesaj"):
            result["mesaj"] = bil["mesaj"]

    # Önceki dönem kaydı
    prev_aktif = state.get("aktif_donem") or {}
    if prev_aktif.get("yil") and prev_aktif.get("donem"):
        if (int(prev_aktif["yil"]), int(prev_aktif["donem"])) != (result["yil"], result["donem"]):
            result["yeni_donem"] = True
            result["onceki"] = {
                "yil": int(prev_aktif["yil"]),
                "donem": int(prev_aktif["donem"]),
                "etiket": f"{prev_aktif['yil']}/{prev_aktif['donem']}",
            }
            # Dönem değişti → eski metrik snapshot'larını "önceki dönem" olarak işaretle
            # (mali_tablo_iyilesme_kontrolu zaten son değerleri saklıyor)
            result["durum"] = result.get("durum") or "geldi"
            result["mesaj"] = (
                f"🆕 Güncel bilanço dönemi: {result['etiket']} "
                f"(önceki: {result['onceki']['etiket']}) — analizler buna göre yenilenmeli"
            )

    # Aktif dönemi kalıcı yaz (analiz/rapor aynı kaynağı kullansın)
    state["aktif_donem"] = {
        "yil": result["yil"],
        "donem": result["donem"],
        "etiket": result["etiket"],
        "aciklama": result["aciklama"],
        "durum": result.get("durum"),
        "son_guncelleme": datetime.now().isoformat(timespec="seconds"),
        "kaynak": result["kaynak"],
    }
    _save_bilanco_state(state)
    # Session'a da koy — UI tek yerden okusun
    try:
        st.session_state["aktif_bilanco"] = result
    except Exception:
        pass
    return result


def _onceki_donem(yil: int, donem: int) -> tuple[int, int]:
    """Önceki çeyrek (3,6,9,12)."""
    order = [3, 6, 9, 12]
    if donem not in order:
        donem = 12
    i = order.index(donem)
    if i == 0:
        return yil - 1, 12
    return yil, order[i - 1]


def bilanco_bildirimleri(probe: bool = False) -> dict:
    """
    FON ve HİSSE bilanço/dönem bildirimlerini ayrı üretir.
    Hisse tarafı her zaman son_bilanco_donemi() (canlı / en güncel dönem).
    """
    aktif = son_bilanco_donemi(probe=probe)
    bek = _beklenen_bilanco_donemi()
    state = _load_bilanco_state()
    now_s = datetime.now().isoformat(timespec="seconds")
    bil_core = {"yil": aktif.get("yil"), "donem": aktif.get("donem"), "durum": aktif.get("durum")}

    # ── HİSSE (BIST / KAP) — her zaman en son gelen dönem ─────────────────
    hisse = {
        "kaynak": "BIST / KAP / İş Yatırım",
        "beklenen_yil": bek["yil"],
        "beklenen_donem": bek["donem"],
        "beklenen_etiket": bek["etiket"],
        "yil": aktif["yil"],
        "donem": aktif["donem"],
        "etiket": aktif["etiket"],
        "aciklama": aktif.get("aciklama"),
        "durum": aktif.get("durum") or "geldi",
        "mesaj": aktif.get("mesaj") or f"Güncel dönem: {aktif['etiket']} ({aktif.get('aciklama')})",
        "yeni_donem": bool(aktif.get("yeni_donem")),
        "probe": aktif.get("probe"),
        "onceki_etiket": (aktif.get("onceki") or {}).get("etiket"),
    }

    # ── FON (TEFAS — NAV günlük; rapor çeyreği hisse ile aynı takvim) ──
    fon = {
        "kaynak": "TEFAS",
        "beklenen_etiket": aktif["etiket"],
        "durum": "gunluk",
        "mesaj": f"NAV günlük güncel · rapor çeyreği: {aktif['etiket']}",
        "yeni_donem": bool(aktif.get("yeni_donem")),
        "yil": aktif["yil"],
        "donem": aktif["donem"],
        "etiket": aktif["etiket"],
    }
    if aktif.get("yeni_donem"):
        fon["durum"] = "yeni_ceyrek"
        fon["mesaj"] = (
            f"📅 Yeni raporlama çeyreği: {fon['etiket']} — "
            "fon dağılım / büyüklük özetlerini tarama ile yenileyin."
        )

    # State güncelle
    state["hisse"] = {
        "yil": hisse.get("yil"), "donem": hisse.get("donem"),
        "durum": hisse.get("durum"), "son_kontrol": now_s,
        "etiket": hisse.get("etiket"),
    }
    state["fon"] = {
        "yil": fon.get("yil"), "donem": fon.get("donem"),
        "durum": fon.get("durum"), "son_kontrol": now_s,
        "etiket": fon.get("etiket"),
    }
    _save_bilanco_state(state)

    return {"hisse": hisse, "fon": fon, "state": state, "core": bil_core}




def isyatirim_mali_tablo(
    kod: str,
    *,
    periods: list[tuple[int, int]] | None = None,
    financial_group: str = "XI_29",
    force: bool = False,
) -> dict:
    """
    İş Yatırım MaliTablo API — bilanço + gelir tablosu.
    Her hisse için API'de *gerçekten dolu olan en yeni dönem* güncel kabul edilir
    (ör. ISDMR 2026/9 geldiyse 2026/6 değil 2026/9 kullanılır).
    Günlük cache: aynı gün aynı hisse tekrar çekilmez; force=True veya gün değişince yenilenir.
    """
    import urllib.request
    import urllib.parse
    import json as _json

    kod = str(kod).upper().replace(".IS", "").strip()
    bugun = datetime.now().strftime("%Y-%m-%d")

    # ── Günlük cache ──────────────────────────────────────────────────────
    cache = st.session_state.setdefault("_isy_mali_cache", {})
    ck = f"{kod}|{financial_group}"
    if not force and ck in cache and cache[ck].get("_cache_day") == bugun and not cache[ck].get("error"):
        return cache[ck]

    # ── Aday dönemler: takvim + geriye doğru (en yeni önce) ───────────────
    if not periods:
        try:
            aktif = _beklenen_bilanco_donemi()
            y, d = int(aktif["yil"]), int(aktif["donem"])
        except Exception:
            now = datetime.now()
            y = now.year
            d = 9 if now.month >= 10 else (6 if now.month >= 7 else (3 if now.month >= 5 else 12))
            if d == 12:
                y = y - 1
        # Beklenen + bir önceki + YoY aynı + önceki yıl sonu + daha eski
        cand = [(y, d)]
        py, pd_ = _onceki_donem(y, d)
        cand.append((py, pd_))
        cand.append((y - 1, d))
        # ara dönemler
        for extra in ((y, 9), (y, 6), (y, 3), (y - 1, 12), (y - 1, 9), (y - 1, 6), (y - 2, 12)):
            if extra not in cand:
                cand.append(extra)
        # Ekim+ ise 9 mutlaka adayda
        now = datetime.now()
        if now.month >= 10 and (now.year, 9) not in cand:
            cand.insert(0, (now.year, 9))
        periods = cand[:8]  # istekte max 5 kullanacağız; aday havuzu geniş

    # API max 5 dönem — en olası 5'i al
    seen = set()
    clean = []
    for p in periods:
        key = (int(p[0]), int(p[1]))
        if key not in seen:
            seen.add(key)
            clean.append(key)
        if len(clean) >= 5:
            break
    periods = clean

    def _call(fg: str, per: list[tuple[int, int]]):
        params = {
            "companyCode": kod,
            "exchange": "TRY",
            "financialGroup": fg,
        }
        for i, (yy, pp) in enumerate(per, 1):
            params[f"year{i}"] = yy
            params[f"period{i}"] = pp
        url = (
            "https://www.isyatirim.com.tr/_layouts/15/IsYatirim.Website/Common/"
            "Data.aspx/MaliTablo?" + urllib.parse.urlencode(params)
        )
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; TEFAS-GUI/2.2)",
                "Accept": "application/json",
            },
        )
        with urllib.request.urlopen(req, timeout=25) as resp:
            return _json.loads(resp.read().decode("utf-8", errors="replace"))

    try:
        raw = _call(financial_group, periods)
    except Exception as e:
        if financial_group != "UFRS":
            return isyatirim_mali_tablo(kod, periods=periods, financial_group="UFRS", force=force)
        out = {"error": str(e), "kod": kod, "_cache_day": bugun}
        cache[ck] = out
        return out

    if not isinstance(raw, dict) or not raw.get("ok", True):
        err = (raw or {}).get("errorDescription") or "API hata"
        if financial_group != "UFRS":
            alt = isyatirim_mali_tablo(kod, periods=periods, financial_group="UFRS", force=force)
            if not alt.get("error"):
                return alt
        out = {"error": err, "kod": kod, "_cache_day": bugun}
        cache[ck] = out
        return out

    rows = raw.get("value") or []
    if not rows:
        out = {"error": "Boş tablo", "kod": kod, "periods": periods, "_cache_day": bugun}
        cache[ck] = out
        return out

    def _num(x):
        if x is None or x == "":
            return None
        try:
            return float(str(x).replace(",", "").replace(" ", ""))
        except (TypeError, ValueError):
            return None

    # Her dönem kolonu için "doluluk" skoru (hasılat / net kâr / özkaynak)
    key_codes = ("3C", "3L", "3J", "2N", "2O")
    col_score = [0] * len(periods)
    col_has_data = [False] * len(periods)
    values_by_code: dict = {}

    for r in rows:
        code = str(r.get("itemCode") or "")
        desc = (r.get("itemDescTr") or r.get("itemDescEng") or code).strip()
        vals = [_num(r.get(f"value{i}")) for i in range(1, len(periods) + 1)]
        values_by_code[code] = {"desc": desc, "values": vals}
        if code in key_codes:
            for i, v in enumerate(vals):
                if v is not None:
                    col_has_data[i] = True
                    col_score[i] += 1

    # En yeni *dolu* dönem = periods listesinde index'i küçük olanlar zaten yeni;
    # ama API sırası istekle aynı. Dolu olanlar içinde (yil,donem) max olanı seç.
    best_i = None
    best_key = (-1, -1)
    for i, (yy, pp) in enumerate(periods):
        if col_has_data[i] and (yy, pp) > best_key:
            best_key = (yy, pp)
            best_i = i
    if best_i is None:
        # hiç key yoksa herhangi dolu kolon
        for i, (yy, pp) in enumerate(periods):
            any_val = any(
                (values_by_code.get(c) or {}).get("values", [None] * len(periods))[i] is not None
                for c in values_by_code
            )
            if any_val and (yy, pp) > best_key:
                best_key = (yy, pp)
                best_i = i
    if best_i is None:
        best_i = 0

    # Kolon sırasını yeniden diz: best_i → 0 (güncel), sonra diğerleri tarihe göre
    order = [best_i] + [i for i in range(len(periods)) if i != best_i]
    # ikincil: YoY aynı dönem tercih
    gunce_y, gunce_d = periods[best_i]
    yoy = (gunce_y - 1, gunce_d)
    rest = [i for i in order[1:]]
    rest_sorted = sorted(rest, key=lambda i: (0 if periods[i] == yoy else 1, -periods[i][0], -periods[i][1]))
    # YoY aynı çeyrek (ör. 2026/9 → 2025/9) ikinci sıraya
    yoy_i = None
    by, bd = periods[best_i]
    for i, (yy, dd) in enumerate(periods):
        if i != best_i and yy == by - 1 and dd == bd:
            yoy_i = i
            break
    rest_sorted = [i for i in rest_sorted if i != yoy_i]
    order = [best_i] + ([yoy_i] if yoy_i is not None else []) + rest_sorted

    periods_ord = [periods[i] for i in order]
    col_labels = [f"{y}/{p}" for y, p in periods_ord]

    def _reval(vals):
        return [vals[i] if i < len(vals) else None for i in order]

    bilanco_rows, gelir_rows, nakit_rows = [], [], []
    for code, meta in values_by_code.items():
        vals = _reval(meta["values"])
        values_by_code[code]["values"] = vals
        rec = {"kod": code, "kalem": meta["desc"]}
        for i, lab in enumerate(col_labels):
            rec[lab] = vals[i] if i < len(vals) else None
        if code.startswith("1") or code.startswith("2"):
            bilanco_rows.append(rec)
        elif code.startswith("3"):
            gelir_rows.append(rec)
        elif code.startswith("4"):
            nakit_rows.append(rec)

    def _pair(*codes):
        for c in codes:
            if c in values_by_code:
                vs = values_by_code[c]["values"]
                a = vs[0] if vs else None
                b = vs[1] if len(vs) > 1 else None
                return a, b, values_by_code[c]["desc"]
        return None, None, None

    mapping = {
        "donen_varliklar": ("1A",),
        "nakit": ("1AA", "1AAA"),
        "toplam_varliklar": ("2ODB", "1", "1G", "1Z"),
        "kisa_vadeli_yukumluluk": ("2A",),
        "uzun_vadeli_yukumluluk": ("2B",),
        "ozkaynaklar": ("2N", "2O", "2C", "2D"),
        "hasilat": ("3C", "3A", "3AA"),
        "brut_kar": ("3D", "3CAB"),
        "esas_faaliyet_kari": ("3DF", "3H", "3E"),
        "net_donem_kari": ("3L", "3J", "3I"),
        "favok": ("3HACA",),
    }
    ozet = {}
    for key, codes in mapping.items():
        a, b, desc = _pair(*codes)
        if a is not None:
            ozet[key] = {
                "deger": a,
                "onceki_yoy": b,
                "degisim_%": round((a / b - 1) * 100, 2) if a and b else None,
                "kalem": desc,
            }

    kv = (ozet.get("kisa_vadeli_yukumluluk") or {}).get("deger")
    uv = (ozet.get("uzun_vadeli_yukumluluk") or {}).get("deger")
    if kv is not None or uv is not None:
        ozet["toplam_yukumluluk"] = {
            "deger": (kv or 0) + (uv or 0),
            "onceki_yoy": None,
            "degisim_%": None,
            "kalem": "Kısa+Uzun Vadeli Yükümlülükler",
        }
    if "favok" not in ozet or ozet["favok"].get("deger") is None:
        ef = (ozet.get("esas_faaliyet_kari") or {}).get("deger")
        if ef is not None:
            ozet["favok"] = {
                "deger": ef,
                "kalem": "Esas faaliyet (FAVÖK proxy)",
                "onceki_yoy": (ozet.get("esas_faaliyet_kari") or {}).get("onceki_yoy"),
                "degisim_%": (ozet.get("esas_faaliyet_kari") or {}).get("degisim_%"),
            }

    tv = (ozet.get("toplam_varliklar") or {}).get("deger")
    oz = (ozet.get("ozkaynaklar") or {}).get("deger")
    ty = (ozet.get("toplam_yukumluluk") or {}).get("deger")
    has = (ozet.get("hasilat") or {}).get("deger")
    net = (ozet.get("net_donem_kari") or {}).get("deger")
    don = (ozet.get("donen_varliklar") or {}).get("deger")
    oranlar = {}
    if net is not None and oz:
        oranlar["ROE_%"] = round(net / oz * 100, 2)
    if net is not None and tv:
        oranlar["ROA_%"] = round(net / tv * 100, 2)
    if net is not None and has:
        oranlar["net_kar_marji_%"] = round(net / has * 100, 2)
    if ty is not None and oz:
        oranlar["borc_ozkaynak"] = round(ty / oz, 3)
    if don is not None and kv:
        oranlar["cari_oran"] = round(don / kv, 3)

    gunce = col_labels[0] if col_labels else None

    # Global aktif dönemi hisse bazlı güncelle (daha yeni ise)
    try:
        state = _load_bilanco_state()
        prev = state.get("aktif_donem") or {}
        gy, gd = gunce_y, gunce_d
        if not prev.get("yil") or (gy, gd) > (int(prev.get("yil") or 0), int(prev.get("donem") or 0)):
            state["aktif_donem"] = {
                "yil": gy, "donem": gd, "etiket": gunce,
                "aciklama": "yıllık" if gd == 12 else f"{gd} aylık",
                "durum": "geldi",
                "son_guncelleme": datetime.now().isoformat(timespec="seconds"),
                "kaynak": f"İş Yatırım:{kod}",
            }
            _save_bilanco_state(state)
            try:
                st.session_state["aktif_bilanco"] = state["aktif_donem"]
            except Exception:
                pass
        # Hisse özel dönem kaydı
        per_hisse = state.setdefault("hisse_donem", {})
        per_hisse[kod] = {
            "yil": gy, "donem": gd, "etiket": gunce,
            "son_guncelleme": bugun, "kaynak": "İş Yatırım",
        }
        state["hisse_donem"] = per_hisse
        _save_bilanco_state(state)
    except Exception:
        pass

    out = {
        "kod": kod,
        "kaynak": "İş Yatırım MaliTablo",
        "financial_group": financial_group,
        "periods": periods_ord,
        "col_labels": col_labels,
        "gunce_donem": gunce,
        "gunce_yil": gunce_y,
        "gunce_donem_no": gunce_d,
        "ozet": ozet,
        "oranlar": oranlar,
        "bilanco_df": pd.DataFrame(bilanco_rows) if bilanco_rows else pd.DataFrame(),
        "gelir_df": pd.DataFrame(gelir_rows) if gelir_rows else pd.DataFrame(),
        "nakit_df": pd.DataFrame(nakit_rows) if nakit_rows else pd.DataFrame(),
        "url": f"https://www.isyatirim.com.tr/tr-tr/analiz/hisse/Sayfalar/sirket-karti.aspx?hisse={kod}",
        "_cache_day": bugun,
        "guncel_mi": True,
    }
    cache[ck] = out
    return out


def mali_tablo_iyilesme_kontrolu(
    sembol: str,
    yeni_metrikler: dict,
    *,
    tur: str = "hisse",
) -> list[str]:
    """
    Son bilanço / metrik setini önceki kayıtla karşılaştırır.
    İyileşme varsa uyarı metinleri döner (kullanıcıya yeşil bildirim).
    tur: 'hisse' | 'fon'
    """
    state = _load_bilanco_state()
    key_store = "hisse_metrik_onceki" if tur == "hisse" else "fon_metrik_onceki"
    onceki_map = state.get(key_store) or {}
    kod = str(sembol).upper().replace(".IS", "")
    onceki = onceki_map.get(kod) or {}
    uyarilar = []

    # Karşılaştırılacak alanlar: (alan, yön, etiket, eşik_abs)
    # yön: 'yukari_iyi' | 'asagi_iyi'
    checks = [
        ("ROE", "yukari_iyi", "ROE", 0.5),
        ("roa", "yukari_iyi", "ROA", 0.3),
        ("net_kar_marjı", "yukari_iyi", "Net kâr marjı", 0.3),
        ("F/K", "asagi_iyi", "F/K", 0.5),  # düşüş = ucuzlama (dikkatli yorum)
        ("PD/DD", "asagi_iyi", "PD/DD", 0.05),
        ("borç/özkaynak", "asagi_iyi", "Borç/Özkaynak", 0.05),
        ("getiri_1y", "yukari_iyi", "1Y getiri", 1.0),
        ("sharpe", "yukari_iyi", "Sharpe", 0.05),
        ("max_drawdown", "yukari_iyi", "MaxDD", 1.0),  # -20 → -10 iyileşme
        ("reel_1y", "yukari_iyi", "Reel 1Y", 1.0),
        ("vs_bist100", "yukari_iyi", "vs BIST100", 1.0),
        ("cagr", "yukari_iyi", "CAGR", 1.0),
    ]

    def _num(v):
        try:
            if v is None:
                return None
            x = float(v)
            if x != x:  # nan
                return None
            return x
        except (TypeError, ValueError):
            return None

    for alan, yon, etiket, esik in checks:
        # yeni metriklerde çeşitli anahtar varyasyonları
        nv = _num(yeni_metrikler.get(alan))
        if nv is None:
            nv = _num(yeni_metrikler.get(alan.lower()))
        if nv is None and alan == "ROE":
            nv = _num(yeni_metrikler.get("returnOnEquity"))
        ov = _num(onceki.get(alan))
        if nv is None or ov is None:
            continue
        delta = nv - ov
        if yon == "yukari_iyi" and delta >= esik:
            uyarilar.append(
                f"⬆️ **{kod}** {etiket}: {ov:.2f} → **{nv:.2f}** (+{delta:.2f}) — önceki bilançoya göre iyileşme"
            )
        elif yon == "asagi_iyi" and delta <= -esik:
            # F/K veya borç düşüşü
            uyarilar.append(
                f"⬇️ **{kod}** {etiket}: {ov:.2f} → **{nv:.2f}** ({delta:.2f}) — önceki döneme göre geriledi (çoğu durumda olumlu)"
            )

    # Kaydı güncelle (her zaman son değerleri sakla)
    snap = {k: yeni_metrikler.get(k) for k in set([c[0] for c in checks]) | set(yeni_metrikler.keys())}
    snap["_kayit"] = datetime.now().isoformat(timespec="seconds")
    try:
        _ad = son_bilanco_donemi(probe=False)
        snap["_bilanco_donem"] = _ad.get("etiket")
        snap["_bilanco_yil"] = _ad.get("yil")
        snap["_bilanco_donem_no"] = _ad.get("donem")
    except Exception:
        pass
    onceki_map[kod] = snap
    state[key_store] = onceki_map
    _save_bilanco_state(state)
    return uyarilar


def render_bilanco_bildirim_paneli():
    """Ana sayfa: FON ve HİSSE bilanço bildirimlerini ayrı kartlarda gösterir."""
    st.subheader("📋 Bilanço & dönem bildirimleri")
    st.caption("Hisse (BIST/KAP) ve Fon (TEFAS) bildirimleri birbirinden bağımsızdır.")

    col_h, col_f = st.columns(2)

    # Yenile / probe
    p1, p2 = st.columns([1, 3])
    with p1:
        do_probe = st.button("🔄 Dönemi kontrol et", key="btn_bilanco_check")
    with p2:
        do_probe_live = st.button("Canlı probe (İş Yatırım — hisse)", key="btn_bilanco_probe")

    bild = bilanco_bildirimleri(probe=bool(do_probe_live))
    hisse, fon = bild["hisse"], bild["fon"]

    with col_h:
        st.markdown("#### 🏦 Hisse (BIST)")
        st.caption(hisse.get("kaynak") or "BIST")
        etiket = hisse.get("etiket") or hisse.get("beklenen_etiket") or "—"
        if hisse.get("yeni_donem") or hisse.get("durum") == "geldi":
            st.success(f"✅ Son dönem: **{etiket}**")
            if hisse.get("yeni_donem"):
                st.warning(hisse.get("mesaj") or "Yeni bilanço dönemi")
        elif hisse.get("durum") == "bekleniyor":
            st.warning(f"⏳ {hisse.get('mesaj') or 'Bekleniyor'}")
        else:
            st.info(hisse.get("mesaj") or f"Dönem: {etiket}")
        if hisse.get("probe") is True:
            st.success("Probe: bilançolar erişilebilir")
        elif hisse.get("probe") is False:
            st.warning("Probe: henüz tam yayınlanmamış olabilir")
        st.caption(f"Beklenen: {hisse.get('beklenen_etiket')}")

    with col_f:
        st.markdown("#### 📊 Fon (TEFAS)")
        st.caption(fon.get("kaynak") or "TEFAS")
        if fon.get("yeni_donem"):
            st.warning(fon.get("mesaj") or "Yeni çeyrek")
        else:
            st.info(fon.get("mesaj") or "Günlük NAV")
        st.caption(f"Raporlama çeyreği (takvim): {fon.get('etiket')}")
        st.caption("Fonlarda KAP bilançosu yok; portföy dağılımı / AUM tarama ile güncellenir.")

    # İyileşme uyarıları (önceki snapshot)
    iyilesme = st.session_state.get("_bilanco_iyilesme") or []
    if iyilesme:
        st.markdown("#### 🟢 Önceki bilançoya göre iyileşenler")
        for msg in iyilesme[:15]:
            st.success(msg)

    # Yeni dönem geldiyse izleme listesi otomatik yorum
    st.markdown("#### 🤖 Yeni bilanço — otomatik yorum")
    st.caption(
        "İzleme listesindeki hisselerde İş Yatırım güncel dönem çekilir; "
        "satış/kâr/borç/F/K kurallarıyla yorum üretilir. Tavsiye değildir."
    )
    izleme = st.text_input(
        "İzleme listesi (virgülle)",
        value=st.session_state.get("bilanco_izleme", "ISDMR,THYAO,GARAN,EREGL,ASELS,TUPRS"),
        key="bilanco_izleme",
    )
    if st.button("Bilanço yorumlarını üret", type="primary", key="btn_bilanco_yorum"):
        kodlar = [k.strip().upper().replace(".IS", "") for k in izleme.split(",") if k.strip()]
        yorumlar = []
        bar = st.progress(0.0)
        for i, k in enumerate(kodlar):
            try:
                isy = isyatirim_mali_tablo(k, force=False)
                deg = {}
                try:
                    ic = deg_from_isyatirim(k)
                    if not ic.get("_error"):
                        deg = ic
                except Exception:
                    pass
                y = mali_tablo_otomatik_yorum(isy if isinstance(isy, dict) else {}, deg)
                y["kod"] = k
                yorumlar.append(y)
            except Exception as e:
                yorumlar.append({"kod": k, "sonuc": f"Hata: {e}", "olumlu": [], "olumsuz": [], "notr": [], "aksiyon": []})
            bar.progress((i + 1) / max(len(kodlar), 1))
        bar.empty()
        st.session_state["_ana_bilanco_yorumlar"] = yorumlar

    for y in (st.session_state.get("_ana_bilanco_yorumlar") or []):
        with st.expander(f"{y.get('kod')} · {y.get('donem') or '—'} · {str(y.get('sonuc') or '')[:80]}", expanded=False):
            render_otomatik_bilanco_yorum(y, baslik=f"{y.get('kod')} bilanço yorumu")

    return bild


def df_to_excel_bytes(df: pd.DataFrame, meta: dict = None) -> bytes:
    buf = io.BytesIO()
    # core.excel_kaydet dosyaya yazar; burada belleğe yazıyoruz
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import ColorScaleRule

        sutun_map = {
            "kod": "Fon Kodu", "ad": "Fon Adı", "skor": "Skor", "puan": "Puan",
            "sira": "Sıra", "kurucu": "Kurucu", "tefas_durum": "TEFAS",
            "risk_degeri": "Risk", "faiz": "Faiz", "kategori": "Kategori",
            "getiri_1a": "1A %", "getiri_3a": "3A %", "getiri_6a": "6A %",
            "getiri_1y": "1Y %", "getiri_3y": "3Y %",
            "reel_1y": "Reel 1Y %", "vs_mevduat": "vs Mevduat",
            "vs_bist100": "vs BIST100", "vs_altin": "vs Altın",
            "volatilite": "Vol %", "max_drawdown": "MaxDD %",
            "sharpe": "Sharpe", "sortino": "Sortino", "calmar": "Calmar",
            "cagr": "CAGR %", "kazanma_orani": "Kazanma %",
            "var95": "VaR %95", "cvar95": "cVaR %95",
            "risk_skor": "Risk skor", "likidite_skor": "Likidite skor",
            "likidite_seviye": "Likidite seviye", "risk_bayrak": "Risk bayrakları",
            "yatirim": "Portföy yatırımı", "baskin_tur": "Baskın Tür",
            "akran_grubu": "Akran Grubu", "getiri_z": "Getiri z-skor",
            "rapor_skor": "Rapor skor", "mesafe_skoru": "Dağılım Sapması",
            "d_hisse": "Hisse %", "d_tahvil": "Tahvil %", "d_repo": "Repo %",
            "d_mevduat": "Mevduat %", "d_altin": "Altın %", "d_doviz": "Döviz %",
            "buyukluk": "Büyüklük", "yatirimci": "Yatırımcı",
            "giris_kom": "Giriş Kom.", "cikis_kom": "Çıkış Kom.",
        }
        mevcut = [c for c in sutun_map if c in df.columns]
        out = df[mevcut].copy()
        for c in out.columns:
            if pd.api.types.is_numeric_dtype(out[c]):
                out[c] = out[c].replace([np.inf, -np.inf], np.nan)
        out = out.rename(columns={c: sutun_map[c] for c in mevcut})

        wb = Workbook()
        ws = wb.active
        ws.title = "Fonlar"
        NAVY, WHITE = "1B2838", "FFFFFF"
        header_fill = PatternFill("solid", fgColor=NAVY)
        header_font = Font(name="Calibri", bold=True, color=WHITE, size=11)
        thin = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        for j, col in enumerate(out.columns, 1):
            cell = ws.cell(1, j, col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin
        for i, row in enumerate(out.itertuples(index=False), 2):
            for j, val in enumerate(row, 1):
                cell = ws.cell(i, j, None if (isinstance(val, float) and val != val) else val)
                cell.border = thin
                cell.font = Font(name="Calibri", size=10)
                if isinstance(val, (int, float)) and val == val:
                    cell.number_format = "0.00"
        for j, col in enumerate(out.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(col)) + 3, 10), 36)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(out.columns))}{len(out)+1}"
        ws.freeze_panes = "A2"

        if meta:
            ws2 = wb.create_sheet("Parametreler")
            ws2["A1"], ws2["B1"] = "Alan", "Değer"
            for cell in (ws2["A1"], ws2["B1"]):
                cell.fill = header_fill
                cell.font = header_font
            for i, (k, v) in enumerate(meta.items(), 2):
                ws2[f"A{i}"] = str(k)
                ws2[f"B{i}"] = str(v)
            ws2.column_dimensions["A"].width = 28
            ws2.column_dimensions["B"].width = 40

        wb.save(buf)
    except Exception:
        out2 = df.copy()
        out2.to_excel(buf, index=False)
    buf.seek(0)
    return buf.read()


def show_df(df: pd.DataFrame, cols: list, height: int = 420):
    if df is None or df.empty:
        st.info("Gösterilecek veri yok.")
        return
    use = [c for c in cols if c in df.columns]
    if not use:
        st.dataframe(df, use_container_width=True, height=height)
        return
    view = df[use].copy()
    pct_cols = {
        "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y", "getiri_3y",
        "reel_1y", "vs_mevduat", "vs_bist100", "vs_altin",
        "volatilite", "max_drawdown", "cagr", "kazanma_orani",
        "var95", "cvar95", "en_iyi_gun", "en_kotu_gun",
        "likidite_skor", "risk_skor", "secim_skor", "skor", "puan",
        "mesafe_skoru", "d_hisse", "d_tahvil", "d_repo", "d_mevduat",
        "d_altin", "d_doviz", "likit_varlik_pct",
    }
    ratio_cols = {"sharpe", "sortino", "calmar", "getiri_z", "rapor_skor", "bist_korelasyon"}
    for c in view.columns:
        if c in pct_cols and pd.api.types.is_numeric_dtype(view[c]):
            view[c] = pd.to_numeric(view[c], errors="coerce").round(2)
        elif c in ratio_cols and pd.api.types.is_numeric_dtype(view[c]):
            view[c] = pd.to_numeric(view[c], errors="coerce").round(2)
    col_cfg = {}
    try:
        for c in view.columns:
            if c in ("ad", "risk_bayrak", "yatirim"):
                col_cfg[c] = st.column_config.TextColumn(c, width="medium")
            elif c in pct_cols or c in ratio_cols:
                col_cfg[c] = st.column_config.NumberColumn(c, format="%.2f")
            elif c in ("kod", "sira"):
                col_cfg[c] = st.column_config.TextColumn(c, width="small")
    except Exception:
        col_cfg = None
    st.dataframe(
        view, use_container_width=True, height=height, hide_index=True,
        column_config=col_cfg if col_cfg else None,
    )


def save_excel_to_disk(df: pd.DataFrame, filename: str, meta: dict = None) -> Path | None:
    """Sonuçları programa ait klasöre Excel olarak yazar."""
    if df is None or df.empty:
        return None
    path = ROOT / filename
    try:
        # core.excel_kaydet varsa onu kullan (daha zengin format)
        core.excel_kaydet(df, str(path), meta=meta)
        return path
    except Exception:
        try:
            data = df_to_excel_bytes(df, meta)
            path.write_bytes(data)
            return path
        except Exception as e:
            st.warning(f"Excel diske yazılamadı: {e}")
            return None


def download_excel_button(df: pd.DataFrame, filename: str, meta: dict = None, key: str = "xl", auto_save: bool = True):
    """Hem indirme butonu gösterir hem (isteğe bağlı) diske kaydeder."""
    if df is None or df.empty:
        return
    if auto_save:
        saved = save_excel_to_disk(df, filename, meta)
        if saved:
            st.success(f"💾 Excel kaydedildi: `{saved}`")
    data = df_to_excel_bytes(df, meta)
    st.download_button(
        "📥 Excel indir",
        data=data,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=key,
    )


# ── Kategori seçimi (grup + tam ad) ─────────────────────────────────────────
KATEGORI_GRUPLARI = {
    "Hisse": ["hisse", "yabanci hisse", "endeks"],
    "Borçlanma araçları": ["borclanma", "tahvil", "bono", "eurobond", "dis borclanma"],
    "Para piyasası / Likit": ["para piyasasi", "likit", "kisa vadeli"],
    "Altın / Kıymetli maden": ["altin", "kiymetli maden", "gumus"],
    "Karma / Değişken": ["karma", "degisken", "dengeli"],
    "Serbest": ["serbest"],
    "Fon sepeti": ["fon sepeti"],
    "Katılım (faizsiz)": ["katilim"],
    "Döviz / Yabancı": ["doviz", "yabanci", "dolar", "euro"],
    "Girişim / Gayrimenkul": ["girisim", "gayrimenkul", "gyf"],
}


def kategori_maske(seri: pd.Series, gruplar=None, tam_adlar=None) -> pd.Series:
    """Grup (anahtar kelime) VEYA tam ad seçimlerinden biri tutan satırları işaretler."""
    if not gruplar and not tam_adlar:
        return pd.Series(True, index=seri.index)
    norm = seri.astype(str).apply(lambda x: core._normalize_tr(x).lower())
    m = pd.Series(False, index=seri.index)
    for g in gruplar or []:
        for kw in KATEGORI_GRUPLARI.get(g, []):
            m |= norm.str.contains(kw, na=False, regex=False)
    if tam_adlar:
        m |= seri.astype(str).isin(tam_adlar)
    return m


def kategori_secici(df, key: str, container=None):
    """Kategori grubu + (varsa) tam TEFAS kategorisi çoklu seçimi. (gruplar, tam_adlar) döner."""
    c = container or st
    gruplar = c.multiselect(
        "Kategori grubu", list(KATEGORI_GRUPLARI.keys()), key=f"{key}_grp",
        help="Birden fazla seçilebilir; seçilenlerden herhangi biri eşleşen fon kalır. Boş = hepsi",
    )
    tam = []
    if df is not None and not df.empty and "kategori" in df.columns:
        vc = df["kategori"].dropna().astype(str).str.strip()
        vc = vc[vc != ""].value_counts()
        if len(vc):
            tam = c.multiselect(
                "TEFAS kategorisi (tam ad)", vc.index.tolist(), key=f"{key}_tam",
                format_func=lambda k: f"{k} ({vc[k]})",
            )
    return gruplar, tam


def risk_badge(seviye: str) -> str:
    s = (seviye or "").lower()
    if s == "yüksek":
        return '<span class="badge badge-ok">Likidite yüksek</span>'
    if s == "orta":
        return '<span class="badge badge-warn">Likidite orta</span>'
    if s == "düşük":
        return '<span class="badge badge-bad">Likidite düşük</span>'
    return ""


# ── Sidebar ─────────────────────────────────────────────────────────────────
# ── Sayfa listesi (navigasyon) ─────────────────────────────────────────────
SAYFA_LISTESI = [
    "🏠 Ana Sayfa",
    "🔍 Fon Tarama",
    "📋 Tüm Fon Sıralaması",
    "🧪 Filtre Stüdyosu",
    "⚖️ Karşılaştırma",
    "📦 Rapor Paketleri",
    "📁 Portföy",
    "💼 Portföyüm",
    "🎯 Optimizasyon",
    "🔎 Fon Detay",
    "📈 Hisse & Karma Portföy",
    "🔎 Hisse Detay",
    "🏦 BIST Değerleme",
    "⭐ Favoriler",
    "⚙️ Ayarlar",
]


def _goto(sayfa_adi: str, *, auto_sepet: bool = False):
    """Programatik sayfa geçişi (sidebar butonları / sepet analizi)."""
    if sayfa_adi in SAYFA_LISTESI:
        st.session_state.sayfa = sayfa_adi
    if auto_sepet:
        st.session_state.sepet_auto_analiz = True


# ── Sidebar ─────────────────────────────────────────────────────────────────
with st.sidebar:
    st.title("📊 M.Durmaz")
    st.caption("Fon tarama · risk · portföy · rapor")

    # key="sayfa" → radio doğrudan session_state.sayfa yazar (_goto ile uyumlu)
    _cur = st.session_state.get("sayfa", SAYFA_LISTESI[0])
    if _cur not in SAYFA_LISTESI:
        st.session_state.sayfa = SAYFA_LISTESI[0]
        _cur = SAYFA_LISTESI[0]
    sayfa = st.radio(
        "Menü",
        SAYFA_LISTESI,
        index=SAYFA_LISTESI.index(_cur),
        label_visibility="collapsed",
        key="sayfa",
    )

    try:
        tefas_portfoyum.sidebar_rozet(st)
    except Exception:
        pass
    st.divider()
    tip = st.selectbox("Fon tipi", ["YAT", "EMK", "BYF"], index=["YAT", "EMK", "BYF"].index(st.session_state.get("fon_tipi", "YAT")) if st.session_state.get("fon_tipi") in ("YAT", "EMK", "BYF") else 0)
    if tip != st.session_state.fon_tipi:
        st.session_state.fon_tipi = tip
        st.session_state.analiz = None  # tip değişince motoru yenile

    cfg = st.session_state.cfg
    st.caption(
        f"Risksiz faiz: %{cfg.get('risk_free', 0.4)*100:.0f} · "
        f"TÜFE: %{cfg.get('tufe_yillik')} · "
        f"Mevduat: %{cfg.get('mevduat_yillik')}"
    )
    try:
        st.caption(f"Cache: `{core.CACHE_DIR.name}/` · v{st.session_state.get('app_version', '2.1')}")
    except Exception:
        st.caption(f"v{st.session_state.get('app_version', '2.1')}")
    sepet = st.session_state.get("portfoy_sepet") or []
    st.divider()
    st.markdown(f"### 🛒 Sepet · {len(sepet)}")
    if sepet:
        st.caption(", ".join(x["kod"] for x in sepet[:15]) + ("…" if len(sepet) > 15 else ""))
        b1, b2 = st.columns(2)
        with b1:
            if st.button("Analiz et", key="sepet_analiz_side", use_container_width=True):
                _goto("📁 Portföy", auto_sepet=True)
                st.rerun()
        with b2:
            if st.button("Temizle", key="sepet_temiz_side", use_container_width=True):
                st.session_state.portfoy_sepet = []
                st.session_state.sepet_analiz = None
                st.rerun()
    else:
        st.caption("Tarama / sıralama / detay → Portföye ekle")

# Aktif sayfa (state kaynaklı)
sayfa = st.session_state.get("sayfa", SAYFA_LISTESI[0])


# ═══════════════════════════════════════════════════════════════════════════
# ANA SAYFA
# ═══════════════════════════════════════════════════════════════════════════


def tefas_fon_gunluk_getiriler(top_n: int = 15) -> tuple[pd.DataFrame, str]:
    """
    TEFAS kaynaklı fon günlük getirileri.
    İş Yatırım kullanılmaz. Sıra: core → analiz.api → TEFAS HTTP → boş.
    """
    import urllib.request
    import urllib.parse
    import json as _json

    def _norm_df(df: pd.DataFrame, kaynak: str) -> pd.DataFrame | None:
        if df is None or not isinstance(df, pd.DataFrame) or df.empty:
            return None
        work = df.copy()
        colmap = {}
        for c in work.columns:
            cl = str(c).lower().replace(" ", "").replace("_", "")
            if cl in ("kod", "fonkod", "fonkodu", "code", "fonkodu"):
                colmap[c] = "kod"
            elif cl in ("ad", "fonad", "fonadi", "unvan", "title", "name", "fonunvan"):
                colmap[c] = "ad"
            elif any(x in cl for x in ("gunlukgetiri", "getiri1g", "dailyreturn", "gunluk", "degisim1g", "1g")):
                colmap[c] = "degisim"
            elif cl in ("getiri", "return", "degisim", "pct"):
                if "degisim" not in colmap.values():
                    colmap[c] = "degisim"
            elif "kategori" in cl or "fontur" in cl:
                colmap[c] = "kategori"
            elif "kurucu" in cl or "founder" in cl:
                colmap[c] = "kurucu"
            elif "fiyat" in cl or "price" in cl or "nav" in cl:
                colmap[c] = "fiyat"
        work = work.rename(columns=colmap)
        if "kod" not in work.columns:
            return None
        if "degisim" not in work.columns:
            return None
        work["kod"] = work["kod"].astype(str).str.strip().str.upper()
        work["degisim"] = pd.to_numeric(work["degisim"], errors="coerce")
        # TEFAS bazen 0.10 (% değil oran) döner
        med = work["degisim"].abs().median()
        if pd.notna(med) and 0 < med < 0.5:
            work["degisim"] = work["degisim"] * 100.0
        work = work.dropna(subset=["degisim"])
        work = work[(work["degisim"] > -25) & (work["degisim"] < 25)]
        work = work.drop_duplicates("kod", keep="first")
        if work.empty:
            return None
        work.attrs["kaynak"] = kaynak
        return work

    # ── 1) core TEFAS fonksiyonları ───────────────────────────────────────
    for fn_name in (
        "tefas_gunluk_degisim",
        "fon_gunluk_getiriler",
        "tum_fon_gunluk",
        "tefas_performans_ozet",
        "fon_listesi_getiri",
        "tefas_en_cok_kazandiran",
        "bind_comparison_returns",
    ):
        fn = getattr(core, fn_name, None)
        if not callable(fn):
            continue
        try:
            data = fn()
            if isinstance(data, pd.DataFrame):
                d = _norm_df(data, f"core.{fn_name}")
            elif isinstance(data, dict):
                d = _norm_df(pd.DataFrame(data.get("rows") or data.get("data") or data.get("value") or []), f"core.{fn_name}")
            elif isinstance(data, list):
                d = _norm_df(pd.DataFrame(data), f"core.{fn_name}")
            else:
                d = None
            if d is not None and len(d) >= 5:
                return d, d.attrs.get("kaynak") or f"core.{fn_name}"
        except Exception:
            continue

    # ── 2) get_analiz().api (tefas_core motor) ────────────────────────────
    try:
        analiz = get_analiz()
        api = getattr(analiz, "api", None)
        for meth in ("gunluk_getiriler", "fon_getirileri", "comparison_returns", "bind_fund_returns", "tum_fonlar"):
            fn = getattr(api, meth, None) if api is not None else None
            if not callable(fn):
                continue
            try:
                data = fn()
                if isinstance(data, pd.DataFrame):
                    d = _norm_df(data, f"analiz.api.{meth}")
                elif isinstance(data, list):
                    d = _norm_df(pd.DataFrame(data), f"analiz.api.{meth}")
                else:
                    d = None
                if d is not None and len(d) >= 5:
                    return d, d.attrs.get("kaynak")
            except Exception:
                continue
        # profil + bilinen kod listesinden son 2 NAV → günlük %
        try:
            codes = []
            profil = getattr(analiz, "profil", None) or {}
            if not profil:
                for loader in ("load_profil", "profil_yukle", "get_profil"):
                    fn = getattr(analiz, loader, None) or getattr(getattr(analiz, "api", None), loader, None)
                    if callable(fn):
                        try:
                            profil = fn() or {}
                            break
                        except Exception:
                            pass
            if isinstance(profil, dict) and profil:
                codes = list(profil.keys())
            # core fon listesi
            for fn_name in ("fon_kodlari", "tum_fon_kodlari", "list_fon_kod"):
                fn = getattr(core, fn_name, None)
                if callable(fn):
                    try:
                        extra = fn()
                        if isinstance(extra, (list, tuple, set)):
                            codes = list(dict.fromkeys(list(codes) + [str(x).upper() for x in extra]))
                        elif isinstance(extra, pd.DataFrame) and "kod" in extra.columns:
                            codes = list(dict.fromkeys(list(codes) + extra["kod"].astype(str).str.upper().tolist()))
                    except Exception:
                        pass
            # session tarama kodları öncelikli
            for key in ("son_tarama", "son_siralama"):
                dfx = st.session_state.get(key)
                if dfx is not None and hasattr(dfx, "columns") and "kod" in getattr(dfx, "columns", []):
                    codes = list(dict.fromkeys(
                        dfx["kod"].astype(str).str.upper().tolist() + list(codes)
                    ))
            codes = [c for c in codes if c and len(c) <= 5][:120]
            rows = []
            api = getattr(analiz, "api", None)
            fiyat_fn = None
            if api is not None:
                for name in ("fiyat", "get_fiyat", "nav_series", "fon_fiyat"):
                    if callable(getattr(api, name, None)):
                        fiyat_fn = getattr(api, name)
                        break
            if fiyat_fn is None:
                for name in ("fon_fiyat", "get_fiyat_df"):
                    if callable(getattr(analiz, name, None)):
                        fiyat_fn = getattr(analiz, name)
                        break
                    if callable(getattr(core, name, None)):
                        fiyat_fn = getattr(core, name)
                        break
            if fiyat_fn is not None and codes:
                for kod in codes:
                    try:
                        try:
                            fd = fiyat_fn(kod, periyod=1)
                        except TypeError:
                            try:
                                fd = fiyat_fn(kod)
                            except Exception:
                                continue
                        if fd is None or getattr(fd, "empty", True):
                            continue
                        s = fd.copy()
                        if "tarih" in s.columns:
                            s = s.sort_values("tarih")
                        pxcol = "fiyat" if "fiyat" in s.columns else (
                            "Fiyat" if "Fiyat" in s.columns else (
                                "close" if "close" in s.columns else None
                            )
                        )
                        if not pxcol:
                            continue
                        px = pd.to_numeric(s[pxcol], errors="coerce").dropna()
                        if len(px) < 2:
                            continue
                        last, prev = float(px.iloc[-1]), float(px.iloc[-2])
                        if prev == 0:
                            continue
                        pct = (last / prev - 1.0) * 100.0
                        ad = ""
                        if isinstance(profil, dict):
                            p = profil.get(kod) or profil.get(kod.lower()) or {}
                            if isinstance(p, dict):
                                ad = p.get("ad") or p.get("unvan") or ""
                        rows.append({"kod": kod, "ad": ad, "degisim": pct, "fiyat": last})
                    except Exception:
                        continue
            d = _norm_df(pd.DataFrame(rows), "TEFAS NAV son 2 gün")
            if d is not None and len(d) >= 5:
                return d, f"TEFAS NAV son 2 gün · {len(d)} fon"
        except Exception:
            pass
    except Exception:
        pass

    # ── 3) Resmi TEFAS HTTP (BindComparisonFundReturns — 1 gün) ───────────
    def _tefas_post(path: str, payload: dict):
        url = "https://www.tefas.gov.tr" + path
        body = _json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            url,
            data=body,
            method="POST",
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; TEFAS-GUI/2.3)",
                "Content-Type": "application/json; charset=UTF-8",
                "Accept": "application/json, text/plain, */*",
                "Origin": "https://www.tefas.gov.tr",
                "Referer": "https://www.tefas.gov.tr/",
            },
        )
        with urllib.request.urlopen(req, timeout=30) as resp:
            return _json.loads(resp.read().decode("utf-8", errors="replace"))

    today = datetime.now()
    # son iş gününe yakın aralık
    for fontip in ("YAT", "EMK", ""):
        payload = {
            "calismatipi": "2",
            "fontip": fontip,
            "sfontur": "",
            "kurucukod": "",
            "fongrup": "",
            "bastarih": (today.replace(day=1)).strftime("%d.%m.%Y") if False else (today - __import__("datetime").timedelta(days=10)).strftime("%d.%m.%Y"),
            "bittarih": today.strftime("%d.%m.%Y"),
            "fonturkod": "",
            "fonunvan": "",
            "strperiod": "1 gün",
            "periodeks": "1",
        }
        # alternatif alan adları
        payloads = [
            payload,
            {
                "fontip": fontip or "YAT",
                "sfontur": "",
                "fonunvan": "",
                "kurucukod": "",
                "fongrup": "",
                "bastarih": (today - __import__("datetime").timedelta(days=7)).strftime("%d.%m.%Y"),
                "bittarih": today.strftime("%d.%m.%Y"),
                "fonturkod": "",
                "fonperiod": "1",
                "strperiod": "1 gün",
                "periodeks": "01",
            },
        ]
        for pl in payloads:
            for path in (
                "/api/DB/BindComparisonFundReturns",
                "/api/DB/BindFundReturns",
                "/api/DB/BindComparisonFundInfo",
            ):
                try:
                    raw = _tefas_post(path, pl)
                    data = raw.get("data") or raw.get("Data") or raw.get("value") or raw
                    if isinstance(data, str):
                        try:
                            data = _json.loads(data)
                        except Exception:
                            continue
                    if not data:
                        continue
                    if isinstance(data, dict):
                        data = data.get("data") or data.get("result") or []
                    df = pd.DataFrame(data if isinstance(data, list) else [])
                    if df.empty:
                        continue
                    d = _norm_df(df, f"TEFAS {path}")
                    if d is not None and len(d) >= 5:
                        return d, f"TEFAS HTTP {path}"
                except Exception:
                    continue

    # ── 4) Son tarama (TEFAS taramasından gelmiş metrik — 1g yoksa 1a) ───
    for key, label in (("son_tarama", "tarama"), ("son_siralama", "sıralama")):
        df = st.session_state.get(key)
        if df is None or not hasattr(df, "empty") or df.empty:
            continue
        work = df.copy()
        metric = None
        for c in ("getiri_1g", "getiri_1h", "getiri_1a"):
            if c in work.columns and pd.to_numeric(work[c], errors="coerce").notna().sum() >= 5:
                metric = c
                break
        if not metric:
            continue
        work["degisim"] = pd.to_numeric(work[metric], errors="coerce")
        d = _norm_df(work, f"TEFAS session {label} · {metric}")
        if d is not None and len(d) >= 3:
            return d, d.attrs.get("kaynak")

    return pd.DataFrame(), "TEFAS günlük veri yok"


def _fon_movers_df(top_n: int = 10) -> tuple[pd.DataFrame, pd.DataFrame, str]:
    """
    Fon en çok artan / düşen — yalnızca TEFAS kaynaklı (İş Yatırım yok).
    """
    work, kaynak = tefas_fon_gunluk_getiriler(top_n=max(top_n * 3, 30))
    if work is None or not isinstance(work, pd.DataFrame) or work.empty:
        return (
            pd.DataFrame(),
            pd.DataFrame(),
            kaynak or "TEFAS günlük getiri alınamadı — Fon Tarama çalıştırın",
        )
    cols = [c for c in ("kod", "ad", "degisim", "fiyat", "kategori", "kurucu") if c in work.columns]
    artan = work.nlargest(top_n, "degisim")[cols]
    dusen = work.nsmallest(top_n, "degisim")[cols]
    return artan, dusen, kaynak

def _hisse_movers_df(top_n: int = 10) -> tuple[pd.DataFrame, pd.DataFrame, str]:
    """
    BIST hisse günlük en çok artan / düşen.
    Öncelik: core BIST movers → yfinance izleme listesi.
    """
    # 1) core
    for fn_name in ("bist_gunluk_movers", "bist_artan_dusen", "hisse_gunluk_degisim", "bist_piyasa_ozet"):
        fn = getattr(core, fn_name, None)
        if not callable(fn):
            continue
        try:
            data = fn()
            if isinstance(data, dict):
                up = pd.DataFrame(data.get("artan") or data.get("yukselen") or [])
                dn = pd.DataFrame(data.get("dusen") or data.get("gerileyen") or [])
                if not up.empty or not dn.empty:
                    return up.head(top_n), dn.head(top_n), f"core.{fn_name}"
            if isinstance(data, pd.DataFrame) and not data.empty:
                df = data.copy()
                colmap = {}
                for c in df.columns:
                    cl = str(c).lower()
                    if cl in ("kod", "sembol", "symbol", "ticker"):
                        colmap[c] = "kod"
                    elif "degisim" in cl or "change" in cl or cl in ("pct", "yuzde", "%"):
                        colmap[c] = "degisim"
                    elif cl in ("ad", "name", "title"):
                        colmap[c] = "ad"
                    elif "fiyat" in cl or cl == "price" or cl == "close":
                        colmap[c] = "fiyat"
                df = df.rename(columns=colmap)
                if "degisim" in df.columns:
                    df["degisim"] = pd.to_numeric(df["degisim"], errors="coerce")
                    df = df.dropna(subset=["degisim"])
                    cols = [c for c in ("kod", "ad", "fiyat", "degisim") if c in df.columns]
                    return df.nlargest(top_n, "degisim")[cols], df.nsmallest(top_n, "degisim")[cols], f"core.{fn_name}"
        except Exception:
            continue

    # 2) yfinance — likit BIST izleme listesi (XU100 temsili)
    watch = [
        "THYAO", "ASELS", "BIMAS", "EREGL", "FROTO", "GARAN", "HALKB", "ISCTR",
        "KCHOL", "KOZAL", "KRDMD", "PETKM", "PGSUS", "SAHOL", "SISE", "TAVHL",
        "TCELL", "TOASO", "TUPRS", "YKBNK", "AKBNK", "ARCLK", "EKGYO", "ENKAI",
        "GUBRF", "HEKTS", "ISDMR", "KONTR", "MGROS", "OTKAR", "SASA", "SMRTG",
        "SOKM", "TATGD", "ULKER", "VESTL", "YEOTK", "ZOREN", "ALARK", "DOHOL",
    ]
    try:
        import yfinance as yf
    except ImportError:
        return pd.DataFrame(), pd.DataFrame(), "yfinance yok — pip install yfinance"

    rows = []
    # toplu indir (son 5 gün)
    tickers = [f"{k}.IS" for k in watch]
    try:
        data = yf.download(
            tickers, period="5d", group_by="ticker",
            auto_adjust=True, threads=True, progress=False,
        )
    except Exception as e:
        return pd.DataFrame(), pd.DataFrame(), f"yfinance hata: {e}"

    for kod in watch:
        t = f"{kod}.IS"
        try:
            if len(watch) == 1:
                sub = data
            else:
                # multi-index columns
                if isinstance(data.columns, pd.MultiIndex):
                    if t not in data.columns.get_level_values(0):
                        continue
                    sub = data[t]
                else:
                    continue
            if sub is None or sub.empty:
                continue
            close = sub["Close"].dropna() if "Close" in sub.columns else sub.iloc[:, 0].dropna()
            if len(close) < 2:
                continue
            last = float(close.iloc[-1])
            prev = float(close.iloc[-2])
            if prev == 0:
                continue
            pct = (last / prev - 1.0) * 100.0
            # tek günde ±40% üstü veri hatası olabilir
            if abs(pct) > 40:
                continue
            rows.append({"kod": kod, "fiyat": round(last, 2), "degisim": round(pct, 2)})
        except Exception:
            continue

    if not rows:
        return pd.DataFrame(), pd.DataFrame(), "Hisse verisi alınamadı"

    df = pd.DataFrame(rows).dropna(subset=["degisim"])
    artan = df.nlargest(top_n, "degisim")
    dusen = df.nsmallest(top_n, "degisim")
    return artan, dusen, f"yfinance 1g · {len(df)} hisse"


def render_artan_dusen_paneli():
    """Ana Sayfa: fon + hisse en çok artan / düşen (doğru sıralama + filtre)."""
    st.subheader("📊 En çok artan / düşen")
    st.caption(
        "Fon: **yalnızca TEFAS** günlük getiri (core / api / tefas.gov.tr; İş Yatırım değil). "
        "Hisse: BIST günlük % (yfinance / core). Aşırı uçlar elenir."
    )
    c_refresh, _ = st.columns([1, 4])
    with c_refresh:
        do_refresh = st.button("Yenile", key="btn_movers_refresh")

    cache_key = "_movers_cache"
    if do_refresh or cache_key not in st.session_state:
        with st.spinner("Artan / düşen hesaplanıyor…"):
            fa, fd, fsrc = _fon_movers_df(10)
            ha, hd, hsrc = _hisse_movers_df(10)
            st.session_state[cache_key] = {
                "fa": fa, "fd": fd, "fsrc": fsrc,
                "ha": ha, "hd": hd, "hsrc": hsrc,
                "ts": datetime.now().strftime("%H:%M:%S"),
            }
    cache = st.session_state.get(cache_key) or {}
    st.caption(f"Son hesap: {cache.get('ts') or '—'} · Fon: {cache.get('fsrc')} · Hisse: {cache.get('hsrc')}")

    col_f, col_h = st.columns(2)
    with col_f:
        st.markdown("**Fonlar**")
        fa, fd = cache.get("fa"), cache.get("fd")
        t1, t2 = st.tabs(["🟢 Artan", "🔴 Düşen"])
        with t1:
            if fa is not None and hasattr(fa, "empty") and not fa.empty:
                show = fa.copy()
                if "degisim" in show.columns:
                    show["degisim"] = show["degisim"].map(lambda x: f"%{tr_num(x, 2)}" if pd.notna(x) else "—")
                st.dataframe(show, hide_index=True, use_container_width=True)
            else:
                st.caption("Fon artan listesi yok. «Fon Tarama» çalıştırın veya TEFAS API bağlayın.")
        with t2:
            if fd is not None and hasattr(fd, "empty") and not fd.empty:
                show = fd.copy()
                if "degisim" in show.columns:
                    show["degisim"] = show["degisim"].map(lambda x: f"%{tr_num(x, 2)}" if pd.notna(x) else "—")
                st.dataframe(show, hide_index=True, use_container_width=True)
            else:
                st.caption("Fon düşen listesi yok.")
    with col_h:
        st.markdown("**Hisse senetleri**")
        ha, hd = cache.get("ha"), cache.get("hd")
        t3, t4 = st.tabs(["🟢 Artan", "🔴 Düşen"])
        with t3:
            if ha is not None and hasattr(ha, "empty") and not ha.empty:
                show = ha.copy()
                if "degisim" in show.columns:
                    show["degisim"] = show["degisim"].map(lambda x: f"%{tr_num(x, 2)}" if pd.notna(x) else "—")
                st.dataframe(show, hide_index=True, use_container_width=True)
            else:
                st.caption("Hisse artan listesi yok. yfinance kurulu mu?")
        with t4:
            if hd is not None and hasattr(hd, "empty") and not hd.empty:
                show = hd.copy()
                if "degisim" in show.columns:
                    show["degisim"] = show["degisim"].map(lambda x: f"%{tr_num(x, 2)}" if pd.notna(x) else "—")
                st.dataframe(show, hide_index=True, use_container_width=True)
            else:
                st.caption("Hisse düşen listesi yok.")

if sayfa == "🏠 Ana Sayfa":
    st.markdown(
        """
<div class="hero">
  <h1>M.Durmaz · TEFAS Fon Analiz</h1>
  <p>Günlük özet · alarm · track record · net getiri · bilanço bildirimi</p>
</div>
""",
        unsafe_allow_html=True,
    )

    try:
        tefas_panel.render_ana_panel()
    except Exception as e:
        st.caption(f"Panel: {e}")

    sepet_n = len(st.session_state.get("portfoy_sepet") or [])
    sa = st.session_state.get("sepet_analiz") or {}
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Fon tipi", st.session_state.fon_tipi)
    c2.metric("Risksiz faiz", f"%{st.session_state.cfg.get('risk_free', 0.4)*100:.0f}")
    c3.metric("TÜFE (yıllık)", f"%{st.session_state.cfg.get('tufe_yillik', '—')}")
    c4.metric("Sepet", f"{sepet_n} fon")
    if sa.get("diversify_score") is not None:
        c5.metric("Çeşitlendirme", f"{sa.get('diversify_score')}/100")
    else:
        c5.metric("Çeşitlendirme", "—")

    # ── Günlük özet ────────────────────────────────────────────────────────
    st.subheader("📌 Bugünün özeti")
    o1, o2, o3 = st.columns(3)
    with o1:
        st.markdown("**Sepet**")
        if sepet_n == 0:
            st.caption("Sepet boş · Portföy sayfasından ekleyin")
        elif sa and sa.get("port_metrikler"):
            pm = sa["port_metrikler"]
            st.write(
                f"1Y **%{pm.get('getiri_1y') or sa.get('ort_getiri_1y') or '—'}** · "
                f"Sharpe **{pm.get('sharpe') or '—'}** · "
                f"MaxDD **%{pm.get('max_drawdown') or '—'}**"
            )
            if sa.get("ort_tehlike_yazi"):
                st.caption(f"Ort. tehlike: {sa.get('ort_tehlike_yazi')}")
        else:
            st.caption(f"{sepet_n} fon sepette · henüz analiz yok → Portföy’de «analiz et»")
    with o2:
        st.markdown("**Son tarama**")
        meta = st.session_state.get("son_tarama_meta") or {}
        df_t = st.session_state.get("son_tarama")
        if df_t is not None and hasattr(df_t, "empty") and not df_t.empty:
            st.write(f"**{len(df_t)}** fon · {meta.get('risk_paketi') or meta.get('mod') or 'tarama'}")
            if "getiri_1y" in df_t.columns:
                top = df_t.nlargest(3, "getiri_1y") if "getiri_1y" in df_t.columns else df_t.head(3)
                st.caption("1Y lider: " + ", ".join(
                    f"{r.get('kod')} %{r.get('getiri_1y'):.0f}" if pd.notna(r.get("getiri_1y")) else str(r.get("kod"))
                    for _, r in top.iterrows()
                ))
        else:
            st.caption("Henüz tarama yok")
    with o3:
        st.markdown("**Bilanço (özet)**")
        try:
            _bmini = bilanco_bildirimleri(probe=False)
            _h, _f = _bmini["hisse"], _bmini["fon"]
            if _h.get("yeni_donem"):
                st.warning(f"🏦 Hisse: yeni dönem {_h.get('etiket')}")
            elif _h.get("durum") == "geldi":
                st.success(f"🏦 Hisse: {_h.get('etiket')}")
            else:
                st.caption(f"🏦 Hisse: {_h.get('mesaj') or _h.get('etiket')}")
            if _f.get("yeni_donem"):
                st.warning(f"📊 Fon: {_f.get('etiket')} çeyrek")
            else:
                st.caption(f"📊 Fon: günlük NAV · çeyrek {_f.get('etiket')}")
        except Exception as e:
            st.caption(f"Bilanço: {e}")

    # ── Bilanço bildirimleri (FON / HİSSE ayrı paneller) ───────────────────
    try:
        render_bilanco_bildirim_paneli()
    except Exception as e:
        st.caption(f"Bilanço paneli: {e}")

    # ── En çok artan / düşen (doğrulanmış sıralama) ────────────────────────
    try:
        render_artan_dusen_paneli()
    except Exception as e:
        st.caption(f"Artan/düşen paneli: {e}")

    # ── Alarmlar ───────────────────────────────────────────────────────────
    st.subheader("🔔 Alarmlar")
    try:
        alarm_msgs = list(core.sepet_alarmlari(sa, st.session_state.cfg) or [])
    except Exception:
        alarm_msgs = []
    # Bilanço iyileşme uyarılarını alarm listesine ekle
    for msg in (st.session_state.get("_bilanco_iyilesme") or []):
        alarm_msgs.append(msg)
    # Ayar alarmları son tarama üzerinde
    al_cfg = (st.session_state.cfg or {}).get("alarmlar") or {}
    df_t = st.session_state.get("son_tarama")
    if df_t is not None and hasattr(df_t, "empty") and not df_t.empty:
        for _, row in df_t.head(30).iterrows():
            kod = row.get("kod")
            if al_cfg.get("min_getiri_1y") is not None and row.get("getiri_1y") is not None:
                if float(row["getiri_1y"]) < float(al_cfg["min_getiri_1y"]):
                    alarm_msgs.append(f"{kod} 1Y %{row['getiri_1y']:.1f} < min")
            if al_cfg.get("min_sharpe") is not None and row.get("sharpe") is not None:
                if float(row["sharpe"]) < float(al_cfg["min_sharpe"]):
                    alarm_msgs.append(f"{kod} Sharpe {row['sharpe']:.2f} < min")
            if al_cfg.get("max_drawdown") is not None and row.get("max_drawdown") is not None:
                if float(row["max_drawdown"]) < float(al_cfg["max_drawdown"]):
                    alarm_msgs.append(f"{kod} MaxDD %{row['max_drawdown']:.1f}")
    if alarm_msgs:
        for m in alarm_msgs[:12]:
            st.error(m)
    else:
        st.caption("Aktif alarm yok. Eşikler: Ayarlar + sepet MaxDD / çeşitlendirme / tehlike.")

    # ── Track record ───────────────────────────────────────────────────────
    st.subheader("📈 Track record")
    tr1, tr2 = st.columns(2)
    with tr1:
        st.markdown("**Fon tarama arşivi**")
        tr_fon = core.evaluate_fon_tarama_track(analiz=None)
        st.caption(tr_fon.get("note") or "—")
        if st.button("Arşivdeki fonları bugünle karşılaştır", key="btn_tr_fon"):
            with st.spinner("Fon detayları çekiliyor…"):
                tr_fon = core.evaluate_fon_tarama_track(analiz=get_analiz())
            st.session_state["_tr_fon"] = tr_fon
        tr_fon = st.session_state.get("_tr_fon") or tr_fon
        if tr_fon.get("rows"):
            st.dataframe(pd.DataFrame(tr_fon["rows"]), hide_index=True, use_container_width=True)
    with tr2:
        st.markdown("**BIST skor arşivi**")
        tr_b = core.bist_track_record_ozet()
        st.caption(tr_b.get("note") or "—")
        if tr_b.get("perf_df") is not None:
            st.dataframe(tr_b["perf_df"], hide_index=True, use_container_width=True)
        elif tr_b.get("ok"):
            st.caption("Performans Takibi sayfası 6A+ birikim sonrası dolar.")

    # ── Net getiri (örnek sepet / son tarama) ───────────────────────────────
    st.subheader("💸 Net getiri (yönetim ücreti düşülmüş)")
    st.caption(
        "TEFAS profilde ücret alanı varsa brüt → net yaklaşık hesaplanır. "
        "Ücret yoksa satır atlanır. Formül: net ≈ (1+g)/(1+fee·dönem)−1"
    )
    net_rows = []
    try:
        analiz = get_analiz()
        kodlar_net = []
        if sepet_n:
            kodlar_net = [str(x.get("kod") or x).upper() for x in (st.session_state.get("portfoy_sepet") or [])][:8]
        elif df_t is not None and hasattr(df_t, "empty") and not df_t.empty:
            kodlar_net = [str(k).upper() for k in df_t["kod"].head(5).tolist()]
        if kodlar_net and st.button("Sepet/tarama için net getiri hesapla", key="btn_net_getiri"):
            with st.spinner("Profil + metrik…"):
                for kod in kodlar_net:
                    try:
                        det = analiz.fon_detay(kod)
                        p, o, m = det.get("profil") or {}, det.get("ozet") or {}, det.get("metrikler") or {}
                        fee = core.profil_yonetim_ucreti(p, o)
                        mn = core.metrik_net_getiriler(m, fee)
                        net_rows.append({
                            "kod": kod,
                            "ad": (o.get("fonUnvan") or p.get("fonUnvan") or kod)[:40],
                            "ücret_%": fee,
                            "1Y_brüt": m.get("getiri_1y"),
                            "1Y_net": mn.get("net_getiri_1y"),
                            "6A_net": mn.get("net_getiri_6a"),
                            "3A_net": mn.get("net_getiri_3a"),
                        })
                    except Exception:
                        continue
            st.session_state["_net_rows"] = net_rows
        net_rows = st.session_state.get("_net_rows") or []
        if net_rows:
            st.dataframe(pd.DataFrame(net_rows), hide_index=True, use_container_width=True)
        else:
            st.caption("Hesaplamak için butona basın (sepet veya son tarama gerekir).")
    except Exception as e:
        st.caption(f"Net getiri: {e}")

    with st.expander("Önerilen akış & modüller", expanded=False):
        f1, f2, f3, f4 = st.columns(4)
        with f1:
            st.markdown("**1. Tarama** — aday havuzu")
        with f2:
            st.markdown("**2. Akıllı seçim** — skor")
        with f3:
            st.markdown("**3. Sepet** — analiz + HTML")
        with f4:
            st.markdown("**4. Detay** — rapor")
        st.markdown(
            """
| Modül | Ne işe yarar |
|-------|----------------|
| **Fon Tarama / Sıralama** | Aday havuzu + arşive kayıt (track record) |
| **Portföy** | Sepet, alarm girdileri, çeşitlendirme |
| **BIST Değerleme** | Skor arşivi + bilanço dönemi |
| **Ayarlar** | Alarm eşikleri (1Y, Sharpe, MaxDD, vol) |
"""
        )
        st.caption("Geçmiş performans geleceği garanti etmez. Yatırım tavsiyesi değildir.")

    favs = core.load_favs()
    if favs:
        st.caption(f"Favoriler ({len(favs)}): " + ", ".join(favs[:20]) + ("…" if len(favs) > 20 else ""))


elif sayfa == "🔍 Fon Tarama":
    st.header("🔍 Fon Tarama")
    st.caption("Hedef dağılım + risk + kalite kapıları → sıralı adaylar → akıllı seçim / portföy")

    # ── Hızlı senaryo ──────────────────────────────────────────────────────
    st.subheader("⚡ Hızlı senaryo")
    senaryo = st.radio(
        "Senaryo",
        [
            "Geniş tarama (çok fon)",
            "Dengeli",
            "Muhafazakâr",
            "Agresif büyüme",
            "Enflasyon üstü (reel getiri)",
            "Düşük risk / yüksek likidite",
            "Özel (aşağıdan ayarla)",
        ],
        horizontal=True,
        label_visibility="collapsed",
        index=0,
    )
    SENARYO_MAP = {
        # Geniş tarama: dağılım toleransı çok yüksek + neredeyse hiç kalite kapısı
        # → mümkün olduğunca çok fon bulmak için
        "Geniş tarama (çok fon)": {
            "paket": "0", "profil": "2", "gp": "2",
            "min_1y": None, "min_sharpe": None, "max_vol": None, "max_dd": None,
            "max_risk_skor": 40.0, "min_liq": 10.0, "pozitif_donem": False,
            "min_kazanma": None, "min_en_kotu": None, "max_var95": None, "max_cvar95": None,
            "max_fon": 50, "max_aday": 120, "tolerans": 80,
        },
        "Dengeli": {
            "paket": "2", "profil": "2", "gp": "2",
            "min_1y": 0.0, "min_sharpe": 0.0, "max_vol": 30.0, "max_dd": -25.0,
            "max_risk_skor": 22.0, "min_liq": 30.0, "pozitif_donem": False,
            "min_kazanma": 40.0, "min_en_kotu": -12.0, "max_var95": 6.0, "max_cvar95": 10.0,
            "max_fon": 40, "max_aday": 100, "tolerans": 55,
        },
        "Muhafazakâr": {
            "paket": "1", "profil": "3", "gp": "1",
            "min_1y": None, "min_sharpe": 0.0, "max_vol": 15.0, "max_dd": -10.0,
            "max_risk_skor": 15.0, "min_liq": 45.0, "pozitif_donem": False,
            "min_kazanma": 48.0, "min_en_kotu": -6.0, "max_var95": 3.5, "max_cvar95": 5.5,
            "min_sortino": 0.3,
            "max_fon": 25, "max_aday": 80, "tolerans": 45,
        },
        "Agresif büyüme": {
            "paket": "3", "profil": "1", "gp": "3",
            "min_1y": 5.0, "min_sharpe": None, "max_vol": 55.0, "max_dd": -45.0,
            "max_risk_skor": 30.0, "min_liq": 20.0, "pozitif_donem": False,
            "min_kazanma": 40.0, "min_en_kotu": -18.0, "max_var95": 8.0, "max_cvar95": 14.0,
            "max_fon": 40, "max_aday": 100, "tolerans": 60,
        },
        "Enflasyon üstü (reel getiri)": {
            "paket": "2", "profil": "2", "gp": "2",
            "min_1y": float(st.session_state.cfg.get("tufe_yillik") or 30),
            "min_sharpe": 0.0, "max_vol": 35.0, "max_dd": -30.0,
            "max_risk_skor": 22.0, "min_liq": 30.0, "pozitif_donem": False,
            "min_kazanma": 40.0, "min_en_kotu": -12.0, "max_var95": 6.0, "max_cvar95": 10.0,
            "max_fon": 35, "max_aday": 90, "tolerans": 55,
        },
        "Düşük risk / yüksek likidite": {
            "paket": "1", "profil": "4", "gp": "1",
            "min_1y": None, "min_sharpe": 0.0, "max_vol": 12.0, "max_dd": -8.0,
            "max_risk_skor": 12.0, "min_liq": 55.0, "pozitif_donem": False,
            "min_kazanma": 50.0, "min_en_kotu": -5.0, "max_var95": 3.0, "max_cvar95": 5.0,
            "min_sortino": 0.5,
            "max_fon": 25, "max_aday": 80, "tolerans": 45,
        },
        "Özel (aşağıdan ayarla)": None,
    }
    sc = SENARYO_MAP.get(senaryo) or {}

    with st.expander("① Profil & risk paketi", expanded=(senaryo == "Özel (aşağıdan ayarla)")):
        col_a, col_b = st.columns(2)
        with col_a:
            profil_default = sc.get("profil", "2")
            profil_keys = list(core.PROFILLER.keys())
            profil_sec = st.selectbox(
                "Dağılım profili",
                profil_keys,
                format_func=lambda k: core.PROFILLER[k][0],
                index=profil_keys.index(profil_default) if profil_default in profil_keys else 1,
            )
            if profil_sec == "6":
                st.markdown("**Özel dağılım (%)**")
                hedef = {}
                for kat, ac in core.KATEGORI_ACIKLAMA.items():
                    v = st.slider(ac, 0, 100, 0, key=f"h_{kat}")
                    if v > 0:
                        hedef[kat] = float(v)
                if not hedef:
                    hedef = core.PROFILLER["2"][1]
            else:
                hedef = core.PROFILLER[profil_sec][1]
                st.json(hedef)
        with col_b:
            paket_default = sc.get("paket", "2")
            paket_keys = list(core.RISK_PAKETLER.keys())
            paket_sec = st.selectbox(
                "Risk paketi",
                paket_keys,
                format_func=lambda k: f"{core.RISK_PAKETLER[k]['ad']} — {core.RISK_PAKETLER[k]['aciklama']}",
                index=paket_keys.index(paket_default) if paket_default in paket_keys else 1,
            )
            paket = core.RISK_PAKETLER[paket_sec]

    with st.expander("② Performans & risk filtreleri", expanded=(senaryo == "Özel (aşağıdan ayarla)")):
        c1, c2, c3 = st.columns(3)
        _m1y = sc.get("min_1y", None)
        _msh = sc.get("min_sharpe", paket.get("min_sharpe"))
        _mvol = sc.get("max_vol", paket.get("max_vol"))
        min_1y = c1.number_input("Min 1Y getiri %", value=float(_m1y) if _m1y is not None else None,
                                 step=1.0, format="%.1f", placeholder="yok")
        min_sharpe = c2.number_input("Min Sharpe",
            value=float(_msh) if _msh is not None else None, step=0.1, format="%.2f", placeholder="yok")
        max_vol = c3.number_input("Max volatilite %",
            value=float(_mvol) if _mvol is not None else None, step=1.0, format="%.1f", placeholder="yok")
        c4, c5, c6 = st.columns(3)
        _mdd = sc.get("max_dd", paket.get("max_dd"))
        max_dd = c4.number_input("Max DD % (örn -20)",
            value=float(_mdd) if _mdd is not None else None, step=1.0, format="%.1f", placeholder="yok")
        min_risk = c5.number_input("Min risk (1-7)", value=None, min_value=1, max_value=7, step=1, placeholder="yok")
        max_risk = c6.number_input("Max risk (1-7)",
            value=int(paket["max_risk"]) if paket.get("max_risk") is not None else None,
            min_value=1, max_value=7, step=1, placeholder="yok")
        c7, c8, c9 = st.columns(3)
        kurucu = c7.text_input("Kurucu metni", placeholder="örn AK PORTFÖY")
        faiz_sec = c8.selectbox("Faiz filtresi", ["Hepsi", "Faizli", "Katılım (faizsiz)"], index=0)
        tolerans = c9.slider(
            "Dağılım toleransı", 10, 100, int(sc.get("tolerans", 55)),
            help="Yüksek = neredeyse tüm fonlar geçer (80–100 ≈ serbest). Düşük = hedef profile yakın fonlar.",
        )
        c10, c11 = st.columns(2)
        max_fon = c10.slider("Kaç fon gösterilsin", 5, 100, int(sc.get("max_fon", 40)))
        max_aday = c11.slider("Aday havuzu", 20, 200, int(sc.get("max_aday", 100)))
        tur_sec = st.multiselect(
            "Fon türü (baskın varlık)",
            options=list(core.TUR_ACIKLAMA.keys()),
            format_func=lambda k: core.TUR_ACIKLAMA[k],
            default=[],
        )
        tur_filtre = tur_sec or None

    # Geniş tarama seçildiyse kalite kapılarını varsayılan kapalı tut
    _genis = senaryo.startswith("Geniş")
    with st.expander("③ Kalite kapıları & skor (seçimi iyileştirir)", expanded=not _genis):
        st.markdown(
            "Tarama sonrası **ikinci eleme**: anomali, likidite, dönem tutarlılığı, "
            "kazanma oranı, en kötü gün vb. **Çok fon istiyorsan bu bölümü gevşet veya kapalı tut.**"
        )
        k1, k2, k3, k4 = st.columns(4)
        _mrs = sc.get("max_risk_skor")
        max_risk_skor = k1.number_input(
            "Max risk_skor (anomali)",
            value=float(_mrs) if _mrs is not None else 40.0, step=1.0,
            help="Yüksek tut → daha az eleme (40 ≈ neredeyse kapalı)",
        )
        _ml = sc.get("min_liq")
        min_liq = k2.number_input(
            "Min likidite skor",
            value=float(_ml) if _ml is not None else 10.0, step=5.0,
            help="Düşük tut → daha az eleme (10 ≈ neredeyse kapalı)",
        )
        pozitif_donem = k3.checkbox(
            "1A/3A/6A/1Y hepsi ≥ 0",
            value=bool(sc.get("pozitif_donem", False)),
            help="Açık bırakırsanız daha fazla fon geçer. Piyasa düşüşlü dönemlerde çok eleme yapar.",
        )
        min_reel = k4.checkbox(
            "Reel getiri > 0 (TÜFE üstü)",
            value=(senaryo == "Enflasyon üstü (reel getiri)"),
        )

        st.markdown("**Ek risk / tutarlılık filtreleri** (boş = filtre yok)")
        e1, e2, e3, e4 = st.columns(4)
        _mk = sc.get("min_kazanma")
        min_kazanma = e1.number_input(
            "Min kazanma oranı %",
            value=float(_mk) if _mk is not None else None,
            min_value=0.0, max_value=100.0, step=1.0, placeholder="yok",
            help="Boş bırakırsan uygulanmaz. 50 ≈ yarıdan fazla gün yeşil.",
        )
        _mek = sc.get("min_en_kotu")
        min_en_kotu = e2.number_input(
            "En kötü gün alt sınırı %",
            value=float(_mek) if _mek is not None else None,
            step=0.5, format="%.1f", placeholder="yok",
            help="Boş = filtre yok. Örn. -5 → tek günde %5'ten fazla düşenler elenir.",
        )
        min_sortino = e3.number_input(
            "Min Sortino",
            value=float(sc.get("min_sortino")) if sc.get("min_sortino") is not None else None,
            step=0.1, format="%.2f", placeholder="yok",
            help="Sadece olumsuz oynaklığa göre risk-ayarlı getiri.",
        )
        _mvar = sc.get("max_var95")
        max_var95 = e4.number_input(
            "Max |VaR %95| (örn 3 = -3%)",
            value=float(_mvar) if _mvar is not None else None,
            min_value=0.5, max_value=20.0, step=0.5, placeholder="yok",
            help="Boş = filtre yok. Günlük kayıp üst sınırı.",
        )
        e5, e6, e7, e8 = st.columns(4)
        min_cagr = e5.number_input(
            "Min CAGR %", value=None, step=1.0, format="%.1f", placeholder="yok", key="min_cagr_kalite",
        )
        min_en_iyi = e6.number_input(
            "Min en iyi gün % (opsiyonel)", value=None, step=0.5, format="%.1f", placeholder="yok",
            help="Aşırı düşük tavan getirili fonları elemek için nadiren kullanılır.",
        )
        _mcvar = sc.get("max_cvar95")
        max_cvar95 = e7.number_input(
            "Max |cVaR %95| (örn 5)",
            value=float(_mcvar) if _mcvar is not None else None,
            min_value=0.5, max_value=30.0, step=0.5, placeholder="yok",
            help="Boş = filtre yok. En kötü %5 günlerin ortalaması.",
        )
        zorunlu_kazanma = e8.checkbox(
            "Kazanma oranı filtresini uygula",
            value=False if _genis else (min_kazanma is not None),
            help="Kapalıyken min kazanma oranı uygulanmaz → daha geniş sonuç.",
        )

        st.markdown("**Maliyet · büyüklük · likidite süresi · diğer**")
        x1, x2, x3, x4 = st.columns(4)
        max_toplam_kom = x1.number_input(
            "Max giriş+çıkış kom. %", value=None, step=0.1, format="%.2f", placeholder="yok",
            help="Giriş ve çıkış komisyonu toplamı üst sınırı",
            key="max_kom_kalite",
        )
        min_buyukluk_mn = x2.number_input(
            "Min portföy büyüklüğü (Mn TL)", value=None, step=10.0, format="%.0f", placeholder="yok",
            help="Örn. 100 = en az 100 milyon TL AUM",
            key="min_buyukluk_kalite",
        )
        min_yatirimci = x3.number_input(
            "Min yatırımcı sayısı", value=None, step=10, format="%d", placeholder="yok",
            key="min_yat_kalite",
        )
        max_valor = x4.number_input(
            "Max geri alım valör (gün)", value=None, min_value=0, max_value=30, step=1, placeholder="yok",
            help="T+0 / T+1 tercih için düşük tutun",
            key="max_valor_kalite",
        )
        x5, x6, x7, x8 = st.columns(4)
        min_calmar = x5.number_input(
            "Min Calmar", value=None, step=0.1, format="%.2f", placeholder="yok", key="min_calmar_kalite",
        )
        min_3y = x6.number_input(
            "Min 3Y getiri %", value=None, step=1.0, format="%.1f", placeholder="yok", key="min_3y_kalite",
        )
        min_6a = x7.number_input(
            "Min 6A getiri %", value=None, step=1.0, format="%.1f", placeholder="yok", key="min_6a_kalite",
        )
        min_vs_mev = x8.number_input(
            "Min vs mevduat (puan)", value=None, step=1.0, format="%.1f", placeholder="yok",
            help="1Y getiri − mevduat farkı",
            key="min_vs_mev_kalite",
        )
        x9, x10, x11 = st.columns(3)
        max_mesafe = x9.number_input(
            "Max dağılım sapması", value=None, step=5.0, format="%.0f", placeholder="yok",
            help="Hedef profile uzaklık; düşük = daha uyumlu",
            key="max_mesafe_kalite",
        )
        max_jump = x10.number_input(
            "Max sıçrama sayısı (>3σ)", value=None, step=1, format="%d", placeholder="yok",
            key="max_jump_kalite",
        )
        sadece_katilim = x11.checkbox("Sadece katılım (faizsiz eğilim)", value=False, key="sadece_katilim")

        st.markdown("**Kategori & BIST korelasyonu**")
        y1, y2, y3 = st.columns(3)
        kategori_filtre = y1.text_input(
            "Kategori içerir (örn Hisse, Borçlanma, Karma)",
            value="",
            help="TEFAS kategori adında geçen metin; boş = hepsi",
            key="kategori_filtre_txt",
        ).strip()
        kategori_gruplari = y1.multiselect(
            "Kategori grubu (çoklu)", list(KATEGORI_GRUPLARI.keys()),
            help="Seçilenlerden biri eşleşen fonlar kalır; boş = hepsi",
            key="kategori_grup_sec",
        )
        max_bist_korr = y2.number_input(
            "Max BIST100 korelasyonu",
            value=None, min_value=-1.0, max_value=1.0, step=0.05, format="%.2f",
            placeholder="yok",
            help="Örn. 0.70 → endekse çok bağlı fonları ele (çeşitlendirme)",
            key="max_bist_korr",
        )
        hesapla_bist = y3.checkbox(
            "BIST korelasyonunu hesapla",
            value=False if senaryo.startswith("Geniş") else True,
            help="Tarama sonrası her aday için BIST100 korelasyonu (biraz sürer). Geniş taramada varsayılan kapalı.",
        )

        skor_sistemi = st.radio(
            "Skor sistemi",
            ["gelismis", "klasik"],
            format_func=lambda x: "Gelişmiş (percentile + akran)" if x == "gelismis" else "Klasik",
            horizontal=True,
        )
        skor_agirlik = None
        if skor_sistemi == "gelismis":
            gp_default = sc.get("gp", "2")
            gp_keys = list(core.GELISMIS_PROFILLER.keys())
            gp = st.selectbox(
                "Gelişmiş skor profili",
                gp_keys,
                format_func=lambda k: core.GELISMIS_PROFILLER[k][0],
                index=gp_keys.index(gp_default) if gp_default in gp_keys else 1,
            )
            if gp == "4":
                skor_agirlik = {}
                for b, ac in core.GELISMIS_BILESEN_ACIKLAMA.items():
                    skor_agirlik[b] = st.slider(ac, 0, 100, 15, key=f"ga_{b}")
            else:
                skor_agirlik = core.GELISMIS_PROFILLER[gp][1]

    faiz_f = {"Faizli": "faiz", "Katılım (faizsiz)": "faizsiz"}.get(faiz_sec)

    if st.button("🚀 Taramayı başlat", type="primary", use_container_width=True):
        analiz = get_analiz()
        # Daha geniş havuz çek, kalite kapısını sonra uygula
        # Daha geniş havuz: max_fon'un en az 2 katı, max_aday kadar, üst sınır 250
        cek_fon = min(max(max_fon * 2, max_fon + 20, max_aday), 250)
        with st.spinner("Fonlar taranıyor (API + cache)… Birkaç dakika sürebilir."):
            try:
                df = analiz.tara(
                    hedef,
                    min_1y=min_1y,
                    min_sharpe=min_sharpe,
                    max_vol=max_vol,
                    max_dd=max_dd,
                    max_fon=cek_fon,
                    tolerans=tolerans,
                    max_aday=max_aday,
                    kurucu_filtre=kurucu or None,
                    faiz_filtre=faiz_f,
                    min_risk=int(min_risk) if min_risk else None,
                    max_risk=int(max_risk) if max_risk else None,
                    tur_filtre=tur_filtre,
                    skor_sistemi=skor_sistemi,
                    skor_agirlik=skor_agirlik,
                )
                st.session_state.son_tarama = df
                try:
                    core.save_tarama_snapshot(df, meta={
                        "kaynak": "tarama", "risk_paketi": paket.get("ad"), "senaryo": senaryo,
                    })
                except Exception:
                    pass
                st.session_state.son_tarama_meta = {
                    "hedef": str(hedef), "risk_paketi": paket["ad"],
                    "skor": skor_sistemi, "senaryo": senaryo,
                    "max_risk_skor": max_risk_skor, "min_liq": min_liq,
                    "pozitif_donem": pozitif_donem, "min_reel": min_reel,
                    "max_fon": max_fon,
                    "min_kazanma": min_kazanma if zorunlu_kazanma else None,
                    "min_en_kotu": min_en_kotu,
                    "min_sortino": min_sortino,
                    "max_var95": max_var95,
                    "max_cvar95": max_cvar95,
                    "min_cagr": min_cagr,
                    "min_en_iyi": min_en_iyi,
                    "max_toplam_kom": max_toplam_kom,
                    "min_buyukluk_mn": min_buyukluk_mn,
                    "min_yatirimci": min_yatirimci,
                    "max_valor": max_valor,
                    "min_calmar": min_calmar,
                    "min_3y": min_3y,
                    "min_6a": min_6a,
                    "min_vs_mev": min_vs_mev,
                    "max_mesafe": max_mesafe,
                    "max_jump": max_jump,
                    "sadece_katilim": sadece_katilim,
                    "kategori_filtre": kategori_filtre or None,
                    "kategori_gruplari": kategori_gruplari or None,
                    "max_bist_korr": max_bist_korr,
                    "hesapla_bist": hesapla_bist,
                }
            except Exception as e:
                st.error(f"Tarama hatası: {e}")
                st.session_state.son_tarama = pd.DataFrame()

    df_raw = st.session_state.son_tarama
    meta = st.session_state.get("son_tarama_meta") or {}

    def _kalite_filtre(df_in: pd.DataFrame) -> pd.DataFrame:
        if df_in is None or df_in.empty:
            return pd.DataFrame()
        w = df_in.copy()
        mrs = meta.get("max_risk_skor", max_risk_skor)
        ml = meta.get("min_liq", min_liq)
        if "risk_skor" in w.columns:
            w = w[pd.to_numeric(w["risk_skor"], errors="coerce").fillna(999) <= mrs]
        if "likidite_skor" in w.columns:
            w = w[pd.to_numeric(w["likidite_skor"], errors="coerce").fillna(0) >= ml]
        if meta.get("pozitif_donem", pozitif_donem):
            for col in ("getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y"):
                if col in w.columns:
                    w = w[pd.to_numeric(w[col], errors="coerce").fillna(-1) >= 0]
        if meta.get("min_reel", min_reel) and "reel_1y" in w.columns:
            w = w[pd.to_numeric(w["reel_1y"], errors="coerce").fillna(-999) > 0]
        # Kazanma oranı
        mkaz = meta.get("min_kazanma")
        if mkaz is not None and "kazanma_orani" in w.columns:
            w = w[pd.to_numeric(w["kazanma_orani"], errors="coerce").fillna(0) >= float(mkaz)]
        # En kötü gün (alt sınır: -5 → -8'likleri ele)
        mek = meta.get("min_en_kotu")
        if mek is not None and "en_kotu_gun" in w.columns:
            w = w[pd.to_numeric(w["en_kotu_gun"], errors="coerce").fillna(-999) >= float(mek)]
        # Sortino
        mso = meta.get("min_sortino")
        if mso is not None and "sortino" in w.columns:
            w = w[pd.to_numeric(w["sortino"], errors="coerce").fillna(-999) >= float(mso)]
        # VaR95: değer negatif; |VaR| <= max_var95 → var95 >= -max_var95
        mvar = meta.get("max_var95")
        if mvar is not None and "var95" in w.columns:
            w = w[pd.to_numeric(w["var95"], errors="coerce").fillna(-999) >= -abs(float(mvar))]
        # cVaR95
        mcvar = meta.get("max_cvar95")
        if mcvar is not None and "cvar95" in w.columns:
            w = w[pd.to_numeric(w["cvar95"], errors="coerce").fillna(-999) >= -abs(float(mcvar))]
        # CAGR
        mcagr = meta.get("min_cagr")
        if mcagr is not None and "cagr" in w.columns:
            w = w[pd.to_numeric(w["cagr"], errors="coerce").fillna(-999) >= float(mcagr)]
        # En iyi gün (opsiyonel tavan değil taban)
        mei = meta.get("min_en_iyi")
        if mei is not None and "en_iyi_gun" in w.columns:
            w = w[pd.to_numeric(w["en_iyi_gun"], errors="coerce").fillna(-999) >= float(mei)]
        # Komisyon toplamı
        mkom = meta.get("max_toplam_kom")
        if mkom is not None:
            g = pd.to_numeric(w.get("giris_kom"), errors="coerce").fillna(0) if "giris_kom" in w.columns else 0
            c = pd.to_numeric(w.get("cikis_kom"), errors="coerce").fillna(0) if "cikis_kom" in w.columns else 0
            try:
                w = w[(g + c) <= float(mkom)]
            except Exception:
                pass
        # Büyüklük (Mn TL)
        mbuy = meta.get("min_buyukluk_mn")
        if mbuy is not None and "buyukluk" in w.columns:
            w = w[pd.to_numeric(w["buyukluk"], errors="coerce").fillna(0) >= float(mbuy) * 1_000_000]
        # Yatırımcı
        myat = meta.get("min_yatirimci")
        if myat is not None and "yatirimci" in w.columns:
            w = w[pd.to_numeric(w["yatirimci"], errors="coerce").fillna(0) >= float(myat)]
        # Valör
        mval = meta.get("max_valor")
        if mval is not None and "valor_geri_alis" in w.columns:
            vv = pd.to_numeric(w["valor_geri_alis"], errors="coerce")
            w = w[vv.isna() | (vv <= float(mval))]
        # Calmar
        mcal = meta.get("min_calmar")
        if mcal is not None and "calmar" in w.columns:
            w = w[pd.to_numeric(w["calmar"], errors="coerce").fillna(-999) >= float(mcal)]
        # 3Y / 6A
        m3 = meta.get("min_3y")
        if m3 is not None and "getiri_3y" in w.columns:
            w = w[pd.to_numeric(w["getiri_3y"], errors="coerce").fillna(-999) >= float(m3)]
        m6 = meta.get("min_6a")
        if m6 is not None and "getiri_6a" in w.columns:
            w = w[pd.to_numeric(w["getiri_6a"], errors="coerce").fillna(-999) >= float(m6)]
        # vs mevduat
        mvm = meta.get("min_vs_mev")
        if mvm is not None and "vs_mevduat" in w.columns:
            w = w[pd.to_numeric(w["vs_mevduat"], errors="coerce").fillna(-999) >= float(mvm)]
        # dağılım sapması
        mmes = meta.get("max_mesafe")
        if mmes is not None and "mesafe_skoru" in w.columns:
            w = w[pd.to_numeric(w["mesafe_skoru"], errors="coerce").fillna(999) <= float(mmes)]
        # sıçrama
        mjump = meta.get("max_jump")
        if mjump is not None and "jump_sayisi" in w.columns:
            w = w[pd.to_numeric(w["jump_sayisi"], errors="coerce").fillna(999) <= float(mjump)]
        # katılım
        if meta.get("sadece_katilim"):
            if "faiz" in w.columns:
                fl = w["faiz"].astype(str).str.lower()
                w = w[fl.str.contains("katılım|katilim|faizsiz", na=False)]
            elif "ad" in w.columns:
                adl = w["ad"].astype(str).str.lower()
                w = w[adl.str.contains("katılım|katilim", na=False)]
        # kategori metin
        kfil = meta.get("kategori_filtre")
        if kfil and "kategori" in w.columns:
            nt = core._normalize_tr(kfil).lower()
            w = w[w["kategori"].astype(str).apply(
                lambda x: nt in core._normalize_tr(x).lower()
            )]
        # kategori grubu (çoklu)
        kgr = meta.get("kategori_gruplari")
        if kgr and "kategori" in w.columns:
            w = w[kategori_maske(w["kategori"], kgr, None)]
        # BIST korelasyon üst sınırı (sütun varsa)
        mbk = meta.get("max_bist_korr")
        if mbk is not None and "bist_korelasyon" in w.columns:
            bk = pd.to_numeric(w["bist_korelasyon"], errors="coerce")
            w = w[bk.isna() | (bk <= float(mbk))]
        # Rapor uyumlu skor ile yeniden sırala
        if not w.empty:
            w = w.copy()
            w["secim_skor"] = w.apply(core.rapor_uyumlu_skor, axis=1)
            sort_cols = [c for c in ["secim_skor", "skor", "sharpe", "getiri_1y"] if c in w.columns]
            w = w.sort_values(sort_cols, ascending=[False] * len(sort_cols))
            n = int(meta.get("max_fon", max_fon) or 20)
            w = w.head(n).reset_index(drop=True)
            if "sira" in w.columns:
                w = w.drop(columns=["sira"])
            w.insert(0, "sira", range(1, len(w) + 1))
        return w

    if df_raw is not None and not df_raw.empty:
        # BIST korelasyonu (kalite öncesi geniş havuzda, sonra filtre)
        if meta.get("hesapla_bist", True) and "bist_korelasyon" not in df_raw.columns:
            analiz_tmp = get_analiz()
            with st.spinner("BIST100 korelasyonları hesaplanıyor…"):
                try:
                    df_raw = core.df_bist_korelasyon_ekle(analiz_tmp, df_raw)
                    st.session_state.son_tarama = df_raw
                except Exception as e:
                    st.warning(f"BIST korelasyonu hesaplanamadı: {e}")

        df = _kalite_filtre(df_raw)
        elenen = len(df_raw) - len(df)
        st.success(f"Ham aday: **{len(df_raw)}** · Kalite sonrası: **{len(df)}**" + (f" · Elenen: {elenen}" if elenen else ""))

        if df.empty:
            st.warning("Kalite kapıları tüm adayları eledi — eşikleri gevşetin veya senaryoyu değiştirin.")
        else:
            # Özet metrikler
            m1, m2, m3, m4, m5 = st.columns(5)
            m1.metric("Ort. 1Y %", fmt_num(df["getiri_1y"].mean() if "getiri_1y" in df else None))
            m2.metric("Ort. Sharpe", fmt_num(df["sharpe"].mean() if "sharpe" in df else None, 2))
            m3.metric("Ort. MaxDD %", fmt_num(df["max_drawdown"].mean() if "max_drawdown" in df else None))
            m4.metric("Ort. Likidite", fmt_num(df["likidite_skor"].mean() if "likidite_skor" in df else None, 0))
            m5.metric(
                "Ort. BIST korr.",
                fmt_num(df["bist_korelasyon"].mean() if "bist_korelasyon" in df else None, 2),
            )

            # ── Akıllı öneri (top 5) ──────────────────────────────────────
            st.subheader("🏆 Akıllı öneri (en iyi 5)")
            top5 = df.head(5)
            for _, row in top5.iterrows():
                nedenler = []
                if row.get("sharpe") is not None and float(row["sharpe"]) >= 1:
                    nedenler.append(f"Sharpe {float(row['sharpe']):.2f}")
                if row.get("reel_1y") is not None and float(row["reel_1y"]) > 0:
                    nedenler.append(f"reel +{float(row['reel_1y']):.1f}%")
                if row.get("likidite_seviye") == "yüksek":
                    nedenler.append("yüksek likidite")
                if row.get("max_drawdown") is not None and float(row["max_drawdown"]) >= -10:
                    nedenler.append(f"sınırlı DD {float(row['max_drawdown']):.1f}%")
                if row.get("vs_mevduat") is not None and float(row["vs_mevduat"]) > 0:
                    nedenler.append(f"mevduat üstü +{float(row['vs_mevduat']):.1f}")
                if row.get("risk_skor") is not None and float(row["risk_skor"]) < 10:
                    nedenler.append("düşük anomali")
                if row.get("bist_korelasyon") is not None:
                    try:
                        nedenler.append(f"BIST ρ={float(row['bist_korelasyon']):.2f}")
                    except (TypeError, ValueError):
                        pass
                if row.get("kategori"):
                    nedenler.append(str(row.get("kategori"))[:28])
                neden_txt = " · ".join(nedenler) if nedenler else "genel skor yüksek"
                with st.container():
                    cL, cR = st.columns([4, 1])
                    cL.markdown(
                        f"**{int(row.get('sira', 0))}. `{row.get('kod')}`** — {str(row.get('ad') or '')[:55]}  \n"
                        f"1Y **{fmt_num(row.get('getiri_1y'))}%** · Sharpe **{fmt_num(row.get('sharpe'), 2)}** · "
                        f"MaxDD **{fmt_num(row.get('max_drawdown'))}%** · Likidite **{fmt_num(row.get('likidite_skor'), 0)}** · "
                        f"Seçim skoru **{fmt_num(row.get('secim_skor'), 1)}**  \n"
                        f"<span style='color:#94a3b8;font-size:0.85rem'>Neden: {neden_txt}</span>",
                        unsafe_allow_html=True,
                    )
                    bayrak = str(row.get("risk_bayrak") or "")
                    if bayrak and bayrak not in ("", "nan"):
                        cR.caption(f"⚠ {bayrak[:60]}")

            # ── Akıllı 10 fon portföy ────────────────────────────────────
            st.subheader("🎯 Akıllı portföy (en iyi 10 fon)")
            st.caption(
                "Seçim skoru + tür çeşitliliği (aynı baskın türden en fazla 3) + "
                "opsiyonel düşük BIST korelasyonu → Max Sharpe / Risk Parity ağırlıkları."
            )
            with st.expander("📖 Akıllı seçim kriterleri — detay", expanded=False):
                st.markdown(core.AKILLI_SECIM_ACIKLAMA)
            ap1, ap2, ap3, ap4 = st.columns(4)
            n_akilli = ap1.slider("Fon sayısı", 5, 15, 10, key="n_akilli_port")
            ym_akilli = ap2.selectbox(
                "Ağırlık yöntemi",
                ["max_sharpe", "risk_parity", "score", "equal"],
                format_func=lambda x: {
                    "max_sharpe": "Max Sharpe", "risk_parity": "Risk Parity",
                    "score": "Skor orantılı", "equal": "Eşit",
                }[x],
                key="ym_akilli",
            )
            max_w_ak = ap3.slider("Tek fon max %", 10, 40, 20, key="maxw_akilli") / 100.0
            max_tur = ap4.slider("Aynı türden max fon", 1, 5, 3, key="max_tur_akilli")
            bt1, bt2 = st.columns(2)
            with bt1:
                btn_a10 = st.button("🚀 En iyi 10 fonluk portföyü oluştur", type="primary", key="btn_akilli10")
            with bt2:
                btn_kar_t = st.button("🔀 Karışık (Muhafazakâr+Dengeli+Agresif)", key="btn_karisik_tar")

            if btn_kar_t:
                kc1, kc2, kc3 = st.columns(3)
                n_muh = kc1.slider("Muhafazakâr", 1, 6, 3, key="n_muh_tar")
                n_den = kc2.slider("Dengeli", 1, 6, 4, key="n_den_tar")
                n_agr = kc3.slider("Agresif", 1, 6, 3, key="n_agr_tar")
                analiz = get_analiz()
                with st.spinner("Karışık portföy hesaplanıyor…"):
                    sonuc = core.akilli_karisik_portfoy(
                        analiz, df,
                        n_muhafazakar=n_muh, n_dengeli=n_den, n_agresif=n_agr,
                        yontem=ym_akilli, max_w=max_w_ak,
                    )
                if sonuc.get("error"):
                    st.error(sonuc["error"])
                else:
                    st.session_state.akilli_portfoy = sonuc

            if btn_a10:
                analiz = get_analiz()
                with st.spinner("Akıllı portföy hesaplanıyor…"):
                    sonuc = core.akilli_portfoy_sec(
                        analiz, df,
                        n_fon=n_akilli,
                        yontem=ym_akilli,
                        max_w=max_w_ak,
                        max_ayni_tur=max_tur,
                        max_bist_korr=meta.get("max_bist_korr"),
                        diversify=True,
                    )
                if sonuc.get("error"):
                    st.error(sonuc["error"])
                else:
                    st.session_state.akilli_portfoy = sonuc
            sonuc = st.session_state.get("akilli_portfoy")
            if sonuc and not sonuc.get("error") and sonuc.get("secili") is not None:
                secili = sonuc["secili"]
                st.success(
                    f"**{sonuc.get('n')} fon** seçildi · yöntem: `{sonuc.get('method')}` · "
                    f"tür dağılımı: {sonuc.get('tur_dagilim')}"
                )
                if sonuc.get("port_sharpe") is not None:
                    p1, p2, p3 = st.columns(3)
                    p1.metric("Beklenen μ", f"%{sonuc.get('port_mu')}")
                    p2.metric("Vol", f"%{sonuc.get('port_vol')}")
                    p3.metric("Sharpe", f"{sonuc.get('port_sharpe')}")
                if sonuc.get("ozet"):
                    with st.expander("Neden bu fonlar? (seçim gerekçeleri)", expanded=True):
                        for line in sonuc["ozet"]:
                            st.markdown(f"- {line}")
                wmap = sonuc.get("weights") or {}
                show_cols = [
                    c for c in [
                        "kod", "ad", "kategori", "baskin_tur", "akilli_skor",
                        "temel_skor", "bonus_kazanma", "bonus_en_kotu_gun",
                        "bonus_likidite", "bonus_bist", "ceza_anomali",
                        "getiri_1y", "sharpe", "sortino", "max_drawdown",
                        "bist_korelasyon", "likidite_skor", "risk_skor", "secim_gerekce",
                    ] if c in secili.columns
                ]
                view_p = secili[show_cols].copy() if show_cols else secili.copy()
                view_p["agirlik_%"] = view_p["kod"].astype(str).map(
                    lambda k: round(float(wmap.get(k, 0)) * 100, 1)
                )
                show_df(view_p, list(view_p.columns), height=400)
                with st.expander("Skor kırılımı (her fon)"):
                    for _, r in secili.iterrows():
                        st.caption(
                            f"**{r.get('kod')}** akilli={r.get('akilli_skor')} = "
                            f"temel {r.get('temel_skor')} + kazanma {r.get('bonus_kazanma')} + "
                            f"kötü gün {r.get('bonus_en_kotu_gun')} + likidite {r.get('bonus_likidite')} + "
                            f"BIST {r.get('bonus_bist')} + anomali {r.get('ceza_anomali')}"
                        )
                edf = secili.copy()
                edf["portfoy_agirlik_%"] = edf["kod"].astype(str).map(
                    lambda k: round(float(wmap.get(k, 0)) * 100, 2)
                )
                download_excel_button(
                    edf,
                    f"tefas_akilli_portfoy10_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={
                        "mod": "akilli_portfoy",
                        "n": sonuc.get("n"),
                        "yontem": sonuc.get("method"),
                        "agirliklar": wmap,
                        "tur_dagilim": sonuc.get("tur_dagilim"),
                    },
                    key="xl_akilli10",
                )
                fa, fb = st.columns(2)
                with fa:
                    if st.button("⭐ Favorilere ekle", key="fav_akilli10"):
                        kodlar = secili["kod"].astype(str).tolist()
                        favs = list(set(core.load_favs()) | set(kodlar))
                        core.save_favs(favs)
                        st.success(f"{len(kodlar)} favori")
                with fb:
                    if st.button("🛒 Portföy sepetine ekle", key="sepet_akilli10"):
                        ad_map = {str(r["kod"]): str(r.get("ad") or r["kod"]) for _, r in secili.iterrows()}
                        n = sepet_ekle(secili["kod"].astype(str).tolist(), ad_map=ad_map, kaynak="tarama_akilli")
                        st.success(f"{n} fon sepete (toplam {len(st.session_state.portfoy_sepet)})")

            # ── Canlı filtre (sonuç üzerinde) ─────────────────────────────
            st.subheader("📋 Sonuç listesi")
            with st.expander("🔎 Liste üzerinde ek filtre", expanded=True):
                f1, f2, f3, f4 = st.columns(4)
                live_min_1y = f1.number_input("Min 1Y %", value=None, step=1.0, format="%.1f",
                                              placeholder="yok", key="live_min_1y")
                live_min_1a = f2.number_input("Min 1A %", value=None, step=0.5, format="%.1f",
                                              placeholder="yok", key="live_min_1a",
                                              help="Son 1 ay — kısa vade kazanan")
                live_min_3a = f3.number_input("Min 3A %", value=None, step=1.0, format="%.1f",
                                              placeholder="yok", key="live_min_3a")
                live_min_sh = f4.number_input("Min Sharpe", value=None, step=0.1, format="%.2f",
                                              placeholder="yok", key="live_min_sh")
                f5, f6, f7, f8 = st.columns(4)
                live_max_dd = f5.number_input("Max DD (örn -15)", value=None, step=1.0, format="%.1f",
                                              placeholder="yok", key="live_max_dd")
                live_min_liq = f6.number_input("Min likidite", value=None, step=5.0, format="%.0f",
                                               placeholder="yok", key="live_min_liq")
                live_max_vol = f7.number_input("Max vol %", value=None, step=1.0, format="%.1f",
                                               placeholder="yok", key="live_max_vol")
                live_max_rs = f8.number_input("Max risk_skor", value=None, step=1.0, format="%.1f",
                                              placeholder="yok", key="live_max_rs")
                f9, f10, f11, f12 = st.columns(4)
                live_min_kaz = f9.number_input("Min kazanma %", value=None, step=1.0, format="%.0f",
                                               placeholder="yok", key="live_min_kaz",
                                               help="Günlük pozitif gün oranı")
                live_min_ekotu = f10.number_input("En kötü gün ≥ %", value=None, step=0.5, format="%.1f",
                                                  placeholder="yok", key="live_min_ekotu",
                                                  help="Örn. -5 → tek günde %5'ten fazla düşenleri ele")
                live_min_so = f11.number_input("Min Sortino", value=None, step=0.1, format="%.2f",
                                               placeholder="yok", key="live_min_so")
                live_max_var = f12.number_input("Max |VaR95|", value=None, step=0.5, format="%.1f",
                                                placeholder="yok", key="live_max_var",
                                                help="3 yazarsanız VaR ≤ -3% olanlar elenir")
                f13a, f13b = st.columns(2)
                live_max_cvar = f13a.number_input("Max |cVaR95|", value=None, step=0.5, format="%.1f",
                                                  placeholder="yok", key="live_max_cvar")
                live_min_cagr = f13b.number_input("Min CAGR %", value=None, step=1.0, format="%.1f",
                                                  placeholder="yok", key="live_min_cagr")
                f13, f14, f15, f16 = st.columns(4)
                live_min_calmar = f13.number_input("Min Calmar", value=None, step=0.1, format="%.2f",
                                                   placeholder="yok", key="live_min_calmar")
                live_min_3y = f14.number_input("Min 3Y %", value=None, step=1.0, format="%.1f",
                                               placeholder="yok", key="live_min_3y")
                live_min_buy = f15.number_input("Min büyüklük Mn", value=None, step=10.0, format="%.0f",
                                                placeholder="yok", key="live_min_buy")
                live_max_kom = f16.number_input("Max kom. toplam %", value=None, step=0.1, format="%.2f",
                                                placeholder="yok", key="live_max_kom")
                f17, f18, f19 = st.columns(3)
                live_text = f17.text_input("Ad/kod/kurucu içerir", "", key="live_text")
                live_tur = f18.multiselect(
                    "Baskın tür",
                    options=sorted(df["baskin_tur"].dropna().unique().tolist()) if "baskin_tur" in df.columns else [],
                    key="live_tur",
                )
                live_sadece_temiz = f19.checkbox(
                    "Risk bayrağı olanları gizle",
                    value=False, key="live_temiz",
                )
                f20, f21 = st.columns(2)
                live_kat = f20.text_input("Kategori içerir", "", key="live_kat")
                live_kgr, live_ktam = kategori_secici(df, "live_kat_sec")
                live_max_bist = f21.number_input(
                    "Max BIST ρ", value=None, min_value=-1.0, max_value=1.0,
                    step=0.05, format="%.2f", placeholder="yok", key="live_max_bist",
                )

            df_goster = df.copy()
            if live_min_1y is not None and "getiri_1y" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["getiri_1y"], errors="coerce").fillna(-999) >= live_min_1y]
            if live_min_1a is not None and "getiri_1a" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["getiri_1a"], errors="coerce").fillna(-999) >= live_min_1a]
            if live_min_3a is not None and "getiri_3a" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["getiri_3a"], errors="coerce").fillna(-999) >= live_min_3a]
            if live_min_sh is not None and "sharpe" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["sharpe"], errors="coerce").fillna(-999) >= live_min_sh]
            if live_max_dd is not None and "max_drawdown" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["max_drawdown"], errors="coerce").fillna(-999) >= live_max_dd]
            if live_min_liq is not None and "likidite_skor" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["likidite_skor"], errors="coerce").fillna(0) >= live_min_liq]
            if live_max_vol is not None and "volatilite" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["volatilite"], errors="coerce").fillna(999) <= live_max_vol]
            if live_max_rs is not None and "risk_skor" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["risk_skor"], errors="coerce").fillna(999) <= live_max_rs]
            if live_min_kaz is not None and "kazanma_orani" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["kazanma_orani"], errors="coerce").fillna(0) >= live_min_kaz]
            if live_min_ekotu is not None and "en_kotu_gun" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["en_kotu_gun"], errors="coerce").fillna(-999) >= live_min_ekotu]
            if live_min_so is not None and "sortino" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["sortino"], errors="coerce").fillna(-999) >= live_min_so]
            if live_max_var is not None and "var95" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["var95"], errors="coerce").fillna(-999) >= -abs(live_max_var)]
            if live_max_cvar is not None and "cvar95" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["cvar95"], errors="coerce").fillna(-999) >= -abs(live_max_cvar)]
            if live_min_cagr is not None and "cagr" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["cagr"], errors="coerce").fillna(-999) >= live_min_cagr]
            if live_min_calmar is not None and "calmar" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["calmar"], errors="coerce").fillna(-999) >= live_min_calmar]
            if live_min_3y is not None and "getiri_3y" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["getiri_3y"], errors="coerce").fillna(-999) >= live_min_3y]
            if live_min_buy is not None and "buyukluk" in df_goster.columns:
                df_goster = df_goster[pd.to_numeric(df_goster["buyukluk"], errors="coerce").fillna(0) >= live_min_buy * 1_000_000]
            if live_max_kom is not None:
                g = pd.to_numeric(df_goster.get("giris_kom"), errors="coerce").fillna(0) if "giris_kom" in df_goster.columns else 0
                c = pd.to_numeric(df_goster.get("cikis_kom"), errors="coerce").fillna(0) if "cikis_kom" in df_goster.columns else 0
                try:
                    df_goster = df_goster[(g + c) <= live_max_kom]
                except Exception:
                    pass
            if live_text:
                nt = core._normalize_tr(live_text).lower()
                mask = (
                    df_goster.get("kod", pd.Series(dtype=str)).astype(str).str.upper().str.contains(live_text.upper(), na=False)
                    | df_goster.get("ad", pd.Series(dtype=str)).astype(str).apply(lambda x: nt in core._normalize_tr(x).lower())
                    | df_goster.get("kurucu", pd.Series(dtype=str)).astype(str).apply(lambda x: nt in core._normalize_tr(x).lower())
                )
                df_goster = df_goster[mask]
            if live_tur and "baskin_tur" in df_goster.columns:
                df_goster = df_goster[df_goster["baskin_tur"].isin(live_tur)]
            if live_sadece_temiz and "risk_bayrak" in df_goster.columns:
                rb = df_goster["risk_bayrak"].astype(str).str.lower()
                df_goster = df_goster[~rb.str.contains("aşırı|akran-dışı|sıçrama|yüksek fiyat", na=False)]
            if (live_kgr or live_ktam) and "kategori" in df_goster.columns:
                df_goster = df_goster[kategori_maske(df_goster["kategori"], live_kgr, live_ktam)]
            if live_kat and "kategori" in df_goster.columns:
                nt = core._normalize_tr(live_kat).lower()
                df_goster = df_goster[df_goster["kategori"].astype(str).apply(
                    lambda x: nt in core._normalize_tr(x).lower()
                )]
            if live_max_bist is not None and "bist_korelasyon" in df_goster.columns:
                bk = pd.to_numeric(df_goster["bist_korelasyon"], errors="coerce")
                df_goster = df_goster[bk.isna() | (bk <= live_max_bist)]

            df_goster = TE.panel(st, core, get_analiz, df_goster, "live")
            st.caption(f"Gösterilen: **{len(df_goster)}** / {len(df)} fon")
            cols = [
                "sira", "secim_skor", "skor", "kod", "ad",
                "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y", "getiri_3y",
                "reel_1y", "vs_mevduat", "vs_bist100",
                "sharpe", "sortino", "calmar", "cagr", "kazanma_orani",
                "en_kotu_gun", "en_iyi_gun", "var95", "cvar95",
                "volatilite", "max_drawdown",
                "risk_degeri", "risk_skor", "likidite_skor", "likidite_seviye",
                "buyukluk", "yatirimci", "giris_kom", "cikis_kom", "valor_geri_alis",
                "mesafe_skoru", "jump_sayisi", "bist_korelasyon", "kategori",
                "baskin_tur", "kurucu", "risk_bayrak",
            ]
            show_df(df_goster, cols + [c for c in st.session_state.get("live_ek_cols", []) if c not in cols], height=420)

            if skor_sistemi == "gelismis" or "pct_getiri" in df_goster.columns:
                with st.expander("Bileşen percentile'ları (akran grubu içinde 0–100)"):
                    pct_cols = [c for c in [
                        "kod", "pct_getiri", "pct_risk_ayarli", "pct_maxdd",
                        "pct_risk_anomali", "pct_maliyet", "pct_akran_sapma", "akran_grubu",
                    ] if c in df_goster.columns]
                    show_df(df_goster, pct_cols, height=280)

            # ── Seçimden portföy / favori ──────────────────────────────────
            st.subheader("✅ Seçim yap")
            kod_list = df_goster["kod"].astype(str).tolist() if "kod" in df_goster.columns else []
            labels = {
                str(r["kod"]): f"{r['kod']} — 1Y {fmt_num(r.get('getiri_1y'))}% · Sh {fmt_num(r.get('sharpe'), 2)}"
                for _, r in df_goster.iterrows()
            }
            secilen = st.multiselect(
                "Portföye / favoriye eklenecek fonlar",
                options=kod_list,
                default=kod_list[: min(5, len(kod_list))],
                format_func=lambda k: labels.get(k, k),
            )
            # Portföy için secili_df kaynağı
            df = df_goster
            a1, a2, a3 = st.columns(3)
            with a1:
                cfa, cfb = st.columns(2)
                with cfa:
                    if st.button("⭐ Favori") and secilen:
                        favs = list(set(core.load_favs()) | set(secilen))
                        core.save_favs(favs)
                        st.success(f"{len(secilen)} favori")
                with cfb:
                    if st.button("🛒 Portföye ekle", key="sepet_tarama") and secilen:
                        ad_map = {str(r["kod"]): str(r.get("ad") or r["kod"]) for _, r in df_goster.iterrows()}
                        n = sepet_ekle(secilen, ad_map=ad_map, kaynak="tarama")
                        st.success(f"{n} fon sepete eklendi (toplam {len(st.session_state.portfoy_sepet)})")
            with a2:
                yontem_p = st.selectbox(
                    "Portföy ağırlığı",
                    ["max_sharpe", "score", "equal", "risk_parity"],
                    format_func=lambda x: {
                        "max_sharpe": "Max Sharpe", "score": "Seçim skoru orantılı",
                        "equal": "Eşit", "risk_parity": "Risk Parity",
                    }[x],
                    key="tar_port_ym",
                )
            with a3:
                max_w_p = st.slider("Tek fon max %", 10, 100, 35, key="tar_maxw") / 100.0

            if st.button("📁 Seçilenlerden portföy kur", type="primary") and secilen:
                secili_df = df[df["kod"].astype(str).isin(secilen)].copy()
                if "rapor_skor" not in secili_df.columns and "secim_skor" in secili_df.columns:
                    secili_df["rapor_skor"] = secili_df["secim_skor"]
                analiz = get_analiz()
                with st.spinner("Ağırlıklar hesaplanıyor…"):
                    out = core.siralamadan_agirlik_hesapla(
                        analiz, secili_df, yontem=yontem_p, max_w=max_w_p,
                    )
                if out.get("error"):
                    st.error(out["error"])
                else:
                    st.markdown(f"**Ağırlıklar ({out.get('method')})**")
                    wdf = pd.DataFrame([
                        {"kod": k, "ağırlık %": round(v * 100, 1),
                         "risk katkısı %": out.get("risk_contrib_pct", {}).get(k, 0)}
                        for k, v in sorted(out.get("weights", {}).items(), key=lambda x: -x[1])
                        if v >= 0.001
                    ])
                    st.dataframe(wdf, hide_index=True, use_container_width=True)
                    if out.get("port_sharpe") is not None:
                        p1, p2, p3 = st.columns(3)
                        p1.metric("Beklenen μ", f"%{out.get('port_mu')}")
                        p2.metric("Vol", f"%{out.get('port_vol')}")
                        p3.metric("Sharpe", f"{out.get('port_sharpe')}")
                    # Excel
                    edf = secili_df.copy()
                    edf["portfoy_agirlik_%"] = edf["kod"].map(
                        lambda k: round(out.get("weights", {}).get(str(k), 0) * 100, 2)
                    )
                    download_excel_button(
                        edf,
                        f"tefas_tarama_portfoy_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        meta={"mod": "tarama_portfoy", "yontem": yontem_p, "agirliklar": out.get("weights")},
                        key="xl_tar_port",
                    )
                    st.session_state.son_portfoy = out

            st.divider()
            cdl1, cdl2 = st.columns(2)
            with cdl1:
                download_excel_button(
                    df,
                    f"tefas_tarama_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta=meta,
                    key="xl_tarama",
                )
            with cdl2:
                if core.HAS_MPL and st.button("📊 PNG grafik rapor", key="gr_tarama"):
                    path = core.grafik_rapor(df, baslik=f"Tarama — {meta.get('risk_paketi', '')} | {senaryo}")
                    if path:
                        st.image(path)
                        with open(path, "rb") as f:
                            st.download_button(
                                "PNG indir", f, file_name=Path(path).name,
                                mime="image/png", key="dl_png_t",
                            )
    elif df_raw is not None and df_raw.empty:
        st.warning("Sonuç yok — filtreleri gevşetin veya senaryo değiştirin.")


# ═══════════════════════════════════════════════════════════════════════════
# TÜM FON SIRALAMASI
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "📋 Tüm Fon Sıralaması":
    st.header("📋 Tüm Fonlar — Dönem Sıralaması")
    st.caption("Dağılım filtresi yok. 1A · 3A · 6A · 1Y · 3Y getiri + risk metrikleri")

    with st.expander("Parametreler", expanded=True):
        c1, c2, c3 = st.columns(3)
        poz = c1.checkbox("Son 1Y her dönem pozitif (1A,3A,6A,1Y)", value=True)
        poz3 = c2.checkbox("3Y de pozitif olsun", value=False)
        max_liste = c3.slider("Kaç fon listelensin", 10, 200, 40)
        lim = st.number_input(
            "En fazla kaç açık fon taransın? (0 = hepsi)",
            value=80, min_value=0, max_value=5000, step=10,
            help="İlk denemede 80–150 önerilir; tüm fonlar 10–40 dk sürebilir.",
        )
        max_tara = int(lim) if lim else None

        tur_sec = st.multiselect(
            "Fon türü filtresi",
            options=list(core.TUR_ACIKLAMA.keys()),
            format_func=lambda k: core.TUR_ACIKLAMA[k],
        )
        tur_filtre = tur_sec or None

        skor_sistemi = st.radio(
            "Skor sistemi", ["gelismis", "klasik"],
            format_func=lambda x: "Gelişmiş" if x == "gelismis" else "Klasik",
            horizontal=True, key="sira_skor",
        )
        skor_agirlik = core.GELISMIS_PROFILLER["2"][1] if skor_sistemi == "gelismis" else None
        if skor_sistemi == "gelismis":
            gp = st.selectbox(
                "Gelişmiş profil",
                list(core.GELISMIS_PROFILLER.keys()),
                format_func=lambda k: core.GELISMIS_PROFILLER[k][0],
                index=1, key="sira_gp",
            )
            if gp != "4":
                skor_agirlik = core.GELISMIS_PROFILLER[gp][1]

    st.info("⏳ Tüm fon taraması uzun sürebilir. Cache sonraki koşuları hızlandırır.")

    if st.button("🚀 Sıralamayı başlat", type="primary", use_container_width=True):
        analiz = get_analiz()
        with st.spinner("Açık fonlar taranıyor…"):
            try:
                df = analiz.tum_fon_siralama(
                    sadece_pozitif_1y=poz,
                    pozitif_3y=poz3,
                    max_liste=max_liste,
                    max_tara=max_tara,
                    tur_filtre=tur_filtre,
                    skor_sistemi=skor_sistemi,
                    skor_agirlik=skor_agirlik,
                )
                st.session_state.son_siralama = df
                try:
                    core.save_tarama_snapshot(df, meta={"kaynak": "siralama"})
                except Exception:
                    pass
            except Exception as e:
                st.error(f"Hata: {e}")
                df = pd.DataFrame()

    df = st.session_state.son_siralama
    if df is not None and not df.empty:
        # Akıllı skor + tutarlılık skoru ekle
        work = df.copy()
        if "secim_skor" not in work.columns:
            work["secim_skor"] = work.apply(core.rapor_uyumlu_skor, axis=1)
        if "tutarlilik_skor" not in work.columns:
            work = core.df_tutarlilik_skoru_ekle(work)

        st.success(f"{len(work)} fon listelendi")

        # ── Akıllı seçim kriterleri ──────────────────────────────────────
        st.subheader("🎯 Akıllı seçim kriterleri")
        with st.expander("📖 Kriterler nasıl çalışır? (detaylı açıklama)", expanded=False):
            st.markdown(core.AKILLI_SECIM_ACIKLAMA)

        preset_keys = list(core.AKILLI_FILTRE_PRESET.keys())
        preset_sec = st.radio(
            "Akıllı filtre profili",
            preset_keys,
            index=preset_keys.index("Dengeli") if "Dengeli" in preset_keys else 0,
            horizontal=True,
            format_func=lambda k: core.AKILLI_FILTRE_PRESET[k]["ad"].split("—")[0].strip(),
            key="sira_akilli_preset",
            help="Muhafazakâr / Dengeli / Agresif… tüm eşikleri otomatik doldurur",
        )
        pr = core.AKILLI_FILTRE_PRESET[preset_sec]
        st.caption(pr.get("ad", preset_sec))
        pk = preset_sec.replace(" ", "_").replace("/", "_")

        with st.expander("Kalite & risk filtreleri (profil eşikleri — ince ayar yapabilirsiniz)", expanded=(preset_sec == "Özel")):
            s1, s2, s3, s4 = st.columns(4)
            s_max_rs = s1.number_input("Max risk_skor", value=float(pr.get("max_risk_skor") or 25), step=1.0,
                                       key=f"sira_max_rs_{pk}", help="Anomali / ani sıçrama skoru üst sınırı")
            s_min_liq = s2.number_input("Min likidite", value=float(pr.get("min_liq") or 20), step=5.0,
                                        key=f"sira_min_liq_{pk}")
            _mkaz = pr.get("min_kazanma")
            s_min_kaz = s3.number_input(
                "Min kazanma %",
                value=float(_mkaz) if _mkaz is not None else None,
                step=1.0, placeholder="yok", key=f"sira_min_kaz_{pk}",
            )
            _mek = pr.get("min_en_kotu")
            s_min_ekotu = s4.number_input(
                "En kötü gün ≥ %",
                value=float(_mek) if _mek is not None else None,
                step=0.5, format="%.1f", placeholder="yok", key=f"sira_min_ekotu_{pk}",
            )
            s5, s6, s7, s8 = st.columns(4)
            _msh = pr.get("min_sharpe")
            s_min_sh = s5.number_input("Min Sharpe", value=float(_msh) if _msh is not None else None,
                                       step=0.1, format="%.2f", placeholder="yok", key=f"sira_min_sh_{pk}")
            _mso = pr.get("min_sortino")
            s_min_so = s6.number_input("Min Sortino", value=float(_mso) if _mso is not None else None,
                                       step=0.1, format="%.2f", placeholder="yok", key=f"sira_min_so_{pk}")
            s_max_dd = s7.number_input("Max DD %", value=float(pr.get("max_dd") if pr.get("max_dd") is not None else -40),
                                       step=1.0, key=f"sira_max_dd_{pk}")
            s_max_vol = s8.number_input("Max vol %", value=float(pr.get("max_vol") if pr.get("max_vol") is not None else 50),
                                        step=1.0, key=f"sira_max_vol_{pk}")
            s9, s10, s11, s12 = st.columns(4)
            _mvar = pr.get("max_var95")
            s_max_var = s9.number_input(
                "Max |VaR95|",
                value=float(_mvar) if _mvar is not None else None,
                step=0.5, placeholder="yok", key=f"sira_max_var_{pk}",
            )
            _mcvar = pr.get("max_cvar95")
            s_max_cvar = s10.number_input(
                "Max |cVaR95|",
                value=float(_mcvar) if _mcvar is not None else None,
                step=0.5, placeholder="yok", key=f"sira_max_cvar_{pk}",
            )
            _mcagr = pr.get("min_cagr")
            s_min_cagr = s11.number_input("Min CAGR %", value=float(_mcagr) if _mcagr is not None else None,
                                          step=1.0, format="%.1f", placeholder="yok", key=f"sira_min_cagr_{pk}")
            _mcal = pr.get("min_calmar")
            s_min_calmar = s12.number_input("Min Calmar", value=float(_mcal) if _mcal is not None else None,
                                            step=0.1, format="%.2f", placeholder="yok", key=f"sira_min_cal_{pk}")
            s13, s14, s15, s16 = st.columns(4)
            s_kategori = s13.text_input("Kategori içerir", "", key=f"sira_kat_{pk}")
            _mbist = pr.get("max_bist_korr")
            s_max_bist = s14.number_input(
                "Max BIST ρ",
                value=float(_mbist) if _mbist is not None else None,
                min_value=-1.0, max_value=1.0, step=0.05, format="%.2f",
                placeholder="yok", key=f"sira_max_bist_{pk}",
            )
            _rejim_preset = preset_sec in (
                "BIST düşünce artan", "Kriz dayanıklı", "Asimetrik (yükselir-korur)",
            ) or (pr.get("sirala") or "") in (
                "dususte_artan_skor", "kriz_dayanik_skor", "asimetri_skor", "bist_korelasyon_asc",
            )
            s_hesapla_bist = s15.checkbox(
                "BIST / rejim metrikleri hesapla",
                value=(_mbist is not None) or _rejim_preset or ("BIST" in preset_sec),
                key=f"sira_bist_hesap_{pk}",
                help="Korelasyon + upside/downside capture + kriz pencereleri. Biraz sürer.",
            )
            s_temiz = s16.checkbox(
                "Risk bayrağı olanları gizle",
                value=bool(pr.get("temiz_bayrak", True)),
                key=f"sira_temiz_{pk}",
            )
            s_min_reel = st.checkbox(
                "Reel getiri > 0 (TÜFE üstü)",
                value=bool(pr.get("min_reel", False)),
                key=f"sira_min_reel_{pk}",
            )
            st.markdown("**Kategoriye göre seçim**")
            kc1, kc2 = st.columns(2)
            s_kgr = kc1.multiselect(
                "Kategori grubu", list(KATEGORI_GRUPLARI.keys()), key=f"sira_kgr_{pk}",
                help="Birden fazla seçilebilir; boş = hepsi",
            )
            _vc = (work["kategori"].dropna().astype(str).str.strip().value_counts()
                   if "kategori" in work.columns else pd.Series(dtype=int))
            s_ktam = kc2.multiselect(
                "TEFAS kategorisi (tam ad)", _vc.index.tolist(), key=f"sira_ktam_{pk}",
                format_func=lambda k: f"{k} ({_vc[k]})",
            )
            st.markdown("**Kısa vade & rejim skor eşikleri**")
            kv1, kv2, kv3 = st.columns(3)
            _m1a = pr.get("min_1a")
            _m3a = pr.get("min_3a")
            s_min_1a = kv1.number_input(
                "Min 1A getiri %",
                value=float(_m1a) if _m1a is not None else None,
                step=0.5, format="%.1f", placeholder="yok",
                key=f"sira_min_1a_{pk}",
                help="Son 1 aylık getiri tabanı",
            )
            s_min_3a = kv2.number_input(
                "Min 3A getiri %",
                value=float(_m3a) if _m3a is not None else None,
                step=1.0, format="%.1f", placeholder="yok",
                key=f"sira_min_3a_{pk}",
                help="Son 3 aylık getiri tabanı",
            )
            _sirala_opts = [
                "akilli_skor", "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y",
                "sharpe", "bist_korelasyon", "bist_korelasyon_asc",
                "dususte_artan_skor", "kriz_dayanik_skor", "asimetri_skor",
                "upside_capture", "downside_capture", "capture_ratio", "kriz_getiri",
            ]
            _sirala_def = pr.get("sirala", "akilli_skor")
            if _sirala_def not in _sirala_opts:
                _sirala_opts = [_sirala_def] + _sirala_opts
            s_sirala = kv3.selectbox(
                "Sıralama ölçütü",
                _sirala_opts,
                index=_sirala_opts.index(_sirala_def) if _sirala_def in _sirala_opts else 0,
                format_func=lambda x: {
                    "akilli_skor": "Akıllı skor",
                    "getiri_1a": "1A getiri (kısa vade)",
                    "getiri_3a": "3A getiri",
                    "getiri_6a": "6A getiri",
                    "getiri_1y": "1Y getiri",
                    "sharpe": "Sharpe",
                    "bist_korelasyon": "BIST ρ (yüksek → düşük)",
                    "bist_korelasyon_asc": "BIST düşünce artan (düşük ρ önce)",
                    "dususte_artan_skor": "📉 Düşünce artan skor (0–100)",
                    "kriz_dayanik_skor": "🛡️ Kriz dayanıklılık skoru (0–100)",
                    "asimetri_skor": "⚖️ Asimetri skor (yükselir-korur)",
                    "upside_capture": "Upside capture %",
                    "downside_capture": "Downside capture % (düşük iyi)",
                    "capture_ratio": "Capture ratio (up/down)",
                    "kriz_getiri": "Kriz dönemi getiri %",
                }.get(x, x),
                key=f"sira_sirala_{pk}",
            )
            # Rejim skor min eşikleri (preset'ten)
            rr1, rr2, rr3 = st.columns(3)
            _mda = pr.get("min_dususte_artan")
            s_min_dususte = rr1.number_input(
                "Min düşünce-artan skor",
                value=float(_mda) if _mda is not None else None,
                min_value=0.0, max_value=100.0, step=5.0, placeholder="yok",
                key=f"sira_min_dususte_{pk}",
            )
            _mkd = pr.get("min_kriz_dayanik")
            s_min_kriz = rr2.number_input(
                "Min kriz-dayanıklılık skor",
                value=float(_mkd) if _mkd is not None else None,
                min_value=0.0, max_value=100.0, step=5.0, placeholder="yok",
                key=f"sira_min_kriz_{pk}",
            )
            _mas = pr.get("min_asimetri")
            s_min_asim = rr3.number_input(
                "Min asimetri skor",
                value=float(_mas) if _mas is not None else None,
                min_value=0.0, max_value=100.0, step=5.0, placeholder="yok",
                key=f"sira_min_asim_{pk}",
            )

        # BIST korelasyonu + rejim metrikleri (isteğe bağlı)
        _need_rejim = s_hesapla_bist or _rejim_preset or s_sirala in (
            "dususte_artan_skor", "kriz_dayanik_skor", "asimetri_skor",
            "upside_capture", "downside_capture", "capture_ratio", "kriz_getiri",
        )
        if _need_rejim and "dususte_artan_skor" not in work.columns:
            analiz_tmp = get_analiz()
            with st.spinner("BIST rejim metrikleri hesaplanıyor (korelasyon + kriz + capture)… Birkaç dk sürebilir."):
                try:
                    work = core.df_rejim_metrikleri_ekle(analiz_tmp, work, periyod=36)
                    st.session_state.son_siralama = work
                except Exception as e:
                    st.warning(f"Rejim metrikleri: {e}")
                    # fallback sadece korelasyon
                    if "bist_korelasyon" not in work.columns:
                        try:
                            work = core.df_bist_korelasyon_ekle(analiz_tmp, work)
                            st.session_state.son_siralama = work
                        except Exception as e2:
                            st.warning(f"BIST korelasyonu: {e2}")
        elif s_hesapla_bist and "bist_korelasyon" not in work.columns:
            analiz_tmp = get_analiz()
            with st.spinner("BIST100 korelasyonları hesaplanıyor…"):
                try:
                    work = core.df_bist_korelasyon_ekle(analiz_tmp, work)
                    st.session_state.son_siralama = work
                except Exception as e:
                    st.warning(f"BIST korelasyonu: {e}")

        filt = work.copy()
        if "risk_skor" in filt.columns and s_max_rs is not None:
            filt = filt[pd.to_numeric(filt["risk_skor"], errors="coerce").fillna(999) <= s_max_rs]
        if "likidite_skor" in filt.columns and s_min_liq is not None:
            filt = filt[pd.to_numeric(filt["likidite_skor"], errors="coerce").fillna(0) >= s_min_liq]
        if s_min_kaz is not None and "kazanma_orani" in filt.columns:
            filt = filt[pd.to_numeric(filt["kazanma_orani"], errors="coerce").fillna(0) >= s_min_kaz]
        if s_min_ekotu is not None and "en_kotu_gun" in filt.columns:
            filt = filt[pd.to_numeric(filt["en_kotu_gun"], errors="coerce").fillna(-999) >= s_min_ekotu]
        if s_min_sh is not None and "sharpe" in filt.columns:
            filt = filt[pd.to_numeric(filt["sharpe"], errors="coerce").fillna(-999) >= s_min_sh]
        if s_min_so is not None and "sortino" in filt.columns:
            filt = filt[pd.to_numeric(filt["sortino"], errors="coerce").fillna(-999) >= s_min_so]
        if s_max_dd is not None and "max_drawdown" in filt.columns:
            filt = filt[pd.to_numeric(filt["max_drawdown"], errors="coerce").fillna(-999) >= s_max_dd]
        if s_max_vol is not None and "volatilite" in filt.columns:
            filt = filt[pd.to_numeric(filt["volatilite"], errors="coerce").fillna(999) <= s_max_vol]
        if s_max_var is not None and "var95" in filt.columns:
            filt = filt[pd.to_numeric(filt["var95"], errors="coerce").fillna(-999) >= -abs(s_max_var)]
        if s_max_cvar is not None and "cvar95" in filt.columns:
            filt = filt[pd.to_numeric(filt["cvar95"], errors="coerce").fillna(-999) >= -abs(s_max_cvar)]
        if s_min_cagr is not None and "cagr" in filt.columns:
            filt = filt[pd.to_numeric(filt["cagr"], errors="coerce").fillna(-999) >= s_min_cagr]
        if s_min_calmar is not None and "calmar" in filt.columns:
            filt = filt[pd.to_numeric(filt["calmar"], errors="coerce").fillna(-999) >= s_min_calmar]
        if s_min_reel and "reel_1y" in filt.columns:
            filt = filt[pd.to_numeric(filt["reel_1y"], errors="coerce").fillna(-999) > 0]
        if (s_kgr or s_ktam) and "kategori" in filt.columns:
            filt = filt[kategori_maske(filt["kategori"], s_kgr, s_ktam)]
        if s_kategori and "kategori" in filt.columns:
            nt = core._normalize_tr(s_kategori).lower()
            filt = filt[filt["kategori"].astype(str).apply(lambda x: nt in core._normalize_tr(x).lower())]
        if s_max_bist is not None and "bist_korelasyon" in filt.columns:
            bk = pd.to_numeric(filt["bist_korelasyon"], errors="coerce")
            filt = filt[bk.isna() | (bk <= s_max_bist)]
        if s_temiz and "risk_bayrak" in filt.columns:
            rb = filt["risk_bayrak"].astype(str).str.lower()
            filt = filt[~rb.str.contains("aşırı|akran-dışı|sıçrama|yüksek fiyat", na=False)]
        if s_min_1a is not None and "getiri_1a" in filt.columns:
            filt = filt[pd.to_numeric(filt["getiri_1a"], errors="coerce").fillna(-999) >= s_min_1a]
        if s_min_3a is not None and "getiri_3a" in filt.columns:
            filt = filt[pd.to_numeric(filt["getiri_3a"], errors="coerce").fillna(-999) >= s_min_3a]
        # Rejim skor eşikleri
        if s_min_dususte is not None and "dususte_artan_skor" in filt.columns:
            filt = filt[pd.to_numeric(filt["dususte_artan_skor"], errors="coerce").fillna(0) >= s_min_dususte]
        if s_min_kriz is not None and "kriz_dayanik_skor" in filt.columns:
            filt = filt[pd.to_numeric(filt["kriz_dayanik_skor"], errors="coerce").fillna(0) >= s_min_kriz]
        if s_min_asim is not None and "asimetri_skor" in filt.columns:
            filt = filt[pd.to_numeric(filt["asimetri_skor"], errors="coerce").fillna(0) >= s_min_asim]
        # Ek filtreler (yaş, süreklilik, kategori içi sıra, alfa/beta, yoğunlaşma, akış, valör, maliyet)
        filt = TE.panel(st, core, get_analiz, filt, "sira")
        # Sıralama (kısa vade / skor / BIST / rejim)
        _sort_col = s_sirala
        _asc = False
        if s_sirala == "bist_korelasyon_asc":
            _sort_col = "bist_korelasyon"
            _asc = True  # düşük / negatif ρ önce = BIST düşünce artmaya aday
        elif s_sirala in ("downside_capture", "bist_korelasyon"):
            # düşük downside capture / düşük ρ daha iyi → küçükten büyüğe
            if s_sirala == "downside_capture":
                _asc = True
        if _sort_col in filt.columns and not filt.empty:
            filt = filt.copy()
            filt["_sort"] = pd.to_numeric(filt[_sort_col], errors="coerce")
            filt = filt.sort_values("_sort", ascending=_asc, na_position="last").drop(columns=["_sort"]).reset_index(drop=True)
            if "sira" in filt.columns:
                filt["sira"] = range(1, len(filt) + 1)
            else:
                filt.insert(0, "sira", range(1, len(filt) + 1))

        st.caption(
            f"Akıllı filtre sonrası: **{len(filt)}** / {len(work)} fon"
            + (f" · sıralama: `{s_sirala}`" if s_sirala else "")
        )

        cols = [
            "sira", "secim_skor", "tutarlilik_skor", "tutarlilik_n", "puan", "kod", "ad", "yatirim", "kategori",
            "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y", "getiri_3y",
            "reel_1y", "vs_mevduat", "vs_bist100",
            "sharpe", "sortino", "calmar", "cagr", "kazanma_orani",
            "en_kotu_gun", "en_iyi_gun", "var95", "cvar95",
            "volatilite", "max_drawdown",
            "risk_skor", "likidite_skor", "likidite_seviye",
            "bist_korelasyon", "beta", "beta_up", "beta_down",
            "upside_capture", "downside_capture", "capture_ratio",
            "bist_dusuk_gun_fon", "kriz_getiri",
            "dususte_artan_skor", "kriz_dayanik_skor", "asimetri_skor",
            "baskin_tur", "risk_bayrak",
        ]
        show_df(filt if not filt.empty else work, cols + [c for c in st.session_state.get("sira_ek_cols", []) if c not in cols], height=480)

        download_excel_button(
            filt if not filt.empty else work,
            f"tefas_donem_sira_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            meta={"mod": "tum_fon_siralama", "pozitif_1y": poz, "pozitif_3y": poz3},
            key="xl_sira",
        )

        # ── Akıllı portföy (10 fon) ──────────────────────────────────────
        st.subheader("🎯 Akıllı portföy (sıralamadan en iyi fonlar)")
        st.caption(
            f"Profil: **{preset_sec}** · çok kriterli skor + tür çeşitliliği. "
            "Yukarıdaki filtreler uygulanmış liste üzerinden seçer."
        )
        ap1, ap2, ap3, ap4 = st.columns(4)
        n_akilli_s = ap1.slider("Fon sayısı", 5, 15, int(pr.get("n_fon", 10)), key=f"n_akilli_sira_{pk}")
        _ym_opts = ["max_sharpe", "risk_parity", "score", "equal"]
        _ym_def = pr.get("yontem", "max_sharpe")
        ym_akilli_s = ap2.selectbox(
            "Ağırlık yöntemi",
            _ym_opts,
            index=_ym_opts.index(_ym_def) if _ym_def in _ym_opts else 0,
            format_func=lambda x: {
                "max_sharpe": "Max Sharpe", "risk_parity": "Risk Parity",
                "score": "Skor orantılı", "equal": "Eşit",
            }[x],
            key=f"ym_akilli_sira_{pk}",
        )
        max_w_s = ap3.slider(
            "Tek fon max %", 10, 40, int(round(float(pr.get("max_w", 0.2)) * 100)),
            key=f"maxw_akilli_sira_{pk}",
        ) / 100.0
        max_tur_s = ap4.slider("Aynı türden max fon", 1, 5, int(pr.get("max_ayni_tur", 3)), key=f"max_tur_sira_{pk}")

        b_col1, b_col2, b_col3 = st.columns(3)
        with b_col1:
            btn_tek = st.button("🚀 Akıllı seçim ile portföy oluştur", type="primary", key="btn_akilli_sira")
        with b_col2:
            btn_kar = st.button("🔀 Karışık (Muh+Den+Agr)", key="btn_karisik_sira")
        with b_col3:
            btn_cu = st.button("🏛️ Çekirdek–Uydu portföy", key="btn_cekirdek_uydu")

        if btn_cu:
            kaynak = filt if not filt.empty else work
            if kaynak.empty:
                st.warning("Filtrelenmiş liste boş.")
            else:
                analiz = get_analiz()
                with st.spinner("Çekirdek–uydu portföy hesaplanıyor…"):
                    sonuc = core.cekirdek_uydu_portfoy(
                        analiz, kaynak,
                        n_cekirdek=6, n_uydu=4,
                        cekirdek_agirlik=0.70,
                        yontem=ym_akilli_s,
                        max_w_cekirdek=max_w_s,
                        max_w_uydu=min(0.15, max_w_s),
                        max_ayni_tur=max_tur_s,
                    )
                if sonuc.get("error"):
                    st.error(sonuc["error"])
                else:
                    st.session_state.akilli_portfoy_sira = sonuc

        if btn_kar:
            with st.expander("Karışık portföy ayarları", expanded=True):
                kc1, kc2, kc3 = st.columns(3)
                n_muh = kc1.slider("Muhafazakâr fon", 1, 6, 3, key="n_muh_kar")
                n_den = kc2.slider("Dengeli fon", 1, 6, 4, key="n_den_kar")
                n_agr = kc3.slider("Agresif fon", 1, 6, 3, key="n_agr_kar")
            kaynak = filt if not filt.empty else work
            if kaynak.empty:
                st.warning("Filtrelenmiş liste boş.")
            else:
                # BIST düşünce artan için korelasyon gerekir
                analiz = get_analiz()
                with st.spinner("Karışık portföy (3 profil) hesaplanıyor…"):
                    sonuc = core.akilli_karisik_portfoy(
                        analiz, kaynak,
                        n_muhafazakar=n_muh,
                        n_dengeli=n_den,
                        n_agresif=n_agr,
                        yontem=ym_akilli_s,
                        max_w=max_w_s,
                    )
                if sonuc.get("error"):
                    st.error(sonuc["error"])
                else:
                    st.session_state.akilli_portfoy_sira = sonuc

        if btn_tek:
            kaynak = filt if not filt.empty else work
            if kaynak.empty:
                st.warning("Filtrelenmiş liste boş.")
            else:
                analiz = get_analiz()
                with st.spinner("Akıllı portföy hesaplanıyor…"):
                    sonuc = core.akilli_portfoy_sec(
                        analiz, kaynak,
                        n_fon=n_akilli_s,
                        yontem=ym_akilli_s,
                        max_w=max_w_s,
                        max_ayni_tur=max_tur_s,
                        max_bist_korr=s_max_bist,
                        diversify=True,
                    )
                if sonuc.get("error"):
                    st.error(sonuc["error"])
                else:
                    st.session_state.akilli_portfoy_sira = sonuc

        sonuc = st.session_state.get("akilli_portfoy_sira")
        if sonuc and not sonuc.get("error") and sonuc.get("secili") is not None:
            secili = sonuc["secili"]
            st.success(
                f"**{sonuc.get('n')} fon** seçildi · yöntem: `{sonuc.get('method')}` · "
                f"tür: {sonuc.get('tur_dagilim')}"
            )
            if sonuc.get("port_sharpe") is not None:
                p1, p2, p3 = st.columns(3)
                p1.metric("Beklenen μ", f"%{sonuc.get('port_mu')}")
                p2.metric("Vol", f"%{sonuc.get('port_vol')}")
                p3.metric("Sharpe", f"{sonuc.get('port_sharpe')}")
            if sonuc.get("ozet"):
                with st.expander("Neden bu fonlar?", expanded=True):
                    for line in sonuc["ozet"]:
                        st.markdown(f"- {line}")
            wmap = sonuc.get("weights") or {}
            show_cols = [
                c for c in [
                    "kod", "ad", "kategori", "baskin_tur", "akilli_skor",
                    "temel_skor", "bonus_kazanma", "bonus_en_kotu_gun",
                    "bonus_likidite", "bonus_bist", "ceza_anomali",
                    "getiri_1y", "sharpe", "sortino", "max_drawdown",
                    "bist_korelasyon", "likidite_skor", "profil_kaynak", "secim_gerekce",
                ] if c in secili.columns
            ]
            view_p = secili[show_cols].copy() if show_cols else secili.copy()
            view_p["agirlik_%"] = view_p["kod"].astype(str).map(
                lambda k: round(float(wmap.get(k, 0)) * 100, 1)
            )
            show_df(view_p, list(view_p.columns), height=400)
            with st.expander("Skor kırılımı (her fon)"):
                for _, r in secili.iterrows():
                    st.caption(
                        f"**{r.get('kod')}** akilli={r.get('akilli_skor')} = "
                        f"temel {r.get('temel_skor')} + kazanma {r.get('bonus_kazanma')} + "
                        f"kötü gün {r.get('bonus_en_kotu_gun')} + likidite {r.get('bonus_likidite')} + "
                        f"BIST {r.get('bonus_bist')} + anomali {r.get('ceza_anomali')}"
                    )
            edf = secili.copy()
            edf["portfoy_agirlik_%"] = edf["kod"].astype(str).map(
                lambda k: round(float(wmap.get(k, 0)) * 100, 2)
            )
            download_excel_button(
                edf,
                f"tefas_sira_akilli_portfoy_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                meta={"mod": "siralama_akilli_portfoy", "yontem": sonuc.get("method"), "agirliklar": wmap},
                key="xl_sira_akilli",
            )
            fc1, fc2 = st.columns(2)
            with fc1:
                if st.button("⭐ Favorilere ekle", key="fav_sira_akilli"):
                    kodlar = secili["kod"].astype(str).tolist()
                    favs = list(set(core.load_favs()) | set(kodlar))
                    core.save_favs(favs)
                    st.success(f"{len(kodlar)} favori")
            with fc2:
                if st.button("🛒 Portföy sepetine ekle", key="sepet_sira_akilli"):
                    ad_map = {str(r["kod"]): str(r.get("ad") or r["kod"]) for _, r in secili.iterrows()}
                    n = sepet_ekle(secili["kod"].astype(str).tolist(), ad_map=ad_map, kaynak="siralama")
                    st.success(f"{n} fon sepete eklendi (toplam {len(st.session_state.portfoy_sepet)})")

        if core.HAS_MPL and st.button("📊 PNG grafik rapor", key="gr_sira"):
            path = core.grafik_rapor(
                filt if not filt.empty else work,
                baslik=f"Dönem sıralaması | {datetime.now().strftime('%d.%m.%Y')}",
            )
            if path:
                st.image(path)
    elif df is not None and df.empty:
        st.warning("Kriterlere uyan fon yok.")


# ═══════════════════════════════════════════════════════════════════════════
# KARŞILAŞTIRMA
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "🧪 Filtre Stüdyosu":
    try:
        tefas_filtre.render(
            core=core, get_analiz=get_analiz, sepet_ekle=sepet_ekle, show_df=show_df,
            download_excel_button=download_excel_button, fmt_num=fmt_num,
        )
    except Exception as e:
        st.error(f"Filtre Stüdyosu yüklenemedi: {e}")
        st.caption("tefas_filtre modülünü kontrol edin.")

    st.markdown("---")
    st.markdown("### Hazır profil — birleşik skor + filtre")
    st.caption(
        "**Dengeli** · **Değer** (F/K≤15) · **Kalite** · **Momentum** · **Taze bilanço**. Tavsiye değildir."
    )
    fs_prof = st.radio(
        "Hazır profil",
        ["Dengeli", "Değer", "Kalite", "Momentum", "Taze bilanço", "Düşük kaldıraç"],
        horizontal=True,
        key="fs_hazir_profil",
    )
    fs_list = st.text_input(
        "Hisse kodları (virgülle)",
        value="ISDMR,THYAO,GARAN,EREGL,ASELS,SISE,KCHOL,TUPRS,ARCLK,TCELL",
        key="fs_birlesik_list",
    )
    if st.button("Profili uygula / sırala", type="primary", key="btn_fs_birlesik"):
        kodlar = [k.strip().upper().replace(".IS", "") for k in fs_list.split(",") if k.strip()]
        rows = []
        bar = st.progress(0.0)
        for i, k in enumerate(kodlar):
            try:
                rows.append(hisse_birlesik_satir(k))
            except Exception as e:
                rows.append({"kod": k, "ad": k, "birlesik_ortalama": None, "hata": str(e)})
            bar.progress((i + 1) / max(len(kodlar), 1))
        bar.empty()
        dfb = hazir_profil_filtrele(pd.DataFrame(rows), fs_prof)
        st.session_state["fs_birlesik_df"] = dfb
        st.session_state["fs_birlesik_profil"] = fs_prof
    dfb = st.session_state.get("fs_birlesik_df")
    if dfb is not None and not getattr(dfb, "empty", True):
        show_cols = [
            c for c in (
                "kod", "ad", "bilanco_donem", "kalite_skor", "model_sinyal",
                "sinyal_norm", "ucuzluk_skor", "fiyat_persentil_1y",
                "birlesik_ortalama", "seviye",
                "F/K", "PD/DD", "ROE", "getiri_1y",
            ) if c in dfb.columns
        ]
        show_df(dfb, show_cols, height=360)
        download_excel_button(
            dfb[show_cols] if show_cols else dfb,
            f"filtre_birlesik_skor_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            meta={"mod": "birlesik_kalite_sinyal", "siralama": "yuksek_dusuk"},
            key="xl_fs_birlesik",
        )
        st.caption(f"Profil: **{st.session_state.get('fs_birlesik_profil', 'Dengeli')}**. Tavsiye değildir.")

elif sayfa == "⚖️ Karşılaştırma":
    st.header("⚖️ Karşılaştırma — Fon + Hisse")
    st.caption("Fon (TEFAS) ve hisse (BIST) yan yana; getiri, risk, çarpan, bilanço, kalite ve model sinyal.")
    c1, c2 = st.columns(2)
    with c1:
        fon_str = st.text_input("Fon kodları (virgülle)", value="AAL,AFO", key="kar_fon")
    with c2:
        hisse_str = st.text_input("Hisse kodları (virgülle)", value="ISDMR,THYAO", key="kar_hisse")
    if st.button("Karşılaştır", type="primary", key="btn_kar_mix"):
        rows_f, rows_h = [], []
        with st.spinner("Fonlar (TEFAS)…"):
            try:
                analiz = get_analiz()
            except Exception as e:
                analiz = None
                st.warning(f"TEFAS analiz: {e}")
            if analiz is not None:
                for kod in [k.strip().upper() for k in fon_str.split(",") if k.strip()]:
                    try:
                        d = analiz.fon_detay(kod)
                        m, o, p = d["metrikler"], d["ozet"], d["profil"]
                        sig = model_sinyal_skoru(m=m)
                        rows_f.append({
                            "tur": "FON", "kod": kod,
                            "ad": (o.get("fonUnvan") or p.get("fonUnvan") or kod)[:48],
                            "risk": p.get("riskDegeri"),
                            "getiri_1a": m.get("getiri_1a"), "getiri_3a": m.get("getiri_3a"),
                            "getiri_1y": m.get("getiri_1y"), "volatilite": m.get("volatilite"),
                            "max_drawdown": m.get("max_drawdown"), "sharpe": m.get("sharpe"),
                            "fiyat": m.get("fiyat"), "F/K": None, "PD/DD": None, "ROE": None,
                            "bilanco_donem": None, "kalite_skor": None,
                            "model_sinyal": sig.get("skor"),
                        })
                    except Exception as e:
                        st.warning(f"Fon {kod}: {e}")
        with st.spinner("Hisseler…"):
            for kod in [k.strip().upper().replace(".IS", "") for k in hisse_str.split(",") if k.strip()]:
                row = {
                    "tur": "HISSE", "kod": kod, "ad": kod, "risk": None,
                    "getiri_1a": None, "getiri_3a": None, "getiri_1y": None,
                    "volatilite": None, "max_drawdown": None, "sharpe": None, "fiyat": None,
                    "F/K": None, "PD/DD": None, "ROE": None,
                    "bilanco_donem": None, "kalite_skor": None, "model_sinyal": None,
                }
                try:
                    import yfinance as yf
                    t = yf.Ticker(f"{kod}.IS")
                    hist = t.history(period="1y", auto_adjust=True)
                    if hist is not None and not hist.empty:
                        px = hist["Close"].dropna()
                        row["fiyat"] = float(px.iloc[-1])
                        if len(px) > 22:
                            row["getiri_1a"] = float((px.iloc[-1] / px.iloc[-22] - 1) * 100)
                        if len(px) > 66:
                            row["getiri_3a"] = float((px.iloc[-1] / px.iloc[-66] - 1) * 100)
                        if len(px) > 2:
                            row["getiri_1y"] = float((px.iloc[-1] / px.iloc[0] - 1) * 100)
                        rets = px.pct_change().dropna()
                        if len(rets) > 10:
                            row["volatilite"] = float(rets.std() * (252 ** 0.5) * 100)
                            cum = (1 + rets).cumprod()
                            row["max_drawdown"] = float((cum / cum.cummax() - 1).min() * 100)
                    try:
                        info = dict(t.info or {})
                    except Exception:
                        info = {}
                    if info.get("returnOnEquity") is not None:
                        row["ROE"] = float(info["returnOnEquity"]) * 100
                    row["ad"] = (info.get("longName") or info.get("shortName") or kod)[:48]
                    try:
                        _ic = deg_from_isyatirim(kod)
                        if not _ic.get("_error"):
                            row["F/K"] = _ic.get("F/K")
                            row["PD/DD"] = _ic.get("PD/DD")
                            if _ic.get("fiyat") is not None:
                                row["fiyat"] = _ic["fiyat"]
                    except Exception:
                        pass
                except Exception as e:
                    st.caption(f"{kod} fiyat: {e}")
                try:
                    isy = isyatirim_mali_tablo(kod)
                    if isy and not isy.get("error"):
                        row["bilanco_donem"] = isy.get("gunce_donem")
                        orans = isy.get("oranlar") or {}
                        if orans.get("ROE_%") is not None:
                            row["ROE"] = orans["ROE_%"]
                        kal = bilanco_kalite_skoru(isy, {"F/K": row["F/K"], "PD/DD": row["PD/DD"], "ROE": row["ROE"]})
                        row["kalite_skor"] = kal.get("skor")
                        sig = model_sinyal_skoru(
                            m={"getiri_1y": row["getiri_1y"], "volatilite": row["volatilite"],
                               "max_drawdown": row["max_drawdown"], "sharpe": row["sharpe"]},
                            deg={"F/K": row["F/K"], "PD/DD": row["PD/DD"], "ROE": row["ROE"]},
                            kalite=kal,
                        )
                        row["model_sinyal"] = sig.get("skor")
                except Exception as e:
                    st.caption(f"{kod} bilanço: {e}")
                rows_h.append(row)
        st.session_state.son_karsilastir = pd.DataFrame(rows_f + rows_h)

    df = st.session_state.get("son_karsilastir")
    if df is not None and not getattr(df, "empty", True):
        st.subheader("Ortak metrik tablosu")
        cols = [
            "tur", "kod", "ad", "bilanco_donem", "fiyat",
            "getiri_1a", "getiri_3a", "getiri_1y", "volatilite", "max_drawdown", "sharpe",
            "F/K", "PD/DD", "ROE", "kalite_skor", "model_sinyal", "risk",
        ]
        show_df(df, cols, height=380)
        download_excel_button(
            df, f"karsilastir_fon_hisse_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            meta={"mod": "karsilastirma_fon_hisse"}, key="xl_kar_mix",
        )
        st.caption("Model sinyal (−100…+100) kural tabanlıdır; al/sat tavsiyesi değildir.")


# ═══════════════════════════════════════════════════════════════════════════
# RAPOR PAKETLERİ
# ═══════════════════════════════════════════════════════════════════════════

elif sayfa == "📦 Rapor Paketleri":
    st.header("📦 Rapor Paketleri")
    st.caption(
        "Kurumsal HTML rapor: F/K kontrolü, nakit/PD, satış ve kâr tabloları, grafikler. "
        "Yatırım tavsiyesi değildir."
    )
    rp1, rp2 = st.columns(2)
    with rp1:
        rp_tur = st.radio("Tür", ["Hisse", "Fon"], horizontal=True, key="rp_tur")
    with rp2:
        rp_kod = st.text_input(
            "Kod",
            value="ISDMR" if rp_tur == "Hisse" else "AAL",
            key="rp_kod",
        )
    if st.button("Rapor paketi oluştur", type="primary", key="btn_rp"):
        kod = (rp_kod or "").strip().upper().replace(".IS", "")
        bloklar, tables, charts = [], [], []
        meta, kpi = {"Kod": kod, "Tür": rp_tur}, {}
        ad, donem, fk_yorum = kod, "", ""
        with st.spinner("Kurumsal rapor hazırlanıyor…"):
            if rp_tur == "Hisse":
                isy, info = {}, {}
                m, deg = {}, {}
                try:
                    isy = isyatirim_mali_tablo(kod, force=True) or {}
                except Exception as e:
                    st.warning(f"İş Yatırım: {e}")
                    isy = {}
                donem = isy.get("gunce_donem") or ""
                meta["Bilanço dönemi"] = donem or "—"
                oz = isy.get("ozet") or {}
                oran = isy.get("oranlar") or {}

                def _oz_val(key):
                    item = oz.get(key) or {}
                    return item.get("deger") if isinstance(item, dict) else None

                def _oz_yoy(key):
                    item = oz.get(key) or {}
                    return item.get("degisim_%") if isinstance(item, dict) else None

                nakit = _oz_val("nakit")
                hasilat = _oz_val("hasilat")
                brut = _oz_val("brut_kar")
                faaliyet = _oz_val("esas_faaliyet_kari")
                net = _oz_val("net_donem_kari")
                ozk = _oz_val("ozkaynaklar")

                # Çarpanlar: YALNIZCA İş Yatırım (Yahoo F/K kullanılmaz)
                try:
                    deg = deg_from_isyatirim(kod, force=True)
                    if deg.get("fiyat") is not None:
                        m["fiyat"] = deg["fiyat"]
                    mcap = deg.get("piyasa_değeri")
                    meta["Çarpan kaynağı"] = "İş Yatırım şirket kartı"
                except Exception as e:
                    st.caption(f"İş Yatırım çarpan: {e}")
                    mcap = None
                # Fiyat / 1Y getiri için yfinance yalnızca OHLCV (çarpan değil)
                try:
                    import yfinance as yf
                    t = yf.Ticker(f"{kod}.IS")
                    hist = t.history(period="1y", auto_adjust=True)
                    if hist is not None and not hist.empty:
                        px = hist["Close"].dropna()
                        if m.get("fiyat") is None and len(px):
                            m["fiyat"] = float(px.iloc[-1])
                        if len(px) > 2:
                            m["getiri_1y"] = float((px.iloc[-1] / px.iloc[0] - 1) * 100)
                    try:
                        info = dict(t.info or {})
                        ad = info.get("longName") or info.get("shortName") or ad
                    except Exception:
                        pass
                except Exception as e:
                    st.caption(f"Fiyat serisi: {e}")
                if oran.get("ROE_%") is not None:
                    deg["ROE"] = oran["ROE_%"]

                fk = deg.get("F/K")
                pddd = deg.get("PD/DD")
                # F/K kontrol
                try:
                    fk_f = float(fk) if fk is not None else None
                except (TypeError, ValueError):
                    fk_f = None
                if fk_f is None:
                    fk_yorum = (
                        "F/K: veri yok veya negatif kâr nedeniyle tanımsız. "
                        "Çarpan yorumu için net kâr ve sektör ortalamasını kontrol edin."
                    )
                elif fk_f < 0:
                    fk_yorum = f"F/K ({fk_f:.1f}) negatif — zarar dönemi veya muhasebe etkisi; klasik F/K yorumu geçersiz."
                elif fk_f < 8:
                    fk_yorum = f"F/K = {fk_f:.1f} — düşük çarpan bandı (sektör ile kıyaslayın; ucuz görünüm)."
                elif fk_f <= 15:
                    fk_yorum = f"F/K = {fk_f:.1f} — orta / makul çarpan bandı."
                elif fk_f <= 25:
                    fk_yorum = f"F/K = {fk_f:.1f} — görece yüksek; büyüme beklentisi fiyatlanmış olabilir."
                else:
                    fk_yorum = f"F/K = {fk_f:.1f} — yüksek çarpan; değerleme riski artmış olabilir."

                nakit_pd = None
                if nakit is not None and mcap:
                    try:
                        nakit_pd = float(nakit) / float(mcap) * 100.0
                    except Exception:
                        nakit_pd = None

                kpi = {
                    "Fiyat": _fmt_tl(m.get("fiyat")) if m.get("fiyat") is not None else "—",
                    "F/K": f"{fk_f:.1f}" if fk_f is not None else "—",
                    "PD/DD": tr_num(pddd, 2) if pddd is not None else "—",
                    "Piyasa değeri": _fmt_tl(mcap) if mcap else "—",
                    "Dönem sonu nakit": _fmt_tl(nakit) if nakit is not None else "—",
                    "Nakit / PD": f"%{nakit_pd:.1f}" if nakit_pd is not None else "—",
                }
                meta.update({
                    "Şirket": ad,
                    "F/K": tr_num(fk_f, 2) if fk_f is not None else "yok",
                    "PD/DD": tr_num(pddd, 2) if pddd is not None else "—",
                    "ROE %": tr_num(deg.get("ROE"), 1) if deg.get("ROE") is not None else "—",
                    "1Y getiri %": f"{m['getiri_1y']:.1f}" if m.get("getiri_1y") is not None else "—",
                    "Nakit/PD %": tr_num(nakit_pd, 2) if nakit_pd is not None else "—",
                })


                # ── Çok dönemli satış / kârlılık (İş Yatırım kolonları) ──
                col_labels = list(isy.get("col_labels") or [])
                gelir_df = isy.get("gelir_df")
                bilanco_df = isy.get("bilanco_df")

                def _series_from_df(df, codes, prefer_desc=None):
                    """itemCode listesinden ilk eşleşen satırın dönem değerleri."""
                    if df is None or getattr(df, "empty", True) or not col_labels:
                        return None, None
                    work = df.copy()
                    if "kod" not in work.columns:
                        return None, None
                    for code in codes:
                        hit = work[work["kod"].astype(str) == code]
                        if hit.empty:
                            continue
                        row0 = hit.iloc[0]
                        vals = []
                        for lab in col_labels:
                            v = row0.get(lab)
                            try:
                                vals.append(float(v) if v is not None and v == v else None)
                            except (TypeError, ValueError):
                                vals.append(None)
                        return row0.get("kalem") or code, vals
                    return None, None

                def _pct(a, b):
                    try:
                        if a is None or b is None or b == 0:
                            return None
                        return round((float(a) / float(b) - 1.0) * 100.0, 1)
                    except Exception:
                        return None

                def _marj(kar, sat):
                    try:
                        if kar is None or sat is None or sat == 0:
                            return None
                        return round(float(kar) / float(sat) * 100.0, 1)
                    except Exception:
                        return None

                sat_lab, sat_vals = _series_from_df(gelir_df, ("3C", "3A", "3AA"))
                brut_lab, brut_vals = _series_from_df(gelir_df, ("3D", "3CAB"))
                faal_lab, faal_vals = _series_from_df(gelir_df, ("3DF", "3H", "3E"))
                net_lab, net_vals = _series_from_df(gelir_df, ("3L", "3J", "3I"))
                nakit_lab, nakit_vals = _series_from_df(bilanco_df, ("1AA", "1AAA"))
                ozk_lab, ozk_vals = _series_from_df(bilanco_df, ("2N", "2O"))

                # Tek dönem yedek (ozet)
                if not sat_vals or all(v is None for v in (sat_vals or [])):
                    sat_vals = [_oz_val("hasilat")] + [None] * max(0, len(col_labels) - 1)
                    col_labels = col_labels or [donem or "Güncel"]
                if not net_vals or all(v is None for v in (net_vals or [])):
                    net_vals = [_oz_val("net_donem_kari")] + [None] * max(0, len(col_labels) - 1)
                if not brut_vals:
                    brut_vals = [_oz_val("brut_kar")] + [None] * max(0, len(col_labels) - 1)
                if not faal_vals:
                    faal_vals = [_oz_val("esas_faaliyet_kari")] + [None] * max(0, len(col_labels) - 1)

                # Çok dönem satış + kâr tablosu
                multi_headers = ["Kalem"] + list(col_labels)
                multi_rows = []
                for label, vals in [
                    ("Satışlar (hasılat)", sat_vals),
                    ("Brüt kâr", brut_vals),
                    ("Esas faaliyet kârı", faal_vals),
                    ("Net dönem kârı", net_vals),
                    ("Nakit", nakit_vals),
                    ("Özkaynaklar", ozk_vals),
                ]:
                    if not vals:
                        continue
                    multi_rows.append(
                        [label] + [_fmt_tl(v) if v is not None else "—" for v in vals]
                    )
                if multi_rows:
                    tables.append((
                        f"Çok dönemli karşılaştırma (İş Yatırım) — güncel: {donem or col_labels[0] if col_labels else '—'}",
                        multi_headers,
                        multi_rows,
                    ))

                # Dönemsel kârlılık marjları
                marj_rows = []
                for i, lab in enumerate(col_labels):
                    s = sat_vals[i] if sat_vals and i < len(sat_vals) else None
                    b = brut_vals[i] if brut_vals and i < len(brut_vals) else None
                    f = faal_vals[i] if faal_vals and i < len(faal_vals) else None
                    n = net_vals[i] if net_vals and i < len(net_vals) else None
                    marj_rows.append([
                        lab,
                        f"%{_marj(b, s)}" if _marj(b, s) is not None else "—",
                        f"%{_marj(f, s)}" if _marj(f, s) is not None else "—",
                        f"%{_marj(n, s)}" if _marj(n, s) is not None else "—",
                        _fmt_tl(s) if s is not None else "—",
                        _fmt_tl(n) if n is not None else "—",
                    ])
                if marj_rows:
                    tables.append((
                        "Dönemsel kârlılık (marjlar)",
                        ["Dönem", "Brüt marj %", "Faaliyet marj %", "Net marj %", "Satış", "Net kâr"],
                        marj_rows,
                    ))

                # Satış büyümesi vs enflasyon
                tufe = None
                try:
                    tufe = float((st.session_state.cfg or {}).get("tufe_yillik") or 0) or None
                except Exception:
                    tufe = None
                if tufe is None:
                    tufe = 40.0  # yedek gösterim; meta'da belirtilir
                    tufe_src = "varsayılan (Ayarlar'da TÜFE yok)"
                else:
                    tufe_src = "Ayarlar · tufe_yillik"

                buyume_rows = []
                # Ardışık dönemler ve YoY aynı çeyrek (col0 vs col1 genelde YoY)
                if sat_vals and len(sat_vals) >= 2:
                    for i in range(len(sat_vals) - 1):
                        a, b = sat_vals[i], sat_vals[i + 1]
                        g = _pct(a, b)
                        if g is None:
                            continue
                        reel = round(g - float(tufe), 1) if tufe is not None else None
                        buyume_rows.append([
                            f"{col_labels[i]} vs {col_labels[i+1]}",
                            f"%{g:.1f}",
                            f"%{tufe:.1f}" if tufe is not None else "—",
                            f"%{reel:.1f}" if reel is not None else "—",
                            "Reel büyüme (+)" if reel is not None and reel > 0 else (
                                "Enflasyon altı" if reel is not None else "—"
                            ),
                        ])
                # Net kâr büyümesi de
                net_buyume = []
                if net_vals and len(net_vals) >= 2:
                    for i in range(min(3, len(net_vals) - 1)):
                        g = _pct(net_vals[i], net_vals[i + 1])
                        if g is None:
                            continue
                        net_buyume.append([
                            f"Net {col_labels[i]} vs {col_labels[i+1]}",
                            f"%{g:.1f}",
                        ])
                if buyume_rows:
                    tables.append((
                        f"Satış büyümesi vs enflasyon (TÜFE %{tufe:.0f} · {tufe_src})",
                        ["Karşılaştırma", "Satış büyüme %", "TÜFE %", "Fark (satış−TÜFE)", "Yorum"],
                        buyume_rows,
                    ))
                    meta["TÜFE %"] = tufe
                    meta["Satış−TÜFE (ilk satır)"] = buyume_rows[0][3] if buyume_rows else "—"

                # Tek dönem özet (geriye uyum)
                gel_rows = []
                for label, key in [
                    ("Satışlar (hasılat)", "hasilat"),
                    ("Brüt kâr", "brut_kar"),
                    ("Esas faaliyet kârı", "esas_faaliyet_kari"),
                    ("Net dönem kârı", "net_donem_kari"),
                ]:
                    v = _oz_val(key)
                    y = _oz_yoy(key)
                    gel_rows.append([
                        label,
                        _fmt_tl(v) if v is not None else "—",
                        f"%{tr_num(y, 1)}" if y is not None else "—",
                    ])
                tables.append((
                    f"Gelir özeti (güncel dönem {donem or '—'})",
                    ["Kalem", "Tutar (TL)", "YoY %"],
                    gel_rows,
                ))

                bil_rows = []
                for label, key in [
                    ("Nakit ve benzerleri", "nakit"),
                    ("Dönen varlıklar", "donen_varliklar"),
                    ("Toplam varlıklar", "toplam_varliklar"),
                    ("Kısa vadeli yükümlülük", "kisa_vadeli_yukumluluk"),
                    ("Özkaynaklar", "ozkaynaklar"),
                ]:
                    v = _oz_val(key)
                    y = _oz_yoy(key)
                    bil_rows.append([
                        label,
                        _fmt_tl(v) if v is not None else "—",
                        f"%{tr_num(y, 1)}" if y is not None else "—",
                    ])
                if nakit_pd is not None:
                    bil_rows.append(["Nakit / Piyasa değeri", f"%{nakit_pd:.2f}", "—"])
                tables.append((
                    f"Bilanço özeti — {donem or 'güncel'}",
                    ["Kalem", "Tutar / oran", "YoY %"],
                    bil_rows,
                ))

                tables.append((
                    "Değerleme çarpanları (İş Yatırım — F/K kontrolü)",
                    ["Çarpan", "Değer", "Not"],
                    [
                        ["F/K", tr_num(fk_f, 2) if fk_f is not None else "—", (fk_yorum or "")[:90]],
                        ["PD/DD", tr_num(pddd, 2) if pddd is not None else "—", "İş Yatırım"],
                        ["FD/FAVÖK", tr_num(deg.get("FD/FAVÖK"), 2) if deg.get("FD/FAVÖK") is not None else "—", "İş Yatırım"],
                        ["ROE %", tr_num(deg.get("ROE"), 1) if deg.get("ROE") is not None else "—", "Bilanço / oran"],
                        ["Nakit/PD %", tr_num(nakit_pd, 2) if nakit_pd is not None else "—", "Dönem sonu nakit ÷ PD"],
                    ],
                ))

                # Grafik: çok dönem satış
                if sat_vals and col_labels and any(v is not None for v in sat_vals):
                    labs = [str(x) for x in col_labels]
                    charts.append((
                        "Satışlar — dönemler",
                        labs,
                        [v if v is not None else 0 for v in sat_vals],
                    ))
                if net_vals and col_labels and any(v is not None for v in net_vals):
                    charts.append((
                        "Net kâr — dönemler",
                        [str(x) for x in col_labels],
                        [v if v is not None else 0 for v in net_vals],
                    ))
                # Güncel dönem kalem kırılımı
                chart_labels, chart_vals = [], []
                for lab, key in [
                    ("Satış", "hasilat"),
                    ("Brüt", "brut_kar"),
                    ("Faaliyet", "esas_faaliyet_kari"),
                    ("Net", "net_donem_kari"),
                ]:
                    v = _oz_val(key)
                    if v is not None:
                        chart_labels.append(lab)
                        chart_vals.append(v)
                if chart_labels:
                    charts.append((f"Satış ve kârlar ({donem})", chart_labels, chart_vals))
                # Marj grafiği (güncel)
                if marj_rows:
                    # her dönem net marj
                    m_labs, m_vals = [], []
                    for r in marj_rows:
                        try:
                            mv = float(str(r[3]).replace("%", "").replace("—", ""))
                            m_labs.append(str(r[0]))
                            m_vals.append(mv)
                        except Exception:
                            pass
                    if m_labs:
                        charts.append(("Net marj % — dönemler", m_labs, m_vals))
                # Satış büyüme vs TÜFE
                if buyume_rows:
                    blabs, bsat, btufe = [], [], []
                    for r in buyume_rows[:4]:
                        try:
                            blabs.append(str(r[0])[:18])
                            bsat.append(float(str(r[1]).replace("%", "")))
                            btufe.append(float(str(r[2]).replace("%", "")))
                        except Exception:
                            pass
                    if blabs:
                        charts.append(("Satış büyüme % (dönem karşılaştırmaları)", blabs, bsat))
                        charts.append((f"TÜFE referans %{tufe:.0f}", blabs, btufe))

                kal = bilanco_kalite_skoru(isy, deg)
                tu = tarihsel_ucuzluk_analizi(kod, deg=deg)
                sig = model_sinyal_skoru(m=m, deg=deg, kalite=kal, tu=tu)
                bir = birlesik_profil_skoru(kalite=kal, sinyal=sig, tu=tu)
                meta["Kalite skoru"] = kal.get("skor")
                meta["Model sinyal"] = sig.get("skor")
                meta["Ucuzluk skoru"] = tu.get("skor")
                meta["Birleşik ortalama"] = bir.get("ortalama")

                bloklar.append((
                    "F/K kontrol notu",
                    fk_yorum or "F/K verisi yok.",
                ))
                bloklar.append((
                    "Bilanço kalite",
                    f"Skor: {kal.get('skor')}/100 ({kal.get('seviye')})\n{kal.get('yorum')}\n"
                    f"Artılar: {', '.join(kal.get('artilar') or ['—'])}\n"
                    f"Eksiler: {', '.join(kal.get('eksiler') or ['—'])}",
                ))
                bloklar.append((
                    "Model sinyal (kural tabanlı — tavsiye değildir)",
                    f"Skor: {sig.get('skor')} · {sig.get('etiket')}\n{sig.get('disclaimer')}",
                ))
                bloklar.append((
                    "Birleşik profil (eşit ağırlık)",
                    bir.get("not") or "",
                ))
            else:
                try:
                    analiz = get_analiz()
                    d = analiz.fon_detay(kod)
                    m, o, p = d["metrikler"], d["ozet"], d["profil"]
                    ad = (o.get("fonUnvan") or p.get("fonUnvan") or kod)
                    meta.update({
                        "Fon ad": ad,
                        "1Y %": m.get("getiri_1y"),
                        "Sharpe": m.get("sharpe"),
                        "MaxDD": m.get("max_drawdown"),
                        "Risk": p.get("riskDegeri"),
                    })
                    kpi = {
                        "1Y getiri": f"%{m.get('getiri_1y')}" if m.get("getiri_1y") is not None else "—",
                        "Sharpe": f"{m.get('sharpe')}" if m.get("sharpe") is not None else "—",
                        "MaxDD": f"%{m.get('max_drawdown')}" if m.get("max_drawdown") is not None else "—",
                        "Volatilite": f"%{m.get('volatilite')}" if m.get("volatilite") is not None else "—",
                    }
                    tables.append((
                        "Fon performans özeti",
                        ["Metrik", "Değer"],
                        [[str(k), str(v)] for k, v in m.items() if v is not None and not str(k).startswith("_")][:20],
                    ))
                    sig = model_sinyal_skoru(m=m)
                    meta["Model sinyal"] = sig.get("skor")
                    bloklar.append((
                        "Model sinyal",
                        f"Skor: {sig.get('skor')} · {sig.get('etiket')}\n{sig.get('disclaimer')}",
                    ))
                except Exception as e:
                    st.error(f"Fon raporu: {e}")
                    bloklar.append(("Hata", str(e)))

        html = rapor_paketi_html(
            baslik="M.Durmaz Kurumsal Analiz Raporu",
            kod=kod, ad=ad or "", donem=donem or "—",
            bloklar=bloklar, meta=meta,
            tables=tables, charts=charts, kpi=kpi,
            fk_yorum=fk_yorum,
        )
        st.session_state["_rapor_html"] = html
        st.session_state["_rapor_meta"] = meta
        st.success("Kurumsal rapor hazır.")

    if st.session_state.get("_rapor_html"):
        meta0 = st.session_state.get("_rapor_meta") or {}
        st.download_button(
            "HTML rapor indir",
            data=st.session_state["_rapor_html"].encode("utf-8"),
            file_name=f"rapor_{meta0.get('Kod', meta0.get('kod', 'x'))}_{datetime.now().strftime('%Y%m%d_%H%M')}.html",
            mime="text/html",
            key="dl_rp_html",
        )
        with st.expander("Önizleme", expanded=True):
            try:
                import streamlit.components.v1 as components
                components.html(st.session_state["_rapor_html"], height=720, scrolling=True)
            except Exception:
                st.code(st.session_state["_rapor_html"][:3000])

# ═══════════════════════════════════════════════════════════════════════════
# PORTFÖY
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "📁 Portföy":
    st.header("📁 Portföy Analizi")
    # Sidebar "Analiz et" gelmişse sepet moduna zorla
    _force_sepet = bool(st.session_state.get("sepet_auto_analiz"))
    _mod_opts = ["🛒 Sepet (portföye ekle)", "Elle ağırlık", "Otomatik (en yüksek puan)"]
    mod = st.radio(
        "Mod",
        _mod_opts,
        index=0 if _force_sepet else 0,
        horizontal=True,
        key="portfoy_mod",
    )

    if mod == "🛒 Sepet (portföye ekle)":
        st.caption(
            "Tarama, sıralama, akıllı portföy veya fon detaydan **Portföye ekle** ile biriktirdiğiniz fonlar. "
            "Burada hepsini analiz edip özet rapor alırsınız."
        )
        sepet = st.session_state.get("portfoy_sepet") or []
        st.markdown("**İstediğiniz kadar fon ekleyin** (virgül / satır; limit yok)")
        c1, c2, c3 = st.columns([3, 1, 1])
        with c1:
            manuel = st.text_area(
                "Kodlar",
                placeholder="AAL, TTE, AFO\nMAC, TTE",
                key="sepet_manuel",
                height=70,
                label_visibility="collapsed",
            )
        with c2:
            st.write("")
            if st.button("Sepete ekle", key="sepet_manuel_btn", use_container_width=True) and manuel:
                import re as _re
                kodlar = [x.strip() for x in _re.split(r"[,;\s]+", manuel) if x.strip()]
                n = sepet_ekle(kodlar, kaynak="manuel")
                st.success(f"{n} eklendi (toplam {len(st.session_state.get('portfoy_sepet') or [])})")
                st.rerun()
        with c3:
            st.write("")
            if st.button("Sepeti temizle", key="sepet_temiz_page", use_container_width=True):
                st.session_state.portfoy_sepet = []
                st.session_state.sepet_analiz = None
                st.rerun()

        if not sepet:
            st.info("Sepet boş. Tarama / sıralama sonuçlarından **🛒 Portföye ekle** kullanın veya yukarıya kod yazın.")
        else:
            st.markdown(f"**Sepetteki fonlar ({len(sepet)})**")
            sdf = pd.DataFrame(sepet)
            st.dataframe(sdf, hide_index=True, use_container_width=True)
            sil = st.multiselect("Sepetten çıkar", [x["kod"] for x in sepet], key="sepet_sil")
            if st.button("Seçilenleri çıkar") and sil:
                st.session_state.portfoy_sepet = [x for x in sepet if x["kod"] not in sil]
                st.rerun()

            st.subheader("Ağırlıklar")
            use_equal = st.checkbox("Eşit ağırlık", value=True, key="sepet_eq_w")
            custom_w = {}
            if not use_equal:
                cols_w = st.columns(min(len(sepet), 4))
                for i, item in enumerate(sepet):
                    with cols_w[i % len(cols_w)]:
                        custom_w[item["kod"]] = st.number_input(
                            f"{item['kod']} %",
                            value=round(100 / len(sepet)),
                            min_value=0, max_value=100,
                            key=f"sepet_w_{item['kod']}",
                        )
            else:
                custom_w = {x["kod"]: 1.0 for x in sepet}

            # Sidebar "Analiz et" → otomatik tetik
            _auto = bool(st.session_state.pop("sepet_auto_analiz", False))
            if st.button("📊 Sepeti analiz et + derin rapor", type="primary", key="btn_sepet_analiz") or _auto:
                with st.spinner("Fonlar, korelasyon ve birleşik metrikler hesaplanıyor…"):
                    st.session_state.sepet_analiz = sepet_analiz_et(weights=custom_w)

            sa = st.session_state.get("sepet_analiz")
            if sa and sa.get("error"):
                st.error(sa["error"])
            elif sa and sa.get("df") is not None:
                pm = sa.get("port_metrikler") or {}
                # KPI şeridi — risk/getiri + tehlike ortalaması
                k1, k2, k3, k4, k5, k6 = st.columns(6)
                k1.metric("Fon", sa.get("n_fon") or len(sa["df"]))
                k2.metric("Birleşik 1Y %", fmt_num(sa.get("ort_getiri_1y") if sa.get("ort_getiri_1y") is not None else pm.get("getiri_1y")))
                k3.metric("Sharpe", fmt_num(pm.get("sharpe"), 2))
                k4.metric("MaxDD %", fmt_num(pm.get("max_drawdown")))
                k5.metric("Vol %", fmt_num(pm.get("volatilite")))
                k6.metric("Çeşitlendirme", f"{sa.get('diversify_score') or '—'}/100")
                t1, t2, t3, t4 = st.columns(4)
                t1.metric(
                    "Ort. tehlike (1–7)",
                    sa.get("ort_tehlike_yazi") or "—",
                    help="Ağırlıklı ortalama SPK risk değeri (1=çok düşük, 7=çok yüksek)",
                )
                t2.metric("Ort. getiri 1A %", fmt_num(sa.get("ort_getiri_1a")))
                t3.metric("Ort. getiri 3A %", fmt_num(sa.get("ort_getiri_3a")))
                t4.metric("Ort. getiri 6A %", fmt_num(sa.get("ort_getiri_6a")))

                tab_o, tab_t, tab_k, tab_d, tab_g, tab_c, tab_r = st.tabs([
                    "Özet rapor", "Fon tablosu", "Korelasyon", "Dağılım", "Grafik",
                    "Çeşitlendirme önerisi", "HTML / QuantStats",
                ])
                with tab_o:
                    st.markdown(sa.get("ozet_md") or "")
                    if sa.get("liderler"):
                        st.caption("Öne çıkanlar yukarıdaki metinde listelenmiştir.")
                with tab_c:
                    oner = sa.get("oneriler") or []
                    st.markdown("**Çeşitlendirme skorunu yükseltmek için eklenebilecek fonlar**")
                    st.caption(
                        "Yalnızca **TEFAS’ta açık** (işlem gören) fonlar önerilir; "
                        "tasfiye / işleme kapalı / yasaklı kurucu elenir. "
                        "Aday kaynağı: son tarama/sıralama veya TEFAS açık evren."
                    )
                    if not oner:
                        st.info(
                            "Öneri yok. Önce **Tüm Fon Sıralaması** veya **Fon Tarama** çalıştırıp "
                            "yeniden analiz edin; en az 2 fon ve fiyat serisi gerekir. "
                            "Öneriler yalnızca TEFAS açık fonlardan seçilir."
                        )
                    else:
                        odf = pd.DataFrame(oner)
                        st.dataframe(odf, hide_index=True, use_container_width=True)
                        ekle_kod = st.multiselect(
                            "Sepete ekle",
                            [o["kod"] for o in oner],
                            default=[o["kod"] for o in oner[:3]],
                            key="sepet_oneri_ekle",
                        )
                        if st.button("🛒 Önerilenleri sepete ekle", key="btn_oneri_sepet") and ekle_kod:
                            ad_map = {o["kod"]: o.get("ad", o["kod"]) for o in oner}
                            n = sepet_ekle(ekle_kod, ad_map=ad_map, kaynak="cesitlendirme")
                            st.success(f"{n} fon sepete eklendi — yeniden analiz edin.")
                            st.rerun()
                with tab_r:
                    st.markdown("**Türkçe HTML portföy raporu** (+ quantstats varsa ayrı risk raporu)")
                    if st.button("📄 HTML rapor üret", type="primary", key="btn_sepet_html"):
                        with st.spinner("HTML rapor yazılıyor…"):
                            try:
                                _pm = dict(pm or {})
                                _pm["ort_tehlike"] = sa.get("ort_tehlike")
                                _pm["ort_tehlike_yazi"] = sa.get("ort_tehlike_yazi")
                                if sa.get("ort_getiri_1y") is not None:
                                    _pm["getiri_1y"] = sa.get("ort_getiri_1y")
                                path = core.portfoy_html_raporu(
                                    baslik=f"Sepet Portföy ({sa.get('n_fon')} fon)",
                                    metrikler=_pm,
                                    tablo=sa.get("df"),
                                    diversify_score=sa.get("diversify_score"),
                                    corr_df=sa.get("corr"),
                                    oneriler=sa.get("oneriler"),
                                    ozet_md=sa.get("ozet_md") or "",
                                    fiyat_df=sa.get("port_fd"),
                                    quantstats=True,
                                )
                                if path:
                                    st.session_state["sepet_html_path"] = path
                                    st.success(f"Rapor: {path}")
                                else:
                                    st.warning("Rapor üretilemedi.")
                            except Exception as e:
                                st.error(f"Rapor hatası: {e}")
                    hp = st.session_state.get("sepet_html_path")
                    if hp and Path(hp).exists():
                        data = Path(hp).read_bytes()
                        st.download_button(
                            "⬇️ HTML indir",
                            data=data,
                            file_name=Path(hp).name,
                            mime="text/html",
                            key="dl_sepet_html",
                        )
                        st.caption(str(hp))
                        # quantstats dosyası aynı klasörde qs_rapor_* olabilir
                        try:
                            qs_files = sorted(Path(hp).parent.glob("qs_rapor_*.html"), key=lambda p: p.stat().st_mtime, reverse=True)
                            if qs_files:
                                qf = qs_files[0]
                                st.download_button(
                                    "⬇️ QuantStats HTML indir",
                                    data=qf.read_bytes(),
                                    file_name=qf.name,
                                    mime="text/html",
                                    key="dl_sepet_qs",
                                )
                        except Exception:
                            pass
                with tab_t:
                    show_df(sa["df"], list(sa["df"].columns), height=420)
                    # Ağırlık kolonu
                    wmap = sa.get("weights") or {}
                    if wmap:
                        wdf = pd.DataFrame([
                            {"kod": k, "ağırlık %": round(v * 100, 1)}
                            for k, v in sorted(wmap.items(), key=lambda x: -x[1])
                        ])
                        st.markdown("**Kullanılan ağırlıklar**")
                        st.dataframe(wdf, hide_index=True, use_container_width=True)
                with tab_k:
                    corr = sa.get("corr")
                    if corr is not None and not corr.empty:
                        st.caption("Günlük getiri korelasyonu — düşük değerler daha iyi çeşitlendirme anlamına gelir.")
                        st.dataframe(corr.round(2), use_container_width=True)
                        try:
                            st.dataframe(
                                corr.style.background_gradient(cmap="RdYlGn_r", vmin=-1, vmax=1),
                                use_container_width=True,
                            )
                        except Exception:
                            pass
                    else:
                        st.info("Korelasyon için en az 2 fon ve yeterli ortak fiyat serisi gerekir.")
                with tab_d:
                    bd = sa.get("birlesik_dag") or {}
                    kd = sa.get("kategori_dag") or {}
                    cda, cdb = st.columns(2)
                    with cda:
                        st.markdown("**Birleşik varlık kırılımı**")
                        if bd:
                            labels = {
                                "hisse": "Hisse", "tahvil": "Tahvil", "repo": "Repo",
                                "mevduat": "Mevduat", "altin": "Altın", "doviz": "Döviz",
                                "fon": "Fon", "diger": "Diğer",
                            }
                            ser = pd.Series({labels.get(k, k): v for k, v in bd.items()})
                            st.bar_chart(ser)
                        else:
                            st.caption("Dağılım verisi yok")
                    with cdb:
                        st.markdown("**Kategori ağırlıkları**")
                        if kd:
                            st.bar_chart(pd.Series(kd))
                        else:
                            st.caption("Kategori verisi yok")
                with tab_g:
                    pret = sa.get("port_ret")
                    if pret is not None and len(pret) > 5:
                        cum = (1 + pret).cumprod()
                        st.markdown("**Birleşik portföy kümülatif getiri (normalize)**")
                        st.line_chart(cum.rename("Portföy"))
                        # Drawdown
                        peak = cum.cummax()
                        dd = (cum / peak - 1.0) * 100
                        st.markdown("**Düşüş (drawdown) %**")
                        st.area_chart(dd.rename("DD %"))
                    else:
                        st.info("Grafik için yeterli birleşik seri yok.")

                download_excel_button(
                    sa["df"],
                    f"tefas_sepet_derin_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={
                        "mod": "portfoy_sepet_derin",
                        "weights": sa.get("weights"),
                        "port_metrikler": pm,
                        "diversify_score": sa.get("diversify_score"),
                        "birlesik_dag": sa.get("birlesik_dag"),
                        "kategori_dag": sa.get("kategori_dag"),
                    },
                    key="xl_sepet",
                )
                if st.button("TEFAS portföy motoru (ağırlıklı)", key="sepet_equal_engine"):
                    w = sa.get("weights") or {}
                    if w:
                        analiz = get_analiz()
                        with st.spinner("portfoy_analiz…"):
                            s = analiz.portfoy_analiz(w)
                        st.session_state.son_portfoy = s
                        mm = s.get("metrikler") or {}
                        st.json({k: mm.get(k) for k in (
                            "getiri_1y", "reel_1y", "sharpe", "max_drawdown", "volatilite", "vs_bist100"
                        )})

    elif mod == "Elle ağırlık":
        n = st.slider("Kaç fon", 2, 8, 3)
        weights = {}
        cols = st.columns(min(n, 4))
        for i in range(n):
            with cols[i % len(cols)]:
                kod = st.text_input(f"Fon {i+1} kodu", key=f"pk_{i}").strip().upper()
                pct = st.number_input(f"Ağırlık %", value=round(100 / n), min_value=1, max_value=100, key=f"pw_{i}")
                if kod:
                    weights[kod] = float(pct)
        if st.button("Portföyü analiz et", type="primary") and weights:
            analiz = get_analiz()
            with st.spinner("Hesaplanıyor…"):
                s = analiz.portfoy_analiz(weights)
            m = s.get("metrikler") or {}
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("1Y %", fmt_num(m.get("getiri_1y")))
            c2.metric("Reel 1Y", fmt_num(m.get("reel_1y")))
            c3.metric("Sharpe", fmt_num(m.get("sharpe"), 2))
            c4.metric("MaxDD %", fmt_num(m.get("max_drawdown")))
            st.markdown("**Ağırlıklar**")
            st.json({k: f"%{v*100:.0f}" for k, v in (s.get("weights") or {}).items()})
            dag = {k: v for k, v in (s.get("dagilim") or {}).items() if v and v > 0.5}
            if dag:
                st.markdown("**Birleşik dağılım**")
                st.json(dag)
            st.session_state.son_portfoy = s
            # Excel: ağırlık + metrik özeti
            rows = []
            for k, v in (s.get("weights") or {}).items():
                rows.append({"kod": k, "agirlik_%": round(v * 100, 2), **{mk: m.get(mk) for mk in (
                    "getiri_1y", "reel_1y", "sharpe", "max_drawdown", "volatilite", "vs_mevduat", "vs_bist100"
                )}})
            if rows:
                edf = pd.DataFrame(rows)
                download_excel_button(
                    edf,
                    f"tefas_portfoy_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={"mod": "elle_portfoy", "agirliklar": s.get("weights"), "metrikler": m},
                    key="xl_port_elle",
                )

    else:
        st.caption("Risk paketine göre tarama yapıp en yüksek skorlu fonlardan portföy kurar.")
        paket_sec = st.selectbox(
            "Risk paketi",
            list(core.RISK_PAKETLER.keys()),
            format_func=lambda k: core.RISK_PAKETLER[k]["ad"],
            index=1, key="oto_paket",
        )
        paket = core.RISK_PAKETLER[paket_sec]
        hedef = core.PROFILLER.get(paket.get("profil") or "2", core.PROFILLER["2"])[1]
        n = st.slider("Portföye kaç fon", 2, 10, 5)
        yontem = st.radio("Ağırlık", ["Puan-orantılı", "Eşit"], horizontal=True)
        max_w = st.slider("Tek fon max ağırlık", 0.1, 1.0, 0.5, 0.05)

        if st.button("Otomatik portföy oluştur", type="primary"):
            analiz = get_analiz()
            with st.spinner("En yüksek puanlı fonlar aranıyor…"):
                df = analiz.tara(
                    hedef,
                    min_sharpe=paket.get("min_sharpe"),
                    max_vol=paket.get("max_vol"),
                    max_dd=paket.get("max_dd"),
                    max_fon=n,
                    max_aday=max(40, n * 6),
                    max_risk=paket.get("max_risk"),
                )
            if df.empty:
                st.warning("Uygun fon bulunamadı.")
            else:
                kodlar = df["kod"].tolist()
                if yontem == "Eşit":
                    w = {k: 1.0 / len(kodlar) for k in kodlar}
                else:
                    skorlar = np.array([float(v) if v is not None else 0.0 for v in df["skor"].tolist()])
                    w_arr = core.optimize_score_proportional(skorlar, max_w=max_w)
                    w = {kodlar[i]: float(w_arr[i]) for i in range(len(kodlar))}
                show_df(df, ["skor", "kod", "ad", "getiri_1y", "sharpe", "volatilite", "max_drawdown", "risk_bayrak"], 300)
                st.markdown("**Ağırlıklar**")
                for k, v in sorted(w.items(), key=lambda x: -x[1]):
                    st.write(f"**{k}**: %{v*100:.1f}")
                s = analiz.portfoy_analiz(w)
                m = s.get("metrikler") or {}
                c1, c2, c3 = st.columns(3)
                c1.metric("1Y %", fmt_num(m.get("getiri_1y")))
                c2.metric("Sharpe", fmt_num(m.get("sharpe"), 2))
                c3.metric("MaxDD %", fmt_num(m.get("max_drawdown")))
                st.session_state.son_portfoy = s
                edf = df.copy()
                edf["portfoy_agirlik_%"] = edf["kod"].map(lambda k: round(w.get(k, 0) * 100, 2))
                download_excel_button(
                    edf,
                    f"tefas_oto_portfoy_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={"mod": "otomatik_portfoy", "risk_paketi": paket["ad"], "agirliklar": w, "metrikler": m},
                    key="xl_port_oto",
                )


# ═══════════════════════════════════════════════════════════════════════════
# OPTİMİZASYON
# ═══════════════════════════════════════════════════════════════════════════



elif sayfa == "💼 Portföyüm":
    # Fonaly tarzı takip + Excel yükleme (fon.xlsx Özet/Lotlar) + düzenleme


    if "sanal_portfoy" not in st.session_state:
        st.session_state.sanal_portfoy = {
            "ad": "Portföyüm",
            "nakit": 0.0,
            "pozisyonlar": {},
            "lotlar": [],
            "islemler": [],
            "kaynak": None,
        }
    sp = st.session_state.sanal_portfoy

    st.markdown(
        f'<div class="fy-hero"><h2>💼 {sp.get("ad") or "Portföyüm"}</h2>'
        f'<p>Excel (Özet / Lotlar) yükle · düzenle · canlı fiyat ile K/Z. Al/sat emri yok. Yatırım tavsiyesi değildir.</p></div>',
        unsafe_allow_html=True,
    )

    def _sp_fiyatlar(kodlar):
        out = {}
        try:
            analiz = get_analiz()
        except Exception:
            analiz = None
        for kod in kodlar:
            px = None
            if analiz is not None:
                try:
                    fd = analiz.api.fiyat(kod, periyod=1)
                    if fd is not None and not getattr(fd, "empty", True) and "fiyat" in fd.columns:
                        px = float(pd.to_numeric(fd["fiyat"], errors="coerce").dropna().iloc[-1])
                except Exception:
                    pass
            if px is None:
                try:
                    import yfinance as yf
                    h = yf.Ticker(f"{kod}.IS").history(period="5d")
                    if h is not None and not h.empty:
                        px = float(h["Close"].iloc[-1])
                except Exception:
                    pass
            out[kod] = px
        return out

    def _norm_col(c):
        s = str(c or "").strip().lower()
        for a, b in (("ı", "i"), ("ğ", "g"), ("ü", "u"), ("ş", "s"), ("ö", "o"), ("ç", "c")):
            s = s.replace(a, b)
        return s

    def _excel_to_portfoy(file_bytes: bytes) -> dict:
        """fon.xlsx: Özet + Lotlar sayfaları → pozisyonlar + lotlar."""
        from io import BytesIO
        xls = pd.ExcelFile(BytesIO(file_bytes))
        sheets = { _norm_col(n): n for n in xls.sheet_names }
        ozet_name = sheets.get("ozet") or sheets.get("özet") or list(xls.sheet_names)[0]
        ozet = pd.read_excel(xls, sheet_name=ozet_name)
        ozet.columns = [str(c).strip() for c in ozet.columns]
        colmap = {_norm_col(c): c for c in ozet.columns}

        def col(*names):
            for n in names:
                if _norm_col(n) in colmap:
                    return colmap[_norm_col(n)]
            return None

        c_kod = col("Fon kodu", "kod", "fon")
        c_ad = col("Fon adı", "ad", "unvan")
        c_adet = col("Adet", "adet", "pay")
        c_maliyet = col("Ortalama maliyet", "ortalama maliyet", "maliyet", "alis")
        c_tm = col("Toplam maliyet (TL)", "toplam maliyet")
        c_td = col("Toplam değer (TL)", "toplam deger", "toplam değer")
        if not c_kod:
            raise ValueError("Excel'de 'Fon kodu' kolonu bulunamadı (Özet).")

        poz = {}
        for _, r in ozet.iterrows():
            kod = str(r.get(c_kod) or "").strip().upper()
            if not kod or kod in ("NAN", "NONE"):
                continue
            adet = float(pd.to_numeric(r.get(c_adet), errors="coerce") or 0)
            if adet <= 0:
                continue
            mal = float(pd.to_numeric(r.get(c_maliyet), errors="coerce") or 0)
            if mal <= 0 and c_tm:
                tm = float(pd.to_numeric(r.get(c_tm), errors="coerce") or 0)
                if adet > 0 and tm > 0:
                    mal = tm / adet
            poz[kod] = {
                "adet": adet,
                "maliyet": mal,
                "ad": str(r.get(c_ad) or kod)[:80],
                "excel_deger": float(pd.to_numeric(r.get(c_td), errors="coerce") or 0) if c_td else None,
                "excel_maliyet_toplam": float(pd.to_numeric(r.get(c_tm), errors="coerce") or 0) if c_tm else None,
            }

        lotlar = []
        lot_name = sheets.get("lotlar") or sheets.get("lot")
        if lot_name:
            lot_df = pd.read_excel(xls, sheet_name=lot_name)
            lot_df.columns = [str(c).strip() for c in lot_df.columns]
            lc = {_norm_col(c): c for c in lot_df.columns}
            def lcol(*names):
                for n in names:
                    if _norm_col(n) in lc:
                        return lc[_norm_col(n)]
                return None
            lk = lcol("Fon kodu", "kod")
            la = lcol("Adet")
            lf = lcol("Alış fiyatı", "alis fiyati", "fiyat")
            lt = lcol("Gerçekleşme tarihi", "tarih")
            ln = lcol("Lot #", "lot")
            for _, r in lot_df.iterrows():
                kod = str(r.get(lk) or "").strip().upper() if lk else ""
                if not kod:
                    continue
                lotlar.append({
                    "kod": kod,
                    "lot": r.get(ln) if ln else None,
                    "tarih": str(r.get(lt) or "")[:16] if lt else "",
                    "adet": float(pd.to_numeric(r.get(la), errors="coerce") or 0) if la else 0,
                    "alis": float(pd.to_numeric(r.get(lf), errors="coerce") or 0) if lf else 0,
                })
        return {"pozisyonlar": poz, "lotlar": lotlar, "ozet_rows": len(poz)}

    # ── Excel yükle ──
    with st.expander("📥 Excel’den yükle (Özet / Lotlar)", expanded=not bool(sp.get("pozisyonlar"))):
        st.caption(
            "Beklenen format: **Özet** sayfası (Fon kodu, Adet, Ortalama maliyet, Toplam maliyet/değer…) "
            "ve isteğe bağlı **Lotlar**. Örnek: fon.xlsx"
        )
        up = st.file_uploader("fon.xlsx / portföy Excel", type=["xlsx", "xls"], key="fy_excel_up")
        if up is not None and st.button("Excel’i portföye aktar", type="primary", key="fy_excel_btn"):
            try:
                data = _excel_to_portfoy(up.getvalue())
                sp["pozisyonlar"] = data["pozisyonlar"]
                sp["lotlar"] = data.get("lotlar") or []
                sp["kaynak"] = up.name
                sp["nakit"] = 0.0
                st.session_state.sanal_portfoy = sp
                # sepete de yansıt
                sepet = [{"kod": k, "ad": v.get("ad") or k, "kaynak": "excel"} for k, v in data["pozisyonlar"].items()]
                st.session_state.portfoy_sepet = sepet
                st.session_state["_fy_analiz_run"] = True
                st.session_state["_fy_excel_just_loaded"] = True
                warns = []
                for k, v in data["pozisyonlar"].items():
                    if float(v.get("maliyet") or 0) <= 0:
                        warns.append(f"{k}: maliyet eksik")
                msg = f"**{data['ozet_rows']} fon** yüklendi"
                if sp.get("lotlar"):
                    msg += f" · **{len(sp['lotlar'])} lot**"
                msg += " · analiz paneli açılacak"
                st.success(msg)
                if warns:
                    st.warning("Excel uyarıları: " + "; ".join(warns[:10]))
                st.rerun()
            except Exception as e:
                st.error(f"Excel okunamadı: {e}")

    # ── Düzenleme ──
    with st.expander("✏️ Portföyü düzenle", expanded=False):
        poz0 = dict(sp.get("pozisyonlar") or {})
        if not poz0:
            st.info("Önce Excel yükleyin veya kod listesi ekleyin.")
        else:
            edit_rows = []
            for kod, p in poz0.items():
                edit_rows.append({
                    "kod": kod,
                    "ad": p.get("ad") or kod,
                    "adet": float(p.get("adet") or 0),
                    "maliyet": float(p.get("maliyet") or 0),
                })
            edf = pd.DataFrame(edit_rows)
            edited = st.data_editor(
                edf,
                num_rows="dynamic",
                use_container_width=True,
                key="fy_editor",
                column_config={
                    "kod": st.column_config.TextColumn("Fon kodu", required=True),
                    "ad": st.column_config.TextColumn("Ad"),
                    "adet": st.column_config.NumberColumn("Adet", min_value=0.0, format="%.4f"),
                    "maliyet": st.column_config.NumberColumn("Ort. maliyet", min_value=0.0, format="%.6f"),
                },
            )
            if st.button("Düzenlemeyi kaydet", type="primary", key="fy_edit_save"):
                poz2 = {}
                for _, r in edited.iterrows():
                    kod = str(r.get("kod") or "").strip().upper()
                    if not kod:
                        continue
                    adet = float(r.get("adet") or 0)
                    if adet <= 0:
                        continue
                    poz2[kod] = {
                        "adet": adet,
                        "maliyet": float(r.get("maliyet") or 0),
                        "ad": str(r.get("ad") or kod)[:80],
                    }
                sp["pozisyonlar"] = poz2
                st.session_state.sanal_portfoy = sp
                st.session_state.portfoy_sepet = [
                    {"kod": k, "ad": v.get("ad") or k, "kaynak": "manuel"} for k, v in poz2.items()
                ]
                st.success(f"{len(poz2)} pozisyon kaydedildi.")
                st.rerun()

        st.markdown("**Kod listesi ile ekle / birleştir**")
        kod_text = st.text_input("Kodlar (virgülle)", placeholder="GPG, NKP, TI1", key="fy_kod_extra")
        if st.button("Kodları ekle (eşit ağırlık, mevcut bozulmaz)", key="fy_kod_add") and kod_text.strip():
            import re as _re
            kodlar = [x.strip().upper() for x in _re.split(r"[,;\s]+", kod_text) if x.strip()]
            pxm = _sp_fiyatlar(kodlar)
            poz2 = dict(sp.get("pozisyonlar") or {})
            for k in kodlar:
                if k in poz2:
                    continue
                px = pxm.get(k) or 1.0
                poz2[k] = {"adet": 1000.0, "maliyet": px, "ad": k}
            sp["pozisyonlar"] = poz2
            st.session_state.sanal_portfoy = sp
            st.rerun()

        if st.button("Portföyü temizle", key="fy_clear"):
            st.session_state.sanal_portfoy = {
                "ad": "Portföyüm", "nakit": 0.0, "pozisyonlar": {}, "lotlar": [], "islemler": [], "kaynak": None,
            }
            st.rerun()

    poz = dict(sp.get("pozisyonlar") or {})
    kodlar = list(poz.keys())
    fiyatlar = _sp_fiyatlar(kodlar) if kodlar else {}

    maliyet_toplam = 0.0
    deger_toplam = 0.0
    for kod, p in poz.items():
        adet = float(p.get("adet") or 0)
        maliyet = float(p.get("maliyet") or 0)
        maliyet_toplam += adet * maliyet
        px = fiyatlar.get(kod)
        if px is None and p.get("excel_deger") and adet:
            # canlı fiyat yoksa excel değerinden birim
            try:
                px = float(p["excel_deger"]) / adet
            except Exception:
                px = maliyet
        deger_toplam += adet * (px if px is not None else maliyet)
    portfoy_deger = deger_toplam
    acik_kz = deger_toplam - maliyet_toplam
    acik_kz_pct = (acik_kz / maliyet_toplam * 100.0) if maliyet_toplam > 0 else 0.0

    left, right = st.columns([1.05, 2.45])
    with left:
        kz_cls = "fy-pos" if acik_kz >= 0 else "fy-neg"
        st.markdown(
            f'<div class="fy-card"><div class="lbl">Portföy değeri</div>'
            f'<div class="val">₺{tr_num(portfoy_deger, 2)}</div>'
            f'<div class="sub">{sp.get("kaynak") or "manuel"}</div></div>',
            unsafe_allow_html=True,
        )
        st.markdown(
            f'<div class="fy-card"><div class="lbl">Toplam maliyet</div>'
            f'<div class="val">₺{tr_num(maliyet_toplam, 2)}</div></div>',
            unsafe_allow_html=True,
        )
        st.markdown(
            f'<div class="fy-card"><div class="lbl">Toplam kâr/zarar</div>'
            f'<div class="val {kz_cls}">₺{tr_num(acik_kz, 2)}</div>'
            f'<div class="sub {kz_cls}">%{acik_kz_pct:+.2f}</div></div>',
            unsafe_allow_html=True,
        )
        st.markdown(
            f'<div class="fy-card"><div class="lbl">Fon sayısı</div>'
            f'<div class="val">{len(poz)}</div></div>',
            unsafe_allow_html=True,
        )
        if st.button("Fiyatları yenile", key="fy_yenile", use_container_width=True):
            st.rerun()

    with right:
        tab_g, tab_l, tab_h, tab_r, tab_d = st.tabs([
            "Genel", "Lotlar", "Getiri", "Risk", "Dağılım",
        ])
        with tab_g:
            if not poz:
                empty_state("Portföy boş", "Excel (Özet / Lotlar) yükleyin veya kod listesi ekleyin.", "Örnek dosya: fon.xlsx")
            else:
                rows = []
                for kod, p in poz.items():
                    adet = float(p.get("adet") or 0)
                    mal = float(p.get("maliyet") or 0)
                    px = fiyatlar.get(kod)
                    if px is None and p.get("excel_deger") and adet:
                        try:
                            px = float(p["excel_deger"]) / adet
                        except Exception:
                            px = mal
                    deg = (px or mal) * adet
                    kz = (px - mal) * adet if px is not None else 0.0
                    kz_p = ((px / mal) - 1) * 100 if px and mal else 0.0
                    rows.append({
                        "kod": kod,
                        "ad": (p.get("ad") or "")[:40],
                        "adet": round(adet, 4),
                        "maliyet": round(mal, 6),
                        "fiyat": round(px, 6) if px is not None else None,
                        "değer": round(deg, 2),
                        "K/Z": round(kz, 2),
                        "K/Z %": round(kz_p, 2),
                        "ağırlık %": round(deg / deger_toplam * 100, 2) if deger_toplam else 0,
                    })
                pdf = pd.DataFrame(rows).sort_values("değer", ascending=False)
                pdf_show = pdf.copy()
                for col in ("adet", "maliyet", "fiyat", "değer", "K/Z", "K/Z %", "ağırlık %"):
                    if col in pdf_show.columns:
                        pdf_show[col] = pdf_show[col].map(
                            lambda x: tr_num(x, 2 if col != "adet" else 4) if x is not None else "—"
                        )
                st.dataframe(pdf_show, hide_index=True, use_container_width=True, height=380)
                download_excel_button(
                    pdf,
                    f"portfoyum_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={"mod": "portfoyum", "kaynak": sp.get("kaynak")},
                    key="xl_fy_out",
                )

        with tab_l:
            lots = list(sp.get("lotlar") or [])
            if not lots:
                st.caption("Lot satırı yok (Excel Lotlar sayfasından gelir).")
            else:
                st.dataframe(pd.DataFrame(lots), hide_index=True, use_container_width=True, height=360)

        with tab_h:
            st.markdown("#### Aylık getiri & performans")
            if not poz:
                st.info("Pozisyon yok.")
            else:
                # Maliyet bazlı anlık K/Z
                heat = []
                for kod, p in poz.items():
                    mal = float(p.get("maliyet") or 0)
                    px = fiyatlar.get(kod) or mal
                    pct = ((px / mal) - 1) * 100 if mal else 0.0
                    heat.append({"kod": kod, "maliyet_getiri_%": round(pct, 2)})
                hdf = pd.DataFrame(heat).sort_values("maliyet_getiri_%", ascending=False)
                st.caption("Alış maliyetine göre toplam getiri %")
                try:
                    import plotly.express as px
                    fig = px.bar(
                        hdf, x="kod", y="maliyet_getiri_%", color="maliyet_getiri_%",
                        color_continuous_scale=["#ff6b6b", "#888", "#3ddc97"],
                        title="Pozisyon getirileri (maliyet bazlı) %",
                    )
                    fig.update_layout(
                        template="plotly_dark", height=320,
                        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                    )
                    st.plotly_chart(fig, use_container_width=True)
                except Exception:
                    st.dataframe(hdf, hide_index=True)

                # Aylık portföy getirisi (ağırlıklı fiyat serisi)
                if st.button("Aylık getiri serisini hesapla", key="fy_aylik_btn"):
                    st.session_state["_fy_aylik_calc"] = True
                if st.session_state.get("_fy_aylik_calc"):
                    with st.spinner("TEFAS fiyat serileri…"):
                        try:
                            analiz = get_analiz()
                        except Exception:
                            analiz = None
                        series = {}
                        wmap = {}
                        for kod, p in poz.items():
                            adet = float(p.get("adet") or 0)
                            px0 = fiyatlar.get(kod) or float(p.get("maliyet") or 0)
                            wmap[kod] = adet * px0
                            fd = None
                            if analiz is not None:
                                try:
                                    fd = analiz.api.fiyat(kod, periyod=12)
                                except Exception:
                                    fd = None
                            if fd is None or getattr(fd, "empty", True):
                                continue
                            try:
                                tcol = "tarih" if "tarih" in fd.columns else fd.columns[0]
                                fcol = "fiyat" if "fiyat" in fd.columns else [c for c in fd.columns if c != tcol][0]
                                s = fd[[tcol, fcol]].copy()
                                s[tcol] = pd.to_datetime(s[tcol])
                                s = s.set_index(tcol)[fcol].astype(float).sort_index()
                                series[kod] = s
                            except Exception:
                                continue
                        if len(series) < 1:
                            st.warning("Yeterli fiyat serisi yok.")
                        else:
                            tw = sum(wmap.get(k, 0) for k in series) or 1.0
                            # hizala
                            all_idx = None
                            for s in series.values():
                                all_idx = s.index if all_idx is None else all_idx.union(s.index)
                            all_idx = all_idx.sort_values()
                            port = None
                            for kod, s in series.items():
                                w = (wmap.get(kod, 0) / tw)
                                aligned = s.reindex(all_idx).ffill().bfill()
                                part = aligned * w
                                port = part if port is None else port + part
                            # normalize to 100
                            port = port.dropna()
                            if len(port) < 5:
                                st.warning("Seri kısa.")
                            else:
                                port_nav = port / float(port.iloc[0]) * 100.0
                                # aylık getiri
                                monthly = port_nav.resample("ME").last().dropna()
                                if len(monthly) < 2:
                                    monthly = port_nav.resample("M").last().dropna()
                                mret = monthly.pct_change().dropna() * 100.0
                                # Aylık tablo: getiri + NAV + kümülatif
                                mnav = monthly.reindex(mret.index)
                                cum = (1 + mret / 100.0).cumprod()
                                mdf = pd.DataFrame({
                                    "Ay": [d.strftime("%Y-%m") for d in mret.index],
                                    "Aylık getiri %": [float(x) for x in mret.values],
                                    "NAV (100=başlangıç)": [
                                        float(mnav.loc[i]) if i in mnav.index else None
                                        for i in mret.index
                                    ],
                                    "Kümülatif getiri %": [
                                        float((cum.loc[i] - 1) * 100) for i in mret.index
                                    ],
                                })
                                # Gösterim: TR ondalık ayırıcı
                                mdf_show = mdf.copy()
                                for col in ("Aylık getiri %", "NAV (100=başlangıç)", "Kümülatif getiri %"):
                                    mdf_show[col] = mdf_show[col].map(
                                        lambda x: tr_num(x, 2) if x is not None else "—"
                                    )
                                st.markdown("**Aylık portföy getirisi (tablo)**")
                                st.dataframe(mdf_show, hide_index=True, use_container_width=True)
                                # grafik için ham sayı
                                mdf_plot = mdf.rename(columns={"Ay": "ay", "Aylık getiri %": "getiri_%"})
                                try:
                                    import plotly.express as px
                                    fig2 = px.bar(
                                        mdf_plot, x="ay", y="getiri_%", color="getiri_%",
                                        color_continuous_scale=["#ff6b6b", "#888", "#3ddc97"],
                                        title="Aylık getiri %",
                                    )
                                    fig2.update_layout(
                                        template="plotly_dark", height=340,
                                        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                                    )
                                    st.plotly_chart(fig2, use_container_width=True)
                                    fig3 = px.line(
                                        x=port_nav.index, y=port_nav.values,
                                        labels={"x": "Tarih", "y": "NAV (100=başlangıç)"},
                                        title="Portföy NAV endeksi",
                                    )
                                    fig3.update_layout(
                                        template="plotly_dark", height=320,
                                        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                                    )
                                    st.plotly_chart(fig3, use_container_width=True)
                                except Exception as e:
                                    st.caption(str(e))
                                # özet
                                pos = int((mret > 0).sum())
                                neg = int((mret < 0).sum())
                                c1, c2, c3, c4 = st.columns(4)
                                c1.metric("Ort. aylık %", tr_num(float(mret.mean()), 2))
                                c2.metric("En iyi ay %", tr_num(float(mret.max()), 2))
                                c3.metric("En kötü ay %", tr_num(float(mret.min()), 2))
                                c4.metric("Pozitif ay", f"{pos}/{pos+neg}")

                                st.session_state["_fy_aylik_tablo"] = mdf
                                st.session_state["_fy_port_nav"] = port_nav
                                st.session_state["_fy_mret"] = mret
                                st.session_state["_fy_series"] = series
                                st.session_state["_fy_wmap"] = wmap

                                # ── Fonaly tarzı: periyot + aylık ısı ızgarası + günlük istatistik ──
                                st.markdown("---")
                                st.markdown("##### Getiri haritası (aylık ızgara)")
                                periyot = st.radio(
                                    "Periyot",
                                    ["Günlük", "Haftalık", "Aylık", "Çeyreklik"],
                                    index=2,
                                    horizontal=True,
                                    key="fy_periyot_harita",
                                )
                                pret_d = port_nav.pct_change().dropna() * 100.0
                                if periyot == "Günlük":
                                    series_r = pret_d
                                elif periyot == "Haftalık":
                                    series_r = port_nav.resample("W").last().pct_change().dropna() * 100.0
                                elif periyot == "Çeyreklik":
                                    series_r = port_nav.resample("QE").last().pct_change().dropna() * 100.0
                                else:
                                    series_r = mret

                                # Aylık ısı tablosu (yıl × ay)
                                ay_ad = ["Oca", "Şub", "Mar", "Nis", "May", "Haz",
                                         "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]
                                heat_rows = []
                                if len(mret) > 0:
                                    for y in sorted(set(mret.index.year)):
                                        row = {"YIL": str(y)}
                                        year_rets = []
                                        for m in range(1, 13):
                                            cell = mret[(mret.index.year == y) & (mret.index.month == m)]
                                            if len(cell):
                                                v = float(cell.iloc[-1])
                                                row[ay_ad[m - 1]] = v
                                                year_rets.append(v)
                                            else:
                                                row[ay_ad[m - 1]] = None
                                        row["Toplam"] = float((pd.Series([1 + x/100 for x in year_rets]).prod() - 1) * 100) if year_rets else None
                                        heat_rows.append(row)
                                if heat_rows:
                                    hgrid = pd.DataFrame(heat_rows)
                                    # HTML renkli ızgara
                                    def _cell_html(v):
                                        if v is None or (isinstance(v, float) and (v != v)):
                                            return '<td class="hm-empty">—</td>'
                                        cls = "hm-pos" if v >= 0 else "hm-neg"
                                        return f'<td class="{cls}">{tr_num(v, 2)}%</td>'
                                    html = [
                                        '<style>',
                                        '.hm-wrap{overflow-x:auto;margin:8px 0 16px;}',
                                        '.hm{border-collapse:collapse;width:100%;font-size:0.85rem;}',
                                        '.hm th{color:#8fa3b8;padding:6px 8px;font-weight:600;}',
                                        '.hm td{text-align:center;padding:8px 6px;border-radius:8px;margin:2px;}',
                                        '.hm-pos{background:#1a3d32;color:#3ddc97;font-weight:600;}',
                                        '.hm-neg{background:#3d1a1a;color:#ff6b6b;font-weight:600;}',
                                        '.hm-empty{color:#556;}',
                                        '.hm-y{color:#c5d0dc;font-weight:700;text-align:left;}',
                                        '</style><div class="hm-wrap"><table class="hm"><tr><th>YIL</th>',
                                    ]
                                    for a in ay_ad:
                                        html.append(f"<th>{a}</th>")
                                    html.append("<th>Toplam</th></tr>")
                                    for _, rr in hgrid.iterrows():
                                        html.append(f'<tr><td class="hm-y">{rr["YIL"]}</td>')
                                        for a in ay_ad:
                                            html.append(_cell_html(rr.get(a)))
                                        html.append(_cell_html(rr.get("Toplam")))
                                        html.append("</tr>")
                                    html.append("</table></div>")
                                    st.markdown("".join(html), unsafe_allow_html=True)

                                # Günlük getiri istatistik kartları
                                st.markdown("##### Günlük getiri özeti")
                                dret = pret_d
                                if len(dret) >= 2:
                                    pos_d = dret[dret > 0]
                                    neg_d = dret[dret < 0]
                                    n_d = len(dret)
                                    n_pos_d = len(pos_d)
                                    n_neg_d = len(neg_d)
                                    pct_pos = n_pos_d / n_d * 100 if n_d else 0
                                    pct_neg = n_neg_d / n_d * 100 if n_d else 0
                                    g1, g2, g3, g4 = st.columns(4)
                                    g1.metric("Ort. pozitif gün", f"%{tr_num(float(pos_d.mean()) if len(pos_d) else 0, 2)}")
                                    g2.metric("En iyi gün", f"%{tr_num(float(dret.max()), 2)}")
                                    g3.metric("Ort. negatif gün", f"%{tr_num(float(neg_d.mean()) if len(neg_d) else 0, 2)}")
                                    g4.metric("En kötü gün", f"%{tr_num(float(dret.min()), 2)}")
                                    st.markdown(
                                        f"**Günlük getiri dağılımı** · {n_d} gün · "
                                        f'<span style="color:#3ddc97">● %{tr_num(pct_pos,1)} Pozitif — {n_pos_d}</span> · '
                                        f'<span style="color:#ff6b6b">● %{tr_num(pct_neg,1)} Negatif — {n_neg_d}</span>',
                                        unsafe_allow_html=True,
                                    )
                                    # progress bars
                                    st.progress(min(1.0, pct_pos / 100), text=f"Pozitif gün %{tr_num(pct_pos,1)}")
                                    st.progress(min(1.0, pct_neg / 100), text=f"Negatif gün %{tr_num(pct_neg,1)}")

                                    # Periyot seçimine göre özet tablo
                                    if len(series_r) > 0:
                                        st.markdown(f"**{periyot} getiri tablosu**")
                                        if periyot == "Aylık":
                                            pass  # already shown
                                        else:
                                            sdf = pd.DataFrame({
                                                "Dönem": [
                                                    (i.strftime("%Y-%m-%d") if hasattr(i, "strftime") else str(i))
                                                    for i in series_r.index
                                                ],
                                                "Getiri %": [tr_num(float(x), 2) for x in series_r.values],
                                            })
                                            st.dataframe(sdf.tail(24), hide_index=True, use_container_width=True)

        with tab_r:
            st.markdown("#### Risk analizi & değerlendirme")
            if len(poz) < 1:
                st.info("Risk için pozisyon gerekli.")
            else:
                weights = []
                labels = []
                for kod, p in poz.items():
                    px = fiyatlar.get(kod) or float(p.get("maliyet") or 0)
                    weights.append(px * float(p.get("adet") or 0))
                    labels.append(kod)
                ssum = sum(weights) or 1
                w = [x / ssum for x in weights]
                herf = sum(x * x for x in w)
                max_w = max(w) if w else 0
                max_kod = labels[int(w.index(max_w))] if w else "—"
                n_eff = (1.0 / herf) if herf > 0 else 0

                # Seri tabanlı risk (varsa hesaplanmış veya şimdi çek)
                vol_y = sharpe = maxdd = cvar = corr_mean = None
                port_daily = None
                series = st.session_state.get("_fy_series") or {}
                wmap = st.session_state.get("_fy_wmap") or {}
                if not series and st.button("Risk metriklerini hesapla (fiyat serisi)", key="fy_risk_calc"):
                    st.session_state["_fy_aylik_calc"] = True  # aynı seriyi kullan
                    st.info("Önce **Getiri** sekmesinde «Aylık getiri serisini hesapla»ya basın veya aşağıyı tekrar deneyin.")
                if series:
                    try:
                        tw = sum(wmap.get(k, 0) for k in series) or 1.0
                        all_idx = None
                        for s in series.values():
                            all_idx = s.index if all_idx is None else all_idx.union(s.index)
                        all_idx = all_idx.sort_values()
                        port = None
                        rets_list = []
                        for kod, s in series.items():
                            ww = (wmap.get(kod, 0) / tw)
                            aligned = s.reindex(all_idx).ffill().bfill()
                            part = aligned * ww
                            port = part if port is None else port + part
                            r = aligned.pct_change().dropna()
                            rets_list.append(r.rename(kod))
                        port = port.dropna()
                        pret = port.pct_change().dropna()
                        port_daily = pret
                        if len(pret) > 5:
                            vol_y = float(pret.std() * (252 ** 0.5) * 100)
                            mean_d = float(pret.mean())
                            rf_d = 0.0
                            try:
                                rf_d = float(get_analiz().rf or 0) / 252.0
                            except Exception:
                                pass
                            if pret.std() > 0:
                                sharpe = float(((mean_d - rf_d) / pret.std()) * (252 ** 0.5))
                            cum = (1 + pret).cumprod()
                            peak = cum.cummax()
                            dd = (cum / peak - 1)
                            maxdd = float(dd.min() * 100)
                            # CVaR %5
                            q = float(pret.quantile(0.05))
                            tail = pret[pret <= q]
                            cvar = float(tail.mean() * 100) if len(tail) else float(q * 100)
                        if len(rets_list) >= 2:
                            rdf = pd.concat(rets_list, axis=1).dropna(how="all").fillna(0)
                            if rdf.shape[1] >= 2:
                                cm = rdf.corr()
                                # ortalama çift korelasyon
                                vals = []
                                cols = list(cm.columns)
                                for i in range(len(cols)):
                                    for j in range(i + 1, len(cols)):
                                        vals.append(float(cm.iloc[i, j]))
                                corr_mean = float(sum(vals) / len(vals)) if vals else None
                    except Exception as e:
                        st.caption(f"Seri risk: {e}")

                r1, r2, r3, r4 = st.columns(4)
                r1.metric("Max ağırlık", f"%{max_w*100:.1f}", max_kod)
                r2.metric("Herfindahl", f"{herf:.3f}")
                r3.metric("Efektif fon sayısı", f"{n_eff:.1f}")
                r4.metric("Fon adedi", len(poz))

                r5, r6, r7, r8 = st.columns(4)
                r5.metric("Yıllık vol %", f"{vol_y:.1f}" if vol_y is not None else "—")
                r6.metric("Sharpe", f"{sharpe:.2f}" if sharpe is not None else "—")
                r7.metric("Max DD %", f"{maxdd:.1f}" if maxdd is not None else "—")
                r8.metric("CVaR %5 (gün)", f"{cvar:.2f}" if cvar is not None else "—")

                if corr_mean is not None:
                    st.metric("Ort. çift korelasyon", f"{corr_mean:.2f}")

                # Değerlendirme metni
                notes = []
                score = 50
                if max_w > 0.25:
                    notes.append(f"⚠️ Tek fon ağırlığı yüksek (**{max_kod} %{max_w*100:.0f}**) — yoğunlaşma riski.")
                    score -= 12
                elif max_w <= 0.15:
                    notes.append("✓ Ağırlıklar dağınık; tek fon baskın değil.")
                    score += 8
                if n_eff < 4:
                    notes.append(f"⚠️ Efektif fon sayısı düşük (**{n_eff:.1f}**) — çeşitlendirme zayıf.")
                    score -= 10
                elif n_eff >= 8:
                    notes.append(f"✓ Efektif çeşitlilik iyi (**{n_eff:.1f}**).")
                    score += 8
                if vol_y is not None:
                    if vol_y >= 25:
                        notes.append(f"⚠️ Yıllık volatilite yüksek (**%{vol_y:.0f}**).")
                        score -= 10
                    elif vol_y <= 12:
                        notes.append(f"✓ Volatilite ılımlı (**%{vol_y:.0f}**).")
                        score += 6
                if maxdd is not None:
                    if maxdd <= -25:
                        notes.append(f"⚠️ Maksimum düşüş derin (**%{maxdd:.0f}**).")
                        score -= 10
                    elif maxdd > -12:
                        notes.append(f"✓ Drawdown sığ (**%{maxdd:.0f}**).")
                        score += 6
                if sharpe is not None:
                    if sharpe >= 1:
                        notes.append(f"✓ Sharpe **{sharpe:.2f}** — risk-ayarlı getiri iyi.")
                        score += 10
                    elif sharpe < 0:
                        notes.append(f"⚠️ Sharpe **{sharpe:.2f}** negatif.")
                        score -= 8
                if corr_mean is not None:
                    if corr_mean >= 0.7:
                        notes.append(f"⚠️ Fonlar birbirine çok bağlı (ort. corr **{corr_mean:.2f}**) — krizde birlikte düşebilir.")
                        score -= 8
                    elif corr_mean <= 0.35:
                        notes.append(f"✓ Düşük korelasyon (**{corr_mean:.2f}**) çeşitlendirmeyi destekler.")
                        score += 8
                score = max(0, min(100, score))
                if score >= 70:
                    seviye = "İyi risk profili"
                elif score >= 50:
                    seviye = "Orta — iyileştirilebilir"
                else:
                    seviye = "Yüksek risk / zayıf çeşitlendirme"

                st.markdown(f"##### Risk skoru: **{score}/100** — {seviye}")
                for n in notes:
                    st.markdown(f"- {n}")
                st.caption(
                    "Skor kural tabanlıdır; yatırım tavsiyesi değildir. "
                    "Vol/Sharpe/MaxDD için **Getiri → Aylık getiri serisini hesapla** ile fiyat serisi gerekir."
                )

                # Ağırlık tablosu
                wdf = pd.DataFrame({
                    "kod": labels,
                    "ağırlık_%": [round(x * 100, 2) for x in w],
                }).sort_values("ağırlık_%", ascending=False)
                st.dataframe(wdf, hide_index=True, use_container_width=True)


        with tab_d:
            st.markdown("#### Portföy dağılımı · kurum · kategori · risk · başarı")
            if not poz or deger_toplam <= 0:
                st.info("Dağılım için pozisyon gerekli.")
            else:
                def _kategori(ad: str, kod: str) -> str:
                    t = f"{ad} {kod}".upper()
                    if any(x in t for x in ("PARA PIYASASI", "PARA PİYASASI", "LIKIT", "LİKİT")):
                        return "Para piyasası"
                    if "BORÇLANMA" in t or "TAHVIL" in t or "BONO" in t:
                        return "Borçlanma"
                    if "KATILIM" in t:
                        return "Katılım"
                    if "HISSE" in t or "HİSSE" in t:
                        return "Hisse senedi"
                    if "FON SEPETI" in t or "FON SEPETİ" in t or "SEPET" in t:
                        return "Fon sepeti"
                    if "EMTIA" in t or "EMTİA" in t or "ALTIN" in t or "GUMUS" in t:
                        return "Emtia"
                    if "DEGISKEN" in t or "DEĞİŞKEN" in t:
                        return "Değişken"
                    if "SERBEST" in t:
                        return "Serbest"
                    if "KAR PAYI" in t or "KÂR PAYI" in t:
                        return "Kâr payı / temettü"
                    if "TEKNOLOJI" in t or "TEKNOLOJİ" in t or "INOVA" in t:
                        return "Teknoloji"
                    if "ISTATISTIKSEL" in t or "İSTATİSTİKSEL" in t or "ARBITRAJ" in t:
                        return "Arbitraj / kantitatif"
                    return "Diğer"

                def _kurum(ad: str) -> str:
                    t = (ad or "").upper()
                    for key in (
                        "İŞ PORTFÖY", "IS PORTFOY", "INVEO", "NEO PORTFÖY", "NEO PORTFOY",
                        "TEB PORTFÖY", "STATECH", "FİBA", "FIBA", "ASTRA", "ATLAS",
                        "ATA PORTFÖY", "KUVEYT", "HEDEF PORTFÖY", "NUROL", "TERA",
                        "AZİMUT", "AZIMUT", "AK PORTFÖY", "GARANTİ", "YAPI KREDİ",
                        "ZİRAAT", "HALK", "QNB", "HSBC", "OSMANLI",
                    ):
                        if key in t:
                            return key.title().replace("Portfoy", "Portföy")
                    # ilk iki kelime
                    parts = (ad or "").split()
                    if len(parts) >= 2:
                        return f"{parts[0]} {parts[1]}"
                    return parts[0] if parts else "Bilinmiyor"

                def _risk_band(kat: str, getiri_pct: float | None) -> str:
                    if kat in ("Para piyasası",):
                        return "Düşük"
                    if kat in ("Borçlanma", "Katılım") and "Hisse" not in kat:
                        return "Düşük-Orta"
                    if kat in ("Hisse senedi", "Serbest", "Arbitraj / kantitatif", "Teknoloji"):
                        return "Yüksek"
                    if kat in ("Değişken", "Fon sepeti", "Emtia", "Kâr payı / temettü"):
                        return "Orta"
                    return "Orta"

                rows_d = []
                for kod, p in poz.items():
                    adet = float(p.get("adet") or 0)
                    mal = float(p.get("maliyet") or 0)
                    px = fiyatlar.get(kod)
                    if px is None and p.get("excel_deger") and adet:
                        try:
                            px = float(p["excel_deger"]) / adet
                        except Exception:
                            px = mal
                    px = px if px is not None else mal
                    deg = adet * px
                    kz_p = ((px / mal) - 1) * 100 if mal else 0.0
                    ad = p.get("ad") or kod
                    kat = _kategori(ad, kod)
                    rows_d.append({
                        "kod": kod,
                        "ad": ad[:50],
                        "değer": deg,
                        "ağırlık_%": deg / deger_toplam * 100,
                        "K/Z_%": kz_p,
                        "kategori": kat,
                        "kurum": _kurum(ad),
                        "risk": _risk_band(kat, kz_p),
                        "basari": "Pozitif" if kz_p >= 0 else "Negatif",
                    })
                ddf = pd.DataFrame(rows_d)

                # Başarı oranı
                n_pos = int((ddf["K/Z_%"] >= 0).sum())
                n_neg = int((ddf["K/Z_%"] < 0).sum())
                n_all = len(ddf)
                basari_adet = (n_pos / n_all * 100) if n_all else 0
                # değer ağırlıklı başarı
                pos_val = float(ddf.loc[ddf["K/Z_%"] >= 0, "değer"].sum())
                basari_deger = (pos_val / deger_toplam * 100) if deger_toplam else 0
                # kâr eden pozisyonların toplam K/Z payı
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Başarı oranı (adet)", f"%{basari_adet:.0f}", f"{n_pos}/{n_all} pozitif")
                c2.metric("Başarı (değer ağırlıklı)", f"%{basari_deger:.0f}")
                c3.metric("Ort. K/Z %", f"{float(ddf['K/Z_%'].mean()):+.1f}")
                c4.metric("Medyan K/Z %", f"{float(ddf['K/Z_%'].median()):+.1f}")

                try:
                    import plotly.express as px
                    # Fon bazlı
                    col_f, col_p = st.columns([1.1, 1.2])
                    with col_f:
                        st.markdown("**Fon dağılımı**")
                        st.caption("Portföyünüzdeki fonların toplam değer içindeki ağırlığı.")
                        fig0 = px.pie(ddf, names="kod", values="değer", hole=0.55, title=None)
                        fig0.update_layout(
                            template="plotly_dark", height=300,
                            paper_bgcolor="rgba(0,0,0,0)", margin=dict(t=10, b=10, l=10, r=10),
                            showlegend=False,
                        )
                        st.plotly_chart(fig0, use_container_width=True)
                        # sıralı liste (Fonaly tarzı)
                        ranked = ddf.sort_values("ağırlık_%", ascending=False)
                        list_html = ['<div style="font-size:0.92rem;line-height:1.85">']
                        for _, rr in ranked.iterrows():
                            list_html.append(
                                f'<div style="display:flex;justify-content:space-between;border-bottom:1px solid #1e2a3a;padding:4px 0">'
                                f'<span><b>{rr["kod"]}</b></span>'
                                f'<span>{tr_num(rr["ağırlık_%"], 1)}%</span></div>'
                            )
                        list_html.append("</div>")
                        st.markdown("".join(list_html), unsafe_allow_html=True)
                    with col_p:
                        st.markdown("**Portföy / kategori dağılımı**")
                        st.caption(f"Toplam: ₺{tr_num(deger_toplam, 2)} · Portföy %100")
                        cdf = ddf.groupby("kategori", as_index=False)["değer"].sum().sort_values("değer", ascending=False)
                        cdf["pay_%"] = cdf["değer"] / cdf["değer"].sum() * 100
                        for _, rr in cdf.iterrows():
                            pay = float(rr["pay_%"])
                            st.markdown(
                                f"**{rr['kategori']}** · %{tr_num(pay, 2)} · ₺{tr_num(rr['değer'], 2)}"
                            )
                            st.progress(min(1.0, max(0.0, pay / 100.0)))

                    st.markdown("##### Aracı kurum / portföy şirketi payı")
                    kdf = ddf.groupby("kurum", as_index=False)["değer"].sum().sort_values("değer", ascending=False)
                    kdf["pay_%"] = kdf["değer"] / kdf["değer"].sum() * 100
                    fig1 = px.pie(kdf, names="kurum", values="değer", hole=0.4, title="Kurum dağılımı")
                    fig1.update_layout(template="plotly_dark", height=360, paper_bgcolor="rgba(0,0,0,0)")
                    st.plotly_chart(fig1, use_container_width=True)
                    kshow = kdf.copy()
                    kshow["değer"] = kshow["değer"].map(lambda x: tr_num(x, 2))
                    kshow["pay_%"] = kshow["pay_%"].map(lambda x: tr_num(x, 2))
                    st.dataframe(kshow, hide_index=True, use_container_width=True)

                    # Risk bandı
                    st.markdown("##### Risk bandı dağılımı")
                    rdf = ddf.groupby("risk", as_index=False)["değer"].sum()
                    order = ["Düşük", "Düşük-Orta", "Orta", "Yüksek"]
                    rdf["risk"] = pd.Categorical(rdf["risk"], categories=order, ordered=True)
                    rdf = rdf.sort_values("risk")
                    fig3 = px.pie(rdf, names="risk", values="değer", hole=0.4, title="Risk bandı (kategori tahmini)")
                    fig3.update_layout(template="plotly_dark", height=340, paper_bgcolor="rgba(0,0,0,0)")
                    st.plotly_chart(fig3, use_container_width=True)

                    # Başarı: pozitif vs negatif
                    st.markdown("##### Başarı (K/Z işareti)")
                    bdf = ddf.groupby("basari", as_index=False).agg(adet=("kod", "count"), değer=("değer", "sum"))
                    fig4 = px.bar(bdf, x="basari", y="değer", color="basari",
                                  color_discrete_map={"Pozitif": "#3ddc97", "Negatif": "#ff6b6b"},
                                  title="Pozitif / negatif pozisyon değeri")
                    fig4.update_layout(template="plotly_dark", height=320, paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                    st.plotly_chart(fig4, use_container_width=True)

                except Exception as e:
                    st.caption(f"Grafik: {e}")
                    st.dataframe(ddf, hide_index=True)

                st.markdown("##### Detay tablo")
                show_cols = ["kod", "ad", "kategori", "kurum", "risk", "ağırlık_%", "K/Z_%", "değer", "basari"]
                dshow = ddf[show_cols].sort_values("ağırlık_%", ascending=False).copy()
                dshow["ağırlık_%"] = dshow["ağırlık_%"].map(lambda x: tr_num(x, 2))
                dshow["K/Z_%"] = dshow["K/Z_%"].map(lambda x: tr_num(x, 2, force_sign=True))
                dshow["değer"] = dshow["değer"].map(lambda x: tr_num(x, 2))
                st.dataframe(dshow, hide_index=True, use_container_width=True, height=320)
                st.session_state["_fy_dagilim_df"] = ddf

                # Kısa analiz
                top_kat = ddf.groupby("kategori")["değer"].sum().idxmax() if len(ddf) else "—"
                top_kur = ddf.groupby("kurum")["değer"].sum().idxmax() if len(ddf) else "—"
                yuksek_risk_pay = float(ddf.loc[ddf["risk"] == "Yüksek", "değer"].sum() / deger_toplam * 100) if deger_toplam else 0
                st.info(
                    f"**Dağılım özeti:** En büyük kategori **{top_kat}** · en büyük kurum **{top_kur}** · "
                    f"yüksek risk bandı payı **%{yuksek_risk_pay:.0f}** · başarı (adet) **%{basari_adet:.0f}**. "
                    "Kategori/risk ad tahminidir (fon unvanından); resmi TEFAS şeması değildir."
                )


    # ── Benchmark + Rebalance ──
    if poz and deger_toplam > 0:
        st.markdown(
            '<div class="bm-panel"><h3>📐 Benchmark & yeniden dengeleme</h3>'
            '<p style="margin:0 0 12px;color:#94a3b8;font-size:0.9rem">'
            'Portföy getirinizi referanslarla kıyaslayın; hedef ağırlığa göre sapma ve rebalance önerisi alın. '
            'Yatırım tavsiyesi değildir.</p></div>',
            unsafe_allow_html=True,
        )
        bc1, bc2, bc3 = st.columns([1.2, 1, 1])
        with bc1:
            bm_sec = st.multiselect(
                "Benchmark",
                ["BIST100", "Mevduat ~TÜFE+", "Nakit (0%)"],
                default=["BIST100"],
                key="fy_bm_sel",
            )
        with bc2:
            rb_mod = st.selectbox(
                "Rebalance hedefi",
                ["Eşit ağırlık", "Mevcut ağırlığı koru (sapma göster)", "Max % tavan"],
                key="fy_rb_mod",
            )
        with bc3:
            max_w_rb = st.slider("Tek fon tavan %", 10, 40, 20, key="fy_rb_cap")

        if st.button("Benchmark & rebalance hesapla", type="primary", key="fy_bm_btn"):
            st.session_state["_fy_bm_run"] = True
            st.session_state["_fy_aylik_calc"] = True

        if st.session_state.get("_fy_bm_run"):
            # Fiili ağırlıklar
            w_act = {}
            for kod, p in poz.items():
                adet = float(p.get("adet") or 0)
                px = fiyatlar.get(kod) or float(p.get("maliyet") or 0)
                w_act[kod] = (adet * px) / deger_toplam if deger_toplam else 0

            # Hedef ağırlık
            n = len(w_act) or 1
            if rb_mod == "Eşit ağırlık":
                w_tgt = {k: 1.0 / n for k in w_act}
            elif rb_mod == "Max % tavan":
                cap = max_w_rb / 100.0
                raw = {k: min(w_act[k], cap) for k in w_act}
                s = sum(raw.values()) or 1
                # tavan üstü kalanı eşit dağıt
                overflow = 1.0 - s
                low = [k for k in w_act if w_act[k] < cap]
                if low and overflow > 0:
                    add = overflow / len(low)
                    for k in low:
                        raw[k] = min(cap, raw[k] + add)
                    s2 = sum(raw.values()) or 1
                    w_tgt = {k: raw[k] / s2 for k in raw}
                else:
                    w_tgt = {k: raw[k] / s for k in raw}
            else:
                w_tgt = dict(w_act)

            # Sapma tablosu
            rows_rb = []
            for kod in w_act:
                a, t = w_act[kod] * 100, w_tgt.get(kod, 0) * 100
                sap = a - t
                rows_rb.append({
                    "kod": kod,
                    "fiili_%": round(a, 2),
                    "hedef_%": round(t, 2),
                    "sapma_%": round(sap, 2),
                    "aksiyon": (
                        "Azalt" if sap > 1.5 else ("Artır" if sap < -1.5 else "Tut")
                    ),
                    "₺ fark": round(deger_toplam * (t - a) / 100.0, 0),
                })
            rbdf = pd.DataFrame(rows_rb).sort_values("sapma_%", key=lambda s: s.abs(), ascending=False)

            st.markdown("##### Ağırlık sapması")
            rb_show = rbdf.copy()
            for c in ("fiili_%", "hedef_%", "sapma_%", "₺ fark"):
                rb_show[c] = rb_show[c].map(lambda x: tr_num(x, 2 if c != "₺ fark" else 0, force_sign=(c == "sapma_%")))
            st.dataframe(rb_show, hide_index=True, use_container_width=True, height=280)

            # progress bars for top drifts
            st.markdown("##### Sapma çubukları")
            for _, rr in rbdf.head(8).iterrows():
                st.markdown(
                    f'<div class="rb-label"><span><b>{rr["kod"]}</b> · {rr["aksiyon"]}</span>'
                    f'<span>fiili %{tr_num(rr["fiili_%"],1)} → hedef %{tr_num(rr["hedef_%"],1)}</span></div>',
                    unsafe_allow_html=True,
                )
                st.progress(min(1.0, abs(float(rr["fiili_%"])) / 100.0))

            # Benchmark getiri (seri varsa)
            port_nav = st.session_state.get("_fy_port_nav")
            chips = []
            if port_nav is not None and len(port_nav) > 5:
                pret = float(port_nav.iloc[-1] / port_nav.iloc[0] * 100 - 100)
                chips.append(f'<span class="bm-chip">Portföy <strong>%{tr_num(pret, 1)}</strong> (seri)</span>')
            else:
                chips.append(
                    f'<span class="bm-chip">Portföy maliyet K/Z <strong>%{tr_num(acik_kz_pct, 1)}</strong></span>'
                )
                pret = acik_kz_pct

            # BIST100
            if "BIST100" in bm_sec:
                bist_ret = None
                try:
                    import yfinance as yf
                    bh = yf.Ticker("XU100.IS").history(period="1y", auto_adjust=True)
                    if bh is not None and not bh.empty and len(bh) > 5:
                        bist_ret = float(bh["Close"].iloc[-1] / bh["Close"].iloc[0] * 100 - 100)
                except Exception:
                    pass
                if bist_ret is not None:
                    fark = pret - bist_ret
                    chips.append(
                        f'<span class="bm-chip">BIST100 <strong>%{tr_num(bist_ret, 1)}</strong></span>'
                    )
                    chips.append(
                        f'<span class="bm-chip">Fark <strong style="color:{"#3ddc97" if fark>=0 else "#ff6b6b"}">%{tr_num(fark, 1, force_sign=True)}</strong></span>'
                    )
                else:
                    chips.append('<span class="bm-chip">BIST100 veri yok</span>')

            if "Mevduat ~TÜFE+" in bm_sec:
                try:
                    tufe = float((st.session_state.cfg or {}).get("tufe_yillik") or 40)
                except Exception:
                    tufe = 40.0
                mev = tufe + 5  # kaba proxy
                chips.append(
                    f'<span class="bm-chip">Mevduat~TÜFE+5 <strong>%{tr_num(mev, 1)}</strong></span>'
                )
                chips.append(
                    f'<span class="bm-chip">Portföy−mevduat <strong>%{tr_num(pret - mev, 1, force_sign=True)}</strong></span>'
                )
            if "Nakit (0%)" in bm_sec:
                chips.append(
                    f'<span class="bm-chip">Nakit 0% · reel ≈ <strong>%{tr_num(pret, 1, force_sign=True)}</strong></span>'
                )

            st.markdown("".join(chips), unsafe_allow_html=True)
            st.caption(
                "BIST100: yfinance XU100.IS ~1Y. Mevduat: Ayarlar TÜFE + 5 puan (yaklaşık). "
                "Rebalance: hedef ağırlığa ₺ fark kadar al/sat simülasyonu — gerçek emir değil."
            )

            # Tek tık eşit ağırlık uygula
            if rb_mod == "Eşit ağırlık" and st.button("Eşit ağırlığı portföye uygula (takip)", key="fy_rb_apply"):
                pxm = fiyatlar
                poz2 = {}
                pay = deger_toplam / n
                for kod in w_act:
                    px = pxm.get(kod) or float(poz[kod].get("maliyet") or 1)
                    if px <= 0:
                        px = 1.0
                    poz2[kod] = {
                        "adet": pay / px,
                        "maliyet": float(poz[kod].get("maliyet") or px),
                        "ad": poz[kod].get("ad") or kod,
                    }
                sp["pozisyonlar"] = poz2
                st.session_state.sanal_portfoy = sp
                st.success("Eşit ağırlık uygulandı (takip tutarı korundu).")
                st.rerun()

        # ── Portföyü analiz et + güçlü / uyarı / akıllı tavsiye ──
    st.markdown("---")
    st.markdown("### 🧠 Portföy analizi · güçlü yönler · uyarılar · akıllı tavsiye")
    cbtn1, cbtn2 = st.columns([1, 1])
    with cbtn1:
        if st.button("Portföyü analiz et", type="primary", key="fy_analiz_btn"):
            st.session_state["_fy_analiz_run"] = True
    with cbtn2:
        auto_show = st.checkbox(
            "Portföy yüklüyken paneli açık tut",
            value=True,
            key="fy_panel_auto",
        )
    if auto_show and poz:
        st.session_state["_fy_analiz_run"] = True
    if st.session_state.get("_fy_excel_just_loaded"):
        st.session_state["_fy_excel_just_loaded"] = False
        st.session_state["_fy_analiz_run"] = True

    if st.session_state.get("_fy_analiz_run") and poz:
        ddf = st.session_state.get("_fy_dagilim_df")
        if ddf is None or getattr(ddf, "empty", True):
            rows_d = []
            for kod, p in poz.items():
                adet = float(p.get("adet") or 0)
                mal = float(p.get("maliyet") or 0)
                px = fiyatlar.get(kod) or mal
                deg = adet * float(px or 0)
                kz_p = ((float(px) / mal) - 1) * 100 if mal and px else 0.0
                rows_d.append({
                    "kod": kod,
                    "değer": deg,
                    "K/Z_%": kz_p,
                    "ağırlık_%": deg / deger_toplam * 100 if deger_toplam else 0,
                })
            ddf = pd.DataFrame(rows_d)

        guclu, uyarılar, tavsiyeler = [], [], []
        n_pos = int((ddf["K/Z_%"] >= 0).sum())
        n_neg = int((ddf["K/Z_%"] < 0).sum())
        n_all = len(ddf)
        basari = n_pos / n_all * 100 if n_all else 0
        pos_val = float(ddf.loc[ddf["K/Z_%"] >= 0, "değer"].sum()) if "değer" in ddf.columns else 0.0
        basari_deger = (pos_val / deger_toplam * 100) if deger_toplam else 0.0

        if basari >= 70:
            guclu.append(f"Başarı oranı yüksek: **%{basari:.0f}** ({n_pos}/{n_all} pozitif).")
        if basari_deger >= 80:
            guclu.append(f"Değerin **%{basari_deger:.0f}**’i kârda.")
        if acik_kz_pct >= 30:
            guclu.append(f"Toplam getiri güçlü: **%{acik_kz_pct:+.1f}** (₺{acik_kz:,.0f}).")
        if "ağırlık_%" in ddf.columns and float(ddf["ağırlık_%"].max()) <= 18:
            guclu.append("Tek fon baskın değil; ağırlıklar dengeli (max ≤ %18).")
        if n_all >= 8:
            guclu.append(f"Çeşitlilik: **{n_all}** fon.")

        if basari < 50:
            uyarılar.append(f"Başarı oranı düşük (**%{basari:.0f}**); {n_neg} pozisyon zarar yazıyor.")
        if acik_kz_pct < 0:
            uyarılar.append(f"Portföy maliyet altında: **%{acik_kz_pct:.1f}**.")
        if "ağırlık_%" in ddf.columns:
            mx = float(ddf["ağırlık_%"].max())
            mxk = str(ddf.loc[ddf["ağırlık_%"].idxmax(), "kod"])
            if mx > 20:
                uyarılar.append(f"**{mxk}** tek başına **%{mx:.0f}** — yoğunlaşma riski.")
            if mx > 30:
                uyarılar.append(f"Kritik yoğunlaşma: {mxk} > %30.")
        deep = ddf[ddf["K/Z_%"] <= -25]
        for _, r in deep.iterrows():
            uyarılar.append(f"**{r['kod']}** derin zarar **%{float(r['K/Z_%']):.0f}**")
        if n_all < 4:
            uyarılar.append("Fon sayısı çok az — idiyosenkratik risk yüksek.")
        if n_all > 30:
            uyarılar.append("30+ fon: izleme zorlaşır.")

        if basari < 50 or n_neg >= 3:
            tavsiyeler.append("Zarar eden fonları gözden geçirin; tez yoksa kademeli azaltın.")
        if "ağırlık_%" in ddf.columns and float(ddf["ağırlık_%"].max()) > 20:
            mxk = str(ddf.loc[ddf["ağırlık_%"].idxmax(), "kod"])
            tavsiyeler.append(f"**{mxk}** payını %15 altına çekin veya düşük vol tampon ekleyin.")
        for _, r in deep.head(3).iterrows():
            tavsiyeler.append(f"**{r['kod']}**: çıkış planı veya pay tavanı (şu an %{float(r['K/Z_%']):.0f}).")
        if acik_kz_pct >= 80:
            tavsiyeler.append("Yüksek kâr: bir kısmını para piyasasına realize etmeyi değerlendirin.")
        if n_all < 6:
            tavsiyeler.append("Para piyasası / hisse / değişken kategorilerinden çeşitlendirin.")
        dfull = st.session_state.get("_fy_dagilim_df")
        if dfull is not None and "kurum" in getattr(dfull, "columns", []):
            kshare = dfull.groupby("kurum")["değer"].sum()
            if len(kshare) and float(kshare.max() / kshare.sum()) > 0.4:
                tavsiyeler.append(
                    f"Tek kurum (**{kshare.idxmax()}**) %{float(kshare.max()/kshare.sum()*100):.0f} — kurum riskini dağıtın."
                )
        if not tavsiyeler:
            tavsiyeler.append("Kritik kırmızı bayrak yok; periyodik rebalance yeterli olabilir.")

        a1, a2, a3, a4 = st.columns(4)
        a1.metric("Başarı (adet) %", f"{basari:.0f}", f"{n_pos}/{n_all}")
        a2.metric("Başarı (değer) %", f"{basari_deger:.0f}")
        a3.metric("Toplam K/Z %", f"{acik_kz_pct:+.1f}")
        a4.metric("Pozisyon", n_all)

        if guclu:
            st.success("### ✓ Güçlü yönler\n" + "\n".join(f"- {x}" for x in guclu))
        else:
            st.info("Güçlü yön henüz işaretlenmedi.")
        if uyarılar:
            st.warning("### ⚠ Uyarılar\n" + "\n".join(f"- {x}" for x in uyarılar))
        else:
            st.success("Kritik uyarı yok.")
        st.info("### 💡 Akıllı tavsiyeler\n" + "\n".join(f"- {x}" for x in tavsiyeler))
        st.caption(
            "Excel sonrası panel otomatik açılır. "
            "Kurallar: başarı, yoğunlaşma, derin zarar, çeşitlilik, kurum payı. **Yatırım tavsiyesi değildir.**"
        )
    elif not poz:
        st.caption("Önce Excel yükleyin — panel o zaman dolar.")

    st.caption("Excel Özet/Lotlar formatı desteklenir. Canlı fiyat TEFAS; yoksa Excel değer kullanılır.")

elif sayfa == "🎯 Optimizasyon":
    st.header("🎯 Portföy Optimizasyonu")
    kodlar_str = st.text_input("Fon kodları (virgülle)", value="AAL,AFO,TTE", key="opt_kod")
    method = st.selectbox(
        "Yöntem",
        ["max_sharpe", "min_var", "risk_parity", "cvar", "equal", "score", "all"],
        format_func=lambda x: {
            "max_sharpe": "Max Sharpe",
            "min_var": "Min Variance",
            "risk_parity": "Risk Parity",
            "cvar": "Min CVaR (%5)",
            "equal": "Eşit ağırlık",
            "score": "Skor orantılı",
            "all": "Hepsi (karşılaştır)",
        }[x],
    )
    max_w = st.slider("Tek fon max ağırlık", 0.05, 1.0, 0.40, 0.05)
    max_risk = st.number_input("Max risk_skor filtresi (boş=yok)", value=None, step=1.0, placeholder="yok")

    if st.button("Optimize et", type="primary"):
        analiz = get_analiz()
        kodlar = [k.strip().upper() for k in kodlar_str.split(",") if k.strip()]
        series, scores = {}, {}
        with st.spinner("Fiyat serileri ve metrikler…"):
            for kod in kodlar:
                p = analiz.api.profil(kod)
                ad = p.get("fonUnvan") or kod
                ok, neden = core.fon_yatirima_uygun_mu(ad, p.get("tefasDurum") or "")
                if not ok:
                    st.warning(f"✗ {kod}: {neden}")
                    continue
                fd = analiz.api.fiyat(kod, periyod=12)
                if fd.empty:
                    st.warning(f"✗ {kod}: fiyat yok")
                    continue
                an = core.anomaly_from_prices(fd)
                o = analiz.api.ozet(kod)
                m = core.full_metrics(fd, analiz.rf)
                liq = core.likidite_risk(o, m, profil=p, anomaly=an)
                risk_skor = (an.get("anomaly_skor") or 0) + (liq.get("likidite_skor_risk") or 0)
                if max_risk is not None and risk_skor > max_risk:
                    st.warning(f"✗ {kod}: risk_skor={risk_skor:.1f}")
                    continue
                series[kod] = fd.set_index("tarih")["fiyat"]
                scores[kod] = float(m.get("sharpe") or 0) * 10 + float(m.get("getiri_1y") or 0) - risk_skor
                st.caption(f"✓ {kod}  risk_skor={risk_skor:.1f}  sharpe={m.get('sharpe')}")

        if len(series) < 2:
            st.error("En az 2 uygun fon gerekli.")
        else:
            methods = (
                ["max_sharpe", "min_var", "risk_parity", "cvar", "equal", "score"]
                if method == "all" else [method]
            )
            rows = []
            for met in methods:
                out = core.run_optimizer(met, series, rf=analiz.rf, max_w=max_w, scores=scores)
                if out.get("error"):
                    st.warning(out["error"])
                    continue
                st.subheader(met)
                c1, c2, c3 = st.columns(3)
                c1.metric("μ (yıllık)", f"%{out['port_mu']}")
                c2.metric("Vol", f"%{out['port_vol']}")
                c3.metric("Sharpe", f"{out['port_sharpe']}")
                wdf = pd.DataFrame([
                    {"kod": k, "ağırlık %": round(v * 100, 1),
                     "risk katkısı %": out["risk_contrib_pct"].get(k, 0)}
                    for k, v in sorted(out["weights"].items(), key=lambda x: -x[1]) if v >= 0.001
                ])
                st.dataframe(wdf, hide_index=True, use_container_width=True)
                hist = core.portfolio_metrics(out["weights"], series, rf=analiz.rf)
                if hist:
                    st.caption(
                        f"Tarihsel 1Y: %{hist.get('getiri_1y')} · "
                        f"MaxDD: %{hist.get('max_drawdown')} · Sharpe: {hist.get('sharpe')}"
                    )
                rows.append({"yontem": met, "mu_%": out["port_mu"], "vol_%": out["port_vol"],
                             "sharpe": out["port_sharpe"], **{f"w_{k}": v for k, v in out["weights"].items()}})
            if rows:
                st.session_state.opt_sonuc = pd.DataFrame(rows)
                download_excel_button(
                    st.session_state.opt_sonuc,
                    f"tefas_opt_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    key="xl_opt",
                )


# ═══════════════════════════════════════════════════════════════════════════
# FON DETAY (kod veya ad ile arama)
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "🔎 Fon Detay":
    st.header("🔎 Fon Detay & Risk Raporu")
    st.caption("Fon **kodu** (örn. AAL) veya **adının bir kısmını** yazın (örn. Ak Portföy, altın, para piyasası)")

    arama = st.text_input(
        "Fon kodu veya adı",
        value=st.session_state.get("detay_arama", ""),
        placeholder="örn. AAL  veya  İş Portföy Hisse  veya  altın",
        key="detay_arama_input",
    ).strip()

    rapor_tur = st.radio(
        "HTML rapor türü",
        ["Türkçe (önerilen)", "Quantstats"] if core.HAS_QUANTSTATS else ["Türkçe (önerilen)"],
        horizontal=True,
    )
    bm_ok = st.checkbox("BIST100 ile karşılaştır", value=True)

    def _normalize_search(s: str) -> str:
        return core._normalize_tr(s or "").lower()

    def _fon_ara(analiz: core.FonAnaliz, q: str) -> pd.DataFrame:
        """Kod tam eşleşme veya unvan içinde arama."""
        q = (q or "").strip()
        if not q:
            return pd.DataFrame()
        q_up = q.upper().replace(" ", "")
        # Kısa ve alfanümerik → muhtemel kod
        if len(q) <= 6 and q_up.isalnum():
            try:
                p = analiz.api.profil(q_up)
                if p:
                    ad = p.get("fonUnvan") or q_up
                    return pd.DataFrame([{"kod": q_up, "ad": ad}])
            except Exception:
                pass
        # İsim / kısmi arama — dağılım listesi
        with st.spinner("Fon listesinde aranıyor…"):
            dag = analiz.api.dagilimlar()
        if dag is None or dag.empty:
            return pd.DataFrame()
        nq = _normalize_search(q)
        rows = []
        for _, r in dag.iterrows():
            kod = str(r.get("kod") or "").upper()
            ad = str(r.get("ad") or "")
            if nq in _normalize_search(ad) or nq in _normalize_search(kod) or q_up == kod:
                rows.append({"kod": kod, "ad": ad})
        if not rows:
            return pd.DataFrame()
        out = pd.DataFrame(rows).drop_duplicates("kod")
        # Daha iyi eşleşenler üste (başta geçen / daha kısa ad)
        def _rank(row):
            ad_n = _normalize_search(row["ad"])
            if ad_n.startswith(nq):
                return (0, len(ad_n))
            if nq in ad_n:
                return (1, len(ad_n))
            return (2, len(ad_n))
        out["_r"] = out.apply(_rank, axis=1)
        out = out.sort_values("_r").drop(columns=["_r"]).reset_index(drop=True)
        return out.head(40)

    def _detay_yukle(analiz: core.FonAnaliz, kod: str):
        d = analiz.fon_detay(kod)
        m, o, p = d["metrikler"], d["ozet"], d["profil"]
        ad = o.get("fonUnvan") or p.get("fonUnvan") or kod
        # Anomali + likidite özeti
        fd = d.get("fiyat_df")
        an = core.anomaly_from_prices(fd) if fd is not None else {}
        liq = core.likidite_risk(o, m, profil=p, anomaly=an)
        return {
            "d": d, "m": m, "o": o, "p": p, "ad": ad, "kod": kod,
            "an": an, "liq": liq,
        }

    col_ara, col_btn = st.columns([3, 1])
    with col_btn:
        st.write("")
        st.write("")
        ara_tik = st.button("🔍 Ara / Analiz et", type="primary", use_container_width=True)

    if ara_tik and arama:
        st.session_state.detay_arama = arama
        analiz = get_analiz()
        bulunan = _fon_ara(analiz, arama)
        st.session_state.detay_arama_sonuc = bulunan
        if bulunan is None or bulunan.empty:
            st.warning(f"«{arama}» için fon bulunamadı. Kodu veya adın bir parçasını deneyin.")
            st.session_state.detay_sonuc = None
        elif len(bulunan) == 1:
            kod = str(bulunan.iloc[0]["kod"])
            with st.spinner(f"{kod} analiz ediliyor…"):
                try:
                    st.session_state.detay_sonuc = _detay_yukle(analiz, kod)
                except Exception as e:
                    st.error(str(e))
                    st.session_state.detay_sonuc = None
        else:
            st.session_state.detay_sonuc = None  # seçim bekleniyor

    # Birden fazla sonuç → seçim
    bulunan = st.session_state.get("detay_arama_sonuc")
    if (
        st.session_state.get("detay_sonuc") is None
        and bulunan is not None
        and not getattr(bulunan, "empty", True)
        and len(bulunan) > 1
    ):
        st.info(f"**{len(bulunan)}** fon bulundu — analiz için birini seçin:")
        sec_etiket = [
            f"{r['kod']} — {str(r['ad'])[:70]}"
            for _, r in bulunan.iterrows()
        ]
        secim = st.radio("Fon seçin", sec_etiket, key="detay_secim_radio")
        if st.button("Seçileni analiz et", type="primary", key="detay_sec_analiz"):
            kod = secim.split(" — ")[0].strip().upper()
            analiz = get_analiz()
            with st.spinner(f"{kod} analiz ediliyor…"):
                try:
                    st.session_state.detay_sonuc = _detay_yukle(analiz, kod)
                except Exception as e:
                    st.error(str(e))
                    st.session_state.detay_sonuc = None

    det = st.session_state.detay_sonuc
    if det:
        m, o, p, ad, kod = det["m"], det["o"], det["p"], det["ad"], det["kod"]
        an, liq = det.get("an") or {}, det.get("liq") or {}
        st.subheader(f"`{kod}` — {ad}")
        try:
            _bd_f = son_bilanco_donemi(probe=False)
            _donem_not = f" · Rapor çeyreği: {_bd_f.get('etiket')}"
        except Exception:
            _donem_not = ""
        st.caption(
            f"TEFAS: {p.get('tefasDurum')} · Risk: {p.get('riskDegeri')} · "
            f"Faiz: {core.faiz_durum(p, ad)} · Kategori: {o.get('fonKategori')} · "
            f"İşlem: {p.get('basIsSaat') or '?'}-{p.get('sonIsSaat') or '?'}"
            f"{_donem_not} · NAV günlük güncel"
        )

        c1, c2, c3, c4, c5, c6 = st.columns(6)
        c1.metric("1Y %", fmt_num(m.get("getiri_1y")))
        c2.metric("Reel 1Y", fmt_num(m.get("reel_1y")))
        c3.metric("Sharpe", fmt_num(m.get("sharpe"), 2))
        c4.metric("Sortino", fmt_num(m.get("sortino"), 2))
        c5.metric("MaxDD %", fmt_num(m.get("max_drawdown")))
        c6.metric("Vol %", fmt_num(m.get("volatilite")))

        c7, c8, c9, c10, c11, c12 = st.columns(6)
        c7.metric("vs Mevduat", fmt_num(m.get("vs_mevduat")))
        c8.metric("vs BIST100", fmt_num(m.get("vs_bist100")))
        c9.metric("CAGR %", fmt_num(m.get("cagr")))
        c10.metric("Fiyat", fmt_num(m.get("fiyat"), 4))
        c11.metric("Likidite", fmt_num(liq.get("likidite_skor"), 0))
        c12.metric("Risk skor", fmt_num(an.get("anomaly_skor"), 1))

        if liq.get("likidite_seviye") or an.get("risk_bayrak") or liq.get("likidite_bayrak"):
            parts = []
            if liq.get("likidite_seviye"):
                parts.append(f"Likidite: {liq.get('likidite_seviye')}")
            if liq.get("likidite_bayrak"):
                parts.append(str(liq.get("likidite_bayrak")))
            if an.get("risk_bayrak"):
                parts.append(str(an.get("risk_bayrak")))
            st.caption(" · ".join(parts))

        more = {
            "1A %": m.get("getiri_1a"), "3A %": m.get("getiri_3a"), "6A %": m.get("getiri_6a"),
            "3Y %": m.get("getiri_3y"), "Calmar": m.get("calmar"),
            "VaR95": m.get("var95"), "cVaR95": m.get("cvar95"),
            "Kazanma %": m.get("kazanma_orani"),
            "En iyi gün %": m.get("en_iyi_gun"), "En kötü gün %": m.get("en_kotu_gun"),
            "Yatırımcı": o.get("yatirimciSayi"), "Büyüklük": o.get("portBuyukluk"),
            "Giriş kom.": p.get("girisKomisyonu"), "Çıkış kom.": p.get("cikisKomisyonu"),
        }
        st.json({k: v for k, v in more.items() if v is not None})

        # Portföy dağılımı (detay ekranı)
        try:
            analiz_d = get_analiz()
            dag = analiz_d.api.dagilimlar()
            if dag is not None and not dag.empty:
                row_d = dag[dag["kod"].astype(str).str.upper() == kod.upper()]
                if not row_d.empty:
                    oz = analiz_d.api.dagilim_ozet(row_d.iloc[0])
                    oz = {core.KATEGORI_ACIKLAMA.get(k, k): round(v, 1) for k, v in oz.items() if v and v >= 0.5}
                    if oz:
                        st.subheader("Portföy varlık dağılımı")
                        st.bar_chart(pd.Series(oz, name="Pay %"))
                        st.caption(" · ".join(f"{k}: %{v}" for k, v in sorted(oz.items(), key=lambda x: -x[1])))
        except Exception:
            pass

        # Önceki fon metrik kaydı ile karşılaştır
        try:
            _fy = mali_tablo_iyilesme_kontrolu(kod, dict(m), tur="fon")
            if _fy:
                st.session_state["_bilanco_iyilesme"] = (
                    list(st.session_state.get("_bilanco_iyilesme") or []) + _fy
                )[-30:]
                for _msg in _fy:
                    st.success(_msg)
        except Exception:
            pass

        # Sade Türkçe değerlendirme (plain markdown — Streamlit HTML'i bazen boş gösteriyordu)
        st.subheader("Sade dilde değerlendirme")
        metin_deg = core.turkce_degerlendirme(
            sharpe=m.get("sharpe"), sortino=m.get("sortino"),
            max_dd=m.get("max_drawdown"), vol=m.get("volatilite"),
            cagr=m.get("cagr"), calmar=m.get("calmar"),
            getiri_1y=m.get("getiri_1y"), reel_1y=m.get("reel_1y"),
            vs_mevduat=m.get("vs_mevduat"), vs_bist100=m.get("vs_bist100"),
            kazanma_orani=m.get("kazanma_orani"),
            en_kotu_gun=m.get("en_kotu_gun"),
            likidite_skor=(liq or {}).get("likidite_skor"),
            plain=True,
        )
        # BIST korelasyonu detayda ayrıca hesaplanmış olabilir
        try:
            if det.get("d") and det["d"].get("fiyat_df") is not None:
                bk = core.fon_bist_korelasyon(det["d"]["fiyat_df"])
                if bk is not None:
                    ekstra = core.turkce_degerlendirme(bist_korelasyon=bk, plain=True)
                    # sadece BIST satırını ekle
                    for line in ekstra.splitlines():
                        if "Korelasyon" in line or "korelasyon" in line:
                            metin_deg = metin_deg.rstrip() + "\n\n" + line + "\n"
        except Exception:
            pass
        st.markdown(metin_deg)
        # HTML kutu da (rapor stili) — yedek görünüm
        with st.expander("HTML kutu görünümü", expanded=False):
            st.components.v1.html(
                core.turkce_degerlendirme(
                    sharpe=m.get("sharpe"), sortino=m.get("sortino"),
                    max_dd=m.get("max_drawdown"), vol=m.get("volatilite"),
                    cagr=m.get("cagr"), calmar=m.get("calmar"),
                    getiri_1y=m.get("getiri_1y"), reel_1y=m.get("reel_1y"),
                    vs_mevduat=m.get("vs_mevduat"), vs_bist100=m.get("vs_bist100"),
                    kazanma_orani=m.get("kazanma_orani"),
                    en_kotu_gun=m.get("en_kotu_gun"),
                    likidite_skor=(liq or {}).get("likidite_skor"),
                    plain=False,
                ),
                height=420,
                scrolling=True,
            )

        b1, b2, b3 = st.columns(3)
        with b1:
            if st.button("⭐ Favorilere ekle", key="detay_fav"):
                favs = list(set(core.load_favs()) | {kod})
                core.save_favs(favs)
                st.success(f"{kod} eklendi.")
            if st.button("🛒 Portföy sepetine ekle", key="detay_sepet"):
                n = sepet_ekle([kod], ad_map={kod: ad}, kaynak="detay")
                st.success(f"{kod} sepete eklendi" if n else f"{kod} zaten sepette")
        with b2:
            # Mini Excel
            row = {"kod": kod, "ad": ad, **{k: m.get(k) for k in (
                "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y", "getiri_3y",
                "reel_1y", "vs_mevduat", "vs_bist100", "volatilite", "max_drawdown",
                "sharpe", "sortino", "calmar", "cagr", "fiyat",
            )}}
            download_excel_button(
                pd.DataFrame([row]),
                f"tefas_detay_{kod}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                meta={"kod": kod, "ad": ad},
                key="xl_detay",
            )
        with b3:
            pass

        if st.button("📄 HTML risk raporu oluştur", type="primary"):
            analiz = get_analiz()
            fd = det["d"].get("fiyat_df")
            if fd is None or len(fd) < 100:
                fd = analiz.api.fiyat(kod, periyod=36)
            bm = core.yahoo_daily_returns("XU100.IS") if bm_ok else None
            # Portföy varlık dağılımı
            dagilim_fon = {}
            try:
                dag = analiz.api.dagilimlar()
                if dag is not None and not dag.empty:
                    row = dag[dag["kod"].astype(str).str.upper() == kod.upper()]
                    if not row.empty:
                        dagilim_fon = analiz.api.dagilim_ozet(row.iloc[0])
            except Exception:
                pass
            ek_m = dict(m)
            for k, v in (dagilim_fon or {}).items():
                ek_m[f"d_{k}"] = v
            with st.spinner("Rapor üretiliyor…"):
                if rapor_tur.startswith("Quant") and core.HAS_QUANTSTATS:
                    path = core.quantstats_rapor(fd, kod, ad=ad, benchmark_returns=bm, ekstra_metrikler=ek_m)
                else:
                    path = core.turkce_risk_raporu(
                        fd, kod, ad=ad, rf=analiz.rf,
                        benchmark_returns=bm, benchmark_ad="BIST100",
                        ekstra_metrikler=ek_m, dagilim=dagilim_fon,
                    )
            if path:
                st.success(f"Rapor: {path}")
                with open(path, "rb") as f:
                    st.download_button(
                        "HTML indir", f, file_name=Path(path).name,
                        mime="text/html", key="dl_html",
                    )
                try:
                    html = Path(path).read_text(encoding="utf-8")
                    st.components.v1.html(html, height=800, scrolling=True)
                except Exception:
                    pass


# ═══════════════════════════════════════════════════════════════════════════
# HİSSE & KARMA PORTFÖY
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "📈 Hisse & Karma Portföy":
    st.header("📈 Hisse analizi & fon+hisse karma portföy")
    st.caption(
        "BIST hisseleri Yahoo Finance üzerinden çekilir (örn. THYAO, GARAN, ASELS). "
        "Fon kodları TEFAS; hisseler `.IS` soneki ile veya kısa kodla girilir."
    )

    tab_h, tab_k = st.tabs(["Hisse analizi", "Fon + Hisse portföy"])

    with tab_h:
        hisse_str = st.text_input(
            "Hisse kodları (virgülle)",
            value=st.session_state.get("hisse_liste", "THYAO,GARAN,ASELS"),
            placeholder="THYAO, GARAN, EREGL, TUPRS",
            key="hisse_input",
        )
        if st.button("Hisseleri analiz et", type="primary", key="btn_hisse_analiz"):
            semboller = [x.strip() for x in hisse_str.split(",") if x.strip()]
            st.session_state.hisse_liste = hisse_str
            analiz = get_analiz()
            with st.spinner("Hisse fiyatları ve metrikler hesaplanıyor…"):
                hdf = core.hisseleri_analiz_et(semboller, rf=analiz.rf, cfg=st.session_state.cfg)
            st.session_state.hisse_sonuc = hdf

        hdf = st.session_state.get("hisse_sonuc")
        if hdf is not None and not hdf.empty:
            if "hata" in hdf.columns and hdf["hata"].notna().any():
                for _, r in hdf[hdf["hata"].notna()].iterrows():
                    st.warning(f"{r.get('sembol')}: {r.get('hata')}")
            ok = hdf[hdf.get("hata", pd.Series(dtype=object)).isna()] if "hata" in hdf.columns else hdf
            if not ok.empty:
                st.success(f"{len(ok)} hisse analiz edildi")
                cols_h = [
                    c for c in [
                        "kod", "sembol", "fiyat",
                        "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y",
                        "reel_1y", "vs_mevduat", "vs_bist100",
                        "sharpe", "sortino", "calmar", "cagr",
                        "volatilite", "max_drawdown", "kazanma_orani",
                        "en_kotu_gun", "var95", "bist_korelasyon", "risk_skor",
                    ] if c in ok.columns
                ]
                show_df(ok, cols_h, height=360)
                download_excel_button(
                    ok,
                    f"tefas_hisse_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    meta={"mod": "hisse_analiz"},
                    key="xl_hisse",
                )
                # Tek hisse detay değerlendirme
                sec_h = st.selectbox("Sade dilde değerlendirme için hisse", ok["kod"].tolist(), key="hisse_det_sec")
                if sec_h:
                    row = ok[ok["kod"] == sec_h].iloc[0]
                    st.markdown(core.turkce_degerlendirme(
                        sharpe=row.get("sharpe"), sortino=row.get("sortino"),
                        max_dd=row.get("max_drawdown"), vol=row.get("volatilite"),
                        cagr=row.get("cagr"), calmar=row.get("calmar"),
                        getiri_1y=row.get("getiri_1y"), reel_1y=row.get("reel_1y"),
                        vs_mevduat=row.get("vs_mevduat"), vs_bist100=row.get("vs_bist100"),
                        kazanma_orani=row.get("kazanma_orani"),
                        en_kotu_gun=row.get("en_kotu_gun"),
                        bist_korelasyon=row.get("bist_korelasyon"),
                        plain=True,
                    ))
                    if st.button("HTML risk raporu (hisse)", key="hisse_html_rap"):
                        analiz = get_analiz()
                        det = core.hisse_analiz(str(row.get("sembol") or sec_h), rf=analiz.rf, cfg=st.session_state.cfg)
                        if det.get("error"):
                            st.error(det["error"])
                        else:
                            path = core.turkce_risk_raporu(
                                det["fiyat_df"], det["kod"], ad=det["sembol"],
                                rf=analiz.rf,
                                benchmark_returns=core.yahoo_daily_returns("XU100.IS"),
                                ekstra_metrikler=det["metrikler"],
                            )
                            if path:
                                st.success(path)
                                with open(path, "rb") as f:
                                    st.download_button("HTML indir", f, file_name=Path(path).name, mime="text/html", key="dl_hisse_html")

    with tab_k:
        st.markdown("Fon kodları + hisse kodlarını birlikte ağırlıklandırın.")
        n_var = st.slider("Kaç varlık", 2, 12, 4, key="karma_n")
        weights = {}
        cols_w = st.columns(min(n_var, 4))
        for i in range(n_var):
            with cols_w[i % len(cols_w)]:
                kod = st.text_input(
                    f"Varlık {i+1} (fon veya hisse)",
                    value=["AAL", "THYAO", "GARAN", "TTE"][i] if i < 4 else "",
                    key=f"karma_kod_{i}",
                    help="Fon: AAL · Hisse: THYAO veya THYAO.IS",
                ).strip().upper()
                pct = st.number_input(f"Ağırlık %", value=round(100 / n_var), min_value=1, max_value=100, key=f"karma_w_{i}")
                if kod:
                    # Hisse kısa kodlarını .IS ile işaretle eğer 5 harf ve TEFAS değil gibi
                    weights[kod] = float(pct)

        if st.button("Karma portföyü analiz et", type="primary", key="btn_karma"):
            if not weights:
                st.warning("En az bir varlık girin.")
            else:
                analiz = get_analiz()
                with st.spinner("Fon + hisse serileri birleştiriliyor…"):
                    out = core.portfoy_fon_hisse_analiz(analiz, weights)
                st.session_state.karma_sonuc = out

        out = st.session_state.get("karma_sonuc")
        if out and out.get("error"):
            st.error(out["error"])
            if out.get("meta"):
                st.json(out["meta"])
        elif out and out.get("metrikler"):
            m = out["metrikler"]
            st.success(f"Ortak seri: {out.get('n_gun')} gün")
            c1, c2, c3, c4, c5 = st.columns(5)
            c1.metric("1Y %", fmt_num(m.get("getiri_1y")))
            c2.metric("Sharpe", fmt_num(m.get("sharpe"), 2))
            c3.metric("MaxDD %", fmt_num(m.get("max_drawdown")))
            c4.metric("Vol %", fmt_num(m.get("volatilite")))
            c5.metric("vs BIST", fmt_num(m.get("vs_bist100")))
            st.markdown("**Normalize ağırlıklar**")
            wdf = pd.DataFrame([
                {"kod": k, "ağırlık %": round(v * 100, 1),
                 "tip": (out.get("meta") or {}).get(k, {}).get("tip", "?"),
                 "ad": (out.get("meta") or {}).get(k, {}).get("ad", k)}
                for k, v in sorted((out.get("weights") or {}).items(), key=lambda x: -x[1])
            ])
            st.dataframe(wdf, hide_index=True, use_container_width=True)
            st.markdown(core.turkce_degerlendirme(
                sharpe=m.get("sharpe"), sortino=m.get("sortino"),
                max_dd=m.get("max_drawdown"), vol=m.get("volatilite"),
                cagr=m.get("cagr"), calmar=m.get("calmar"),
                getiri_1y=m.get("getiri_1y"), reel_1y=m.get("reel_1y"),
                vs_mevduat=m.get("vs_mevduat"), vs_bist100=m.get("vs_bist100"),
                plain=True,
            ))
            download_excel_button(
                wdf.assign(**{k: m.get(k) for k in ("getiri_1y", "sharpe", "max_drawdown", "volatilite")}),
                f"tefas_karma_portfoy_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                meta={"mod": "karma_fon_hisse", "weights": out.get("weights"), "metrikler": m},
                key="xl_karma",
            )


# ═══════════════════════════════════════════════════════════════════════════
# FAVORİLER
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "🏦 BIST Değerleme":
    st.header("🏦 BIST Değerleme")
    # Güncel bilanço dönemi — raporlar buna göre üretilmeli
    try:
        _bd_bist = son_bilanco_donemi(probe=False)
    except Exception:
        _bd_bist = {"etiket": "?", "aciklama": ""}
    st.caption(
        f"Takvim / aktif bilanço dönemi: **{_bd_bist.get('etiket')}** "
        f"({_bd_bist.get('aciklama') or '—'}). "
        "Her hisse için İş Yatırım’da **dolu en yeni dönem** (ör. ISDMR **2026/9**) rapor ve sıralamaya yazılır; "
        "eski kolon (2026/6) güncel kabul edilmez."
    )
    # tefas_bist modülüne dönem bilgisini ilet (varsa okur)
    try:
        st.session_state["bist_rapor_donem"] = {
            "yil": _bd_bist.get("yil"),
            "donem": _bd_bist.get("donem"),
            "etiket": _bd_bist.get("etiket"),
            "kaynak": _bd_bist.get("kaynak") or "takvim",
        }
        for attr, val in (
            ("RAPOR_DONEM", _bd_bist.get("etiket")),
            ("BILANCO_YIL", _bd_bist.get("yil")),
            ("BILANCO_DONEM", _bd_bist.get("donem")),
        ):
            try:
                setattr(tefas_bist, attr, val)
            except Exception:
                pass
    except Exception:
        pass

    col_a, col_b = st.columns([1, 1])
    with col_a:
        if st.button("Seçili / örnek hisse İş Yatırım dönemini doğrula", key="btn_bist_isy_probe"):
            probe_kod = st.session_state.get("bist_probe_kod") or "ISDMR"
            with st.spinner(f"{probe_kod} İş Yatırım…"):
                try:
                    isy = isyatirim_mali_tablo(probe_kod, force=True)
                    if isy.get("error"):
                        st.warning(isy["error"])
                    else:
                        st.success(
                            f"**{probe_kod}** güncel kolon: **{isy.get('gunce_donem')}** "
                            f"(API dolu en yeni dönem). Raporlar buna göre."
                        )
                        st.session_state["bist_probe_isy"] = {
                            "kod": probe_kod,
                            "donem": isy.get("gunce_donem"),
                            "oranlar": isy.get("oranlar"),
                            "ozet_keys": list((isy.get("ozet") or {}).keys()),
                        }
                except Exception as e:
                    st.error(str(e))
    with col_b:
        st.text_input("Probe hisse kodu", value="ISDMR", key="bist_probe_kod")
        _pi = st.session_state.get("bist_probe_isy")
        if _pi:
            st.caption(f"Son probe: {_pi.get('kod')} → **{_pi.get('donem')}**")

    st.markdown("#### Tarihsel ucuzluk")
    st.caption(
        "Filtre Stüdyosu ile aynı mantık: 1Y/3Y fiyat persentili + F/K & PD/DD. "
        "Yüksek skor = tarihsel olarak daha ucuz fiyat/çarpan bölgesi."
    )
    tu_kod = st.text_input("Hisse kodu (ucuzluk)", value="ISDMR", key="bist_tu_kod")
    if st.button("Tarihsel ucuzluk hesapla", key="btn_bist_tu"):
        kod_u = (tu_kod or "").strip().upper().replace(".IS", "")
        with st.spinner(f"{kod_u}…"):
            fd_u, info_u, err_u = None, {}, None
            try:
                # yfinance fiyat
                import yfinance as yf
                t = yf.Ticker(f"{kod_u}.IS")
                hist = t.history(period="3y", auto_adjust=True)
                if hist is not None and not hist.empty:
                    hist = hist.reset_index()
                    dc = "Date" if "Date" in hist.columns else hist.columns[0]
                    fd_u = hist.rename(columns={dc: "tarih", "Close": "fiyat"})[["tarih", "fiyat"]]
                try:
                    info_u = dict(t.info or {})
                except Exception:
                    info_u = {}
            except Exception as e:
                err_u = str(e)
            deg_u = {}
            if info_u:
                # F/K / PD/DD yalnızca İş Yatırım
                try:
                    _ic = deg_from_isyatirim(kod_u)
                    if not _ic.get("_error"):
                        if _ic.get("F/K") is not None:
                            deg_u["F/K"] = _ic["F/K"]
                        if _ic.get("PD/DD") is not None:
                            deg_u["PD/DD"] = _ic["PD/DD"]
                except Exception:
                    pass
            # İş Yatırım oranları varsa ekle
            try:
                isy_u = isyatirim_mali_tablo(kod_u)
                if isy_u and not isy_u.get("error"):
                    st.caption(f"Bilanço dönemi (İş Yatırım): **{isy_u.get('gunce_donem')}**")
            except Exception:
                pass
            tu = tarihsel_ucuzluk_analizi(kod_u, fd=fd_u, deg=deg_u)
            st.session_state["_bist_tu"] = tu
            if err_u and tu.get("skor") is None:
                st.warning(err_u)
    tu = st.session_state.get("_bist_tu")
    if tu:
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Ucuzluk skoru", f"{tu.get('skor')}/100" if tu.get("skor") is not None else "—")
        c2.metric("1Y fiyat %ile", f"%{tu.get('fiyat_persentil_1y')}" if tu.get("fiyat_persentil_1y") is not None else "—")
        c3.metric("3Y fiyat %ile", f"%{tu.get('fiyat_persentil_3y')}" if tu.get("fiyat_persentil_3y") is not None else "—")
        c4.metric("Ucuz mu?", "Evet" if tu.get("ucuz_mu") else ("Hayır" if tu.get("ucuz_mu") is False else "—"))
        for y in tu.get("yorum") or []:
            st.markdown(f"- {y}")

    st.markdown("#### Hazır profil — birleşik skor + filtre")
    st.caption(
        "**Dengeli** · **Değer** (F/K≤15) · **Kalite** (≥55) · **Momentum** · **Taze bilanço**. "
        "Kalite+sinyal+ucuzluk eşit ağırlık. Tavsiye değildir."
    )
    bist_prof = st.radio(
        "Hazır profil",
        ["Dengeli", "Değer", "Kalite", "Momentum", "Taze bilanço", "Düşük kaldıraç"],
        horizontal=True,
        key="bist_hazir_profil",
    )
    bist_list = st.text_input(
        "Profil hisse listesi",
        value="ISDMR,THYAO,GARAN,EREGL,ASELS,SISE,KCHOL,TUPRS,BIMAS,SAHOL,ARCLK,TCELL",
        key="bist_birlesik_list",
    )
    if st.button("Profili uygula / sırala", type="primary", key="btn_bist_birlesik"):
        kodlar = [k.strip().upper().replace(".IS", "") for k in bist_list.split(",") if k.strip()]
        rows = []
        bar = st.progress(0.0)
        for i, k in enumerate(kodlar):
            try:
                rows.append(hisse_birlesik_satir(k))
            except Exception as e:
                rows.append({"kod": k, "ad": k, "birlesik_ortalama": None, "hata": str(e)})
            bar.progress((i + 1) / max(len(kodlar), 1))
        bar.empty()
        dfb = hazir_profil_filtrele(pd.DataFrame(rows), bist_prof)
        st.session_state["bist_birlesik_df"] = dfb
        st.session_state["bist_birlesik_profil"] = bist_prof
    dfb = st.session_state.get("bist_birlesik_df")
    if dfb is not None and not getattr(dfb, "empty", True):
        show_cols = [
            c for c in (
                "kod", "ad", "bilanco_donem", "kalite_skor", "model_sinyal",
                "sinyal_norm", "ucuzluk_skor", "fiyat_persentil_1y",
                "birlesik_ortalama", "seviye",
                "F/K", "PD/DD", "ROE", "getiri_1y",
            ) if c in dfb.columns
        ]
        show_df(dfb, show_cols, height=360)
        download_excel_button(
            dfb[show_cols] if show_cols else dfb,
            f"bist_birlesik_skor_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            meta={"mod": "bist_birlesik_profil", "siralama": "yuksek_dusuk"},
            key="xl_bist_birlesik",
        )
        st.caption("1. satır = en yüksek birleşik ortalama. **bilanco_donem** sütunu her hissenin İş Yatırım güncel dönemidir (ör. ISDMR → 2026/9).")

    try:
        try:
            tefas_bist.render(
                show_df=show_df,
                download_excel_button=download_excel_button,
                bilanco_donem=_bd_bist.get("etiket"),
                bilanco_yil=_bd_bist.get("yil"),
                bilanco_donem_no=_bd_bist.get("donem"),
            )
        except TypeError:
            tefas_bist.render(show_df=show_df, download_excel_button=download_excel_button)
    except Exception as e:
        st.error(f"BIST Değerleme yüklenemedi: {e}")
        st.caption("tefas_bist modülünü kontrol edin.")
        st.info(
            "Tek hisse için güncel bilanço + SWOT: **🔎 Hisse Detay** "
            "(İş Yatırım, örn. ISDMR → 2026/9)."
        )

elif sayfa == "🔎 Hisse Detay":
    st.header("🔎 Hisse Detay Analizi")
    st.caption(
        "Excel «Hisse Analiz» tarzı: fiyat · getiri · risk · değerleme · grafikler · SWOT. "
        "Mali tablo **İş Yatırım** üzerinden; hisse için API’de dolu **en yeni dönem** "
        "(ör. 2026/9 geldiyse 2026/6 kullanılmaz) tüm rapor ve skora işlenir. "
        "Kod örn. THYAO, GARAN, ISDMR."
    )

    def _hisse_sembol_normalize(raw: str) -> str:
        s = (raw or "").strip().upper().replace(" ", "")
        if not s:
            return ""
        if s.endswith(".IS"):
            return s
        # BIST kısa kod
        if len(s) <= 6 and s.replace(".", "").isalnum():
            return f"{s}.IS"
        return s

    def _df_from_yfinance(sembol: str, period: str = "3y"):
        """Yahoo Finance fallback — fiyat + info."""
        try:
            import yfinance as yf
        except ImportError:
            return None, None, "yfinance yüklü değil: pip install yfinance"
        try:
            t = yf.Ticker(sembol)
            hist = t.history(period=period, auto_adjust=True)
            if hist is None or hist.empty:
                return None, None, f"{sembol}: fiyat verisi yok"
            hist = hist.reset_index()
            # tarih kolonunu normalize
            date_col = "Date" if "Date" in hist.columns else hist.columns[0]
            hist = hist.rename(columns={date_col: "tarih", "Close": "fiyat", "Volume": "hacim"})
            hist["tarih"] = pd.to_datetime(hist["tarih"]).dt.tz_localize(None)
            info = {}
            try:
                info = dict(t.info or {})
            except Exception:
                info = {}
            return hist[["tarih", "fiyat"] + (["hacim"] if "hacim" in hist.columns else [])], info, None
        except Exception as e:
            return None, None, str(e)

    def _metrikler_from_fd(fd: pd.DataFrame, rf: float, cfg: dict) -> dict:
        """core.full_metrics + benchmark varsa attach."""
        m = {}
        try:
            m = core.full_metrics(fd, rf=rf) or {}
        except Exception:
            # minimal fallback
            try:
                s = fd.set_index("tarih")["fiyat"].astype(float).sort_index()
                s = s[~s.index.duplicated(keep="last")]
                rets = s.pct_change().dropna()
                if len(rets) < 5:
                    return {}
                def _g(n):
                    if len(s) <= n:
                        return None
                    return float((s.iloc[-1] / s.iloc[-n] - 1) * 100)
                m = {
                    "fiyat": float(s.iloc[-1]),
                    "getiri_1a": _g(22), "getiri_3a": _g(66),
                    "getiri_6a": _g(132), "getiri_1y": _g(252),
                    "volatilite": float(rets.std() * (252 ** 0.5) * 100),
                    "kazanma_orani": float((rets > 0).mean() * 100),
                }
                cum = (1 + rets).cumprod()
                dd = (cum / cum.cummax() - 1) * 100
                m["max_drawdown"] = float(dd.min())
                excess = rets - (rf / 252)
                vol = rets.std()
                m["sharpe"] = float((excess.mean() / vol) * (252 ** 0.5)) if vol > 0 else None
            except Exception:
                return {}
        try:
            m = core.attach_benchmark(m, core.fetch_benchmarks(cfg))
        except Exception:
            pass
        return m

    def _degerleme_from_info(info: dict) -> dict:
        """Yahoo info → Excel tarzı değerleme alanları."""
        if not info:
            return {}
        map_keys = {
            "longName": "şirket",
            "shortName": "kısa_ad",
            "sector": "sektör",
            "industry": "endüstri",
            "currency": "para_birimi",
            "currentPrice": "fiyat",
            "previousClose": "önceki_kapanış",
            "open": "açılış",
            "dayLow": "gün_düşük",
            "dayHigh": "gün_yüksek",
            "fiftyTwoWeekLow": "52h_düşük",
            "fiftyTwoWeekHigh": "52h_yüksek",
            "volume": "hacim",
            "averageVolume": "ort_hacim",
            "marketCap": "piyasa_değeri",
            "enterpriseValue": "firma_değeri",
            "trailingPE": "F/K",
            "forwardPE": "F/K_ileri",
            "priceToBook": "PD/DD",
            "enterpriseToEbitda": "FD/FAVÖK",
            "enterpriseToRevenue": "FD/Satış",
            "pegRatio": "PEG",
            "profitMargins": "net_kar_marjı",
            "grossMargins": "brüt_marj",
            "operatingMargins": "faaliyet_marjı",
            "returnOnEquity": "ROE",
            "returnOnAssets": "ROA",
            "debtToEquity": "borç/özkaynak",
            "currentRatio": "cari_oran",
            "quickRatio": "asit_test",
            "dividendYield": "temettü_getirisi",
            "payoutRatio": "temettü_ödeme",
            "beta": "beta",
            "trailingEps": "EPS",
            "bookValue": "defter_değeri",
            "targetMeanPrice": "hedef_fiyat",
            "recommendationKey": "tavsiye",
            "numberOfAnalystOpinions": "analist_sayısı",
            "sharesOutstanding": "hisse_adet",
            "floatShares": "serbest_dolaşım",
            "heldPercentInsiders": "içeriden_%",
            "heldPercentInstitutions": "kurumsal_%",
        }
        out = {}
        for k, tr in map_keys.items():
            v = info.get(k)
            if v is None:
                continue
            # oranlar yüzdeye
            if tr in ("net_kar_marjı", "brüt_marj", "faaliyet_marjı", "ROE", "ROA",
                      "temettü_getirisi", "temettü_ödeme", "içeriden_%", "kurumsal_%") and isinstance(v, (int, float)):
                # yahoo çoğu zaten 0-1
                if abs(v) <= 1.5:
                    v = v * 100
            out[tr] = v
        return out


    def _teknik_analiz(fd: pd.DataFrame) -> dict:
        """Fiyat serisinden teknik göstergeler (MA, RSI, trend, destek/direnç)."""
        out = {"error": None}
        if fd is None or getattr(fd, "empty", True) or "fiyat" not in (fd.columns if hasattr(fd, "columns") else []):
            out["error"] = "Fiyat serisi yok"
            return out
        try:
            s = fd.copy()
            s["tarih"] = pd.to_datetime(s["tarih"])
            s = s.sort_values("tarih").drop_duplicates("tarih", keep="last").set_index("tarih")
            px = s["fiyat"].astype(float)
            if len(px) < 30:
                out["error"] = "Yetersiz veri (<30 gün)"
                return out
            rets = px.pct_change().dropna()
            last = float(px.iloc[-1])
            out["fiyat"] = last
            out["n_gun"] = int(len(px))
            # Hareketli ortalamalar
            for w in (5, 22, 50, 100, 200):
                if len(px) >= w:
                    ma = float(px.rolling(w).mean().iloc[-1])
                    out[f"ma{w}"] = ma
                    out[f"ma{w}_uzaklik_%"] = round((last / ma - 1) * 100, 2) if ma else None
            # RSI(14)
            delta = px.diff()
            gain = delta.clip(lower=0).rolling(14).mean()
            loss = (-delta.clip(upper=0)).rolling(14).mean()
            rs = gain / loss.replace(0, np.nan)
            rsi = 100 - (100 / (1 + rs))
            out["rsi14"] = float(rsi.iloc[-1]) if pd.notna(rsi.iloc[-1]) else None
            # Getiriler
            for label, n in [("1h", 5), ("1a", 22), ("3a", 66), ("6a", 132), ("1y", 252)]:
                if len(px) > n:
                    out[f"getiri_{label}"] = float((px.iloc[-1] / px.iloc[-n] - 1) * 100)
            # 52 hafta
            w252 = px.tail(min(252, len(px)))
            out["52h_yuksek"] = float(w252.max())
            out["52h_dusuk"] = float(w252.min())
            out["52h_yuksek_uzak_%"] = round((last / out["52h_yuksek"] - 1) * 100, 2)
            out["52h_dusuk_uzak_%"] = round((last / out["52h_dusuk"] - 1) * 100, 2)
            # Trend yorumu
            rsi_v = out.get("rsi14")
            if rsi_v is not None:
                if rsi_v >= 70:
                    out["rsi_yorum"] = "Aşırı alım bölgesi (≥70)"
                elif rsi_v <= 30:
                    out["rsi_yorum"] = "Aşırı satım bölgesi (≤30)"
                else:
                    out["rsi_yorum"] = "Nötr bölge"
            # MA trend
            ma22 = out.get("ma22")
            ma50 = out.get("ma50")
            ma200 = out.get("ma200")
            sinyaller = []
            if ma22 and last > ma22:
                sinyaller.append("Fiyat > MA22 (kısa vade pozitif)")
            elif ma22:
                sinyaller.append("Fiyat < MA22 (kısa vade zayıf)")
            if ma50 and ma200:
                if ma50 > ma200:
                    sinyaller.append("MA50 > MA200 (orta vade yükseliş yapısı)")
                else:
                    sinyaller.append("MA50 < MA200 (orta vade zayıf yapı)")
            if ma200:
                if last > ma200:
                    sinyaller.append("Fiyat > MA200 (uzun vade trend üstü)")
                else:
                    sinyaller.append("Fiyat < MA200 (uzun vade trend altı)")
            out["sinyaller"] = sinyaller
            # Hacim trendi
            if "hacim" in s.columns and s["hacim"].notna().sum() > 20:
                vol = s["hacim"].astype(float)
                out["hacim_son"] = float(vol.iloc[-1])
                out["hacim_ort20"] = float(vol.tail(20).mean())
                if out["hacim_ort20"]:
                    out["hacim_oran"] = round(out["hacim_son"] / out["hacim_ort20"], 2)
            # Seri (grafik için MA)
            out["_px"] = px
            for w in (22, 50, 200):
                if len(px) >= w:
                    out[f"_ma{w}_seri"] = px.rolling(w).mean()
            return out
        except Exception as e:
            out["error"] = str(e)
            return out

    def _temel_mali_analiz(deg: dict, m: dict, bilanco: dict | None = None) -> dict:
        """
        Güncel bilanço dönemine göre temel + mali tablo değerlendirmesi.
        Yahoo/info + performans metrikleri birleştirilir.
        """
        bilanco = bilanco or {}
        donem = bilanco.get("etiket") or "güncel"
        aciklama = bilanco.get("aciklama") or ""
        notlar = []
        skor = 50  # 0-100 nötr başlangıç
        detay = {
            "donem": donem,
            "donem_aciklama": aciklama,
            "carpanlar": {},
            "karlilik": {},
            "bilanco_saglik": {},
            "notlar": notlar,
            "skor": skor,
            "skor_yorum": "",
        }

        def _f(v):
            try:
                if v is None:
                    return None
                x = float(v)
                return x if x == x else None
            except (TypeError, ValueError):
                return None

        fk = _f(deg.get("F/K"))
        pddd = _f(deg.get("PD/DD"))
        fdf = _f(deg.get("FD/FAVÖK"))
        peg = _f(deg.get("PEG"))
        roe = _f(deg.get("ROE"))
        roa = _f(deg.get("ROA"))
        net_m = _f(deg.get("net_kar_marjı"))
        brut = _f(deg.get("brüt_marj"))
        faal = _f(deg.get("faaliyet_marjı"))
        borc = _f(deg.get("borç/özkaynak"))
        cari = _f(deg.get("cari_oran"))
        beta = _f(deg.get("beta"))
        tem = _f(deg.get("temettü_getirisi"))

        detay["carpanlar"] = {
            "F/K": fk, "PD/DD": pddd, "FD/FAVÖK": fdf, "PEG": peg,
            "Temettü %": tem, "Beta": beta,
        }
        detay["karlilik"] = {
            "ROE %": roe, "ROA %": roa,
            "Net kâr marjı %": net_m, "Brüt marj %": brut, "Faaliyet marjı %": faal,
        }
        detay["bilanco_saglik"] = {
            "Borç/Özkaynak": borc, "Cari oran": cari,
            "Asit test": _f(deg.get("asit_test")),
        }

        # Skor / notlar (basit kural seti — tavsiye değil)
        if fk is not None:
            if fk < 8:
                notlar.append(f"F/K {fk:.1f} düşük çarpan bandında (sektör ile kıyaslayın).")
                skor += 8
            elif fk > 25:
                notlar.append(f"F/K {fk:.1f} yüksek — büyüme fiyatlanmış olabilir.")
                skor -= 6
            else:
                notlar.append(f"F/K {fk:.1f} orta bant.")
        if pddd is not None:
            if pddd < 1:
                notlar.append(f"PD/DD {pddd:.2f} — defter değerinin altında.")
                skor += 6
            elif pddd > 3:
                notlar.append(f"PD/DD {pddd:.2f} — deftere göre primli.")
                skor -= 4
        if roe is not None:
            if roe >= 15:
                notlar.append(f"ROE %{roe:.1f} güçlü özkaynak kârlılığı.")
                skor += 10
            elif roe < 5:
                notlar.append(f"ROE %{roe:.1f} zayıf.")
                skor -= 8
            else:
                notlar.append(f"ROE %{roe:.1f} orta.")
        if net_m is not None:
            if net_m >= 10:
                notlar.append(f"Net kâr marjı %{net_m:.1f} iyi.")
                skor += 6
            elif net_m < 3:
                notlar.append(f"Net kâr marjı %{net_m:.1f} düşük.")
                skor -= 5
        if borc is not None:
            # yahoo debtToEquity bazen yüzde (53 = 0.53)
            b = borc / 100.0 if borc > 5 else borc
            if b <= 0.5:
                notlar.append(f"Borç/özkaynak {b:.2f} düşük kaldıraç.")
                skor += 5
            elif b >= 1.5:
                notlar.append(f"Borç/özkaynak {b:.2f} yüksek kaldıraç — dikkat.")
                skor -= 7
        if cari is not None:
            if cari >= 1.5:
                notlar.append(f"Cari oran {cari:.2f} likidite rahat.")
                skor += 4
            elif cari < 1:
                notlar.append(f"Cari oran {cari:.2f} kısa vadeli baskı riski.")
                skor -= 6
        # Performans katkısı
        g1y = _f(m.get("getiri_1y"))
        if g1y is not None:
            if g1y > 20:
                notlar.append(f"1Y getiri %{g1y:.0f} güçlü.")
                skor += 5
            elif g1y < -20:
                notlar.append(f"1Y getiri %{g1y:.0f} zayıf.")
                skor -= 5
        sh = _f(m.get("sharpe"))
        if sh is not None and sh >= 1:
            notlar.append(f"Sharpe {sh:.2f} risk-ayarlı getiri iyi.")
            skor += 4

        skor = max(0, min(100, skor))
        detay["skor"] = skor
        if skor >= 70:
            detay["skor_yorum"] = "Temel görünüm görece olumlu (kurallar skoru)."
        elif skor >= 45:
            detay["skor_yorum"] = "Temel görünüm karışık / nötr."
        else:
            detay["skor_yorum"] = "Temel görünüm temkinli — çarpan/kârlılık/borç kontrol edin."
        notlar.insert(0, f"Mali tablo referans dönemi: **{donem}** ({aciklama or 'son bilanço'}).")
        detay["notlar"] = notlar
        return detay


    def _swot_yatirim_analizi(
        kod: str,
        m: dict,
        deg: dict,
        tem: dict | None,
        tek: dict | None,
        isy: dict | None,
        bil: dict | None,
    ) -> dict:
        """
        Tüm metrikleri birleştirerek bilanço +/- ve SWOT üretir.
        Yatırım tavsiyesi değildir; yapılandırılmış yorum.
        """
        tem = tem or {}
        tek = tek or {}
        isy = isy or {}
        bil = bil or {}
        S, W, O, T = [], [], [], []
        olumlu, olumsuz = [], []

        def _f(v):
            try:
                if v is None:
                    return None
                x = float(v)
                return x if x == x else None
            except (TypeError, ValueError):
                return None

        donem = (
            isy.get("gunce_donem")
            or bil.get("etiket")
            or tem.get("donem")
            or "?"
        )
        oz = isy.get("ozet") or {}
        oran = isy.get("oranlar") or {}

        # ── Bilanço / gelir (İş Yatırım) ──────────────────────────────────
        net_chg = (oz.get("net_donem_kari") or {}).get("degisim_%")
        has_chg = (oz.get("hasilat") or {}).get("degisim_%")
        oz_chg = (oz.get("ozkaynaklar") or {}).get("degisim_%")
        ef_chg = (oz.get("esas_faaliyet_kari") or {}).get("degisim_%")
        net_v = (oz.get("net_donem_kari") or {}).get("deger")
        has_v = (oz.get("hasilat") or {}).get("deger")
        oz_v = (oz.get("ozkaynaklar") or {}).get("deger")

        if net_chg is not None:
            if net_chg >= 20:
                olumlu.append(f"Net kâr YoY **%{net_chg:.0f}** arttı ({donem}).")
                S.append(f"Güçlü kâr büyümesi (YoY %{net_chg:.0f})")
            elif net_chg <= -15:
                olumsuz.append(f"Net kâr YoY **%{net_chg:.0f}** geriledi.")
                W.append(f"Kâr daralması (YoY %{net_chg:.0f})")
            else:
                olumlu.append(f"Net kâr YoY değişim %{net_chg:.0f} (ılımlı).")
        if has_chg is not None:
            if has_chg >= 10:
                olumlu.append(f"Hasılat YoY **%{has_chg:.0f}** büyüdü.")
                S.append(f"Satış büyümesi %{has_chg:.0f}")
            elif has_chg < 0:
                olumsuz.append(f"Hasılat YoY **%{has_chg:.0f}** küçüldü.")
                W.append(f"Ciro gerilemesi %{has_chg:.0f}")
        if oz_chg is not None and oz_chg >= 5:
            olumlu.append(f"Özkaynak **%{oz_chg:.0f}** güçlendi.")
            S.append("Özkaynak artışı")
        if ef_chg is not None and ef_chg >= 15:
            S.append(f"Esas faaliyet kârı YoY %{ef_chg:.0f}")
            olumlu.append(f"Operasyonel kâr YoY **%{ef_chg:.0f}**.")
        elif ef_chg is not None and ef_chg < -10:
            W.append(f"Operasyonel kâr zayıf (YoY %{ef_chg:.0f})")
            olumsuz.append(f"Esas faaliyet kârı YoY **%{ef_chg:.0f}**.")

        roe = _f(oran.get("ROE_%")) or _f(deg.get("ROE"))
        roa = _f(oran.get("ROA_%")) or _f(deg.get("ROA"))
        marj = _f(oran.get("net_kar_marji_%")) or _f(deg.get("net_kar_marjı"))
        borc = _f(oran.get("borc_ozkaynak")) or _f(deg.get("borç/özkaynak"))
        if borc is not None and borc > 5:
            borc = borc / 100.0
        cari = _f(oran.get("cari_oran")) or _f(deg.get("cari_oran"))

        if roe is not None:
            if roe >= 15:
                olumlu.append(f"ROE **%{roe:.1f}** — güçlü özkaynak kârlılığı.")
                S.append(f"ROE %{roe:.1f}")
            elif roe < 5:
                olumsuz.append(f"ROE **%{roe:.1f}** düşük.")
                W.append(f"Zayıf ROE %{roe:.1f}")
        if marj is not None:
            if marj >= 10:
                S.append(f"Net marj %{marj:.1f}")
                olumlu.append(f"Net kâr marjı **%{marj:.1f}**.")
            elif marj < 3:
                W.append(f"İnce marj %{marj:.1f}")
                olumsuz.append(f"Net marj **%{marj:.1f}** zayıf.")
        if borc is not None:
            if borc <= 0.5:
                S.append(f"Düşük kaldıraç ({borc:.2f})")
                olumlu.append(f"Borç/özkaynak **{borc:.2f}** düşük.")
            elif borc >= 1.5:
                W.append(f"Yüksek kaldıraç {borc:.2f}")
                olumsuz.append(f"Borç/özkaynak **{borc:.2f}** yüksek — faiz/refinansman riski.")
                T.append("Borç yükü / faiz hassasiyeti")
        if cari is not None:
            if cari >= 1.5:
                S.append(f"Cari oran {cari:.2f}")
                olumlu.append(f"Cari oran **{cari:.2f}** — kısa vadeli likidite rahat.")
            elif cari < 1:
                W.append(f"Cari oran {cari:.2f}")
                olumsuz.append(f"Cari oran **{cari:.2f}** — kısa vadeli yükümlülük baskısı.")
                T.append("Likidite / kısa vadeli borç baskısı")

        # ── Çarpanlar ─────────────────────────────────────────────────────
        fk = _f(deg.get("F/K"))
        pddd = _f(deg.get("PD/DD"))
        fdf = _f(deg.get("FD/FAVÖK"))
        temettu = _f(deg.get("temettü_getirisi"))
        if fk is not None:
            if fk < 8:
                O.append(f"F/K {fk:.1f} düşük çarpan — değer fırsatı olabilir (sektör ile kıyaslayın)")
                olumlu.append(f"F/K **{fk:.1f}** görece ucuz band.")
            elif fk > 25:
                T.append(f"Yüksek F/K ({fk:.1f}) — beklenti fiyatlanmış olabilir")
                olumsuz.append(f"F/K **{fk:.1f}** yüksek.")
            else:
                olumlu.append(f"F/K **{fk:.1f}** orta bant.")
        if pddd is not None:
            if pddd < 1:
                O.append(f"PD/DD {pddd:.2f} < 1 — defter altı")
                olumlu.append(f"PD/DD **{pddd:.2f}** defter değerinin altında.")
            elif pddd > 3:
                T.append(f"PD/DD {pddd:.2f} primli")
                olumsuz.append(f"PD/DD **{pddd:.2f}** primli.")
        if temettu is not None and temettu >= 5:
            S.append(f"Temettü verimi %{temettu:.1f}")
            olumlu.append(f"Temettü getirisi **%{temettu:.1f}**.")
            O.append("Temettü geliri destekleyici")

        # ── Performans / risk ─────────────────────────────────────────────
        g1y = _f(m.get("getiri_1y"))
        g3a = _f(m.get("getiri_3a"))
        sh = _f(m.get("sharpe"))
        dd = _f(m.get("max_drawdown"))
        vol = _f(m.get("volatilite"))
        if g1y is not None:
            if g1y >= 30:
                S.append(f"1Y getiri %{g1y:.0f}")
                olumlu.append(f"1 yıllık getiri **%{g1y:.0f}** güçlü.")
            elif g1y <= -20:
                W.append(f"1Y getiri %{g1y:.0f}")
                olumsuz.append(f"1 yıllık getiri **%{g1y:.0f}** zayıf.")
        if sh is not None:
            if sh >= 1:
                S.append(f"Sharpe {sh:.2f}")
                olumlu.append(f"Sharpe **{sh:.2f}** — risk-ayarlı getiri iyi.")
            elif sh < 0:
                W.append(f"Sharpe {sh:.2f}")
                olumsuz.append(f"Sharpe **{sh:.2f}** negatif.")
        if dd is not None:
            if dd <= -40:
                W.append(f"MaxDD %{dd:.0f}")
                olumsuz.append(f"Maksimum düşüş **%{dd:.0f}** — yüksek çizim riski.")
                T.append("Derin drawdown geçmişi")
            elif dd > -15:
                S.append(f"Sığ MaxDD %{dd:.0f}")
        if vol is not None and vol >= 50:
            T.append(f"Yüksek volatilite %{vol:.0f}")
            olumsuz.append(f"Volatilite **%{vol:.0f}** yüksek.")

        # ── Teknik ────────────────────────────────────────────────────────
        rsi = _f(tek.get("rsi14"))
        if rsi is not None:
            if rsi <= 30:
                O.append(f"RSI {rsi:.0f} aşırı satım bölgesi — teknik tepki potansiyeli")
            elif rsi >= 70:
                T.append(f"RSI {rsi:.0f} aşırı alım — kısa vadeli düzeltme riski")
        for sig in (tek.get("sinyaller") or [])[:4]:
            if "pozitif" in sig.lower() or "üst" in sig.lower() or ">" in sig:
                if "zayıf" not in sig.lower() and "altı" not in sig.lower():
                    S.append(f"Teknik: {sig}")
                else:
                    W.append(f"Teknik: {sig}")
            elif "zayıf" in sig.lower() or "altı" in sig.lower() or "<" in sig:
                W.append(f"Teknik: {sig}")

        # ── Sektör / genel tehdit-fırsat ───────────────────────────────────
        O.append("Güncel KAP / İş Yatırım bilançosu ile oranlar yenilendi")
        T.append("Makro: faiz, kur, emtia ve sektör döngüsü fiyatı etkileyebilir")
        T.append("Tek hisse riski — portföy çeşitlendirmesi gerekir")

        # Tekrarları sadeleştir
        def _uniq(xs):
            seen, out = set(), []
            for x in xs:
                k = str(x)
                if k not in seen:
                    seen.add(k)
                    out.append(x)
            return out

        S, W, O, T = map(_uniq, (S, W, O, T))
        olumlu, olumsuz = _uniq(olumlu), _uniq(olumsuz)

        skor = _f(tem.get("skor"))
        if skor is None:
            skor = 50 + 4 * len(S) - 4 * len(W)
            skor = max(0, min(100, skor))

        if skor >= 70:
            sonuc = (
                f"**Özet:** `{kod}` için güncel bilanço (**{donem}**) ve çarpan/performans "
                "birlikte **görece olumlu** bir tablo çiziyor. Yine de sektör ve makro riskler "
                "SWOT’taki tehditlerle birlikte izlenmeli."
            )
            ton = "olumlu"
        elif skor >= 45:
            sonuc = (
                f"**Özet:** `{kod}` için tablo **karışık / nötr**. "
                f"Bilanço dönemi **{donem}**; güçlü ve zayıf yanlar dengeli görünüyor. "
                "Pozisyon boyutu ve giriş seviyesi kritik."
            )
            ton = "notr"
        else:
            sonuc = (
                f"**Özet:** `{kod}` için temel/teknik göstergeler **temkinli** bir görünüm sunuyor "
                f"(dönem **{donem}**). Zayıf yanlar ve tehditler ağır basıyor olabilir; "
                "ek araştırma ve risk yönetimi şart."
            )
            ton = "temkinli"

        sonuc += (
            "\n\n*Bu metin otomatik kural setiyle üretilmiştir; **yatırım tavsiyesi değildir**. "
            "Kendi risk profiliniz ve güncel KAP dipnotlarıyla doğrulayın.*"
        )

        return {
            "donem": donem,
            "olumlu": olumlu,
            "olumsuz": olumsuz,
            "S": S, "W": W, "O": O, "T": T,
            "skor": skor,
            "ton": ton,
            "sonuc": sonuc,
        }

    def _try_core_bilanco(kod: str) -> dict:
        """Önce İş Yatırım MaliTablo; yoksa core/tefas_bist. Günlük güncel dönem."""
        try:
            force = bool(st.session_state.get("_isy_force_refresh"))
            isy = isyatirim_mali_tablo(kod, force=force)
            st.session_state["_isy_force_refresh"] = False
            if isy and not isy.get("error"):
                return isy
        except Exception as e:
            isy = {"error": str(e)}
        for fn_name in ("hisse_bilanco_ozet", "bilanco_ozet", "hisse_mali_tablo"):
            fn = getattr(core, fn_name, None)
            if callable(fn):
                try:
                    r = fn(kod)
                    if isinstance(r, dict) and r:
                        return r
                except Exception:
                    pass
        try:
            import tefas_bist as tb
            for fn_name in ("hisse_bilanco_ozet", "bilanco_ozet", "mali_tablo"):
                fn = getattr(tb, fn_name, None)
                if callable(fn):
                    try:
                        r = fn(kod)
                        if isinstance(r, dict) and r:
                            return r
                    except Exception:
                        pass
        except Exception:
            pass
        return isy if isinstance(isy, dict) else {}

    def _draw_hisse_charts(fd: pd.DataFrame, m: dict, sembol: str, bm_rets=None):
        """Fiyat, kümülatif getiri, drawdown, hacim grafikleri."""
        if fd is None or fd.empty or "fiyat" not in fd.columns:
            st.info("Grafik için fiyat serisi yok.")
            return
        s = fd.copy()
        s["tarih"] = pd.to_datetime(s["tarih"])
        s = s.sort_values("tarih").drop_duplicates("tarih", keep="last")
        s = s.set_index("tarih")
        px = s["fiyat"].astype(float)
        rets = px.pct_change().dropna()
        cum = (1 + rets).cumprod()
        peak = cum.cummax()
        dd = (cum / peak - 1.0) * 100

        g1, g2 = st.columns(2)
        with g1:
            st.markdown("**Fiyat (TL)**")
            st.line_chart(px.rename("Fiyat"))
        with g2:
            st.markdown("**Kümülatif getiri (normalize=1)**")
            chart_df = cum.rename("Hisse").to_frame()
            if bm_rets is not None and len(bm_rets):
                try:
                    bm = bm_rets.reindex(cum.index).fillna(0)
                    chart_df["BIST100"] = (1 + bm).cumprod()
                    # hizala başlangıç
                    chart_df = chart_df / chart_df.iloc[0]
                except Exception:
                    pass
            st.line_chart(chart_df)
        g3, g4 = st.columns(2)
        with g3:
            st.markdown("**Drawdown %**")
            st.area_chart(dd.rename("DD %"))
        with g4:
            if "hacim" in s.columns and s["hacim"].notna().any():
                st.markdown("**İşlem hacmi**")
                vol = s["hacim"].astype(float)
                st.bar_chart(vol.rename("Hacim").tail(120))
            else:
                st.markdown("**Aylık getiri % (son 24 ay)**")
                try:
                    monthly = px.resample("ME").last().pct_change().dropna() * 100
                    st.bar_chart(monthly.tail(24).rename("Aylık %"))
                except Exception:
                    st.caption("Aylık getiri hesaplanamadı")

        # Getiri özeti bar
        st.markdown("**Dönem getirileri %**")
        donem = {}
        for k, label in [
            ("getiri_1a", "1A"), ("getiri_3a", "3A"), ("getiri_6a", "6A"),
            ("getiri_1y", "1Y"), ("cagr", "CAGR"),
        ]:
            if m.get(k) is not None:
                try:
                    donem[label] = float(m[k])
                except (TypeError, ValueError):
                    pass
        if donem:
            st.bar_chart(pd.Series(donem, name="Getiri %"))

    # ── Girdi ──────────────────────────────────────────────────────────────
    c_in, c_btn, c_per = st.columns([3, 1, 1])
    with c_in:
        raw_kod = st.text_input(
            "Hisse kodu veya adı",
            value=st.session_state.get("hisse_detay_kod", "THYAO"),
            placeholder="THYAO  ·  GARAN  ·  ASELS.IS",
            key="hisse_detay_input",
        ).strip()
    with c_per:
        period = st.selectbox("Periyot", ["1y", "2y", "3y", "5y", "max"], index=2, key="hisse_det_per")
    with c_btn:
        st.write("")
        st.write("")
        run = st.button("🔍 Analiz et", type="primary", use_container_width=True, key="btn_hisse_detay")

    # Mali tablo zorla yenile → otomatik yeniden analiz
    if st.session_state.pop("_hisse_auto_analiz", False) and not run:
        run = True
        if st.session_state.get("hisse_kod_input"):
            raw_kod = st.session_state["hisse_kod_input"]

    if run and raw_kod:
        st.session_state.hisse_detay_kod = raw_kod
        sembol = _hisse_sembol_normalize(raw_kod)
        analiz = get_analiz()
        det = None
        err = None
        with st.spinner(f"{sembol} analiz ediliyor…"):
            # 1) core.hisse_analiz
            try:
                det = core.hisse_analiz(sembol, rf=analiz.rf, cfg=st.session_state.cfg)
                if isinstance(det, dict) and det.get("error"):
                    err = det.get("error")
                    det = None
            except Exception as e:
                err = str(e)
                det = None
            # 2) yfinance fallback
            info = {}
            if det is None:
                fd, info, yerr = _df_from_yfinance(sembol, period=period)
                if yerr:
                    st.error(yerr if not err else f"{err} · {yerr}")
                    st.stop()
                m = _metrikler_from_fd(fd, analiz.rf, st.session_state.cfg)
                det = {
                    "kod": sembol.replace(".IS", ""),
                    "sembol": sembol,
                    "fiyat_df": fd,
                    "metrikler": m,
                    "info": info,
                    "kaynak": "yfinance",
                }
            else:
                # core başarılı — info için yine de yf dene
                try:
                    _, info, _ = _df_from_yfinance(sembol, period=period)
                    det["info"] = info or det.get("info") or {}
                except Exception:
                    det["info"] = det.get("info") or {}
                if det.get("fiyat_df") is None or (hasattr(det.get("fiyat_df"), "empty") and det["fiyat_df"].empty):
                    fd, info2, _ = _df_from_yfinance(sembol, period=period)
                    if fd is not None:
                        det["fiyat_df"] = fd
                        if info2:
                            det["info"] = info2
        st.session_state.hisse_detay_sonuc = det

    det = st.session_state.get("hisse_detay_sonuc")
    if not det:
        empty_state("Henüz hisse seçilmedi", "Kod girip Analiz et ile İş Yatırım bilanço ve skorlar gelir.", "Örnek: ISDMR · THYAO · GARAN · EREGL")
    else:
        m = det.get("metrikler") or {}
        info = det.get("info") or {}
        deg = _degerleme_from_info(info)
        kod = det.get("kod") or ""
        # İş Yatırım çarpanları Yahoo üzerine yazar (F/K hatasını önler)
        try:
            _isy_c = deg_from_isyatirim(kod)
            if not _isy_c.get("_error"):
                for _ck in ("F/K", "PD/DD", "FD/FAVÖK", "FD/Satış", "fiyat", "piyasa_değeri"):
                    if _isy_c.get(_ck) is not None:
                        deg[_ck] = _isy_c[_ck]
                deg["_kaynak_carpan"] = "İş Yatırım"
        except Exception:
            pass
        sembol = det.get("sembol") or kod
        ad = deg.get("şirket") or deg.get("kısa_ad") or info.get("longName") or sembol
        sektor = deg.get("sektör") or info.get("sector") or "—"
        endustri = deg.get("endüstri") or info.get("industry") or "—"

        st.subheader(f"`{kod}` — {ad}")
        try:
            _bd = son_bilanco_donemi(probe=False)
            st.caption(
                f"{sembol} · {sektor} · {endustri} · kaynak: {det.get('kaynak') or 'core'} · "
                f"**Mali tablo dönemi (takvim): {_bd.get('etiket')} ({_bd.get('aciklama')})** — "
                "hisseye özel güncel dönem Mali Tablo sekmesinde İş Yatırım’dan gelir"
            )
            if _bd.get("yeni_donem"):
                st.warning(
                    f"Yeni bilanço dönemi: **{_bd.get('etiket')}**"
                    + (f" — önceki {_bd['onceki']['etiket']}" if _bd.get("onceki") else "")
                    + ". Raporlar bu döneme göre güncellenir."
                )
        except Exception:
            st.caption(f"{sembol} · {sektor} · {endustri} · kaynak: {det.get('kaynak') or 'core'}")

        # KPI şeridi
        k1, k2, k3, k4, k5, k6 = st.columns(6)
        k1.metric("Fiyat", fmt_num(m.get("fiyat") or deg.get("fiyat"), 2))
        k2.metric("1Y %", fmt_num(m.get("getiri_1y")))
        k3.metric("Sharpe", fmt_num(m.get("sharpe"), 2))
        k4.metric("MaxDD %", fmt_num(m.get("max_drawdown")))
        k5.metric("Vol %", fmt_num(m.get("volatilite")))
        k6.metric("Beta", fmt_num(deg.get("beta"), 2) if deg.get("beta") is not None else fmt_num(m.get("bist_korelasyon"), 2))

        k7, k8, k9, k10, k11, k12 = st.columns(6)
        k7.metric("3A %", fmt_num(m.get("getiri_3a")))
        k8.metric("6A %", fmt_num(m.get("getiri_6a")))
        k9.metric("Reel 1Y", fmt_num(m.get("reel_1y")))
        k10.metric("vs BIST", fmt_num(m.get("vs_bist100")))
        k11.metric("F/K", fmt_num(deg.get("F/K"), 1))
        k12.metric("PD/DD", fmt_num(deg.get("PD/DD"), 2))

        # Önceki bilanço / metrik kaydı ile karşılaştır → iyileşme uyarısı
        _cmp = dict(m)
        _cmp.update({k: deg.get(k) for k in ("ROE", "ROA", "F/K", "PD/DD", "borç/özkaynak", "net_kar_marjı") if deg.get(k) is not None})
        try:
            _iy = mali_tablo_iyilesme_kontrolu(kod or sembol, _cmp, tur="hisse")
            if _iy:
                st.session_state["_bilanco_iyilesme"] = (
                    list(st.session_state.get("_bilanco_iyilesme") or []) + _iy
                )[-30:]
                for _msg in _iy:
                    st.success(_msg)
            else:
                st.caption("Önceki bilanço kaydı yok veya anlamlı değişiklik yok (ilk analiz kayıt altına alındı).")
        except Exception as _e:
            st.caption(f"İyileşme kontrolü: {_e}")

        # Güncel bilanço dönemi + teknik + temel
        try:
            _bil = son_bilanco_donemi(probe=False)
        except Exception:
            _bil = {"etiket": "?", "aciklama": ""}
        _tek = _teknik_analiz(det.get("fiyat_df"))
        _core_bil = _try_core_bilanco(kod)
        # İş Yatırım oranlarını deg'e işle → temel analiz güncel bilançoya dayansın
        if _core_bil and not _core_bil.get("error"):
            _or = _core_bil.get("oranlar") or {}
            if _or.get("ROE_%") is not None:
                deg["ROE"] = _or["ROE_%"]
            if _or.get("ROA_%") is not None:
                deg["ROA"] = _or["ROA_%"]
            if _or.get("net_kar_marji_%") is not None:
                deg["net_kar_marjı"] = _or["net_kar_marji_%"]
            if _or.get("borc_ozkaynak") is not None:
                deg["borç/özkaynak"] = _or["borc_ozkaynak"]
            if _or.get("cari_oran") is not None:
                deg["cari_oran"] = _or["cari_oran"]
            if _core_bil.get("gunce_donem"):
                _bil = dict(_bil)
                _bil["etiket"] = _core_bil["gunce_donem"]
                _bil["yil"] = _core_bil.get("gunce_yil")
                _bil["donem"] = _core_bil.get("gunce_donem_no")
                _bil["aciklama"] = (
                    "yıllık" if _core_bil.get("gunce_donem_no") == 12
                    else f"{_core_bil.get('gunce_donem_no')} aylık"
                )
                _bil["kaynak"] = "İş Yatırım"
                _bil["guncel_mi"] = True
        _tem = _temel_mali_analiz(deg, m, bilanco=_bil)
        det["teknik"] = {k: v for k, v in _tek.items() if not str(k).startswith("_")}
        det["temel"] = _tem
        det["bilanco_donem"] = _bil
        det["isyatirim"] = {
            k: v for k, v in (_core_bil or {}).items()
            if k not in ("bilanco_df", "gelir_df", "nakit_df")
        }

        # Skorlar + sticky özet şeridi
        try:
            _kal_s = bilanco_kalite_skoru(_core_bil if isinstance(_core_bil, dict) else {}, deg)
            _tu_s = tarihsel_ucuzluk_analizi(kod, fd=det.get("fiyat_df"), deg=deg)
            _sig_s = model_sinyal_skoru(m=m, deg=deg, kalite=_kal_s, tu=_tu_s)
            _bir_s = birlesik_profil_skoru(kalite=_kal_s, sinyal=_sig_s, tu=_tu_s)
            _mom_s = bilanco_momentum(_core_bil if isinstance(_core_bil, dict) else {})
        except Exception:
            _kal_s, _tu_s, _sig_s, _bir_s, _mom_s = {}, {}, {}, {}, {}

        _fk_v = deg.get("F/K")
        _pddd_v = deg.get("PD/DD")
        _don_v = (_bil or {}).get("etiket") or "—"
        _bir_v = (_bir_s or {}).get("birlesik")
        _kal_v = (_kal_s or {}).get("skor")
        _mom_v = (_mom_s or {}).get("skor")

        st.markdown(
            (
                '<div class="sticky-strip">'
                f'<div class="sticky-item"><div class="sk">Kod</div><div class="sv">{kod}</div></div>'
                f'<div class="sticky-item"><div class="sk">Bilanço</div><div class="sv">{_don_v}</div></div>'
                f'<div class="sticky-item"><div class="sk">F/K</div><div class="sv">{tr_num(_fk_v, 1) if _fk_v is not None else "—"}</div></div>'
                f'<div class="sticky-item"><div class="sk">PD/DD</div><div class="sv">{tr_num(_pddd_v, 2) if _pddd_v is not None else "—"}</div></div>'
                f'<div class="sticky-item"><div class="sk">Birleşik</div><div class="sv">{tr_num(_bir_v, 0) if _bir_v is not None else "—"}</div></div>'
                f'<div class="sticky-item"><div class="sk">Kalite</div><div class="sv">{tr_num(_kal_v, 0) if _kal_v is not None else "—"}</div></div>'
                f'<div class="sticky-item"><div class="sk">Momentum</div><div class="sv">{tr_num(_mom_v, 0) if _mom_v is not None else "—"}</div></div>'
                f'<div class="sticky-item">{badge_fk(_fk_v)}</div>'
                '</div>'
            ),
            unsafe_allow_html=True,
        )

        # Bilanço momentum özeti (analiz aracı)
        if _mom_s and _mom_s.get("satirlar"):
            with st.expander(f"Bilanço momentumu · skor {tr_num(_mom_v, 0)}/100 · {_mom_s.get('ozet') or ''}", expanded=False):
                mdf = pd.DataFrame(_mom_s["satirlar"]).copy()
                if "Güncel" in mdf.columns:
                    mdf["Güncel"] = mdf["Güncel"].map(lambda v: _fmt_tl(v))
                if "YoY dönem" in mdf.columns:
                    mdf["YoY dönem"] = mdf["YoY dönem"].map(lambda v: _fmt_tl(v))
                if "YoY %" in mdf.columns:
                    mdf["YoY %"] = mdf["YoY %"].map(lambda v: tr_num(v, 1, force_sign=True) if v is not None else "—")
                st.dataframe(mdf, hide_index=True, use_container_width=True)
                st.caption("YoY = önceki yıl aynı dönem. Yükümlülük artışında yön ters yorumlanır.")

        tab_oz, tab_sk, tab_g, tab_mt, tab_tem, tab_tek, tab_r, tab_d, tab_v, tab_x = st.tabs([
            "🏢 Şirket Özeti",
            "📋 Skor Kart",
            "📈 Grafikler",
            "📑 Mali Tablo",
            "🏛️ Temel Analiz",
            "📉 Teknik Analiz",
            "📊 Getiri & Risk",
            "💰 Değerleme",
            "📝 Değerlendirme",
            "📥 Excel / HTML",
        ])

        # ── Fintables tarzı şirket özeti ──────────────────────────────────
        with tab_oz:
            st.markdown("#### Şirket özeti · Fintables tarzı")
            st.caption(
                "Özet kart, karne (kârlılık / büyüme / borç / değerleme), çarpanlar ve oranlar. "
                "Veri: **İş Yatırım**. [Fintables örnek](https://fintables.com/sirketler/A1CAP) benzeri düzen."
            )
            _karne = fintables_tarzi_karne(_core_bil if isinstance(_core_bil, dict) else {}, deg, m)
            det["fintables_karne"] = _karne

            # Üst KPI şeridi
            k1, k2, k3, k4, k5, k6 = st.columns(6)
            k1.metric("Fiyat", fmt_num(deg.get("fiyat") or m.get("fiyat"), 2))
            k2.metric("F/K", fmt_num(deg.get("F/K"), 1))
            k3.metric("PD/DD", fmt_num(deg.get("PD/DD"), 2))
            k4.metric("FD/FAVÖK", fmt_num(deg.get("FD/FAVÖK"), 1))
            k5.metric("Piyasa değeri", _fmt_tl(deg.get("piyasa_değeri")) if deg.get("piyasa_değeri") else "—")
            k6.metric("Bilanço", _core_bil.get("gunce_donem") or _bil.get("etiket") or "—")
            if deg.get("_kaynak_carpan") or deg.get("_kaynak"):
                st.caption(f"Çarpan kaynağı: **{deg.get('_kaynak_carpan') or deg.get('_kaynak')}** (Yahoo F/K kullanılmaz)")

            st.markdown("##### Fintables tarzı karne")
            c1, c2, c3, c4, c5 = st.columns(5)
            c1.metric("Genel", f"{_karne.get('genel')}/100", _karne.get("harf"))
            c2.metric("Kârlılık", f"{_karne.get('karlilik')}/100")
            c3.metric("Büyüme", f"{_karne.get('buyume')}/100")
            c4.metric("Borçluluk*", f"{_karne.get('borcluluk')}/100")
            c5.metric("Değerleme*", f"{_karne.get('degerleme')}/100")
            st.caption("*Borçluluk: yüksek skor = daha sağlıklı borç/likidite. Değerleme: yüksek = görece ucuz çarpan.")
            st.info(_karne.get("yorum") or "")

            mm = _karne.get("metrikler") or {}
            st.markdown("##### Oran analizi (likidite · kaldıraç · kârlılık)")
            oran_rows = [
                {"Grup": "Kârlılık", "Oran": "ROE %", "Değer": mm.get("ROE")},
                {"Grup": "Kârlılık", "Oran": "Net marj %", "Değer": mm.get("net_marj")},
                {"Grup": "Büyüme", "Oran": "Satış YoY %", "Değer": mm.get("satis_yoy")},
                {"Grup": "Büyüme", "Oran": "Net kâr YoY %", "Değer": mm.get("net_yoy")},
                {"Grup": "Likidite", "Oran": "Cari oran", "Değer": mm.get("cari")},
                {"Grup": "Kaldıraç", "Oran": "Borç / özkaynak", "Değer": mm.get("borc_oz")},
                {"Grup": "Değerleme", "Oran": "F/K", "Değer": mm.get("F/K")},
                {"Grup": "Değerleme", "Oran": "PD/DD", "Değer": mm.get("PD/DD")},
            ]
            st.dataframe(pd.DataFrame(oran_rows), hide_index=True, use_container_width=True)

            # Bilanço yapısı mini
            oz = (_core_bil.get("ozet") if isinstance(_core_bil, dict) else None) or {}
            if oz:
                st.markdown("##### Bilanço / gelir özeti (güncel dönem)")
                orows = []
                for key, lab in [
                    ("hasilat", "Satışlar"),
                    ("brut_kar", "Brüt kâr"),
                    ("esas_faaliyet_kari", "Faaliyet kârı"),
                    ("net_donem_kari", "Net kâr"),
                    ("nakit", "Nakit"),
                    ("ozkaynaklar", "Özkaynak"),
                    ("toplam_varliklar", "Toplam varlık"),
                    ("kisa_vadeli_yukumluluk", "KV yükümlülük"),
                ]:
                    it = oz.get(key) or {}
                    if isinstance(it, dict) and it.get("deger") is not None:
                        orows.append({
                            "Kalem": lab,
                            "Tutar": _fmt_tl(it.get("deger")),
                            "YoY %": it.get("degisim_%"),
                        })
                if orows:
                    st.dataframe(pd.DataFrame(orows), hide_index=True, use_container_width=True)

            # Marj trendi grafik (çok dönem)
            col_labels = list((_core_bil.get("col_labels") if isinstance(_core_bil, dict) else None) or [])
            gdf = _core_bil.get("gelir_df") if isinstance(_core_bil, dict) else None
            if gdf is not None and not getattr(gdf, "empty", True) and col_labels and "kod" in gdf.columns:
                def _row_vals(codes):
                    for c in codes:
                        hit = gdf[gdf["kod"].astype(str) == c]
                        if hit.empty:
                            continue
                        r0 = hit.iloc[0]
                        vals = []
                        for lab in col_labels:
                            try:
                                vals.append(float(r0.get(lab)) if r0.get(lab) is not None else None)
                            except (TypeError, ValueError):
                                vals.append(None)
                        return vals
                    return None
                sat_v = _row_vals(("3C", "3A"))
                net_v = _row_vals(("3L", "3J"))
                if sat_v and net_v:
                    marj_tr = []
                    for i, lab in enumerate(col_labels):
                        s, n = sat_v[i] if i < len(sat_v) else None, net_v[i] if i < len(net_v) else None
                        if s and n is not None and s != 0:
                            marj_tr.append({"dönem": lab, "net_marj_%": round(n / s * 100, 1)})
                    if marj_tr:
                        st.markdown("##### Net marj trendi (dönemler)")
                        st.line_chart(pd.DataFrame(marj_tr).set_index("dönem"))

            # Dış linkler
            st.markdown("##### Kaynaklar")
            lc1, lc2, lc3 = st.columns(3)
            with lc1:
                st.markdown(f"[İş Yatırım şirket kartı](https://www.isyatirim.com.tr/tr-tr/analiz/hisse/Sayfalar/sirket-karti.aspx?hisse={kod})")
            with lc2:
                st.markdown(f"[Fintables {kod}](https://fintables.com/sirketler/{kod})")
            with lc3:
                st.markdown(f"[KAP / BIST](https://www.kap.org.tr/tr/)` · yfinance {sembol}`")

            # Excel: mali tablo + karne
            if st.button("Excel’e aktar (özet + mali tablo)", key="btn_ft_excel"):
                try:
                    from io import BytesIO
                    bio = BytesIO()
                    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
                        pd.DataFrame([_karne]).to_excel(writer, sheet_name="Karne", index=False)
                        pd.DataFrame(oran_rows).to_excel(writer, sheet_name="Oranlar", index=False)
                        if isinstance(_core_bil, dict):
                            bdf = _core_bil.get("bilanco_df")
                            gdf2 = _core_bil.get("gelir_df")
                            if bdf is not None and not getattr(bdf, "empty", True):
                                bdf.to_excel(writer, sheet_name="Bilanço", index=False)
                            if gdf2 is not None and not getattr(gdf2, "empty", True):
                                gdf2.to_excel(writer, sheet_name="Gelir", index=False)
                    st.download_button(
                        "İndir Excel",
                        data=bio.getvalue(),
                        file_name=f"sirket_ozet_{kod}_{datetime.now().strftime('%Y%m%d')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_ft_xlsx",
                    )
                except Exception as e:
                    st.warning(f"Excel: {e}")

        with tab_sk:
            st.markdown("#### Skor Kart · Halk Yatırım Analizim tarzı")
            st.caption(
                "Tablolar: fiyat performansı · piyasa değeri · teknik · temel · fiyat özeti. "
                f"[Halk Yatırım Skor Kart](https://analizim.halkyatirim.com.tr/Financial/ScoreCardDetail?hisseKod={kod}) benzeri düzen. "
                "Veri: İş Yatırım + yfinance OHLCV. Yatırım tavsiyesi değildir."
            )

            # ── 1) Fiyat özeti ──
            st.markdown("##### 1) Fiyat özeti")
            fd = det.get("fiyat_df")
            last_px = deg.get("fiyat") or m.get("fiyat")
            prev = None
            high52 = low52 = None
            if fd is not None and not getattr(fd, "empty", True) and "fiyat" in fd.columns:
                s = pd.to_numeric(fd["fiyat"], errors="coerce").dropna()
                if len(s) >= 2:
                    last_px = float(s.iloc[-1])
                    prev = float(s.iloc[-2])
                if len(s) >= 5:
                    tail = s.tail(min(252, len(s)))
                    high52 = float(tail.max())
                    low52 = float(tail.min())
            gunluk_pct = ((last_px / prev - 1) * 100) if last_px and prev else None
            fy_rows = [
                {"Alan": "Son fiyat", "Değer": fmt_num(last_px, 2)},
                {"Alan": "Günlük değişim %", "Değer": f"%{gunluk_pct:+.2f}" if gunluk_pct is not None else "—"},
                {"Alan": "52h yüksek", "Değer": fmt_num(high52, 2)},
                {"Alan": "52h düşük", "Değer": fmt_num(low52, 2)},
                {"Alan": "52h konumu %", "Değer": (
                    f"%{((last_px - low52) / (high52 - low52) * 100):.0f}"
                    if last_px and high52 and low52 and high52 != low52 else "—"
                )},
                {"Alan": "Kaynak çarpan", "Değer": deg.get("_kaynak_carpan") or deg.get("_kaynak") or "İş Yatırım"},
            ]
            st.dataframe(pd.DataFrame(fy_rows), hide_index=True, use_container_width=True)

            # ── 2) Fiyat performansı ──
            st.markdown("##### 2) Fiyat performansı")
            perf_map = [
                ("1 hafta", "getiri_1h", 5),
                ("1 ay", "getiri_1a", 22),
                ("3 ay", "getiri_3a", 66),
                ("6 ay", "getiri_6a", 132),
                ("1 yıl", "getiri_1y", 252),
                ("Yılbaşından", "getiri_ytd", None),
            ]
            perf_rows = []
            s_px = None
            if fd is not None and not getattr(fd, "empty", True) and "fiyat" in fd.columns:
                s_px = pd.to_numeric(fd["fiyat"], errors="coerce").dropna()
            for label, key, bars in perf_map:
                val = m.get(key)
                if val is None and s_px is not None and bars and len(s_px) > bars:
                    try:
                        val = (float(s_px.iloc[-1]) / float(s_px.iloc[-bars - 1]) - 1) * 100
                    except Exception:
                        val = None
                if val is None and key == "getiri_ytd" and s_px is not None and len(s_px) > 5:
                    try:
                        # approximate: from start of series year if possible
                        if "tarih" in fd.columns:
                            tfd = fd.copy()
                            tfd["tarih"] = pd.to_datetime(tfd["tarih"])
                            y0 = tfd[tfd["tarih"].dt.year == tfd["tarih"].dt.year.max()]
                            if len(y0) > 1:
                                val = (float(y0["fiyat"].iloc[-1]) / float(y0["fiyat"].iloc[0]) - 1) * 100
                    except Exception:
                        pass
                perf_rows.append({
                    "Dönem": label,
                    "Getiri %": round(float(val), 2) if val is not None else None,
                    "Yorum": (
                        "Güçlü" if val is not None and val >= 15 else
                        ("Zayıf" if val is not None and val <= -10 else
                         ("Nötr" if val is not None else "—"))
                    ),
                })
            st.dataframe(pd.DataFrame(perf_rows), hide_index=True, use_container_width=True)

            # ── 3) Piyasa değeri & çarpanlar ──
            st.markdown("##### 3) Piyasa değeri ve çarpanlar")
            mcap = deg.get("piyasa_değeri")
            pd_rows = [
                {"Gösterge": "Piyasa değeri", "Değer": _fmt_tl(mcap) if mcap else "—", "Not": "İş Yatırım"},
                {"Gösterge": "F/K", "Değer": fmt_num(deg.get("F/K"), 2), "Not": "İş Yatırım şirket kartı"},
                {"Gösterge": "PD/DD", "Değer": fmt_num(deg.get("PD/DD"), 2), "Not": "İş Yatırım"},
                {"Gösterge": "FD/FAVÖK", "Değer": fmt_num(deg.get("FD/FAVÖK"), 2), "Not": "İş Yatırım"},
                {"Gösterge": "FD/Satış", "Değer": fmt_num(deg.get("FD/Satış"), 2), "Not": "İş Yatırım"},
                {"Gösterge": "ROE %", "Değer": fmt_num(
                    (deg.get("ROE") if deg.get("ROE") is not None else
                     (_core_bil.get("oranlar") or {}).get("ROE_%")), 1
                ), "Not": "Bilanço / oran"},
            ]
            st.dataframe(pd.DataFrame(pd_rows), hide_index=True, use_container_width=True)

            # ── 4) Teknik veriler ──
            st.markdown("##### 4) Teknik veriler")
            tek_rows = [
                {"Gösterge": "RSI (14)", "Değer": fmt_num(_tek.get("rsi14"), 1),
                 "Sinyal": (
                     "Aşırı satım" if (_tek.get("rsi14") or 50) <= 30 else
                     ("Aşırı alım" if (_tek.get("rsi14") or 50) >= 70 else "Nötr")
                 )},
                {"Gösterge": "MA22", "Değer": fmt_num(_tek.get("ma22"), 2),
                 "Sinyal": "Fiyat üstü" if last_px and _tek.get("ma22") and last_px > _tek["ma22"] else "Fiyat altı"},
                {"Gösterge": "MA50", "Değer": fmt_num(_tek.get("ma50"), 2),
                 "Sinyal": "Fiyat üstü" if last_px and _tek.get("ma50") and last_px > _tek["ma50"] else "Fiyat altı"},
                {"Gösterge": "MA200", "Değer": fmt_num(_tek.get("ma200"), 2),
                 "Sinyal": "Fiyat üstü" if last_px and _tek.get("ma200") and last_px > _tek["ma200"] else "Fiyat altı"},
                {"Gösterge": "Volatilite %", "Değer": fmt_num(m.get("volatilite"), 1), "Sinyal": "—"},
                {"Gösterge": "Max DD %", "Değer": fmt_num(m.get("max_drawdown"), 1), "Sinyal": "—"},
                {"Gösterge": "Sharpe", "Değer": fmt_num(m.get("sharpe"), 2), "Sinyal": "—"},
            ]
            st.dataframe(pd.DataFrame(tek_rows), hide_index=True, use_container_width=True)

            # ── 5) Temel veri analizi ──
            st.markdown("##### 5) Temel veri analizi")
            oz = (_core_bil.get("ozet") if isinstance(_core_bil, dict) else None) or {}
            oran = (_core_bil.get("oranlar") if isinstance(_core_bil, dict) else None) or {}
            tem_rows = []
            for key, lab in [
                ("hasilat", "Satışlar"),
                ("brut_kar", "Brüt kâr"),
                ("esas_faaliyet_kari", "Esas faaliyet kârı"),
                ("net_donem_kari", "Net dönem kârı"),
                ("nakit", "Nakit"),
                ("ozkaynaklar", "Özkaynaklar"),
                ("toplam_varliklar", "Toplam varlıklar"),
                ("kisa_vadeli_yukumluluk", "KV yükümlülük"),
            ]:
                it = oz.get(key) or {}
                if isinstance(it, dict) and it.get("deger") is not None:
                    tem_rows.append({
                        "Kalem": lab,
                        "Tutar": _fmt_tl(it.get("deger")),
                        "YoY %": it.get("degisim_%"),
                        "Dönem": _core_bil.get("gunce_donem") or "—",
                    })
            tem_rows.append({"Kalem": "ROE %", "Tutar": oran.get("ROE_%"), "YoY %": None, "Dönem": _core_bil.get("gunce_donem")})
            tem_rows.append({"Kalem": "Net marj %", "Tutar": oran.get("net_kar_marji_%"), "YoY %": None, "Dönem": _core_bil.get("gunce_donem")})
            tem_rows.append({"Kalem": "Cari oran", "Tutar": oran.get("cari_oran"), "YoY %": None, "Dönem": _core_bil.get("gunce_donem")})
            tem_rows.append({"Kalem": "Borç/özkaynak", "Tutar": oran.get("borc_ozkaynak"), "YoY %": None, "Dönem": _core_bil.get("gunce_donem")})
            if tem_rows:
                st.dataframe(pd.DataFrame(tem_rows), hide_index=True, use_container_width=True)
            else:
                st.info("Temel veri için İş Yatırım mali tablo gerekli.")

            # ── 6) Skor özeti (renkli kartlar) ──
            st.markdown("##### 6) Skor özeti")
            _karne2 = det.get("fintables_karne") or fintables_tarzi_karne(
                _core_bil if isinstance(_core_bil, dict) else {}, deg, m
            )
            s1, s2, s3, s4, s5 = st.columns(5)
            s1.metric("Karne genel", f"{_karne2.get('genel')}/100", _karne2.get("harf"))
            s2.metric("Kârlılık", f"{_karne2.get('karlilik')}/100")
            s3.metric("Büyüme", f"{_karne2.get('buyume')}/100")
            s4.metric("Borçluluk", f"{_karne2.get('borcluluk')}/100")
            s5.metric("Değerleme", f"{_karne2.get('degerleme')}/100")
            s6, s7 = st.columns(2)
            s6.metric("Temel skor", f"{_tem.get('skor')}/100" if _tem.get("skor") is not None else "—")
            s7.metric("Model sinyal", f"{(det.get('model_sinyal') or {}).get('skor', '—')}")

            # Excel tüm skor kart tabloları
            if st.button("Skor Kart Excel indir", key="btn_sk_excel"):
                try:
                    from io import BytesIO
                    bio = BytesIO()
                    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
                        pd.DataFrame(fy_rows).to_excel(writer, sheet_name="Fiyat_ozeti", index=False)
                        pd.DataFrame(perf_rows).to_excel(writer, sheet_name="Fiyat_performans", index=False)
                        pd.DataFrame(pd_rows).to_excel(writer, sheet_name="Piyasa_carpan", index=False)
                        pd.DataFrame(tek_rows).to_excel(writer, sheet_name="Teknik", index=False)
                        pd.DataFrame(tem_rows).to_excel(writer, sheet_name="Temel", index=False)
                        pd.DataFrame([_karne2]).to_excel(writer, sheet_name="Karne", index=False)
                    st.download_button(
                        "İndir",
                        data=bio.getvalue(),
                        file_name=f"skor_kart_{kod}_{datetime.now().strftime('%Y%m%d')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_sk_xlsx",
                    )
                except Exception as e:
                    st.warning(str(e))

        with tab_g:
            bm = None
            try:
                bm = core.yahoo_daily_returns("XU100.IS")
            except Exception:
                pass
            _draw_hisse_charts(det.get("fiyat_df"), m, sembol, bm_rets=bm)
            # MA overlay
            if not _tek.get("error") and _tek.get("_px") is not None:
                st.markdown("**Fiyat + hareketli ortalamalar**")
                ma_df = pd.DataFrame({"Fiyat": _tek["_px"]})
                for w in (22, 50, 200):
                    ser = _tek.get(f"_ma{w}_seri")
                    if ser is not None:
                        ma_df[f"MA{w}"] = ser
                st.line_chart(ma_df.tail(300))

        with tab_mt:
            st.markdown(
                f"#### Mali tablo — **İş Yatırım** · güncel dönem **"
                f"{_core_bil.get('gunce_donem') or _bil.get('etiket')}**"
            )
            st.caption(
                "Kaynak: isyatirim.com.tr MaliTablo API. "
                "Her hisse için API’de **dolu olan en yeni dönem** seçilir (ör. 2026/9 varsa 2026/6 kullanılmaz). "
                "Cache günlük; «Zorla yenile» ile anında çeker. Gelir kalemleri kümülatif."
            )
            if st.button("Güncelle ve raporları yenile (İş Yatırım)", type="primary", key="btn_isy_force"):
                # Tüm ilgili cache temiz — güncel mali tablo (ör. 2026/9) yeniden çekilir
                try:
                    for ck in (f"{kod}|XI_29", f"{kod}|UFRS"):
                        st.session_state.get("_isy_mali_cache", {}).pop(ck, None)
                    st.session_state.get("_isy_carpan_cache", {}).pop(kod, None)
                    # Hisse detay sonucunu sil → Analiz et ile taze rapor
                    st.session_state.pop("hisse_detay_sonuc", None)
                    st.session_state["_isy_force_refresh"] = True
                    st.session_state["_hisse_auto_analiz"] = True
                    st.session_state["hisse_kod_input"] = kod
                except Exception:
                    pass
                st.success(f"{kod}: cache temizlendi — mali tablo ve raporlar İş Yatırım’dan yenilenecek.")
                st.rerun()
            if _core_bil.get("error"):
                st.warning(f"İş Yatırım veri alınamadı: {_core_bil.get('error')}")
            elif _core_bil.get("kaynak"):
                st.success(
                    f"{_core_bil.get('kaynak')} · güncel kolon: **{_core_bil.get('gunce_donem')}** · "
                    f"grup: {_core_bil.get('financial_group')}"
                )
                if _core_bil.get("url"):
                    st.caption(f"Şirket kartı: {_core_bil['url']}")

                # Özet KPI
                def _tl(v):
                    try:
                        return _fmt_tl(v)
                    except NameError:
                        if v is None:
                            return "—"
                        try:
                            x = float(v)
                            if abs(x) >= 1e9:
                                return f"{x/1e9:.2f} mlr"
                            if abs(x) >= 1e6:
                                return f"{x/1e6:.1f} mn"
                            return f"{x:,.0f}"
                        except Exception:
                            return "—"

                oz = _core_bil.get("ozet") or {}
                orans = _core_bil.get("oranlar") or {}
                kpis = st.columns(4)
                labels_kpi = [
                    ("Hasılat", "hasilat"),
                    ("Net dönem kârı", "net_donem_kari"),
                    ("Özkaynaklar", "ozkaynaklar"),
                    ("Toplam varlıklar", "toplam_varliklar"),
                ]
                for i, (lab, key) in enumerate(labels_kpi):
                    item = oz.get(key) or {}
                    delta = item.get("degisim_%")
                    kpis[i].metric(
                        lab,
                        _tl(item.get("deger")),
                        delta=f"%{delta}" if delta is not None else None,
                    )
                if orans:
                    st.markdown("**İş Yatırım kalemlerinden oranlar (güncel dönem)**")
                    st.dataframe(
                        pd.DataFrame([{"Oran": k, "Değer": v} for k, v in orans.items()]),
                        hide_index=True, use_container_width=True,
                    )
                    # Temel skora İş Yatırım ROE/marj enjekte
                    if orans.get("ROE_%") is not None:
                        deg["ROE"] = orans["ROE_%"]
                    if orans.get("ROA_%") is not None:
                        deg["ROA"] = orans["ROA_%"]
                    if orans.get("net_kar_marji_%") is not None:
                        deg["net_kar_marjı"] = orans["net_kar_marji_%"]
                    if orans.get("borc_ozkaynak") is not None:
                        deg["borç/özkaynak"] = orans["borc_ozkaynak"]
                    if orans.get("cari_oran") is not None:
                        deg["cari_oran"] = orans["cari_oran"]

                # Özet değişim tablosu
                oz_rows = []
                for key, item in oz.items():
                    if not isinstance(item, dict):
                        continue
                    oz_rows.append({
                        "Kalem": item.get("kalem") or key,
                        _core_bil.get("gunce_donem") or "Güncel": _tl(item.get("deger")),
                        "YoY aynı dönem": _tl(item.get("onceki_yoy")),
                        "Değişim %": item.get("degisim_%"),
                    })
                if oz_rows:
                    st.markdown("**Özet kalemler (YoY)**")
                    st.dataframe(pd.DataFrame(oz_rows), hide_index=True, use_container_width=True)


                _cols = list(_core_bil.get("col_labels") or [])
                _gunce = _core_bil.get("gunce_donem") or (_cols[0] if _cols else "—")
                _yoy_lab = None
                try:
                    _y, _p = str(_gunce).split("/")
                    _yoy_lab = f"{int(_y)-1}/{int(_p)}"
                except Exception:
                    _yoy_lab = _cols[1] if len(_cols) > 1 else None
                st.markdown(
                    f"**Güncel bilanço dönemi: `{_gunce}`**"
                    + (f" · YoY karşılaştırma: **`{_yoy_lab}`** (önceki yıl aynı dönem)" if _yoy_lab else "")
                )

                # Reel satış / kârlılık
                _reel = reel_satis_karlilik(_core_bil)
                st.markdown(
                    f"##### Reel büyüme (nominal YoY − TÜFE %{_reel.get('tufe'):.0f}, {_reel.get('tufe_kaynak')})"
                )
                if _reel.get("satirlar"):
                    _rdf = pd.DataFrame(_reel["satirlar"]).copy()
                    if "Tutar" in _rdf.columns:
                        _rdf["Tutar"] = _rdf["Tutar"].map(lambda v: _tl(v) if v is not None else "—")
                    if "YoY tutar" in _rdf.columns:
                        _rdf["YoY tutar"] = _rdf["YoY tutar"].map(lambda v: _tl(v) if v is not None else "—")
                    st.dataframe(_rdf, hide_index=True, use_container_width=True)
                    st.caption(
                        "Örnek: satış nominal %+50, TÜFE %40 → reel %+10. "
                        "TÜFE’yi ⚙️ Ayarlar’dan güncelleyin."
                    )
                else:
                    st.caption("YoY değişim verisi yok — reel hesaplanamadı.")

                # Gelişmiş otomatik mali yorum
                try:
                    _yorum_mt = mali_tablo_otomatik_yorum(
                        _core_bil if isinstance(_core_bil, dict) else {},
                        deg if isinstance(deg, dict) else {},
                    )
                    render_otomatik_bilanco_yorum(
                        _yorum_mt,
                        baslik=f"Mali tablo analizi · {_gunce}",
                    )
                except Exception as _ye:
                    st.caption(f"Otomatik yorum: {_ye}")

                mt1, mt2, mt3 = st.tabs(["Bilanço", "Gelir tablosu", "Nakit akım"])
                _birim = st.radio(
                    "Tutar birimi",
                    ["mn (milyon TL)", "mlr (milyar TL)", "tam (ondalıklı)"],
                    horizontal=True,
                    key=f"mali_birim_{kod}",
                    index=0,
                )
                _b = "mn" if _birim.startswith("mn") else ("mlr" if _birim.startswith("mlr") else "tam")
                st.caption(
                    f"Hücre biçimi · birim **{_b}** · güncel **{_gunce}** vs YoY **{_yoy_lab or '—'}**. "
                    "Boş dönem **—**."
                )

                def _show_mali(df0, label):
                    if df0 is None or getattr(df0, "empty", True):
                        st.caption(f"{label} satırı yok")
                        return
                    # YoY karşılaştırma özeti
                    yoy_df = mali_yoy_karsilastirma_df(df0, _cols)
                    if yoy_df is not None and not yoy_df.empty:
                        st.markdown(f"**{_gunce} vs {_yoy_lab or 'önceki'} (YoY %)**")
                        yoy_view = yoy_df.copy()
                        for c in yoy_view.columns:
                            if c in ("kod", "kalem", "YoY %"):
                                continue
                            if _b == "tam":
                                yoy_view[c] = yoy_view[c].map(
                                    lambda v: "—" if v is None or (isinstance(v, float) and v != v)
                                    else f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                                )
                            else:
                                yoy_view[c] = yoy_view[c].map(
                                    lambda v: format_mali_df(
                                        pd.DataFrame({c: [v]}), birim=_b
                                    )[c].iloc[0] if v is not None else "—"
                                )
                        show_df(yoy_view, list(yoy_view.columns), height=280)
                    st.markdown(f"**Tüm dönemler**")
                    if _b == "tam":
                        fmt = df0.copy()
                        for c in fmt.columns:
                            if str(c) in ("kod", "kalem"):
                                continue
                            def _tam(v):
                                if v is None or (isinstance(v, float) and v != v):
                                    return "—"
                                try:
                                    x = float(v)
                                    if x == 0:
                                        return "0"
                                    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                                except Exception:
                                    return "—" if str(v).lower() in ("none", "nan") else str(v)
                            fmt[c] = fmt[c].map(_tam)
                        show_df(fmt, list(fmt.columns), height=360)
                    else:
                        show_df(format_mali_df(df0, birim=_b), list(df0.columns), height=360)
                    with st.expander(f"Ham {label}"):
                        show_df(df0, list(df0.columns), height=200)

                with mt1:
                    _show_mali(_core_bil.get("bilanco_df"), "Bilanço")
                with mt2:
                    _show_mali(_core_bil.get("gelir_df"), "Gelir")
                with mt3:
                    ndf = _core_bil.get("nakit_df")
                    if ndf is not None and not getattr(ndf, "empty", True):
                        _show_mali(ndf, "Nakit akım")
                    else:
                        st.info(
                            "Nakit akım satırı (kod 4…) bu hisse/grup için boş. "
                            "Bilanço nakit kalemi (1AA) özet KPI’da kullanılabilir."
                        )
                        # nakit bilanço satırı yedek
                        bdf = _core_bil.get("bilanco_df")
                        if bdf is not None and not getattr(bdf, "empty", True) and "kod" in bdf.columns:
                            nak = bdf[bdf["kod"].astype(str).str.startswith("1AA")]
                            if not nak.empty:
                                st.markdown("**Bilanço nakit kalemleri (yedek)**")
                                _show_mali(nak, "Nakit (bilanço)")
            else:
                st.info("İş Yatırım’dan veri gelmedi; Yahoo çarpanları yedek olarak aşağıda.")

            c1, c2, c3 = st.columns(3)
            with c1:
                st.markdown("**Çarpanlar (fiyat / yedek)**")
                for k, v in (_tem.get("carpanlar") or {}).items():
                    st.write(f"{k}: **{fmt_num(v, 2) if v is not None else '—'}**")
            with c2:
                st.markdown("**Kârlılık**")
                for k, v in (_tem.get("karlilik") or {}).items():
                    st.write(f"{k}: **{fmt_num(v, 2) if v is not None else '—'}**")
            with c3:
                st.markdown("**Bilanço sağlığı**")
                for k, v in (_tem.get("bilanco_saglik") or {}).items():
                    st.write(f"{k}: **{fmt_num(v, 2) if v is not None else '—'}**")

        with tab_tem:
            st.markdown(f"#### Temel analiz — referans bilanço **{_bil.get('etiket')}**")
            s1, s2, s3 = st.columns(3)
            s1.metric("Temel skor", f"{_tem.get('skor')}/100")
            s2.metric("Dönem", _bil.get("etiket") or "—")
            s3.metric("F/K", fmt_num((_tem.get("carpanlar") or {}).get("F/K"), 1))
            st.markdown(f"**{_tem.get('skor_yorum')}**")
            for n in _tem.get("notlar") or []:
                st.markdown(f"- {n}")
            # Tarihsel ucuzluk (Filtre Stüdyosu ile uyumlu)
            _tu = tarihsel_ucuzluk_analizi(kod, fd=det.get("fiyat_df"), deg=deg)
            det["tarihsel_ucuzluk"] = _tu
            st.markdown("#### Tarihsel ucuzluk")
            u1, u2, u3, u4 = st.columns(4)
            u1.metric("Ucuzluk skoru", f"{_tu.get('skor')}/100" if _tu.get("skor") is not None else "—")
            u2.metric("1Y fiyat %ile", f"%{_tu.get('fiyat_persentil_1y')}" if _tu.get("fiyat_persentil_1y") is not None else "—")
            u3.metric("3Y fiyat %ile", f"%{_tu.get('fiyat_persentil_3y')}" if _tu.get("fiyat_persentil_3y") is not None else "—")
            u4.metric("Ucuz bölge?", "Evet" if _tu.get("ucuz_mu") else ("Hayır" if _tu.get("ucuz_mu") is False else "—"))
            for n in _tu.get("yorum") or []:
                st.markdown(f"- {n}")
            st.caption("Skor kural tabanlıdır; yatırım tavsiyesi değildir. Sektör ortalaması ile birlikte okuyun.")

        with tab_tek:
            st.markdown("#### Teknik analiz")
            if _tek.get("error"):
                st.warning(_tek["error"])
            else:
                t1, t2, t3, t4, t5 = st.columns(5)
                t1.metric("Fiyat", fmt_num(_tek.get("fiyat"), 2))
                t2.metric("RSI(14)", fmt_num(_tek.get("rsi14"), 1))
                t3.metric("MA22 uzaklık %", fmt_num(_tek.get("ma22_uzaklik_%"), 1))
                t4.metric("MA50 uzaklık %", fmt_num(_tek.get("ma50_uzaklik_%"), 1))
                t5.metric("MA200 uzaklık %", fmt_num(_tek.get("ma200_uzaklik_%"), 1))
                if _tek.get("rsi_yorum"):
                    st.caption(f"RSI: {_tek['rsi_yorum']}")
                st.markdown("**Sinyaller**")
                for s in _tek.get("sinyaller") or []:
                    st.markdown(f"- {s}")
                st.markdown("**Getiri (fiyat serisi)**")
                gcols = st.columns(5)
                for i, (lab, key) in enumerate([
                    ("1H", "getiri_1h"), ("1A", "getiri_1a"), ("3A", "getiri_3a"),
                    ("6A", "getiri_6a"), ("1Y", "getiri_1y"),
                ]):
                    gcols[i].metric(lab, f"%{fmt_num(_tek.get(key))}")
                st.markdown("**52 hafta**")
                h1, h2, h3, h4 = st.columns(4)
                h1.metric("Yüksek", fmt_num(_tek.get("52h_yuksek"), 2))
                h2.metric("Düşük", fmt_num(_tek.get("52h_dusuk"), 2))
                h3.metric("Yüksekten %", fmt_num(_tek.get("52h_yuksek_uzak_%"), 1))
                h4.metric("Düşükten %", fmt_num(_tek.get("52h_dusuk_uzak_%"), 1))
                if _tek.get("hacim_oran") is not None:
                    st.caption(
                        f"Hacim: son / 20g ort = **{_tek['hacim_oran']}x** "
                        f"(son {fmt_num(_tek.get('hacim_son'), 0)}, ort {fmt_num(_tek.get('hacim_ort20'), 0)})"
                    )
                st.caption("Teknik göstergeler gecikmeli fiyata dayanır; tek başına al-sat sinyali değildir.")

        with tab_r:
            st.markdown("#### Getiri metrikleri")
            getiri_rows = {
                "1A %": m.get("getiri_1a"), "3A %": m.get("getiri_3a"),
                "6A %": m.get("getiri_6a"), "1Y %": m.get("getiri_1y"),
                "3Y %": m.get("getiri_3y"), "CAGR %": m.get("cagr"),
                "Reel 1Y %": m.get("reel_1y"),
                "vs Mevduat": m.get("vs_mevduat"), "vs BIST100": m.get("vs_bist100"),
            }
            st.dataframe(
                pd.DataFrame([{"Metrik": k, "Değer": v} for k, v in getiri_rows.items() if v is not None]),
                hide_index=True, use_container_width=True,
            )
            st.markdown("#### Risk metrikleri")
            risk_rows = {
                "Volatilite %": m.get("volatilite"), "Max Drawdown %": m.get("max_drawdown"),
                "Sharpe": m.get("sharpe"), "Sortino": m.get("sortino"),
                "Calmar": m.get("calmar"), "VaR %95": m.get("var95"),
                "cVaR %95": m.get("cvar95"), "Kazanma oranı %": m.get("kazanma_orani"),
                "En kötü gün %": m.get("en_kotu_gun"), "En iyi gün %": m.get("en_iyi_gun"),
                "BIST korelasyon": m.get("bist_korelasyon"), "Risk skor": m.get("risk_skor"),
            }
            st.dataframe(
                pd.DataFrame([{"Metrik": k, "Değer": v} for k, v in risk_rows.items() if v is not None]),
                hide_index=True, use_container_width=True,
            )

        with tab_d:
            st.markdown("#### Excel tarzı değerleme alanları")
            if not deg:
                st.info(
                    "Değerleme için Yahoo Finance bilgisi gerekir. "
                    "`pip install yfinance` sonrası yeniden analiz edin."
                )
            else:
                # Gruplar
                groups = {
                    "Piyasa": ["fiyat", "önceki_kapanış", "açılış", "gün_düşük", "gün_yüksek",
                               "52h_düşük", "52h_yüksek", "hacim", "ort_hacim", "piyasa_değeri", "firma_değeri"],
                    "Çarpanlar": ["F/K", "F/K_ileri", "PD/DD", "FD/FAVÖK", "FD/Satış", "PEG", "EPS", "defter_değeri"],
                    "Kârlılık": ["ROE", "ROA", "net_kar_marjı", "brüt_marj", "faaliyet_marjı"],
                    "Bilanço / risk": ["borç/özkaynak", "cari_oran", "asit_test", "beta"],
                    "Temettü / sahiplik": ["temettü_getirisi", "temettü_ödeme", "içeriden_%", "kurumsal_%",
                                           "serbest_dolaşım", "hisse_adet"],
                    "Analist": ["hedef_fiyat", "tavsiye", "analist_sayısı"],
                }
                for gname, keys in groups.items():
                    rows = []
                    for k in keys:
                        if k in deg and deg[k] is not None:
                            v = deg[k]
                            if isinstance(v, float):
                                if k in ("piyasa_değeri", "firma_değeri", "hacim", "ort_hacim",
                                         "serbest_dolaşım", "hisse_adet"):
                                    if abs(v) >= 1e9:
                                        disp = f"{v/1e9:.2f} Mr"
                                    elif abs(v) >= 1e6:
                                        disp = f"{v/1e6:.1f} Mn"
                                    else:
                                        disp = f"{v:,.0f}"
                                else:
                                    disp = f"{v:.2f}"
                            else:
                                disp = str(v)
                            rows.append({"Alan": k, "Değer": disp})
                    if rows:
                        st.markdown(f"**{gname}**")
                        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)

        with tab_v:
            st.markdown("#### Kapsamlı yorum & SWOT")
            st.caption(
                "Bilanço (İş Yatırım güncel dönem), çarpanlar, kârlılık, risk, getiri ve teknik "
                "göstergelerin tamamı değerlendirilir. Yatırım tavsiyesi değildir."
            )
            _sw = _swot_yatirim_analizi(
                kod, m, deg, _tem, _tek,
                isy=_core_bil if isinstance(_core_bil, dict) else {},
                bil=_bil,
            )
            det["swot"] = {k: v for k, v in _sw.items() if k != "sonuc"}

            s1, s2, s3 = st.columns(3)
            s1.metric("Bütüncül skor", f"{_sw.get('skor')}/100")
            s2.metric("Bilanço dönemi", _sw.get("donem") or "—")
            s3.metric("Görünüm", {"olumlu": "Olumlu", "notr": "Nötr", "temkinli": "Temkinli"}.get(_sw.get("ton"), "—"))

            # Olumlu / olumsuz
            c_pos, c_neg = st.columns(2)
            with c_pos:
                st.markdown("##### ✅ Olumlu yanlar")
                if _sw.get("olumlu"):
                    for x in _sw["olumlu"]:
                        st.markdown(f"- {x}")
                else:
                    st.caption("Belirgin olumlu madde yok.")
            with c_neg:
                st.markdown("##### ⚠️ Olumsuz yanlar")
                if _sw.get("olumsuz"):
                    for x in _sw["olumsuz"]:
                        st.markdown(f"- {x}")
                else:
                    st.caption("Belirgin olumsuz madde yok.")

            st.markdown("#### SWOT (yatırım çerçevesi)")
            sw1, sw2 = st.columns(2)
            with sw1:
                st.markdown("**Strengths (Güçlü yönler)**")
                for x in _sw.get("S") or ["—"]:
                    st.markdown(f"- {x}")
                st.markdown("**Opportunities (Fırsatlar)**")
                for x in _sw.get("O") or ["—"]:
                    st.markdown(f"- {x}")
            with sw2:
                st.markdown("**Weaknesses (Zayıf yönler)**")
                for x in _sw.get("W") or ["—"]:
                    st.markdown(f"- {x}")
                st.markdown("**Threats (Tehditler)**")
                for x in _sw.get("T") or ["—"]:
                    st.markdown(f"- {x}")

            st.markdown("#### Sonuç")
            st.markdown(_sw.get("sonuc") or "")

            _kal = bilanco_kalite_skoru(
                _core_bil if isinstance(_core_bil, dict) else {},
                deg,
            )
            _tu2 = det.get("tarihsel_ucuzluk") or tarihsel_ucuzluk_analizi(
                kod, fd=det.get("fiyat_df"), deg=deg
            )
            _sig = model_sinyal_skoru(
                m=m, deg=deg, tek=_tek, tem=_tem, kalite=_kal, tu=_tu2
            )
            det["bilanco_kalite"] = _kal
            det["model_sinyal"] = _sig

            st.markdown("#### Bilanço kalite skoru")
            k1, k2, k3 = st.columns(3)
            k1.metric("Kalite", f"{_kal.get('skor')}/100")
            k2.metric("Seviye", _kal.get("seviye") or "—")
            k3.metric("Dönem", _kal.get("donem") or _bil.get("etiket") or "—")
            st.caption(_kal.get("yorum") or "")
            kc1, kc2 = st.columns(2)
            with kc1:
                st.markdown("**Artılar**")
                for x in _kal.get("artilar") or ["—"]:
                    st.markdown(f"- {x}")
            with kc2:
                st.markdown("**Eksiler**")
                for x in _kal.get("eksiler") or ["—"]:
                    st.markdown(f"- {x}")

            st.markdown("#### Model sinyal (kural tabanlı)")
            st.warning(_sig.get("disclaimer") or "")
            s1, s2 = st.columns(2)
            s1.metric("Sinyal skoru", f"{_sig.get('skor')}")
            s2.metric("Etiket", _sig.get("etiket") or "—")
            if _sig.get("bilesenler"):
                st.dataframe(
                    pd.DataFrame(
                        [{"Bileşen": k, "Katkı": v} for k, v in _sig["bilesenler"].items()]
                    ),
                    hide_index=True, use_container_width=True,
                )

            _tu2 = det.get("tarihsel_ucuzluk") or tarihsel_ucuzluk_analizi(
                kod, fd=det.get("fiyat_df"), deg=deg
            )
            _bir = birlesik_profil_skoru(kalite=_kal, sinyal=_sig, tu=_tu2)
            det["birlesik_skor"] = _bir
            st.markdown("#### Birleşik skor (eşit ağırlık: kalite + sinyal + ucuzluk)")
            st.info(_bir.get("not") or "")
            b1, b2, b3, b4, b5 = st.columns(5)
            b1.metric("Kalite", f"{_bir.get('kalite')}/100" if _bir.get("kalite") is not None else "—")
            b2.metric("Sinyal (norm)", f"{_bir.get('sinyal_norm')}/100" if _bir.get("sinyal_norm") is not None else "—")
            b3.metric("Ucuzluk", f"{_bir.get('ucuzluk')}/100" if _bir.get("ucuzluk") is not None else "—")
            b4.metric("Ortalama", f"{_bir.get('ortalama')}/100" if _bir.get("ortalama") is not None else "—")
            b5.metric("Seviye", _bir.get("seviye") or "—")
            st.caption(
                f"Her bileşen eşit ağırlık (~%{_bir.get('esit_agirlik_%') or '—'}). "
                "Hazır profil sıralaması **ortalama yüksek → düşük**. Yatırım tavsiyesi değildir."
            )

            st.markdown("#### Risk–getiri (sade dil)")
            try:
                st.markdown(core.turkce_degerlendirme(
                    sharpe=m.get("sharpe"), sortino=m.get("sortino"),
                    max_dd=m.get("max_drawdown"), vol=m.get("volatilite"),
                    cagr=m.get("cagr"), calmar=m.get("calmar"),
                    getiri_1y=m.get("getiri_1y"), reel_1y=m.get("reel_1y"),
                    vs_mevduat=m.get("vs_mevduat"), vs_bist100=m.get("vs_bist100"),
                    kazanma_orani=m.get("kazanma_orani"),
                    en_kotu_gun=m.get("en_kotu_gun"),
                    bist_korelasyon=m.get("bist_korelasyon") or deg.get("beta"),
                    plain=True,
                ))
            except Exception as e:
                st.caption(f"Risk metni: {e}")

        with tab_x:
            # Excel: metrik + değerleme tek satır
            row = {
                "kod": kod, "sembol": sembol, "ad": ad, "sektör": sektor, "endüstri": endustri,
                "bilanco_donem": (_bil or {}).get("etiket"),
                "temel_skor": (_tem or {}).get("skor"),
                "rsi14": (_tek or {}).get("rsi14"),
                "ma22": (_tek or {}).get("ma22"),
                "ma50": (_tek or {}).get("ma50"),
                "ma200": (_tek or {}).get("ma200"),
            }
            row.update({k: m.get(k) for k in (
                "fiyat", "getiri_1a", "getiri_3a", "getiri_6a", "getiri_1y", "getiri_3y",
                "reel_1y", "vs_mevduat", "vs_bist100", "cagr",
                "volatilite", "max_drawdown", "sharpe", "sortino", "calmar",
                "var95", "cvar95", "kazanma_orani", "en_kotu_gun", "en_iyi_gun",
                "bist_korelasyon", "risk_skor",
            )})
            for k, v in deg.items():
                if k not in row:
                    row[f"d_{k}"] = v
            edf = pd.DataFrame([row])
            try:
                _bd_x = son_bilanco_donemi(probe=False)
                _meta_x = {
                    "mod": "hisse_detay", "sembol": sembol, "ad": ad,
                    "bilanco_donem": _bd_x.get("etiket"),
                    "bilanco_aciklama": _bd_x.get("aciklama"),
                    "bilanco_kaynak": _bd_x.get("kaynak"),
                }
            except Exception:
                _meta_x = {"mod": "hisse_detay", "sembol": sembol, "ad": ad}
            download_excel_button(
                edf,
                f"hisse_detay_{kod}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                meta=_meta_x,
                key="xl_hisse_detay",
            )
            if st.button("📄 HTML risk raporu", key="btn_hisse_det_html"):
                fd = det.get("fiyat_df")
                if fd is None or len(fd) < 30:
                    st.warning("Yeterli fiyat serisi yok.")
                else:
                    try:
                        path = core.turkce_risk_raporu(
                            fd, kod, ad=ad, rf=get_analiz().rf,
                            benchmark_returns=core.yahoo_daily_returns("XU100.IS"),
                            ekstra_metrikler=m,
                        )
                        if path:
                            st.success(str(path))
                            with open(path, "rb") as f:
                                st.download_button(
                                    "HTML indir", f, file_name=Path(path).name,
                                    mime="text/html", key="dl_hisse_det_html",
                                )
                    except Exception as e:
                        st.error(f"Rapor: {e}")

            if st.button("🛒 Portföy sepetine ekle (fon değil — not)", key="hisse_sepet_info"):
                st.info(
                    "Sepet şu an TEFAS fon kodları içindir. "
                    "Hisseyi **Hisse & Karma Portföy** sekmesinden karma portföye ekleyebilirsiniz."
                )

elif sayfa == "⭐ Favoriler":
    st.header("⭐ Favoriler")
    favs = core.load_favs()

    c1, c2 = st.columns([3, 1])
    with c1:
        yeni = st.text_input("Fon kodu ekle", placeholder="örn AAL").strip().upper()
    with c2:
        st.write("")
        st.write("")
        if st.button("Ekle") and yeni:
            favs = list(set(favs) | {yeni})
            core.save_favs(favs)
            st.rerun()

    if favs:
        st.write("Mevcut:", ", ".join(favs))
        cikar = st.multiselect("Çıkarılacaklar", favs)
        if st.button("Seçilenleri çıkar") and cikar:
            core.save_favs([f for f in favs if f not in cikar])
            st.rerun()

        if st.button("Favorileri güncelle / alarm kontrol", type="primary"):
            analiz = get_analiz()
            with st.spinner("Favoriler çekiliyor…"):
                df = analiz.favorileri_guncelle()
            if not df.empty:
                cols = [c for c in [
                    "kod", "ad", "getiri_1y", "reel_1y", "sharpe", "max_drawdown",
                    "volatilite", "tefas_durum", "risk_degeri",
                ] if c in df.columns]
                show_df(df, cols, 350)
                download_excel_button(
                    df, f"favoriler_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", key="xl_fav",
                )
            if core.ALERT_LOG.exists():
                st.caption(f"Alarm log: {core.ALERT_LOG}")
                try:
                    log = core.ALERT_LOG.read_text(encoding="utf-8").strip().splitlines()[-15:]
                    st.code("\n".join(log) or "(boş)")
                except Exception:
                    pass
    else:
        st.info("Favori listesi boş.")


# ═══════════════════════════════════════════════════════════════════════════
# AYARLAR
# ═══════════════════════════════════════════════════════════════════════════
elif sayfa == "⚙️ Ayarlar":
    st.header("⚙️ Ayarlar")
    cfg = st.session_state.cfg

    c1, c2, c3 = st.columns(3)
    rf = c1.number_input("Risksiz faiz (0.40 = %40)", value=float(cfg.get("risk_free", 0.40)), min_value=0.0, max_value=1.0, step=0.01)
    tufe = c2.number_input("TÜFE yıllık %", value=float(cfg.get("tufe_yillik", 32.0)), min_value=0.0, step=0.5)
    mev = c3.number_input("Mevduat yıllık %", value=float(cfg.get("mevduat_yillik", 40.0)), min_value=0.0, step=0.5)

    st.subheader("Alarm eşikleri")
    al = cfg.get("alarmlar") or {}
    a1, a2, a3, a4 = st.columns(4)
    min_g = a1.number_input("Min 1Y % alarm", value=al.get("min_getiri_1y"), step=1.0, placeholder="yok")
    min_s = a2.number_input("Min Sharpe alarm", value=al.get("min_sharpe"), step=0.1, placeholder="yok")
    max_d = a3.number_input("Max DD % alarm", value=al.get("max_drawdown"), step=1.0, placeholder="yok")
    max_v = a4.number_input("Max vol % alarm", value=al.get("max_volatilite"), step=1.0, placeholder="yok")

    if st.button("Kaydet", type="primary"):
        cfg["risk_free"] = rf
        cfg["tufe_yillik"] = tufe
        cfg["mevduat_yillik"] = mev
        cfg["alarmlar"] = {
            "min_getiri_1y": min_g,
            "min_sharpe": min_s,
            "max_drawdown": max_d,
            "max_volatilite": max_v,
        }
        core.save_config(cfg)
        st.session_state.cfg = cfg
        if st.session_state.analiz is not None:
            st.session_state.analiz.cfg = cfg
            st.session_state.analiz.rf = float(rf)
            st.session_state.analiz.bench = {}
        st.success(f"Kaydedildi → {core.CONFIG_PATH}")

    st.divider()
    st.subheader("Cache")
    cache_files = list(core.CACHE_DIR.glob("*.json")) if core.CACHE_DIR.exists() else []
    st.caption(f"{len(cache_files)} cache dosyası · {core.CACHE_DIR}")
    if st.button("Cache temizle") and cache_files:
        for f in cache_files:
            try:
                f.unlink()
            except Exception:
                pass
        st.success("Cache silindi.")
        st.rerun()


# ── Alt bilgi (her sayfada) ─────────────────────────────────────────────────
st.markdown(
    """
<div class="disclaimer">
  <b>Uyarı:</b> Bu uygulama bilgilendirme amaçlıdır; yatırım tavsiyesi değildir.
  Geçmiş performans geleceği garanti etmez. TEFAS / Yahoo verileri gecikmeli olabilir.
  Kararlarınızı kendi risk tercihinize ve profesyonel danışmana göre verin.
</div>
""",
    unsafe_allow_html=True,
)
